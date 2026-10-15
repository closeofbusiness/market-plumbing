#!/usr/bin/env python3
"""Weekly Commitments-of-Traders ledger for COMEX gold, NYMEX WTI crude and COMEX copper, from the CFTC's free
Disaggregated COT report -- positions by trader class, the zero-sum identity, physical units, dollars, and the weekly
change in managed-money dollar exposure split into a VOLUME and a PRICE effect. Part of E-021 ("where does the money in
commodities come from?"). A futures long funds nothing physical: only margin moves, and every long has a short; this ledger
measures who holds the two sides, not who paid whom. Needs xlrd (EIA .xls) and openpyxl (SPDR .xlsx), as requirements.txt says.

  python3 bin/pull_cftc_cot.py             # pull from the CFTC and the price sources, write every output, print the report
  python3 bin/pull_cftc_cot.py --offline   # rebuild every output from the saved raw files (no network); byte-identical
Read LICENSING below before committing anything this script writes.

SOURCE. CFTC Public Reporting Environment, Socrata API (https://publicreporting.cftc.gov/resource/<id>.json), no key:
  72hh-3qpy  "Disaggregated Futures Only"            (report type 'fo')
  kh3c-gbw2  "Disaggregated Futures and Options Combined"  (report type 'combined'; options are counted on a futures-equivalent
             basis with exchange delta factors, so its figures are rounded and the identities below hold only to rounding)
Anonymous use; the script sends the neutral User-Agent below and makes one request per contract per dataset with a pause
between them. If a server throttles or refuses, the script retries only where the server says to (429/503, honouring
Retry-After), then stops; it never works around a block. Rows are selected by cftc_contract_market_code only (the dataset's
names change over time: CRUDE OIL, LIGHT SWEET became WTI-PHYSICAL, COPPER-GRADE #1 became COPPER- #1; the codes do not).
Both datasets run weekly from 2006-06-13; this ledger starts at the first 2010 report, 2010-01-05, because the brief asks for 2010
onward (the GLD and EIA series start earlier; the Westmetall pages are fetched from 2010). The dataset's own field name for swap-dealer short is 'swap__positions_short_all' (two underscores).

CONTRACTS (code; size read from the data's own contract_units field and asserted, on every row, against this table):
  088691 COMEX GOLD    100 troy ounces        067651 NYMEX WTI (WTI-PHYSICAL)   1,000 barrels       085692 COMEX COPPER #1   25,000 pounds

TRADER CLASSES (CFTC Disaggregated explanatory notes): Producer/Merchant/Processor/User (pm) = commercial hedgers of the physical
commodity; Swap Dealers (sw) = dealers hedging swap books (in gold this is largely bullion-bank hedging, a HYPOTHESIS not a
measurement); Managed Money (mm) = registered CTAs, CPOs and funds the CFTC identifies; Other Reportables (ot) = every other
reportable trader; Non-Reportables (nr) = open interest minus total reportable, i.e. a DERIVED RESIDUAL ("derived by subtracting total
long and short Reportable Positions from the total open interest", CFTC). Classes come from traders' Form 40 responses
(the E-021 map: "by predominant activity"), not from the position taken. Positions are as of Tuesday's close (published the following Friday).
SPREADING ("a computed amount equal to offsetting long and short positions held by a trader", CFTC) is reported once per class
for sw, mm and ot (pm has none) and belongs on BOTH sides: a class's gross long = long + spreading, gross short = short +
spreading. So a class's NET = long - short, and spreading drops out of every net. 'net' below is always long - short, excluding
spreading; a NEGATIVE net is a net short.

THE IDENTITY (stated exactly, and asserted on every row; the script stops if it fails). CFTC explanatory notes: "The aggregate of
all long open interest is equal to the aggregate of all short open interest." With the CFTC's definitions this gives, on each row:
    (1) pm_long + sw_long + sw_spread + mm_long + mm_spread + ot_long + ot_spread  ==  tot_rept_long     (reported)
        pm_short + sw_short + sw_spread + mm_short + mm_spread + ot_short + ot_spread == tot_rept_short  (reported)
    (2) tot_rept_long + nr_long == open_interest == tot_rept_short + nr_short
    =>  sum over the five classes of (long - short)  ==  0       (the zero-sum identity; every long has a short)
  * Futures only: (1), (2) and the zero-sum hold EXACTLY, to the contract, on every row. Asserted exactly.
  * Combined: the same identities hold only to rounding. Options are converted to futures equivalents with delta factors and each
    reported figure is rounded to a whole contract. Each side sums eight rounded class figures against a rounded open interest,
    so the worst case is 9 x 0.5 = 4.5, i.e. at most 4 contracts per side, and at most 8 for the sum of nets. Asserted at those
    bounds; the observed maxima are printed (3 contracts on the net sum in the 2010-2026 data). The cause (rounding) is inferred
    from the size and the symmetric sign pattern, not stated by the CFTC.
  * Non-reportables are a residual by the CFTC's construction, so (2) is true for them by definition. The identity's independent
    content is (1): the reported classes, with spreading counted on both sides, add up to the CFTC's own reportable totals.
  * One-sided only: a class being net long says nothing about who funded it. Zero-sum means the counterparties net out inside the
    market; the dollars that moved were margin and variation margin, which this report does not show.

DATES. Week = the dataset's report date: a Tuesday in 867 of 875 weeks since 2010; eight report dates are the MONDAY instead (2012-12-24,
2012-12-31, 2017-07-03, 2018-12-24, 2018-12-31, 2020-12-21, 2023-07-03, 2025-11-10). Kept as the CFTC dates them. The script checks there
is exactly one report in every ISO week, none missing and none doubled, and that the three contracts share the same dates.

PHYSICAL UNITS (constants below). Gold: 1 troy ounce = 31.1034768 g exactly, so 1 contract = 100 oz = 3.11034768 kg. Copper: 1 pound =
0.45359237 kg exactly, so 1 contract = 25,000 lb = 11.33980925 metric tonnes. WTI: 1 contract = 1,000 barrels (no tonnes: density varies).

LICENSING (supervisor decision, 12 Oct 2026). The repository is public and the principal has not ruled on third-party data. Therefore:
  (1) PUBLIC outputs -- data/cftc_cot/*.csv and report.txt (not raw/), and data/history/cftc_cot_*.csv -- contain NO gold or copper price or
      dollar column derived from GLD or from the Westmetall/LME mirror. They keep contracts, all class positions, open interest and
      physical units (oz, tonnes, bbl, kt). WTI price and dollar columns stay public: the EIA spot series is US-government public domain.
  (2) The gold and copper price-derived outputs -- prices, $bn, volume/price/cross splits, the gold and copper lead-lag correlations (they
      are computed from the licensed prices) and price_crosscheck -- are written to data/vintages/cot_prices/derived/ (gitignored) when
      the sponsor files are present, and skipped with a message when they are absent. The sponsor files are the GLD xlsx (its disclaimer
      prohibits reproduction or redistribution), the Westmetall pages and the World Bank workbook; they are inputs only, never committed.
  (3) report.txt is public and has the gold and copper price and dollar figures stripped; the full report, with them, is
      data/vintages/cot_prices/derived/report_full.txt (local).
  (4) The public outputs are a function of the CFTC and EIA raw files alone, so they are byte-identical whether or not the sponsor files
      are present. Every run ends by checking that no public file carries a gold or copper price or dollar column and stops if one does.

DOLLARS. Exposure at a report date = net contracts x units per contract x that week's price. Price per week, on the report date (Tuesday),
carried back to the last earlier day only where the source has no quote that day (the lag in days is a ledger column; the script stops
if it exceeds 5 days). Gold and copper dollars exist only in the local derived/ files (see LICENSING); WTI dollars are public:
  gold    GLD "Total Net Asset Value in the Trust" / "Total Ounces of Gold in the Trust", $ per troy ounce, from the SPDR Gold Shares
          historical archive (https://api.spdrgoldshares.com/api/v1/historical-archive?product=gld&exchange=NYSE&lang=en, an .xlsx,
          no key). The sponsor removed the LBMA price column at ICE Benchmark Administration's request, so this is the trust's own
          valuation, which the file's notes say is struck from the LBMA Gold Price PM (the London fix before 20 Mar 2015, and the next
          available AM or PM price on days with no PM price). It is a proxy for that benchmark, not the benchmark. The file's disclaimer
          prohibits reproduction or redistribution without written consent: it is kept only under data/vintages/cot_prices/ (gitignored)
          and is NOT committed; the per-week gold price derived from it is local too (see LICENSING).
          Cross-check asserted: NAV-per-share / ounces-per-share agrees with the trust-level ratio within 0.5% (observed max 0.34%).
  wti     EIA "Cushing, OK WTI Spot Price FOB", $ per barrel, the key-free dnav file https://www.eia.gov/dnav/pet/hist_xls/RWTCd.xls
          (US government, public domain; saved under data/cftc_cot/raw/). A spot assessment, not the NYMEX front-month settlement; the
          two track closely but are not identical, and managed money holds the whole curve.
  copper  LME copper cash-settlement ("Official LME-Prices in US Dollar ... Settlement Kasse", $ per metric tonne), from the third-party mirror westmetall.com (one table page per year,
          https://www.westmetall.com/en/markdaten.php?action=table&field=LME_Cu_cash&year=<yyyy>). This is the LME price, NOT the COMEX
          price of the contract being valued. A free, key-free, scripted COMEX series was not found: CME's own pages refused scripted access (HTTP 403,
          per the E-021 map) and stooq.com answered with a JavaScript bot challenge; neither was pursued. Every copper dollar column is therefore
          marked '_lme' and the local copper ledger carries usd_basis = 'LME_cash_via_Westmetall_not_COMEX'. Third-party; kept under
          data/vintages/cot_prices/ (gitignored), not committed.
  check   Each daily series is averaged by month and compared with the World Bank Commodity Price Data ("Pink Sheet"), monthly xlsx,
          which the script finds from https://www.worldbank.org/en/research/commodity-markets: gold (LBMA PM), copper (LME cash), WTI.
          Asserted: mean absolute difference <= 1%, maximum <= 5%; the actual figures are printed in the full report and written to derived/price_crosscheck.csv (local).
  TIMING. COT positions are as of Tuesday's close of business. None of the three price sources is stamped at that moment: the GLD NAV column
  is headed 10:30am New York (so about the London PM fix), the LME cash price is a London official price, and the EIA file states no time. The
  offsets are hours, not days, but they put part of Tuesday's move in the "next week" return and so slightly blur the same-week / next-week test.

THE SPLIT. For managed-money net contracts N, units per contract U and price P (all at report dates t-1 and t):
    total change   = N_t U P_t - N_{t-1} U P_{t-1}
    VOLUME effect  = (N_t - N_{t-1}) U P_{t-1}          change in contracts, at the PRIOR price
    PRICE effect   = N_{t-1} U (P_t - P_{t-1})           prior contracts, on the change in price
    CROSS term     = (N_t - N_{t-1}) U (P_t - P_{t-1})  both moved
  volume + price + cross == total is an algebraic identity; the script checks it every week (max residual printed, in dollars).
  Multi-week changes are the SUM of the weekly effects (chained), which also adds to the total exactly; the endpoint-only split of the
  same window (using the window's first price and first position) differs and is written beside it (wti_usd_windows.csv for WTI, derived/windows_usd.csv for gold and copper). Dollar exposure is
  contracts x units x price: a notional at market, not a cash outlay (a futures long posts margin, not the notional).

LEAD-LAG TEST (2024-01 onward). Pearson and Spearman correlation of the weekly change in managed-money net contracts with (a) the same week's
price change and (b) the NEXT week's price change, price change as a percentage of the prior price; (c) the PREVIOUS week's, an extra, to see
whether positions follow price (trend-following) rather than lead it. WTI is public (leadlag.csv); gold and copper are local (derived/). n is printed; for independent weeks
|r| above about 1.96/sqrt(n) is distinguishable from zero at 5%. Correlation, not cause: trend-following and price pressure cannot be separated.

OUTPUTS (data/history/ files are the repo's disposable cache, overwritten each run; two columns, as_of,value; K in {fo,combined}).
PUBLIC (a function of the CFTC and EIA raw files alone):
  data/history/cftc_cot_<gold|wti|copper>_<prodmerc|swap|mm|other|nonrep>_net_<K>.csv       net contracts (long - short) by class
  data/history/cftc_cot_<c>_open_interest_<K>.csv, _mm_long_<K>.csv, _mm_short_<K>.csv       contracts
  data/history/cftc_cot_<c>_mm_net_<troy_oz|tonnes|bbl>_<K>.csv                             managed-money net, physical units
  data/history/cftc_cot_wti_mm_net_usd_bn_<K>.csv                                           WTI only: managed-money net, $bn at EIA spot
  data/cftc_cot/<gold|copper>_ledger.csv   every position and physical column, both report types side by side (fo_ and co_), one row per week
  data/cftc_cot/wti_ledger.csv             the same plus the WTI price and every dollar column and the volume/price/cross split
  data/cftc_cot/checkpoints.csv, windows.csv   managed money at end-2023, end-2024, end-2025, 2026-06-30 and the latest week; changes over 4, 12,
                                           52 weeks and since end-2023; contracts, physical units and share of open interest, all contracts
  data/cftc_cot/wti_usd_checkpoints.csv, wti_usd_windows.csv   WTI price, dollars and the chained and endpoint volume/price/cross split
  data/cftc_cot/leadlag.csv (WTI only), report.txt (gold and copper dollar figures stripped; without the mode line)
  data/cftc_cot/raw/*.csv, eia_RWTCd.xls, MANIFEST.tsv    the exact CFTC rows and EIA file pulled (public domain; --offline rebuilds from them)
LOCAL ONLY, gitignored, written when the sponsor files are present (data/vintages/cot_prices/):
  *.xlsx, westmetall_*.html                                the third-party price files (inputs)
  derived/<gold|copper>_ledger_usd.csv                     report_date + price, price date, lag, usd_basis, every $bn column, split (join on report_date)
  derived/checkpoints_usd.csv, windows_usd.csv             gold and copper price, dollars and volume/price/cross split
  derived/leadlag_gold_copper.csv, price_crosscheck.csv, report_full.txt
  derived/history/cftc_cot_gold_mm_net_usd_bn_<K>.csv, cftc_cot_copper_mm_net_usd_bn_lme_<K>.csv
"""
import os, re, sys, csv, json, ssl, time, math, bisect, hashlib, argparse, datetime, urllib.request, urllib.parse, urllib.error

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HIST = os.path.join(ROOT, "data", "history")
OUT = os.path.join(ROOT, "data", "cftc_cot")
RAWD = os.path.join(OUT, "raw")
VINT = os.path.join(ROOT, "data", "vintages", "cot_prices")
DER = os.path.join(VINT, "derived")                 # LOCAL, gitignored: everything derived from the non-government price files
DER_HIST = os.path.join(DER, "history")
LICENSED = ("gold", "copper")                       # price series from non-government sources: no price or dollar column in any public output
UA = {"User-Agent": "Mozilla/5.0 (compatible; ThirdDerivativeResearch/1.0; non-commercial research)", "Accept-Encoding": "identity"}
CFTC_URL = "https://publicreporting.cftc.gov/resource/{}.json"
START = "2010-01-01"
LATE_SPLIT = "2024-01-01"            # lead-lag sample: report dates from here
MAX_PRICE_LAG_DAYS = 5

DATASETS = [("fo", "72hh-3qpy", "FutOnly", "cftc_disagg_futonly_72hh-3qpy.csv"),
            ("combined", "kh3c-gbw2", "Combined", "cftc_disagg_combined_kh3c-gbw2.csv")]
KINDS = ["fo", "combined"]
KP = {"fo": "fo_", "combined": "co_"}          # ledger column prefixes

G_PER_TROY_OZ = 31.1034768            # exact by definition
KG_PER_LB = 0.45359237                # exact by definition
GOLD_OZ, WTI_BBL, COPPER_LB = 100, 1000, 25000
COPPER_T = COPPER_LB * KG_PER_LB / 1000.0          # 11.33980925 t per contract

# slug, CFTC contract market code, label, (contract_units unit text, size), price quantity per contract, price column name, usd tag,
# physical conversions [(suffix, per-contract factor)], plausible price range (unit guard), print unit (suffix, scale, label)
CONTRACTS = [
    dict(slug="gold", code="088691", label="COMEX GOLD", units=("TROY OUNCES", GOLD_OZ), qty=float(GOLD_OZ), pcol="price_usd_per_oz", tag="",
         phys=[("troy_oz", float(GOLD_OZ)), ("tonnes", GOLD_OZ * G_PER_TROY_OZ / 1e6)], prange=(800.0, 8000.0), disp=("tonnes", 1.0, "t"),
         price_label="GLD trust NAV/oz (LBMA PM based), $/troy oz"),
    dict(slug="wti", code="067651", label="NYMEX WTI", units=("BARRELS", WTI_BBL), qty=float(WTI_BBL), pcol="price_usd_per_bbl", tag="",
         phys=[("bbl", float(WTI_BBL))], prange=(-50.0, 200.0), disp=("bbl", 1e-6, "Mbbl"),
         price_label="EIA Cushing WTI spot, $/bbl"),
    dict(slug="copper", code="085692", label="COMEX COPPER", units=("POUNDS", COPPER_LB), qty=COPPER_T, pcol="price_usd_per_tonne_lme_cash", tag="_lme",
         phys=[("tonnes", COPPER_T)], prange=(3000.0, 25000.0), disp=("tonnes", 1e-3, "kt"),
         price_label="LME cash via Westmetall (NOT COMEX), $/tonne"),
]
USD_BASIS = {"gold": "GLD_NAV_per_oz_proxy_for_LBMA_PM", "wti": "EIA_Cushing_spot_not_NYMEX_settle", "copper": "LME_cash_via_Westmetall_not_COMEX"}
SLUGS = [c["slug"] for c in CONTRACTS]
CBY = {c["slug"]: c for c in CONTRACTS}

# class key, history slug, label, has spreading
CLASSES = [("pm", "prodmerc", "Producer/Merchant/Processor/User", False), ("sw", "swap", "Swap Dealers", True), ("mm", "mm", "Managed Money", True),
           ("ot", "other", "Other Reportables", True), ("nr", "nonrep", "Non-Reportables (residual)", False)]

# raw column -> API field
FIELDS = [
    ("id", "id"), ("report_date", "report_date_as_yyyy_mm_dd"), ("code", "cftc_contract_market_code"), ("name", "market_and_exchange_names"),
    ("units", "contract_units"), ("kind", "futonly_or_combined"), ("oi", "open_interest_all"),
    ("pm_long", "prod_merc_positions_long"), ("pm_short", "prod_merc_positions_short"),
    ("sw_long", "swap_positions_long_all"), ("sw_short", "swap__positions_short_all"), ("sw_spread", "swap__positions_spread_all"),
    ("mm_long", "m_money_positions_long_all"), ("mm_short", "m_money_positions_short_all"), ("mm_spread", "m_money_positions_spread"),
    ("ot_long", "other_rept_positions_long"), ("ot_short", "other_rept_positions_short"), ("ot_spread", "other_rept_positions_spread"),
    ("totrep_long", "tot_rept_positions_long_all"), ("totrep_short", "tot_rept_positions_short"),
    ("nr_long", "nonrept_positions_long_all"), ("nr_short", "nonrept_positions_short_all"),
]
TEXT_COLS = ("id", "report_date", "code", "name", "units", "kind")
INTS = [c for c, _ in FIELDS if c not in TEXT_COLS]

# combined-report rounding tolerance (contracts); futures only is exact
TOL_SIDE, TOL_NET = 4, 8

EIA_URL = "https://www.eia.gov/dnav/pet/hist_xls/RWTCd.xls"
GLD_URL = "https://api.spdrgoldshares.com/api/v1/historical-archive?product=gld&exchange=NYSE&lang=en"
WM_URL = "https://www.westmetall.com/en/markdaten.php?action=table&field=LME_Cu_cash&year={}"
WB_PAGE = "https://www.worldbank.org/en/research/commodity-markets"
RAW_EIA = os.path.join(RAWD, "eia_RWTCd.xls")
GLD_X = os.path.join(VINT, "gld_US_GLD_Archive_EN.xlsx")
WB_X = os.path.join(VINT, "worldbank_CMO-Historical-Data-Monthly.xlsx")
WM_FIRST = 2010
MANIFEST = os.path.join(RAWD, "MANIFEST.tsv")
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]

MAN = []          # (relpath, url, fetched_utc, sha256, bytes) for every file fetched this run


def fetch(url, tries=4):
    for k in range(tries):
        try:
            r = urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=120, context=ssl.create_default_context())
            return r.read()
        except urllib.error.HTTPError as e:
            if e.code in (429, 503) and k < tries - 1:
                wait = int(e.headers.get("Retry-After") or 0) or 15 * (k + 1)
                print(f"  HTTP {e.code}; the server asked for a pause; waiting {wait}s", file=sys.stderr); time.sleep(wait); continue
            raise SystemExit(f"STOP: {url} returned HTTP {e.code}. Not retrying further and not working around it.")
        except (urllib.error.URLError, TimeoutError, ConnectionError) as e:
            if k < tries - 1: time.sleep(10 * (k + 1)); continue
            raise SystemExit(f"STOP: network error for {url}: {e}")


def save(path, data, url):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "wb") as f: f.write(data)
    MAN.append((os.path.relpath(path, ROOT), url, datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
                hashlib.sha256(data).hexdigest(), len(data)))


def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""): h.update(chunk)
    return h.hexdigest()


def pull_cftc():
    select = ",".join(a for _, a in FIELDS)
    for kind, dsid, label, fname in DATASETS:
        rows = []
        for c in CONTRACTS:
            q = {"$select": select, "$where": f"cftc_contract_market_code = '{c['code']}' AND report_date_as_yyyy_mm_dd >= '{START}T00:00:00.000'",
                 "$order": "report_date_as_yyyy_mm_dd", "$limit": "5000"}
            url = CFTC_URL.format(dsid) + "?" + urllib.parse.urlencode(q)
            d = json.loads(fetch(url))
            assert 0 < len(d) < 5000, f"{dsid} {c['code']}: {len(d)} rows (page limit is 5000; add paging if reached)"
            print(f"  {dsid} {c['code']} {c['label']}: {len(d)} weekly rows")
            for r in d: rows.append({col: (r.get(a) or "").strip() for col, a in FIELDS})
            time.sleep(1.0)
        for r in rows: r["report_date"] = r["report_date"][:10]
        path = os.path.join(RAWD, fname); os.makedirs(RAWD, exist_ok=True)
        with open(path, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=[col for col, _ in FIELDS], lineterminator="\n"); w.writeheader(); w.writerows(rows)
        data = open(path, "rb").read()
        MAN.append((os.path.relpath(path, ROOT), CFTC_URL.format(dsid) + "  [selected fields, contract codes " + ",".join(c["code"] for c in CONTRACTS) + f", report_date >= {START}]",
                    datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"), hashlib.sha256(data).hexdigest(), len(data)))


def pull_prices():
    print("  EIA RWTCd.xls"); save(RAW_EIA, fetch(EIA_URL), EIA_URL); time.sleep(1.0)
    print("  SPDR GLD historical archive"); save(GLD_X, fetch(GLD_URL), GLD_URL); time.sleep(1.0)
    for y in range(WM_FIRST, datetime.date.today().year + 1):
        print(f"  Westmetall LME copper cash {y}")
        save(os.path.join(VINT, f"westmetall_LME_Cu_cash_{y}.html"), fetch(WM_URL.format(y)), WM_URL.format(y)); time.sleep(1.0)
    print("  World Bank commodity-markets page, then the monthly xlsx it links")
    page = fetch(WB_PAGE).decode("utf-8", "replace"); time.sleep(1.0)
    links = sorted(set(re.findall(r"https://thedocs\.worldbank\.org/en/doc/[0-9a-f]+-\d+/related/CMO-Historical-Data-Monthly\.xlsx", page)))
    if len(links) != 1: raise SystemExit(f"STOP: expected exactly one CMO-Historical-Data-Monthly.xlsx link on {WB_PAGE}, found {len(links)}")
    save(WB_X, fetch(links[0]), links[0])


def write_manifest():
    os.makedirs(RAWD, exist_ok=True)
    with open(MANIFEST, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter="\t", lineterminator="\n"); w.writerow(["path", "source_url", "fetched_utc", "sha256", "size_bytes"]); w.writerows(MAN)


def verify_manifest():
    """Offline: every file in the manifest must match its recorded sha256 (the raw snapshot is the vintage). A missing PUBLIC raw file is
    fatal. A missing third-party file under data/vintages/ (gitignored, never committed) is not: it is returned so the caller can skip the
    outputs that need it. Returns (files verified, missing vintage paths, every path the manifest lists)."""
    if not os.path.exists(MANIFEST): raise SystemExit(f"STOP: {os.path.relpath(MANIFEST, ROOT)} missing; run once online first")
    n, missing, listed = 0, [], []
    with open(MANIFEST, newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f, delimiter="\t"):
            p = os.path.join(ROOT, r["path"]); listed.append(r["path"])
            if not os.path.exists(p):
                if r["path"].startswith("data/vintages/"): missing.append(r["path"]); continue
                raise SystemExit(f"STOP: {r['path']} is missing; it is a public raw file, restore it from the repository")
            if sha256_file(p) != r["sha256"]: raise SystemExit(f"STOP: {r['path']} does not match the sha256 recorded in the manifest")
            n += 1
    return n, missing, listed


def availability(missing, listed):
    """Which third-party price inputs are present. Prints a clean SKIPPED line for each that is not."""
    rel = lambda p: os.path.relpath(p, ROOT)
    wm = [x for x in listed if os.path.basename(x).startswith("westmetall_")]
    avail = {"gold": rel(GLD_X) in listed and rel(GLD_X) not in missing,
             "copper": bool(wm) and not any(x in missing for x in wm),
             "wb": rel(WB_X) in listed and rel(WB_X) not in missing}
    why = {"gold": f"gold price and dollar outputs: {rel(GLD_X)} is not present",
           "copper": "copper price and dollar outputs: the Westmetall pages under data/vintages/cot_prices/ are not all present",
           "wb": f"price cross-check and the full report: {rel(WB_X)} is not present"}
    for k in avail:
        if not avail[k]:
            print(f"SKIPPED {why[k]} (third-party sponsor file: gitignored, never committed; run once online to fetch it). Public outputs are unaffected.")
    return avail


# ---------------------------------------------------------------- loaders
def load_raw():
    out = {}
    for kind, dsid, label, fname in DATASETS:
        with open(os.path.join(RAWD, fname), newline="", encoding="utf-8") as f: out[kind] = list(csv.DictReader(f))
    return out


def toint(r, c):
    v = r[c]
    if v == "" or not re.fullmatch(r"-?\d+", v): raise AssertionError(f"{r['code']} {r['report_date']} {c}: not an integer: {v!r}")
    return int(v)


def validate(raw):
    """Hard checks on the pulled rows; raises on any failure. Returns (by[kind][slug][date] -> row with ints, stats)."""
    code2slug = {c["code"]: c["slug"] for c in CONTRACTS}
    by = {k: {s: {} for s in SLUGS} for k in KINDS}
    stat = {k: dict(rows=0, max_long=0, max_short=0, max_net=0, max_tot_long=0, max_tot_short=0, max_nr_long=0, max_nr_short=0) for k in KINDS}
    want_kind = {k: lab for k, _, lab, _ in DATASETS}
    for kind in KINDS:
        for r in raw[kind]:
            assert r["code"] in code2slug, f"unexpected contract code {r['code']}"
            s = code2slug[r["code"]]
            assert r["kind"] == want_kind[kind], f"{kind} {s} {r['report_date']}: futonly_or_combined is {r['kind']!r}, expected {want_kind[kind]!r} (wrong dataset?)"
            for c in INTS: r[c] = toint(r, c)
            assert r["report_date"] not in by[kind][s], f"duplicate {kind} {s} {r['report_date']}"
            by[kind][s][r["report_date"]] = r
    tol = {"fo": (0, 0), "combined": (TOL_SIDE, TOL_NET)}
    for kind in KINDS:
        for s, d in by[kind].items():
            ds = sorted(d); cfg = CBY[s]
            assert ds and ds[0] >= START, (kind, s, ds[:1])
            wd = {x: datetime.date.fromisoformat(x).weekday() for x in ds}
            assert all(v in (0, 1) for v in wd.values()), f"{kind} {s}: a report date is neither a Tuesday nor a Monday"
            mon = lambda x: datetime.date.fromisoformat(x) - datetime.timedelta(days=wd[x])
            gaps = [(a, b) for a, b in zip(ds, ds[1:]) if (mon(b) - mon(a)).days != 7]
            assert not gaps, f"{kind} {s}: gaps or doubled weeks in the weekly series: {gaps[:5]}"
            tol_side, tol_net = tol[kind]
            for x in ds:
                r = d[x]
                # units: the contract size the CFTC reports must be the size we convert with
                m = re.fullmatch(r"\(CONTRACTS OF ([\d,]+) ([A-Z ]+)\)", r["units"])
                assert m and m.group(2) == cfg["units"][0] and int(m.group(1).replace(",", "")) == cfg["units"][1], \
                    f"{kind} {s} {x}: contract_units {r['units']!r} != {cfg['units']}"
                for c in INTS: assert r[c] >= 0, f"{kind} {s} {x}: negative {c}"
                tot = {}
                for side in ("long", "short"):
                    reportable = sum(r[f"{k}_{side}"] for k in ("pm", "sw", "mm", "ot")) + sum(r[f"{k}_spread"] for k in ("sw", "mm", "ot"))   # (1)
                    d1 = reportable - r[f"totrep_{side}"]
                    d2 = r[f"totrep_{side}"] + r[f"nr_{side}"] - r["oi"]                                                                          # (2)
                    tot[side] = reportable + r[f"nr_{side}"] - r["oi"]
                    assert abs(d1) <= tol_side, f"{kind} {s} {x}: classes + spreading != total reportable ({side}): off by {d1}"
                    assert abs(d2) <= tol_side, f"{kind} {s} {x}: total reportable + non-reportable != open interest ({side}): off by {d2}"
                    assert abs(tot[side]) <= tol_side, f"{kind} {s} {x}: sum of all classes ({side}) != open interest: off by {tot[side]}"
                    stat[kind][f"max_tot_{side}"] = max(stat[kind][f"max_tot_{side}"], abs(d1)); stat[kind][f"max_nr_{side}"] = max(stat[kind][f"max_nr_{side}"], abs(d2))
                    stat[kind][f"max_{side}"] = max(stat[kind][f"max_{side}"], abs(tot[side]))
                net = sum(r[f"{k}_long"] - r[f"{k}_short"] for k, *_ in CLASSES)
                assert net == tot["long"] - tot["short"], "internal: net identity algebra"
                assert abs(net) <= tol_net, f"{kind} {s} {x}: the five class nets do not sum to zero: {net}"
                stat[kind]["max_net"] = max(stat[kind]["max_net"], abs(net)); stat[kind]["rows"] += 1
    for kind in KINDS:
        assert len({frozenset(by[kind][s]) for s in SLUGS}) == 1, f"{kind}: the three contracts do not share the same report dates"
    assert set(by["fo"]["gold"]) == set(by["combined"]["gold"]), "futures-only and combined do not share the same report dates"
    for k in ("fo",): assert all(stat[k][m] == 0 for m in stat[k] if m.startswith("max_")), "futures-only identities must be exact"
    return by, stat


def load_gld():
    import openpyxl
    wb = openpyxl.load_workbook(GLD_X, read_only=True, data_only=True)
    ws = wb["US GLD Historical Archive"]; it = ws.iter_rows(values_only=True); hdr = list(next(it))
    want = {"date": "Date", "oz_share": "Ounces of Gold per Share", "nav_share": "NAV/Share at 10:30am NYT", "oz": "Total Ounces of Gold in the Trust", "nav": "Total Net Asset Value in the Trust"}
    ix = {k: hdr.index(v) for k, v in want.items()}                  # ValueError if the sponsor renames a column
    ser, worst = {}, (0.0, None)
    for r in it:
        if r[ix["date"]] is None: continue
        d = r[ix["date"]]; d = datetime.datetime.strptime(d, "%d-%b-%Y").date() if isinstance(d, str) else d.date()
        if r[ix["oz"]] == "US Holiday": continue
        oz, nav = r[ix["oz"]], r[ix["nav"]]
        assert isinstance(oz, (int, float)) and isinstance(nav, (int, float)) and oz > 0 and nav > 0, f"GLD {d}: non-numeric ounces/NAV {oz!r} {nav!r}"
        assert d not in ser, f"GLD duplicate date {d}"
        p = nav / oz; ser[d] = p
        ps, pz = r[ix["nav_share"]], r[ix["oz_share"]]
        if d >= datetime.date.fromisoformat(START) and isinstance(ps, (int, float)) and isinstance(pz, (int, float)) and pz > 0:
            dev = abs(p / (ps / pz) - 1)
            if dev > worst[0]: worst = (dev, d)
    assert worst[0] <= 0.005, f"GLD: trust-level NAV/oz and per-share NAV/oz differ by {worst[0]:.3%} on {worst[1]}"
    return ser, worst


def load_eia():
    import xlrd
    b = xlrd.open_workbook(RAW_EIA); s = b.sheet_by_name("Data 1")
    assert "WTI" in s.cell_value(0, 1) and "Dollars per Barrel" in s.cell_value(0, 1) and s.cell_value(2, 0) == "Date", "EIA RWTCd.xls layout changed"
    ser = {}
    for i in range(3, s.nrows):
        dv, v = s.row_values(i)[:2]
        if not isinstance(v, float): continue
        d = xlrd.xldate_as_datetime(dv, b.datemode).date(); assert d not in ser; ser[d] = v
    return ser


def load_westmetall():
    ser, dashes = {}, []
    WM_YEARS = sorted(int(m.group(1)) for fn in os.listdir(VINT) for m in [re.fullmatch(r"westmetall_LME_Cu_cash_(\d{4})\.html", fn)] if m)
    assert WM_YEARS and WM_YEARS == list(range(WM_FIRST, WM_YEARS[-1] + 1)), f"Westmetall pages are not an unbroken run of years from {WM_FIRST}: {WM_YEARS}"
    for y in WM_YEARS:
        t = open(os.path.join(VINT, f"westmetall_LME_Cu_cash_{y}.html"), encoding="utf-8").read()
        rows = re.findall(r'<td >(\d\d)\. (\w+) (\d{4})</td>\s*<td >([^<]*)</td>\s*<td >([^<]*)</td>\s*<td class="last">([^<]*)</td>', t)
        assert rows and len(rows) == len(re.findall(r"<td >\d\d\. ", t)), f"Westmetall {y}: row pattern did not match every table row"
        if y < WM_YEARS[-1]: assert len(rows) >= 240, f"Westmetall {y}: only {len(rows)} rows"
        for dd, mm, yy, cash, _m3, _stk in rows:
            assert int(yy) == y, f"Westmetall page {y} contains {yy}"
            d = datetime.date(y, MONTHS.index(mm) + 1, int(dd)); assert d not in ser and d not in dashes
            if cash.strip() == "-": dashes.append(d); continue          # no cash price quoted that day
            ser[d] = float(cash.replace(",", ""))
    return ser, dashes


def load_worldbank():
    import openpyxl
    wb = openpyxl.load_workbook(WB_X, read_only=True, data_only=True); rows = list(wb["Monthly Prices"].iter_rows(values_only=True))
    hdr, units = rows[4], rows[5]
    cols = {"wti": ("Crude oil, WTI", "($/bbl)"), "copper": ("Copper", "($/mt)"), "gold": ("Gold", "($/troy oz)")}
    ix = {}
    for k, (name, u) in cols.items():
        j = hdr.index(name); assert units[j] == u, f"World Bank {name}: units {units[j]!r} != {u!r}"; ix[k] = j
    out = {k: {} for k in cols}
    for r in rows[6:]:
        if r[0] and re.fullmatch(r"\d{4}M\d\d", str(r[0])):
            ym = (int(r[0][:4]), int(r[0][5:]))
            for k, j in ix.items():
                if isinstance(r[j], (int, float)): out[k][ym] = float(r[j])
    return out


def load_prices(avail):
    series, aux = {"wti": load_eia()}, {}
    if avail["gold"]: series["gold"], aux["gld_worst"] = load_gld()
    if avail["copper"]: series["copper"], aux["lme_dashes"] = load_westmetall()
    aux["wb"] = load_worldbank() if avail["wb"] else None
    return series, aux


def attach_prices(dates, series, slug):
    """Price at each report date; the last earlier quote where the source has none that day. Returns {date: (price, price_date, lag_days)}."""
    ks = sorted(series); lo, hi = CBY[slug]["prange"]; out = {}
    for d in dates:
        dd = datetime.date.fromisoformat(d); i = bisect.bisect_right(ks, dd) - 1
        assert i >= 0, f"{slug}: no price on or before {d}"
        lag = (dd - ks[i]).days; assert lag <= MAX_PRICE_LAG_DAYS, f"{slug} {d}: nearest earlier price is {lag} days old"
        p = series[ks[i]]; assert lo <= p <= hi, f"{slug} {d}: price {p} outside the plausible range {lo}..{hi} for its unit"
        out[d] = (p, ks[i].isoformat(), lag)
    return out


# ---------------------------------------------------------------- numbers
def f(v, nd):
    """Fixed-decimal text, '' for None, no negative zero, no scientific notation (so the outputs are byte-stable)."""
    if v is None: return ""
    s = format(round(v, nd), f".{nd}f")
    return s[1:] if s.startswith("-") and float(s) == 0 else s


def decompose(N, P, U):
    """Weekly volume / price / cross split of the change in N x U x P. N, P are lists aligned by week. Returns list of dicts (None first)."""
    out = [None]
    for t in range(1, len(N)):
        E1, E0 = N[t] * U * P[t], N[t - 1] * U * P[t - 1]
        dN, dP = N[t] - N[t - 1], P[t] - P[t - 1]
        v, p, x = dN * U * P[t - 1], N[t - 1] * U * dP, dN * U * dP
        res = (E1 - E0) - (v + p + x)
        scale = max(abs(E1), abs(E0), abs(v), abs(p), abs(x))
        assert abs(res) <= 1e-3 + 1e-12 * scale, f"volume + price + cross != total change at week {t}: residual ${res}"
        out.append(dict(E=E1, dE=E1 - E0, vol=v, price=p, cross=x, res=res, dN=dN, dP=dP))
    return out


def build(by, px):
    """All computed quantities. Returns L[kind][slug] = list of row dicts (one per week), the dates, and the largest volume+price+cross residual
    ($) per contract. Contracts with no price (px has no entry) get the position-only quantities; their price/dollar fields are None."""
    dates = sorted(by["fo"]["gold"]); L = {k: {} for k in KINDS}; maxres = {}
    for c in CONTRACTS:
        s = c["slug"]; P = [px[s][d][0] for d in dates] if s in px else None
        for kind in KINDS:
            N = [by[kind][s][d]["mm_long"] - by[kind][s][d]["mm_short"] for d in dates]
            dec = decompose(N, P, c["qty"]) if P else None
            if dec: maxres[s] = max([maxres.get(s, 0.0)] + [abs(x["res"]) for x in dec[1:]])
            rows = []
            for i, d in enumerate(dates):
                r = by[kind][s][d]
                o = {"report_date": d, "price": None if P is None else P[i], "price_date": None if P is None else px[s][d][1], "lag": None if P is None else px[s][d][2]}
                o["oi"] = r["oi"]
                for k, _, _, has_sp in CLASSES:
                    o[f"{k}_long"], o[f"{k}_short"] = r[f"{k}_long"], r[f"{k}_short"]
                    if has_sp: o[f"{k}_spread"] = r[f"{k}_spread"]
                    o[f"{k}_net"] = r[f"{k}_long"] - r[f"{k}_short"]
                o["totrep_long"], o["totrep_short"] = r["totrep_long"], r["totrep_short"]
                o["sum_long"] = sum(r[f"{k}_long"] for k, *_ in CLASSES) + sum(r[f"{k}_spread"] for k, _, _, sp in CLASSES if sp)
                o["sum_short"] = sum(r[f"{k}_short"] for k, *_ in CLASSES) + sum(r[f"{k}_spread"] for k, _, _, sp in CLASSES if sp)
                o["resid_long"], o["resid_short"] = o["sum_long"] - o["oi"], o["sum_short"] - o["oi"]
                o["resid_net"] = sum(o[f"{k}_net"] for k, *_ in CLASSES)
                o["mm_net_pct_oi"], o["mm_long_pct_oi"], o["mm_short_pct_oi"] = 100.0 * o["mm_net"] / o["oi"], 100.0 * o["mm_long"] / o["oi"], 100.0 * o["mm_short"] / o["oi"]
                o["mm_dN"] = None if i == 0 else N[i] - N[i - 1]
                for key in ("dE", "vol", "price", "cross", "res"): o["mm_" + key] = None if dec is None or dec[i] is None else dec[i][key]
                rows.append(o)
            L[kind][s] = rows
    return dates, L, maxres


def window(rows, i0, i1, U):
    """Change in managed-money net, and (if the contract has a price) in its dollar exposure between week i0 and i1: chained (sum of weekly
    effects) and endpoint-only split. Returns (chained, endpoint, total, dN, dP); the first three and dP are None without a price."""
    a, b = rows[i0], rows[i1]; dN = b["mm_net"] - a["mm_net"]
    if a["price"] is None: return None, None, None, dN, None
    ch = {k: sum(rows[t]["mm_" + k] for t in range(i0 + 1, i1 + 1)) for k in ("dE", "vol", "price", "cross")}
    tot = b["mm_net"] * U * b["price"] - a["mm_net"] * U * a["price"]
    tol = 1e-2 + 1e-12 * max(abs(tot), abs(ch["vol"]), abs(ch["price"]), abs(ch["cross"]))
    assert abs(tot - ch["dE"]) <= tol, f"chained weekly effects do not add to the window total: {tot} vs {ch['dE']}"
    assert abs(ch["vol"] + ch["price"] + ch["cross"] - tot) <= tol, "window split does not add to the total"
    dP = b["price"] - a["price"]
    ep = dict(vol=dN * U * a["price"], price=a["mm_net"] * U * dP, cross=dN * U * dP)
    return ch, ep, tot, dN, dP


# ---------------------------------------------------------------- outputs
def hist(name, pairs, nd=None, base=HIST):
    os.makedirs(base, exist_ok=True)
    with open(os.path.join(base, f"cftc_cot_{name}.csv"), "w", encoding="utf-8", newline="") as fh:
        fh.write("as_of,value\n")
        for d, v in pairs: fh.write(f"{d},{v if nd is None else f(v, nd)}\n")


def write_csv(path, rows):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(dict.fromkeys(k for r in rows for k in r)), lineterminator="\n"); w.writeheader(); w.writerows(rows)


def is_usd_col(col, c):
    """A ledger column that carries a price or a dollar figure (or says which price)."""
    return col in ("price_date", "price_lag_days", c["pcol"], "usd_basis") or "_usd_bn" in col or col.endswith("_usd")


def ledger_columns(c):
    cols = ["report_date", "price_date", "price_lag_days", c["pcol"], "usd_basis"]
    for kind in KINDS:
        p = KP[kind]; cols.append(p + "oi")
        for k, _, _, has_sp in CLASSES:
            cols += [f"{p}{k}_long", f"{p}{k}_short"] + ([f"{p}{k}_spread"] if has_sp else []) + [f"{p}{k}_net"]
        cols += [p + "totrep_long", p + "totrep_short", p + "sum_long", p + "sum_short", p + "resid_long", p + "resid_short", p + "resid_net"]
        cols += [f"{p}{k}_net_{u}" for k, *_ in CLASSES for u, _ in c["phys"]] + [f"{p}oi_{u}" for u, _ in c["phys"]]
        cols += [f"{p}{k}_net_usd_bn{c['tag']}" for k, *_ in CLASSES] + [f"{p}oi_usd_bn{c['tag']}"]
        cols += [p + "mm_net_pct_oi", p + "mm_long_pct_oi", p + "mm_short_pct_oi"]
        cols += [p + "mm_net_chg", f"{p}mm_exp_chg_usd_bn{c['tag']}", f"{p}mm_vol_usd_bn{c['tag']}", f"{p}mm_price_usd_bn{c['tag']}",
                 f"{p}mm_cross_usd_bn{c['tag']}", p + "mm_split_resid_usd"]
    return cols


def ledger_row(c, L, i):
    """Every ledger column for week i; the price and dollar columns only if the contract has a price."""
    o0 = L["fo"][c["slug"]][i]; priced = o0["price"] is not None
    out = {"report_date": o0["report_date"]}
    if priced: out.update({"price_date": o0["price_date"], "price_lag_days": str(o0["lag"]), c["pcol"]: f(o0["price"], 4), "usd_basis": USD_BASIS[c["slug"]]})
    for kind in KINDS:
        o = L[kind][c["slug"]][i]; p = KP[kind]; out[p + "oi"] = str(o["oi"])
        for k, _, _, has_sp in CLASSES:
            out[f"{p}{k}_long"], out[f"{p}{k}_short"], out[f"{p}{k}_net"] = str(o[f"{k}_long"]), str(o[f"{k}_short"]), str(o[f"{k}_net"])
            if has_sp: out[f"{p}{k}_spread"] = str(o[f"{k}_spread"])
        for key in ("totrep_long", "totrep_short", "sum_long", "sum_short", "resid_long", "resid_short", "resid_net"): out[p + key] = str(o[key])
        for k, *_ in CLASSES:
            for u, fac in c["phys"]: out[f"{p}{k}_net_{u}"] = f(o[f"{k}_net"] * fac, 4)
            if priced: out[f"{p}{k}_net_usd_bn{c['tag']}"] = f(o[f"{k}_net"] * c["qty"] * o["price"] / 1e9, 6)
        for u, fac in c["phys"]: out[f"{p}oi_{u}"] = f(o["oi"] * fac, 4)
        if priced: out[f"{p}oi_usd_bn{c['tag']}"] = f(o["oi"] * c["qty"] * o["price"] / 1e9, 6)
        for key in ("mm_net_pct_oi", "mm_long_pct_oi", "mm_short_pct_oi"): out[p + key] = f(o[key], 4)
        out[p + "mm_net_chg"] = "" if o["mm_dN"] is None else str(o["mm_dN"])
        if priced:
            for key, nm in (("dE", "exp_chg"), ("vol", "vol"), ("price", "price"), ("cross", "cross")):
                out[f"{p}mm_{nm}_usd_bn{c['tag']}"] = f(None if o["mm_" + key] is None else o["mm_" + key] / 1e9, 6)
            out[p + "mm_split_resid_usd"] = f(o["mm_res"], 2)
    return out


def write_ledgers(L, dates):
    """Public ledgers: WTI in full; gold and copper without any price or dollar column. Local: derived/<gold|copper>_ledger_usd.csv."""
    wrote_local = []
    for c in CONTRACTS:
        s = c["slug"]; cols = ledger_columns(c); priced = L["fo"][s][0]["price"] is not None
        usd_cols = [x for x in cols if is_usd_col(x, c)]
        pub_cols = [x for x in cols if x not in usd_cols] if s in LICENSED else cols
        rows = [ledger_row(c, L, i) for i in range(len(dates))]
        assert all(set(pub_cols) <= set(r) for r in rows), (s, "a public ledger column is missing")
        with open(os.path.join(OUT, f"{s}_ledger.csv"), "w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=pub_cols, extrasaction="ignore", lineterminator="\n"); w.writeheader(); w.writerows(rows)
        if s in LICENSED and priced:
            assert all(set(usd_cols) <= set(r) for r in rows), (s, "a local ledger column is missing")
            os.makedirs(DER, exist_ok=True)
            with open(os.path.join(DER, f"{s}_ledger_usd.csv"), "w", newline="", encoding="utf-8") as fh:
                w = csv.DictWriter(fh, fieldnames=["report_date"] + usd_cols, extrasaction="ignore", lineterminator="\n"); w.writeheader(); w.writerows(rows)
            wrote_local.append(f"{s}_ledger_usd.csv")
    return wrote_local


def write_history(L, dates):
    """Position and physical-unit history is public for all three; the dollar history is public for WTI only, local for gold and copper."""
    n = nl = 0
    for c in CONTRACTS:
        s = c["slug"]
        for kind in KINDS:
            R = L[kind][s]; K = kind
            for k, hs, *_ in CLASSES: hist(f"{s}_{hs}_net_{K}", [(r["report_date"], r[f"{k}_net"]) for r in R]); n += 1
            hist(f"{s}_open_interest_{K}", [(r["report_date"], r["oi"]) for r in R]); hist(f"{s}_mm_long_{K}", [(r["report_date"], r["mm_long"]) for r in R])
            hist(f"{s}_mm_short_{K}", [(r["report_date"], r["mm_short"]) for r in R]); n += 3
            for u, fac in c["phys"]:
                hist(f"{s}_mm_net_{u}_{K}", [(r["report_date"], r["mm_net"] * fac) for r in R], 4); n += 1
            if R[0]["price"] is not None:
                pairs = [(r["report_date"], r["mm_net"] * c["qty"] * r["price"] / 1e9) for r in R]
                if s in LICENSED: hist(f"{s}_mm_net_usd_bn{c['tag']}_{K}", pairs, 6, DER_HIST); nl += 1
                else: hist(f"{s}_mm_net_usd_bn{c['tag']}_{K}", pairs, 6); n += 1
    return n, nl


def checkpoint_dates(dates):
    qe = [("end-2023", "2023-12-31"), ("end-2024", "2024-12-31"), ("end-2025", "2025-12-31"), ("2026-06-30", "2026-06-30")]
    pts = [(lab, q, max(d for d in dates if d <= q)) for lab, q in qe] + [("latest", "", dates[-1])]
    return pts


def write_tables(L, dates):
    """checkpoints.csv and windows.csv (positions and physical units, all contracts, public); wti_usd_* (public) and derived/*_usd.csv (local)
    hold the price and dollar columns. Returns (cp, wd): the merged rows (position + price/dollar where the contract has a price) for printing."""
    pts = checkpoint_dates(dates); idx = {d: i for i, d in enumerate(dates)}
    cp_pub, wd_pub, cp_usd, wd_usd, cp, wd = [], [], [], [], [], []
    for c in CONTRACTS:
        s = c["slug"]; t = c["tag"]
        for kind in KINDS:
            R = L[kind][s]; priced = R[0]["price"] is not None
            for lab, q, d in pts:
                o = R[idx[d]]
                key = {"contract": s, "report_type": kind, "label": lab, "quarter_end": q, "report_date": d}
                pub = dict(key, mm_long=o["mm_long"], mm_short=o["mm_short"], mm_net=o["mm_net"])
                for u, fac in c["phys"]: pub[f"mm_net_{u}"] = f(o["mm_net"] * fac, 4)
                pub["mm_net_pct_oi"], pub["mm_long_pct_oi"], pub["mm_short_pct_oi"] = f(o["mm_net_pct_oi"], 2), f(o["mm_long_pct_oi"], 2), f(o["mm_short_pct_oi"], 2)
                pub["oi"] = o["oi"]
                for u, fac in c["phys"]: pub[f"oi_{u}"] = f(o["oi"] * fac, 2)
                cp_pub.append(pub); full = dict(pub)
                if priced:
                    usd = dict(key, price_date=o["price_date"], price=f(o["price"], 4))
                    usd["mm_net_usd_bn" + t] = f(o["mm_net"] * c["qty"] * o["price"] / 1e9, 4)
                    usd["oi_usd_bn" + t] = f(o["oi"] * c["qty"] * o["price"] / 1e9, 4)
                    cp_usd.append(usd); full.update(usd)
                cp.append(full)
            i1 = len(dates) - 1
            for wlab, i0 in (("4 weeks", i1 - 4), ("12 weeks", i1 - 12), ("52 weeks", i1 - 52), ("since end-2023", idx[pts[0][2]])):
                ch, ep, tot, dN, dP = window(R, i0, i1, c["qty"])
                key = {"contract": s, "report_type": kind, "window": wlab, "weeks": i1 - i0, "from_date": dates[i0], "to_date": dates[i1]}
                pub = dict(key, mm_net_from=R[i0]["mm_net"], mm_net_to=R[i1]["mm_net"], mm_net_chg=dN)
                for u, fac in c["phys"]: pub[f"mm_net_chg_{u}"] = f(dN * fac, 4)
                pub["mm_net_pct_oi_from"], pub["mm_net_pct_oi_to"] = f(R[i0]["mm_net_pct_oi"], 2), f(R[i1]["mm_net_pct_oi"], 2)
                wd_pub.append(pub); full = dict(pub)
                if priced:
                    usd = dict(key, price_from=f(R[i0]["price"], 4), price_to=f(R[i1]["price"], 4), price_chg_pct=f(100.0 * (R[i1]["price"] / R[i0]["price"] - 1), 3))
                    usd["exp_chg_usd_bn" + t] = f(tot / 1e9, 4)
                    usd["chained_vol_usd_bn" + t], usd["chained_price_usd_bn" + t], usd["chained_cross_usd_bn" + t] = f(ch["vol"] / 1e9, 4), f(ch["price"] / 1e9, 4), f(ch["cross"] / 1e9, 4)
                    usd["endpoint_vol_usd_bn" + t], usd["endpoint_price_usd_bn" + t], usd["endpoint_cross_usd_bn" + t] = f(ep["vol"] / 1e9, 4), f(ep["price"] / 1e9, 4), f(ep["cross"] / 1e9, 4)
                    usd["split_check_usd"] = f(ch["vol"] + ch["price"] + ch["cross"] - tot, 2)
                    wd_usd.append(usd); full.update(usd)
                wd.append(full)
    write_csv(os.path.join(OUT, "checkpoints.csv"), cp_pub); write_csv(os.path.join(OUT, "windows.csv"), wd_pub)
    write_csv(os.path.join(OUT, "wti_usd_checkpoints.csv"), [r for r in cp_usd if r["contract"] not in LICENSED])
    write_csv(os.path.join(OUT, "wti_usd_windows.csv"), [r for r in wd_usd if r["contract"] not in LICENSED])
    loc_cp, loc_wd = [r for r in cp_usd if r["contract"] in LICENSED], [r for r in wd_usd if r["contract"] in LICENSED]
    if loc_cp: write_csv(os.path.join(DER, "checkpoints_usd.csv"), loc_cp); write_csv(os.path.join(DER, "windows_usd.csv"), loc_wd)
    return cp, wd


def pearson(x, y):
    n = len(x); mx, my = sum(x) / n, sum(y) / n
    sxy = sum((a - mx) * (b - my) for a, b in zip(x, y)); sxx = sum((a - mx) ** 2 for a in x); syy = sum((b - my) ** 2 for b in y)
    return sxy / math.sqrt(sxx * syy)


def ranks(v):
    order = sorted(range(len(v)), key=lambda i: v[i]); r = [0.0] * len(v); i = 0
    while i < len(v):
        j = i
        while j + 1 < len(v) and v[order[j + 1]] == v[order[i]]: j += 1
        for k in range(i, j + 1): r[order[k]] = (i + j) / 2.0 + 1
        i = j + 1
    return r


def leadlag(L, dates, slugs):
    rows = []
    for c in CONTRACTS:
        if c["slug"] not in slugs: continue
        for kind in KINDS:
            R = L[kind][c["slug"]]
            t0 = next(i for i, d in enumerate(dates) if d >= LATE_SPLIT)
            ts = list(range(t0, len(dates)))
            dn = {t: R[t]["mm_net"] - R[t - 1]["mm_net"] for t in ts}
            ret = {t: R[t]["price"] / R[t - 1]["price"] - 1 for t in range(1, len(dates))}
            dp = {t: R[t]["price"] - R[t - 1]["price"] for t in range(1, len(dates))}
            same = ts; nxt = [t for t in ts if t + 1 < len(dates)]; prv = ts
            row = {"contract": c["slug"], "report_type": kind, "first_week": dates[ts[0]], "last_week": dates[ts[-1]]}
            for tag, tt, shift in (("same", same, 0), ("next", nxt, 1), ("prev", prv, -1)):
                x = [dn[t] for t in tt]; y = [ret[t + shift] for t in tt]; z = [dp[t + shift] for t in tt]
                row[f"n_{tag}"] = len(tt)
                row[f"pearson_{tag}_pct"] = f(pearson(x, y), 3); row[f"spearman_{tag}_pct"] = f(pearson(ranks(x), ranks(y)), 3); row[f"pearson_{tag}_abs"] = f(pearson(x, z), 3)
                row[f"band95_{tag}"] = f(1.96 / math.sqrt(len(tt)), 3)
            rows.append(row)
    return rows


def crosscheck(series, wb):
    """Monthly means of each daily price series against the World Bank Pink Sheet. Asserts gross agreement; LOCAL output (derived/price_crosscheck.csv)."""
    rows, summ = [], []
    for s in SLUGS:
        if s not in series: continue
        m = {}
        for d, v in series[s].items():
            if d >= datetime.date.fromisoformat(START): m.setdefault((d.year, d.month), []).append(v)
        diffs = []
        for ym in sorted(m):
            if ym in wb[s]:
                mine = sum(m[ym]) / len(m[ym]); w = wb[s][ym]; diffs.append((mine / w - 1, ym, mine, w, len(m[ym])))
                rows.append({"series": s, "month": f"{ym[0]}-{ym[1]:02d}", "days": len(m[ym]), "mine": f(mine, 3), "worldbank": f(w, 3), "pct_diff": f(100 * (mine / w - 1), 3)})
        ab = [abs(x[0]) for x in diffs]; mean_abs, mx = sum(ab) / len(ab), max(ab); worst = max(diffs, key=lambda x: abs(x[0]))
        assert mean_abs <= 0.01 and mx <= 0.05, f"{s}: monthly mean disagrees with the World Bank (mean abs {mean_abs:.2%}, max {mx:.2%})"
        summ.append(dict(series=s, n=len(diffs), mean_signed=100 * sum(x[0] for x in diffs) / len(diffs), mean_abs=100 * mean_abs, max_abs=100 * mx, worst=f"{worst[1][0]}-{worst[1][1]:02d}"))
    write_csv(os.path.join(DER, "price_crosscheck.csv"), rows)
    return summ


# ---------------------------------------------------------------- report
def report(by, stat, dates, L, px, aux, cp, wd, ll, cc, public):
    """The printed report. public=True: gold and copper show positions and physical units only (no price, no dollar, no price-derived statistic),
    and nothing depends on a third-party file. public=False: the full report, every contract with prices and dollars (LOCAL)."""
    out = []; P = out.append
    usd = lambda s: (not public) or s not in LICENSED
    n_rows = sum(stat[k]["rows"] for k in KINDS)
    P(f"validated: 3 contracts x {len(dates)} weekly reports ({dates[0]} to {dates[-1]}) x 2 report types = {n_rows:,} rows; one report in every ISO week, none missing; "
      f"contract sizes match the data's contract_units on every row")
    P("IDENTITY asserted on every row (CFTC: aggregate long open interest = aggregate short open interest; spreading counted on both sides; non-reportables = open interest minus reportable):")
    P("   sum(pm,sw,mm,ot long + spreading) = total reportable long; same short;  reportable + non-reportable = open interest on each side;  =>  sum over 5 classes of (long - short) = 0")
    for k, lab in (("fo", "futures only"), ("combined", "futures+options combined")):
        s = stat[k]
        if k == "fo":
            P(f"   {lab}: held EXACTLY on all {s['rows']:,} rows (maximum absolute difference anywhere = 0 contracts)")
        else:
            P(f"   {lab}: held within rounding on all {s['rows']:,} rows; maxima in contracts: classes vs reportable total long/short {s['max_tot_long']}/{s['max_tot_short']}, "
              f"reportable + non-reportable vs OI {s['max_nr_long']}/{s['max_nr_short']}, sum of all classes vs OI {s['max_long']}/{s['max_short']}, "
              f"sum of the five nets {s['max_net']} (allowed {TOL_SIDE} per side, {TOL_NET} for the net sum)")
    P("")
    P("PRICES, one per report date (carried back to the last earlier quote only where the source has none; lag stops the script above 5 days):")
    for c in CONTRACTS:
        s = c["slug"]
        if not usd(s): continue
        lags = [px[s][d][2] for d in dates]; nl = sum(1 for x in lags if x)
        P(f"   {s:6s} {c['price_label']}; {nl} of {len(dates)} weeks carried back (max {max(lags)} day(s)); range {min(px[s][d][0] for d in dates):,.2f} to {max(px[s][d][0] for d in dates):,.2f}")
        if nl: P(f"          carried back on report dates: {', '.join(d for d in dates if px[s][d][2])}")
    if public:
        P("   gold and copper: no price or dollar figure in the public outputs (LICENSING in the script's docstring). They are in data/vintages/cot_prices/derived/ (local, gitignored), with the full report, when the sponsor files are present.")
    else:
        P(f"   gold: GLD trust-level NAV/oz vs per-share NAV/oz, maximum disagreement {100 * aux['gld_worst'][0]:.2f}% (on {aux['gld_worst'][1]}); Westmetall had no cash quote on {len(aux['lme_dashes'])} days ({', '.join(d.isoformat() for d in aux['lme_dashes'])})")
        P("   cross-check of monthly mean of each daily series against the World Bank Pink Sheet (gold LBMA PM, copper LME cash, WTI), months since 2010-01 present in both:")
        for r in cc:
            P(f"      {r['series']:6s} n={r['n']} months  mean signed {r['mean_signed']:+.3f}%  mean absolute {r['mean_abs']:.3f}%  max absolute {r['max_abs']:.2f}% (month {r['worst']})")
    res = max(v for s_, v in aux["maxres"].items() if usd(s_))
    P(f"   split identity (volume + price + cross = total change) held every week; largest residual over all weeks and both report types ({'WTI' if public else 'all three contracts'}) = ${res:.6f}")
    P("")
    for c in CONTRACTS:
        s = c["slug"]; u, sc, ul = c["disp"]; tag = c["tag"]; su = usd(s)
        for kind in KINDS:
            kl = "FUTURES ONLY" if kind == "fo" else "FUTURES+OPTIONS COMBINED"
            P(f"=== {c['label']} ({c['code']}), {kl}: managed money, net = long - short, spreading excluded" + (f"; price = {c['price_label']}" if su else "")
              + ("  [$ columns are at the LME price, NOT the COMEX price]" if su and tag else ""))
            pu = '$/' + ('oz' if s == 'gold' else 'bbl' if s == 'wti' else 't')
            P(f"{'checkpoint':11s} {'report date':11s}" + (f" {'price':>10s}" if su else "") + f" {'MM long':>9s} {'MM short':>9s} {'MM net':>9s} {'MM net ' + ul:>12s}" + (f" {'MM net $bn':>11s}" if su else "")
              + f" {'MM net/OI':>10s} {'MM long/OI':>10s} {'MM sht/OI':>9s} {'OI':>9s} {'OI ' + ul:>11s}" + (f" {'OI $bn':>8s}" if su else ""))
            P(f"{'':11s} {'':11s}" + (f" {pu:>10s}" if su else "") + f" {'contracts':>9s} {'contracts':>9s} {'contracts':>9s} {'(' + ul + ')':>12s}" + (f" {'(at price)':>11s}" if su else "")
              + f" {'(% of OI)':>10s} {'(% of OI)':>10s} {'(% of OI)':>9s} {'contracts':>9s} {'(' + ul + ')':>11s}" + (f" {'($bn)':>8s}" if su else ""))
            for r in [x for x in cp if x["contract"] == s and x["report_type"] == kind]:
                mmp = float(r["mm_net_" + u]) * sc; oip = float(r["oi_" + u]) * sc
                P(f"{r['label']:11s} {r['report_date']:11s}" + (f" {float(r['price']):>10,.2f}" if su else "") + f" {r['mm_long']:>9,d} {r['mm_short']:>9,d} {r['mm_net']:>9,d} {mmp:>12,.1f}"
                  + (f" {float(r['mm_net_usd_bn' + tag]):>11.2f}" if su else "")
                  + f" {float(r['mm_net_pct_oi']):>9.1f}% {float(r['mm_long_pct_oi']):>9.1f}% {float(r['mm_short_pct_oi']):>8.1f}% {r['oi']:>9,d} {oip:>11,.1f}"
                  + (f" {float(r['oi_usd_bn' + tag]):>8.1f}" if su else ""))
            if su:
                for r in [x for x in cp if x["contract"] == s and x["report_type"] == kind and x["price_date"] != x["report_date"]]:
                    P(f"  note: the {r['label']} price ({float(r['price']):,.2f}) is from {r['price_date']}, the last quote before the report date {r['report_date']}")
                P(f"  changes in managed-money net, chained weekly volume/price split (endpoint-only split in brackets); $bn at each week's price" + (" (LME)" if tag else ""))
                P(f"  {'window':>14s} {'from':>10s} {'MM net chg':>11s} {'chg ' + ul:>11s} {'price chg':>9s} {'total $bn':>10s} {'volume $bn':>11s} {'price $bn':>10s} {'cross $bn':>10s}   {'[endpoint vol / price / cross $bn]':s}")
                for r in [x for x in wd if x["contract"] == s and x["report_type"] == kind]:
                    chg = float(r["mm_net_chg_" + u]) * sc
                    P(f"  {r['window']:>14s} {r['from_date']:>10s} {r['mm_net_chg']:>+11,d} {chg:>+11,.1f} {float(r['price_chg_pct']):>+8.1f}% {float(r['exp_chg_usd_bn' + tag]):>+10.2f} "
                      f"{float(r['chained_vol_usd_bn' + tag]):>+11.2f} {float(r['chained_price_usd_bn' + tag]):>+10.2f} {float(r['chained_cross_usd_bn' + tag]):>+10.2f}   "
                      f"[{float(r['endpoint_vol_usd_bn' + tag]):+.2f} / {float(r['endpoint_price_usd_bn' + tag]):+.2f} / {float(r['endpoint_cross_usd_bn' + tag]):+.2f}]")
            else:
                P("  changes in managed-money net (contracts and physical units) and its share of open interest at each end")
                P(f"  {'window':>14s} {'from':>10s} {'MM net chg':>11s} {'chg ' + ul:>11s} {'MM net/OI at from':>18s} {'MM net/OI at to':>16s}")
                P(f"  {'':>14s} {'':>10s} {'contracts':>11s} {'(' + ul + ')':>11s} {'(% of OI)':>18s} {'(% of OI)':>16s}")
                for r in [x for x in wd if x["contract"] == s and x["report_type"] == kind]:
                    chg = float(r["mm_net_chg_" + u]) * sc
                    P(f"  {r['window']:>14s} {r['from_date']:>10s} {r['mm_net_chg']:>+11,d} {chg:>+11,.1f} {float(r['mm_net_pct_oi_from']):>17.1f}% {float(r['mm_net_pct_oi_to']):>15.1f}%")
            P("")
    P(f"LEAD-LAG, weekly change in managed-money net CONTRACTS vs price change (%), report dates from {LATE_SPLIT} (first {ll[0]['first_week']}, last {ll[0]['last_week']}); 'same' = the change in the same week, "
      f"'next' = the following week, 'prev' = the preceding week (extra); r = correlation coefficient (unitless); band = +/-1.96/sqrt(n)"
      + ("; WTI only here, gold and copper are local (LICENSING)" if public else ""))
    P(f"  {'contract':8s} {'report':9s} {'n same':>6s} {'Pearson same':>12s} {'Spearman same':>13s} {'abs $ same':>10s} {'n next':>6s} {'Pearson next':>12s} {'Spearman next':>13s} {'abs $ next':>10s} "
      f"{'Pearson prev':>12s} {'Spearman prev':>13s} {'band (n=same/next)':>20s}")
    for r in ll:
        P(f"  {r['contract']:8s} {r['report_type']:9s} {r['n_same']:>6d} {float(r['pearson_same_pct']):>+12.3f} {float(r['spearman_same_pct']):>+13.3f} {float(r['pearson_same_abs']):>+10.3f} "
          f"{r['n_next']:>6d} {float(r['pearson_next_pct']):>+12.3f} {float(r['spearman_next_pct']):>+13.3f} {float(r['pearson_next_abs']):>+10.3f} "
          f"{float(r['pearson_prev_pct']):>+12.3f} {float(r['spearman_prev_pct']):>+13.3f} {r['band95_same']:>9s}/{r['band95_next']:<9s}")
    return out


def public_files():
    """Every file the licensing rule covers: data/cftc_cot/*.csv and report.txt (not raw/), and data/history/cftc_cot_*.csv."""
    a = [os.path.join(OUT, x) for x in sorted(os.listdir(OUT)) if os.path.isfile(os.path.join(OUT, x))]
    b = [os.path.join(HIST, x) for x in sorted(os.listdir(HIST)) if x.startswith("cftc_cot_") and x.endswith(".csv")]
    return a + b


def check_public():
    """Fail loudly if any public file carries a gold or copper price or dollar column. WTI files (EIA, public domain) are exempt. Checks file names,
    CSV headers, the rows of leadlag.csv, and every gold/copper line and section of report.txt for a '$'. Returns the number of files checked."""
    bad = re.compile(r"price|usd|dollar|lme|\$", re.I); problems = []; files = public_files()
    for p in files:
        name = os.path.basename(p); low = name.lower()
        if name.startswith("wti_") or "_wti_" in low: continue
        if bad.search(low): problems.append(f"{name}: file name carries a price/dollar marker")
        if name == "report.txt":
            sec = False
            for ln in open(p, encoding="utf-8").read().splitlines():
                if ln.startswith("==="): sec = bool(re.search(r"GOLD|COPPER", ln))
                elif ln.startswith("LEAD-LAG"): sec = False
                if (sec or re.search(r"gold|copper", ln, re.I)) and "$" in ln: problems.append(f"report.txt: a gold/copper line carries '$': {ln.strip()[:80]}")
        elif name.endswith(".csv") and p.startswith(OUT):
            with open(p, newline="", encoding="utf-8") as fh:
                rd = csv.reader(fh); hdr = next(rd)
                if (bad.search(",".join(hdr))): problems.append(f"{name}: header carries a price/dollar column: {[h for h in hdr if bad.search(h)]}")
                if name == "leadlag.csv":
                    ci = hdr.index("contract")
                    if any(r[ci] in LICENSED for r in rd): problems.append("leadlag.csv: carries a gold/copper row (price-derived)")
    assert not problems, "LICENSING: public output carries third-party-derived content:\n  " + "\n  ".join(problems)
    return len(files)


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--offline", action="store_true", help="rebuild from the saved raw files; no network")
    a = ap.parse_args()
    head = "CFTC Disaggregated COT ledger: COMEX gold, NYMEX WTI, COMEX copper" + (" (offline, from the saved snapshot)" if a.offline else "")
    print(head)
    if not a.offline:
        pull_cftc(); pull_prices(); write_manifest()
    nfiles, missing, listed = verify_manifest()
    print(f"raw snapshot: {nfiles} files match the sha256 in {os.path.relpath(MANIFEST, ROOT)}" + (f"; {len(missing)} third-party files absent" if missing else ""))
    avail = availability(missing, listed)
    by, stat = validate(load_raw())
    series, aux = load_prices(avail)
    dates = sorted(by["fo"]["gold"])
    px = {s: attach_prices(dates, series[s], s) for s in series}
    dates, L, maxres = build(by, px); aux["maxres"] = maxres
    os.makedirs(OUT, exist_ok=True)
    local = write_ledgers(L, dates); nh, nl = write_history(L, dates); cp, wd = write_tables(L, dates)
    ll_all = leadlag(L, dates, list(px))
    ll_pub, ll_loc = [r for r in ll_all if r["contract"] not in LICENSED], [r for r in ll_all if r["contract"] in LICENSED]
    write_csv(os.path.join(OUT, "leadlag.csv"), ll_pub)
    if ll_loc: write_csv(os.path.join(DER, "leadlag_gold_copper.csv"), ll_loc)
    cc = crosscheck(series, aux["wb"]) if aux["wb"] else None
    pub = report(by, stat, dates, L, px, aux, cp, wd, ll_pub, cc, True)
    with open(os.path.join(OUT, "report.txt"), "w", encoding="utf-8", newline="") as fh: fh.write("\n".join(pub) + "\n")
    npub = check_public()
    full_ok = all(s in px for s in SLUGS) and cc is not None
    if full_ok:
        full = report(by, stat, dates, L, px, aux, cp, wd, ll_all, cc, False)
        with open(os.path.join(DER, "report_full.txt"), "w", encoding="utf-8", newline="") as fh: fh.write("\n".join(full) + "\n")
        print("\n".join(full))
    else:
        print("\n".join(pub)); print("\nSKIPPED the full report (derived/report_full.txt): it needs every third-party price file; the public report above is complete.")
    print(f"\nwrote {nh} cftc_cot_*.csv files in data/history and the public tables, ledgers and report in data/cftc_cot ({npub} public files checked: no gold or copper price or dollar column)")
    print("local, gitignored (data/vintages/cot_prices/derived/): " + (f"{nl} history files, {len(local)} ledgers, checkpoints/windows/leadlag/crosscheck/full report where their inputs were present" if (nl or local) else "nothing written (sponsor files absent)"))


if __name__ == "__main__":
    main()
