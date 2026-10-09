#!/usr/bin/env python3
"""R1: what moved the US 10-year Treasury yield, end-2023 to latest.

Three decompositions, each over the same periods:
  1. Market data (an identity):  nominal 10y par = real 10y par (TIPS) + breakeven inflation.
     Source: US Treasury daily par yield curves (nominal and real), public domain.
  2. Models: fitted 10y zero yield = expected average short rate + term premium, in two models.
     - ACM (Adrian, Crump, Moench), Federal Reserve Bank of New York: ACMY10 = ACMRNY10 + ACMTP10.
     - Kim-Wright (FEDS 2005-33), Federal Reserve Board staff: THREEFY1000 - THREEFYTP1000 = expected.
     The split is a model construct: within each model it sums by construction (an identity), and the two
     models' disagreement is the band. Neither measures what markets expected.
  3. A holder's return on a constant-maturity 10y par bond, month by month: income (y/12), roll-down
     (duration x the 7y-10y slope per year / 12) and price change (-duration x change in yield). That split is
     a first-order approximation and leaves out convexity, so the total is also computed by full revaluation:
     buy a new 10y par bond at the month's start, price it at the month's end as a (10 - elapsed)-year bond on
     that day's curve (interpolated between the 7y and 10y par yields), add the coupon accrued. Added 27 Sep 2026
     after the R12R review; both totals are compounded month by month.

Inputs (not redistributed; fetched into data/vintages/r1_treasury/ if missing, and ignored by git):
  Treasury CSVs by year, the NY Fed ACM workbook (.xls), the Fed Board Kim-Wright CSV.
  4. A survey check that uses neither model: the Philadelphia Fed Survey of Professional Forecasters asks each
     first quarter for the expected 10-year average 3-month bill rate (BILL10) and 10-year bond yield (BOND10),
     and every quarter for annual bill-rate forecasts for the current and next three calendar years (TBILLA-D).
     Compared with the models' expected path on each survey's deadline (the Philadelphia Fed's release-date
     file; corrected 27 Sep 2026 from an approximate 14 February, C-125). Years 1-4 are the mean of TBILLA-D
     (calendar years, an approximation to the four years after the deadline); years 5-10 are implied as
     (10 x BILL10 - 4 x years 1-4) / 6. The models give the same split from their 4- and 10-year expected rates.
Outputs: data/r1_rates/levels.csv, decomp_10y.csv, holder_return_monthly.csv, holder_return_periods.csv,
         survey_check.csv, survey_changes.csv.
Requires: Python 3, xlrd and openpyxl (pip install -r requirements.txt). Run from anywhere:
  python3 bin/r1_treasury_decomposition.py
"""
import csv, datetime as dt, math, os, sys, urllib.request

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW = os.path.join(ROOT, "data", "vintages", "r1_treasury")
OUT = os.path.join(ROOT, "data", "r1_rates")
UA = {"User-Agent": "Mozilla/5.0 (research; free public data)"}
YEARS = [2023, 2024, 2025, 2026]
TSY = ("https://home.treasury.gov/resource-center/data-chart-center/interest-rates/daily-treasury-rates.csv/"
       "{y}/all?type={t}&field_tdr_date_value={y}&page&_format=csv")
ACM_URL = "https://www.newyorkfed.org/medialibrary/media/research/data_indicators/ACMTermPremium.xls"
KW_URL = "https://www.federalreserve.gov/data/yield-curve-tables/feds200533.csv"
SPF_URL = ("https://www.philadelphiafed.org/-/media/frbp/assets/surveys-and-data/survey-of-professional-forecasters/"
           "historical-data/meanlevel.xlsx")
KEY_DATES = [dt.date(2023, 12, 29), dt.date(2024, 12, 31), dt.date(2025, 12, 31), dt.date(2026, 6, 30)]
# SPF response deadlines, from the Philadelphia Fed's spf-release-dates.txt (read 27 Sep 2026). The 2026Q1 survey
# was delayed by the federal government shutdown.
SPF_DEADLINES = {(2024, 1): dt.date(2024, 2, 6), (2024, 2): dt.date(2024, 5, 6), (2024, 3): dt.date(2024, 8, 5),
                 (2024, 4): dt.date(2024, 11, 11), (2025, 1): dt.date(2025, 2, 11), (2025, 2): dt.date(2025, 5, 13),
                 (2025, 3): dt.date(2025, 8, 12), (2025, 4): dt.date(2025, 11, 11), (2026, 1): dt.date(2026, 3, 2),
                 (2026, 2): dt.date(2026, 5, 12), (2026, 3): dt.date(2026, 8, 11)}


def fetch(url, path):
    if os.path.exists(path):
        return
    os.makedirs(os.path.dirname(path), exist_ok=True)
    data = urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=120).read()
    open(path, "wb").write(data)


def load_treasury(kind):
    out = {}
    for y in YEARS:
        p = os.path.join(RAW, f"{kind}_{y}.csv")
        fetch(TSY.format(y=y, t=kind), p)
        for r in csv.DictReader(open(p, encoding="utf-8-sig")):
            d = dt.datetime.strptime(r["Date"], "%m/%d/%Y").date()
            out[d] = {k.strip(): (float(v) if v not in ("", "N/A") else None) for k, v in r.items() if k != "Date"}
    return out


def load_acm():
    import xlrd
    p = os.path.join(RAW, "ACMTermPremium.xls")
    fetch(ACM_URL, p)
    sh = xlrd.open_workbook(p).sheet_by_name("ACM Daily")
    h = [str(c.value) for c in sh.row(0)]
    iy, itp, irn, irn4 = h.index("ACMY10"), h.index("ACMTP10"), h.index("ACMRNY10"), h.index("ACMRNY04")
    out = {}
    for i in range(1, sh.nrows):
        d = dt.datetime.strptime(str(sh.cell_value(i, 0)), "%d-%b-%Y").date()
        if d.year >= 2023:
            out[d] = {"y": sh.cell_value(i, iy), "tp": sh.cell_value(i, itp), "exp": sh.cell_value(i, irn),
                      "exp4": sh.cell_value(i, irn4)}
    return out


def load_kw():
    p = os.path.join(RAW, "feds200533.csv")
    fetch(KW_URL, p)
    rows = list(csv.reader(open(p, encoding="utf-8-sig")))
    h = next(i for i, r in enumerate(rows) if r and r[0] == "Date")
    H = rows[h]; iy, itp = H.index("THREEFY1000.B"), H.index("THREEFYTP1000.B")
    iy4, itp4 = H.index("THREEFY0400.B"), H.index("THREEFYTP0400.B")
    out = {}
    for r in rows[h + 1:]:
        if not r or not r[0][:4].isdigit() or r[iy] in ("", "NA"):
            continue
        d = dt.date.fromisoformat(r[0])
        if d.year >= 2023:
            y, tp = float(r[iy]), float(r[itp])
            out[d] = {"y": y, "tp": tp, "exp": y - tp, "exp4": float(r[iy4]) - float(r[itp4])}
    return out


def load_spf():
    import openpyxl
    p = os.path.join(RAW, "meanLevel.xlsx")
    fetch(SPF_URL, p)
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    out = {}                                                        # keyed (year, quarter)
    for var in ("BILL10", "BOND10"):
        for row in wb[var].iter_rows(min_row=2, values_only=True):
            y, q, v = row[0], row[1], row[2]
            if y and q == 1 and isinstance(v, (int, float)) and y >= 2024:
                out.setdefault((int(y), 1), {})[var] = float(v)
    rows = list(wb["TBILL"].iter_rows(values_only=True)); th = rows[0]
    cols = [th.index(c) for c in ("TBILLA", "TBILLB", "TBILLC", "TBILLD")]
    for row in rows[1:]:
        if row[0] and row[0] >= 2024 and all(isinstance(row[c], (int, float)) for c in cols):
            out.setdefault((int(row[0]), int(row[1])), {})["TBILL_1to4"] = sum(row[c] for c in cols) / 4.0
    return out


def on_or_before(series, d):
    ks = [k for k in series if k <= d]
    return max(ks) if ks else None


def coupon_price(coupon_pct, yield_pct, T):
    """Clean price per 100 of a semiannual-coupon bond, T years to maturity (continuous periods)."""
    c, y, n = coupon_pct / 100.0, yield_pct / 100.0, 2.0 * T
    disc = (1.0 + y / 2.0) ** (-n)
    return 100.0 * (c / y * (1.0 - disc) + disc)


def par_mod_duration(y_pct, n=10):
    y = y_pct / 100.0
    return (1.0 / y) * (1.0 - 1.0 / (1.0 + y / 2.0) ** (2 * n))


def main():
    nom, real = load_treasury("daily_treasury_yield_curve"), load_treasury("daily_treasury_real_yield_curve")
    acm, kw = load_acm(), load_kw()
    latest = min(max(nom), max(real), max(acm), max(kw))            # the last date every source covers
    dates = KEY_DATES + [latest]
    os.makedirs(OUT, exist_ok=True)

    lv = []
    for d in dates:
        dn, dr, da, dk = (on_or_before(s, d) for s in (nom, real, acm, kw))
        n, r, a, k = nom[dn], real[dr], acm[da], kw[dk]
        lv.append({"key_date": d.isoformat(), "treasury_date": dn.isoformat(), "acm_date": da.isoformat(),
                   "kw_date": dk.isoformat(), "par10": n["10 Yr"], "par3m": n["3 Mo"], "par2y": n["2 Yr"],
                   "real10": r["10 YR"], "breakeven10": round(n["10 Yr"] - r["10 YR"], 4), "slope_10y_2y": round(n["10 Yr"] - n["2 Yr"], 4),
                   "acm_y10": round(a["y"], 4), "acm_exp10": round(a["exp"], 4), "acm_tp10": round(a["tp"], 4),
                   "kw_y10": round(k["y"], 4), "kw_exp10": round(k["exp"], 4), "kw_tp10": round(k["tp"], 4)})
    with open(os.path.join(OUT, "levels.csv"), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(lv[0])); w.writeheader(); w.writerows(lv)

    L = {x["key_date"]: x for x in lv}
    k0, k24, k25, kjun, klat = (x.isoformat() for x in dates)
    periods = [("2024", k0, k24), ("2025", k24, k25), ("2026 to 30 Jun", k25, kjun), ("30 Jun to latest", kjun, klat),
               ("end-2023 to 30 Jun 2026 (the equity window)", k0, kjun), ("end-2023 to latest", k0, klat)]
    dec = []
    for name, a, b in periods:
        A, B = L[a], L[b]
        dd = lambda c: round(B[c] - A[c], 4)
        row = {"period": name, "from": a, "to": b, "d_par10": dd("par10"), "d_par2y": dd("par2y"), "d_slope_10y_2y": dd("slope_10y_2y"), "d_real10": dd("real10"),
               "d_breakeven10": dd("breakeven10"), "d_acm_y10": dd("acm_y10"), "d_acm_exp10": dd("acm_exp10"),
               "d_acm_tp10": dd("acm_tp10"), "d_kw_y10": dd("kw_y10"), "d_kw_exp10": dd("kw_exp10"), "d_kw_tp10": dd("kw_tp10")}
        for m in ("acm", "kw"):
            tot = row[f"d_{m}_y10"]
            row[f"{m}_tp_share"] = round(row[f"d_{m}_tp10"] / tot, 3) if abs(tot) >= 0.05 else ""
        dec.append(row)
    with open(os.path.join(OUT, "decomp_10y.csv"), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(dec[0])); w.writeheader(); w.writerows(dec)

    # holder's return, month-end to month-end on the Treasury par curve
    months = sorted({(d.year, d.month) for d in nom if d >= dt.date(2023, 12, 1) and d <= latest})
    ends = [max(d for d in nom if (d.year, d.month) == ym and d <= latest) for ym in months]
    hr = []
    for p, c in zip(ends, ends[1:]):
        y0, y1 = nom[p]["10 Yr"], nom[c]["10 Yr"]
        D = par_mod_duration(y0)
        slope = (nom[p]["10 Yr"] - nom[p]["7 Yr"]) / 3.0                # pp per year of maturity near 10y
        frac = (c - p).days / 365.25
        y_end = nom[c]["10 Yr"] - (nom[c]["10 Yr"] - nom[c]["7 Yr"]) / 3.0 * frac   # yield at 10 - frac years
        full = coupon_price(y0, y_end, 10.0 - frac) - 100.0 + y0 * frac
        hr.append({"from": p.isoformat(), "to": c.isoformat(), "y10_start": y0, "y10_end": y1, "mod_duration": round(D, 3),
                   "income_pct": round(y0 * frac, 4), "rolldown_pct": round(D * slope * frac, 4),
                   "price_pct": round(-D * (y1 - y0), 4), "bill3m_income_pct": round(nom[p]["3 Mo"] * frac, 4),
                   "full_reval_pct": round(full, 4)})
    with open(os.path.join(OUT, "holder_return_monthly.csv"), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(hr[0])); w.writeheader(); w.writerows(hr)
    agg = []
    for name, a, b in periods:
        sub = [x for x in hr if dt.date.fromisoformat(x["from"]) >= on_or_before(nom, dt.date.fromisoformat(a)) - dt.timedelta(days=3)
               and dt.date.fromisoformat(x["to"]) <= dt.date.fromisoformat(b) + dt.timedelta(days=3)]
        s = lambda c: round(sum(x[c] for x in sub), 3)
        comp = lambda f: round(100.0 * (math.prod(1.0 + f(x) / 100.0 for x in sub) - 1.0), 3)   # compounded month by month
        agg.append({"period": name, "months": len(sub), "income_pct": s("income_pct"), "rolldown_pct": s("rolldown_pct"),
                    "price_pct": s("price_pct"), "total_pct": round(s("income_pct") + s("rolldown_pct") + s("price_pct"), 3),
                    "bill3m_income_pct": s("bill3m_income_pct"),
                    "total_compounded_pct": comp(lambda x: x["income_pct"] + x["rolldown_pct"] + x["price_pct"]),
                    "full_reval_compounded_pct": comp(lambda x: x["full_reval_pct"]),
                    "bill3m_compounded_pct": comp(lambda x: x["bill3m_income_pct"])})
    with open(os.path.join(OUT, "holder_return_periods.csv"), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(agg[0])); w.writeheader(); w.writerows(agg)
    spf = load_spf()
    sv, S = [], {}
    split = lambda e10, e4: round((10 * e10 - 4 * e4) / 6, 4) if e10 is not None and e4 is not None else ""
    for key in sorted(SPF_DEADLINES):
        d = SPF_DEADLINES[key]
        if key not in spf or d > latest:
            continue
        dn, da, dk = (on_or_before(s, d) for s in (nom, acm, kw))
        a, k, sp = acm[da], kw[dk], spf[key]
        row = {"survey": f"{key[0]}Q{key[1]}", "deadline": d.isoformat(), "treasury_date": dn.isoformat(),
               "par10": nom[dn]["10 Yr"], "acm_exp_1to4": round(a["exp4"], 4), "acm_exp10": round(a["exp"], 4),
               "acm_exp_5to10": split(a["exp"], a["exp4"]), "kw_exp_1to4": round(k["exp4"], 4),
               "kw_exp10": round(k["exp"], 4), "kw_exp_5to10": split(k["exp"], k["exp4"]),
               "spf_bill_1to4": round(sp["TBILL_1to4"], 4) if "TBILL_1to4" in sp else "",
               "spf_bill10": sp.get("BILL10", ""), "spf_bond10": sp.get("BOND10", ""),
               "spf_bill_5to10": split(sp.get("BILL10"), sp.get("TBILL_1to4"))}
        sv.append(row); S[row["survey"]] = row
    with open(os.path.join(OUT, "survey_check.csv"), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(sv[0])); w.writeheader(); w.writerows(sv)
    # changes in basis points: the two first-quarter surveys two years apart (all horizons), and 2026Q1 to the
    # latest quarterly survey (years 1-4 only; BILL10 and BOND10 are asked in the first quarter only)
    chg = []
    last_q = max((r["survey"] for r in sv), key=lambda x: (int(x[:4]), int(x[-1])))
    for a_q, b_q in (("2024Q1", "2026Q1"), ("2026Q1", last_q)):
        if a_q not in S or b_q not in S:
            continue
        A, B = S[a_q], S[b_q]
        bp = lambda c: round(100 * (B[c] - A[c])) if A[c] != "" and B[c] != "" else ""
        chg.append({"from": a_q, "to": b_q, **{f"d_{c}_bp": bp(c) for c in list(A)[3:]}})
    with open(os.path.join(OUT, "survey_changes.csv"), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(chg[0])); w.writeheader(); w.writerows(chg)
    print(f"[done] latest common date {latest}; wrote levels, decomp_10y, holder_return_monthly, holder_return_periods, "
          f"survey_check, survey_changes")


if __name__ == "__main__":
    main()
