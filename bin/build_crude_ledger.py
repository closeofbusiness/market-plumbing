#!/usr/bin/env python3
"""Crude oil physical ledger: where US barrels came from and went, weekly, and where the world's went, monthly, in barrels and in
dollars, with the supply/disposition balance asserted -- E-021 follow-up ("where does the money in commodities come from", crude).
Standard library plus xlrd (EIA .xls) and openpyxl (EIA STEO .xlsx), both in requirements.txt.

  python3 bin/build_crude_ledger.py             # pull every source, save the raw files, write every output, print the report
  python3 bin/build_crude_ledger.py --offline   # rebuild every output from the saved raw files (no network); byte-identical

SOURCES (all US government, free, no key, scripted with the neutral User-Agent below; one request at a time, 1 s apart; the script
retries only where the server says to (429/503, honouring Retry-After), then stops, and never works around a block). The EIA API needs
a key and is NOT used; no key found in any page is used.
  1. EIA dnav history workbooks, https://www.eia.gov/dnav/pet/hist_xls/<SERIES>w.xls (weekly, week ending Friday). The unit is read from
     each file's own header and asserted:
       stocks, thousand barrels (kb):  WCESTUS1 U.S. ending stocks of crude oil EXCLUDING SPR (commercial, excl. lease stock)
         WCSSTUS1 crude oil in the SPR (Strategic Petroleum Reserve)     WCRSTUS1 crude oil INCLUDING SPR
         W_EPC0_SAX_YCUOK_MBBL Cushing, OK crude oil stocks             WGTSTUS1 total motor gasoline stocks
         WDISTUS1 distillate fuel oil stocks                            WTTSTUS1 crude oil and petroleum products INCLUDING SPR
       flows, thousand barrels per day (kb/d):  WCRFPUS2 field production of crude oil     WCRIMUS2 imports of crude oil
         WCREXUS2 exports of crude oil    WCRNTUS2 net imports of crude oil (including SPR)    WCRRIUS2 refiner net input of crude oil
         WRPUPUS2 product supplied, total petroleum products    WTTNTUS2 net imports of crude oil and petroleum products
  2. EIA Weekly Petroleum Status Report (WPSR) Table 1 in machine-readable form, https://ir.eia.gov/wpsr/psw01.xls ("U.S. Petroleum
     Balance Sheet", full weekly history from 1982; EIA answers with a redirect to a time-limited signed address on the same host, which is
     followed; no login). It is the only free key-free source of the lines the dnav workbooks do not carry (the dnav "weekly supply
     estimates" page has no crude adjustment and no transfers line; their hist_xls files return 404):
         WCRAUUS2 "Unaccounted for crude oil" = Table 1 line (16) "Adjustment"           kb/d, weekly from 2001-11-09
         W_EPC0_TVP_NUS_MBBLD "Transfers to crude oil supply" = Table 1 line (4)         kb/d, weekly from 2023-11-10 ONLY
           (W_EPLLCL2A_TVP_NUS_MBBLD Alaskan NGL and condensate, W_EPLLNGPU_TVP_NUS_MBBLD natural gasoline and unfinished oils are its two parts;
            W_EPC0NG_TVP_NUS_MBBLD is the same amount with the opposite sign on the products side)
         W_EPC0_SCG_NUS_MBBLD "Stock change" (line 13), WCESCUS2 (line 14, commercial), WCSSCUS2 (line 15, SPR)                  kb/d
         W_EPP1_YPT_NUS_MBBLD "Production" of other supply (line 18) = W_EPL0_YPG_NUS_MBBLD natural gas plant liquids + W_EPOOXR_YPT_NUS_MBBLD
           renewable fuels/oxygenates + WPGNPUS2 refinery processing gain (asserted within 1 kb/d; EIA's own title for the first one says "less"
           these three, which its numbers contradict)                                                              kb/d, from 2010-06-04
         WPESCUS2 products stock change (line 27), W_EPP0_VUA_NUS_MBBLD adjustment of total petroleum excluding crude (line 28)    kb/d
     Every series that is in both routes (13 of the 14 above, all but Cushing) is asserted IDENTICAL, date by date, between the dnav file and psw01.xls.
     https://ir.eia.gov/wpsr/table1.csv (the latest release's Table 1: this week, last week, a year earlier) is kept for the spot check.
  3. EIA Short-Term Energy Outlook, https://www.eia.gov/outlooks/steo/xls/STEO_m.xlsx, tab 3atab (world petroleum and other liquid fuels
     production, consumption, inventories), tab 2tab (spot prices) and tab Dates (the workbook's own "Last Historical Month" and per-month
     historical flag). STEO mixes history and forecast: months flagged historical in the file are labelled "historical", the rest "forecast".
     Even a "historical" month is mostly an EIA estimate for the world (consumption outside the OECD is "apparent consumption", and
     "other inventory draws and balance" is a balancing item), so the world withdrawal is BOUNDED, not measured.
  4. EIA daily spot prices, https://www.eia.gov/dnav/pet/hist_xls/RWTCd.xls (Cushing WTI) and RBRTEd.xls (Europe Brent), $ per barrel.
  5. Term structure (looked for briefly, then stopped): https://www.eia.gov/dnav/pet/hist_xls/RCLC1d.xls (NYMEX Cushing contract 1) and the EIA
     futures page https://www.eia.gov/dnav/pet/pet_pri_fut_s1_d.htm are kept as the evidence of where EIA's futures series ends.
  6. LOCAL ONLY (not committed, git-ignored data/vintages/crude/): the World Bank Pink Sheet monthly workbook, found from
     https://www.worldbank.org/en/research/commodity-markets, for a Brent cross-check. If the file is absent the leg is skipped and every
     public output is identical (checked: the public outputs are a function of the saved EIA files alone).

UNITS. kb = thousand barrels; kb/d = thousand barrels per day; mb = million barrels (1 mb = 1,000 kb); mb/d = million barrels per day;
$mn = US$ millions; $bn = US$ billions. A "week" is the Saturday-to-Friday week ending on the Friday "week_end" date; it has 7 days, so a
flow of F kb/d over a week is 7 x F kb. Dollars = barrels x price: kb x $/bbl = $ thousand, so $mn = kb x $/bbl / 1000 and $bn = kb x $/bbl / 1,000,000.

THE WEEKLY BALANCE (the identity, derived from EIA's Table 1; every term is in kb/d, the stock change is the one-week change in the level)
    stock change (commercial + SPR) = 7 x ( field production + transfers to crude supply + imports - exports + unaccounted-for - refiner input )
  * The brief's identity had no transfers term. It is needed from the week ending 2023-11-10, the first week of EIA's Table 1 line (4) "Transfers to
    crude oil supply" (Alaskan NGL, condensate, natural gasoline and unfinished oils that refiners run as crude; about +450 to +730 kb/d; EIA's file has
    no earlier value). Without it the
    identity misses by about 3,700 kb a week from then on (asserted in the printout as a check on the check). Before that week the term is
    ZERO because the line did not exist, and the balancing item then carried the same barrels: unaccounted-for crude averaged +611 kb/d in 2022
    and +859 kb/d in 2023 up to October, against +8 kb/d in 2024. The 2023 break in WCRAUUS2 is therefore a definitional change, not a change in the oil.
  * "Unaccounted for" (EIA's "Adjustment") is the BALANCING ITEM, computed by EIA so that Table 1 closes. So the identity holds by construction for the
    published balance sheet; what it tests is (a) the rounding closure of the published lines, (b) that the published stock change ties to the
    stock LEVELS in WCESTUS1 + WCSSTUS1, and (c) how big the plug is. The plug's size is the finding: it is the part of the barrels that the
    reported production, trade and refinery lines do not explain.
  * Three residuals are written per week (all kb, stock change minus supply-minus-disposition):
      bal_residual_published_kb   7 x published stock change (line 13)  -  7 x terms        asserted |.| <= 24.5 kb in EVERY week from 2001-11-09
                                  (seven published integers in kb/d, each +-0.5, times 7 days = 24.5 kb)
      bal_residual_kb             week-on-week change in WCRSTUS1 (= WCESTUS1 + WCSSTUS1, asserted exact)  -  7 x terms         THE BRIEF'S IDENTITY
                                  asserted |.| <= 22 kb (six published terms x 0.5 x 7 = 21 kb, +1 kb for whole-barrel levels) in every week from
                                  2017-01-06 except REVISION WEEKS
      levels_vs_published_kb      change in the levels  -  7 x published stock change.  A week is a REVISION WEEK if this exceeds 14 kb (twice the
                                  largest value seen in a non-revision week since 2017, 7 kb): the level history was revised after the balance-sheet line was
                                  published. Asserted: revision weeks are at most 1% of the weeks from 2017-01-06 (today: one, 2023-12-22, 203 kb).
  * Before 2017-01-06 the level change and the published stock change disagree by up to 3,100 kb in about a dozen weeks a year, always in the first week of a
    month (2010-2015). The reading offered is that EIA re-benchmarked the weekly stock LEVELS to its monthly survey and did not restate the weekly flow
    lines; that is an inference from the pattern, EIA does not say so in these files. Those weeks are written, shown in the distribution, and NOT asserted.
  * The total-petroleum identity (transfers cancel between crude and products) is also written and asserted on the same terms, from 2017-01-06:
      change in WTTSTUS1 = 7 x ( crude production + other production (line 18) + net imports of crude and products + crude unaccounted-for
                                 + products adjustment (line 28) - product supplied )
    WTTSTUS1 has a one-week glitch on 2016-09-30 (about -42 mb, reversed the next week); it is before the asserted window and is not corrected.

LEDGER COLUMNS (data/crude/us_weekly_ledger.csv; names carry their unit: _kb thousand barrels, _kbd thousand barrels per day, _usd_bbl $/bbl, _usd_mn $ million, _usd_bn $ billion):
  bal_status: asserted | revision_week (excluded from the assertion, listed in the report) | before_2017_not_asserted | unaccounted_for_not_published (nine weeks in
  2005-2007 where EIA's file has no balancing item: no residual is computed). imports_basis: imports_minus_exports where both weekly series exist (imports from
  2010-10-29, exports from 2007-06-22), else the published net imports WCRNTUS2. transfers_in_identity: published from 2023-11-10, else zero_line_did_not_exist.
  tp_* columns: the total-petroleum identity. since_end2025_* : running sums after the week ending 2025-12-26. In checkpoints.csv the start and end of a world
  window are the first days of its first and last months; for US windows start is the earlier checkpoint week (exclusive) and end the later one (inclusive).

DOLLARS. WTI (Cushing spot) is the price for every US dollar column (US crude is priced off WTI); the Brent versions of the stock-change columns are
written beside them. A week's price is the mean of the daily EIA spot prints for the Saturday-to-Friday week (n days is a column; a holiday week has
fewer). Stock-change dollars = kb x that week's price; flow dollars = 7 x kb/d x price (imports and exports are valued at the WTI spot too, a proxy:
actual import and export prices differ by grade and by the Brent-WTI spread). The SPR "release" is the fall in the SPR LEVEL: it nets out sales,
exchanges and their returns, purchases and any loans, so the dollars are barrels at market, NOT proceeds to the Treasury or the DOE's cost basis.
The value of stocks is barrels at the week's spot, not book cost. One WTI print is negative (2020-04-20, -$36.98); it is in the weekly average
for the week ending 2020-04-24 and is not adjusted.
"End-2025" for cumulative windows is the last week ending on or before 31 Dec 2025 (2025-12-26); the week ending 2026-01-02 is the first 2026 week and
includes 27-31 Dec 2025. Checkpoints use the last week ending on or before the date (30 Jun 2026 -> 2026-06-26).

THE FIVE-YEAR RANGE (commercial stocks, levels and the one-week change). Two alignments are written, because they differ by up to six days in the season:
  * date-aligned (columns *_5y_*, the primary): the same week in each of the five preceding calendar years is the weekly observation within 3 days of
    t shifted back by 1..5 years (2021-2025 for a 2026 week); range = minimum and maximum of those five, plus the mean;
  * ISO-week-aligned (columns *_5yiso_*): the Friday with the same ISO week number (week 53 -> the last week of a 52-week year). This is the alignment the
    first-pass map used: for the week ending 2026-10-02 it reproduces the map's 420.3-439.1 mb (the date-aligned range is 414.1-429.2 mb).

KNOWN DATA ISSUES (reported, not corrected):
  * EIA's Europe Brent spot print of $135.51 on 2026-10-02 (from $114.82 the day before, while WTI fell) is analysed in the report, not adjusted.
  * EIA's WTI spot print of $85.23 on 2026-09-25 sits between $95.88 and $99.37; it is in the weekly average for that week, not adjusted.
  * The EIA daily spot files run to 2026-10-06, a week before the pull; the latest stock week is 2026-10-02 (released 7 Oct), so no week lacks a price.
  * psw01.xls and table1.csv are rebuilt by EIA each Wednesday; the saved copies are the vintage. The latest week's figures can be revised in later releases.

TERM STRUCTURE. EIA's NYMEX futures series (RCLC1 ... RCLC4) end on 2024-04-05 and the page says so (RCLC1 and the page are kept as raw evidence). CME refused scripted
access in the first pass (HTTP 403) and was not retried. CFTC data carry no prices. Beyond those two files, EIA's STEO landing page and the "What drives crude oil prices"
page (whose text says updates to its figures were discontinued) were opened and not saved; a guessed STEO "uncertainty" address returned 404. Nothing else was tried; that leg is stopped.

LICENSING (the repository is public). Every committed input and output here is US-government work product (EIA, a part of the US Department of
Energy): in the public domain in the United States (17 U.S.C. 105), so raw files, outputs and the two-column history files may all be public; EIA asks
only that the source be credited ("Source: U.S. Energy Information Administration"). The one non-government input, the World Bank Pink Sheet, is read
only if present in the git-ignored data/vintages/crude/, is never copied into a public file, and does not change any public file.

OUTPUTS
  data/crude/us_weekly_ledger.csv        one row per week from 2001-11-09 (first week of the published balance sheet): stocks, changes, flows, the three
                                         residuals, the five-year range, 4-week averages, prices, dollars, cumulative-since-end-2025 columns
  data/crude/world_monthly_balance.csv   one row per month Jan 2022 - Dec 2027 (STEO 3atab): production, consumption, implied net withdrawal, OECD stocks, Brent, dollars
  data/crude/checkpoints.csv             end-2023, end-2024, end-2025, 2026-06-30 and the latest week; window changes in barrels and dollars; world windows; identity statistics
  data/crude/report.txt                  the printed report (identity statistics, checkpoint tables, Brent check, term-structure finding, spot check)
  data/crude/brent_spread_check.csv      the last 60 daily WTI and Brent prints, the spread, and the one-day changes
  data/history/crude_us_commercial_stocks_kb.csv, crude_us_spr_stocks_kb.csv, crude_us_cushing_stocks_kb.csv, crude_us_field_production_kbd.csv,
  data/history/crude_us_exports_kbd.csv, crude_us_balance_residual_kb.csv        two columns, as_of,value, the full available weekly history
  data/crude/raw/                        every EIA file pulled, and pull_log.tsv (url, UTC time, sha256, size of each)
  LOCAL (git-ignored): data/vintages/crude/worldbank_CMO-Historical-Data-Monthly.xlsx, data/vintages/crude/derived/brent_worldbank_check.csv
"""
import os, re, sys, csv, ssl, math, time, bisect, hashlib, calendar, argparse, datetime, statistics, urllib.request, urllib.error

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HIST = os.path.join(ROOT, "data", "history")
OUT = os.path.join(ROOT, "data", "crude")
RAW = os.path.join(OUT, "raw")
VINT = os.path.join(ROOT, "data", "vintages", "crude")
DER = os.path.join(VINT, "derived")
LOG = os.path.join(RAW, "pull_log.tsv")
WB_X = os.path.join(VINT, "worldbank_CMO-Historical-Data-Monthly.xlsx")
UA = {"User-Agent": "ThirdDerivativeResearch/1.0", "Accept-Encoding": "identity"}
Dt = datetime.date
ONE_WEEK = datetime.timedelta(days=7)

DNAV_W = "https://www.eia.gov/dnav/pet/hist_xls/{}w.xls"
DNAV_D = "https://www.eia.gov/dnav/pet/hist_xls/{}d.xls"
PSW01_URL = "https://ir.eia.gov/wpsr/psw01.xls"
TABLE1_URL = "https://ir.eia.gov/wpsr/table1.csv"
STEO_URL = "https://www.eia.gov/outlooks/steo/xls/STEO_m.xlsx"
FUT_PAGE_URL = "https://www.eia.gov/dnav/pet/pet_pri_fut_s1_d.htm"
WB_PAGE = "https://www.worldbank.org/en/research/commodity-markets"
XLS, ZIP = b"\xd0\xcf\x11\xe0", b"PK\x03\x04"

KB, KBD = "Thousand Barrels", "Thousand Barrels per Day"
WEEKLY = [   # id, unit expected in the file header, role
    ("WCESTUS1", KB, "commercial crude oil stocks, excl. SPR"), ("WCSSTUS1", KB, "crude oil stocks in the SPR"),
    ("WCRSTUS1", KB, "crude oil stocks incl. SPR"), ("W_EPC0_SAX_YCUOK_MBBL", KB, "Cushing, OK crude oil stocks"),
    ("WGTSTUS1", KB, "total motor gasoline stocks"), ("WDISTUS1", KB, "distillate fuel oil stocks"),
    ("WTTSTUS1", KB, "crude oil and petroleum products stocks incl. SPR"),
    ("WCRFPUS2", KBD, "field production of crude oil"), ("WCRIMUS2", KBD, "imports of crude oil"),
    ("WCREXUS2", KBD, "exports of crude oil"), ("WCRNTUS2", KBD, "net imports of crude oil incl. SPR"),
    ("WCRRIUS2", KBD, "refiner net input of crude oil"), ("WRPUPUS2", KBD, "product supplied, total petroleum products"),
    ("WTTNTUS2", KBD, "net imports of crude oil and petroleum products")]
PSW_KEYS = {k: KBD for k in [
    "WCRAUUS2", "W_EPC0_TVP_NUS_MBBLD", "W_EPLLCL2A_TVP_NUS_MBBLD", "W_EPLLNGPU_TVP_NUS_MBBLD", "W_EPC0NG_TVP_NUS_MBBLD",
    "W_EPC0_SCG_NUS_MBBLD", "WCESCUS2", "WCSSCUS2", "W_EPP1_YPT_NUS_MBBLD", "W_EPL0_YPG_NUS_MBBLD", "W_EPOOXR_YPT_NUS_MBBLD",
    "WPGNPUS2", "WPESCUS2", "W_EPP0_VUA_NUS_MBBLD"]}

LEDGER_START = Dt(2001, 11, 9)       # first week with the unaccounted-for series (and so a published balance sheet)
ASSERT_FROM = Dt(2017, 1, 6)         # level-based identity asserted from here (the pre-2017 level history was revised; see docstring)
TOL_PUBLISHED = 24.5                 # kb: seven published integers in kb/d, each +-0.5, x 7 days
TOL_LEVELS = 22.0                    # kb: six published integers x 0.5 x 7 = 21, +1 for whole-barrel stock levels
REV_FLAG = 14.0                      # kb: a revision week if the level change differs from 7 x the published stock change by more than this
MAX_REV_SHARE = 0.01
TRANSFERS_START = Dt(2023, 11, 10)
CHECKPOINTS = [("end-2023", Dt(2023, 12, 31)), ("end-2024", Dt(2024, 12, 31)), ("end-2025", Dt(2025, 12, 31)),
               ("30 Jun 2026", Dt(2026, 6, 30)), ("latest week", None)]


# ------------------------------------------------------------------------------------------------ small helpers
def fmt(x, nd=3):
    if x is None or (isinstance(x, float) and math.isnan(x)): return ""
    if isinstance(x, str): return x
    if isinstance(x, bool): return "1" if x else "0"
    if isinstance(x, int): return str(x)
    return f"{round(x, nd) + 0.0:.{nd}f}"

def colnd(name):
    if name.endswith("_4wk_kbd") or name.endswith("_mean_kb"): return 1
    if name.endswith("_kb") or name.endswith("_kbd"): return 0
    if name.endswith("_usd_bbl") or name.endswith("_usd_bbl_brent"): return 2
    if "_usd_mn" in name: return 1
    if "_usd_bn" in name: return 3
    return 3

def write_csv(path, fields, rows, nd=colnd):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n"); w.writerow(fields)
        for r in rows: w.writerow([fmt(r.get(c), nd(c)) for c in fields])

def write_hist(name, pairs, nd=0):
    os.makedirs(HIST, exist_ok=True)
    with open(os.path.join(HIST, f"crude_{name}.csv"), "w", encoding="utf-8", newline="") as f:
        f.write("as_of,value\n")
        for d, v in pairs: f.write(f"{d.isoformat()},{fmt(v, nd)}\n")

def sub(a, b):
    return None if a is None or b is None else a - b

def mul(a, k):
    return None if a is None else a * k

def usd_mn(kb, price):
    """kb x $/bbl = $ thousand; /1000 = $ million."""
    return None if kb is None or price is None else kb * price / 1000.0

def pctl(xs, q):
    s = sorted(xs); return s[max(0, math.ceil(q / 100.0 * len(s)) - 1)]

def mon(m):
    return f"{m // 100:04d}-{m % 100:02d}"


# ------------------------------------------------------------------------------------------------ network (pull mode)
def get(url, tries=4, magic=None):
    for k in range(tries):
        try:
            r = urllib.request.urlopen(urllib.request.Request(url, headers=dict(UA)), timeout=120, context=ssl.create_default_context())
            data = r.read()
            if magic and not data.startswith(magic):
                raise SystemExit(f"STOP: {url} did not return the expected file type (first bytes {data[:12]!r}). Not retrying.")
            return data
        except urllib.error.HTTPError as e:
            if e.code in (429, 503) and k < tries - 1:
                wait = int(e.headers.get("Retry-After") or 0) or 15 * (k + 1)
                print(f"  HTTP {e.code}; the server asked for a pause; waiting {wait}s", file=sys.stderr); time.sleep(wait); continue
            raise SystemExit(f"STOP: HTTP {e.code} for {url}. Not retrying further and not working around it.")
        except urllib.error.URLError as e:
            if k < tries - 1: time.sleep(10 * (k + 1)); continue
            raise SystemExit(f"STOP: network error for {url}: {e}")

LOG_ROWS = []

def save(path, data, url, log=True):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "wb") as f: f.write(data)
    if log:
        LOG_ROWS.append((os.path.relpath(path, ROOT), url, datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%MZ"),
                         hashlib.sha256(data).hexdigest(), len(data)))
    print(f"  saved {os.path.relpath(path, ROOT)}  {len(data):,} bytes")
    time.sleep(1.0)

def pull():
    print("pulling (one request at a time, 1 s apart)")
    for sid, _, _ in WEEKLY:
        u = DNAV_W.format(sid); save(os.path.join(RAW, f"eia_{sid}w.xls"), get(u, magic=XLS), u)
    save(os.path.join(RAW, "eia_wpsr_psw01.xls"), get(PSW01_URL, magic=XLS), PSW01_URL)
    save(os.path.join(RAW, "eia_wpsr_table1.csv"), get(TABLE1_URL, magic=b'"STUB_1"'), TABLE1_URL)
    for sid in ("RWTC", "RBRTE", "RCLC1"):
        u = DNAV_D.format(sid); save(os.path.join(RAW, f"eia_{sid}d.xls"), get(u, magic=XLS), u)
    save(os.path.join(RAW, "eia_STEO_m.xlsx"), get(STEO_URL, magic=ZIP), STEO_URL)
    save(os.path.join(RAW, "eia_futures_page.htm"), get(FUT_PAGE_URL), FUT_PAGE_URL)
    os.makedirs(RAW, exist_ok=True)
    with open(LOG, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter="\t", lineterminator="\n"); w.writerow(["path", "source_url", "pulled_at_utc", "sha256", "size_bytes"]); w.writerows(LOG_ROWS)
    # local-only leg: World Bank Pink Sheet (Brent cross-check). A stop here ends this leg only.
    try:
        print("  World Bank commodity-markets page, then the monthly xlsx it links (local only)")
        page = get(WB_PAGE).decode("utf-8", "replace"); time.sleep(1.0)
        links = sorted(set(re.findall(r"https://thedocs\.worldbank\.org/en/doc/[0-9a-f]+-\d+/related/CMO-Historical-Data-Monthly\.xlsx", page)))
        if len(links) != 1: raise SystemExit(f"STOP: expected exactly one CMO-Historical-Data-Monthly.xlsx link on {WB_PAGE}, found {len(links)}")
        save(WB_X, get(links[0], magic=ZIP), links[0], log=False)
    except SystemExit as e:
        print(f"  World Bank leg stopped (local cross-check will be skipped): {e}")

def rd(path):
    with open(path, "rb") as f: return f.read()

def verify_raw():
    """Re-hash every raw file against pull_log.tsv (a mismatch means the saved snapshot was edited)."""
    if not os.path.exists(LOG): raise SystemExit("STOP: no pull_log.tsv; run without --offline first")
    n = bad = miss = 0
    with open(LOG, newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f, delimiter="\t"):
            p = os.path.join(ROOT, r["path"])
            if not os.path.exists(p): miss += 1; continue
            n += 1; bad += hashlib.sha256(rd(p)).hexdigest() != r["sha256"]
    print(f"  raw files re-hashed against pull_log.tsv: {n} checked, {bad} changed, {miss} absent from this checkout")
    if bad or miss: raise SystemExit("STOP: a saved raw file does not match its logged sha256 (or is missing)")


# ------------------------------------------------------------------------------------------------ parsers
def xl_date(v, datemode):
    import xlrd
    return xlrd.xldate_as_datetime(v, datemode).date()

def unit_of(title):
    m = re.search(r"\(([^()]*)\)\s*$", title.strip()); return m.group(1).strip() if m else None

def parse_hist(path, sid, unit):
    """One EIA dnav hist_xls workbook: sheet 'Data 1', row 2 Sourcekey, row 3 title with the unit, then (Excel date, value)."""
    import xlrd
    b = xlrd.open_workbook(path); sh = b.sheet_by_name("Data 1")
    assert sh.cell_value(1, 0) == "Sourcekey" and sh.cell_value(1, 1) == sid, (path, sh.cell_value(1, 1))
    title = sh.cell_value(2, 1); assert unit_of(title) == unit, f"{sid}: header unit {unit_of(title)!r} != {unit!r} ({title})"
    out = {}
    for r in range(3, sh.nrows):
        v = sh.cell_value(r, 1)
        if v != "": out[xl_date(sh.cell_value(r, 0), b.datemode)] = float(v)
    return out, title

def parse_psw01(path):
    """psw01.xls: sheets 'Data 1' (stocks) and 'Data 2' (supply); row 2 Sourcekey per column, row 3 title; Data 3 (4-week averages) is not used."""
    import xlrd
    b = xlrd.open_workbook(path); out, titles, release = {}, {}, None
    cs = b.sheet_by_name("Contents")
    for r in range(cs.nrows):
        if cs.cell_value(r, 1) == "Release Date:": release = cs.cell_value(r, 2)
    for name in ("Data 1", "Data 2"):
        sh = b.sheet_by_name(name); assert sh.cell_value(1, 0) == "Sourcekey"
        for c in range(1, sh.ncols):
            k = sh.cell_value(1, c); assert k not in out, f"psw01: duplicate {k}"
            d = {}
            for r in range(3, sh.nrows):
                v = sh.cell_value(r, c)
                if v != "": d[xl_date(sh.cell_value(r, 0), b.datemode)] = float(v)
            out[k], titles[k] = d, sh.cell_value(2, c)
    return out, titles, release

def parse_steo(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    dr = list(wb["Dates"].iter_rows(values_only=True)); info = {}
    for r in dr:
        if r[0] and str(r[0]).strip().startswith("Last Historical Month"): info["last_hist"] = int(r[3])
        if r[0] and str(r[0]).strip().startswith("Forecast Month"): info["release"] = str(r[3])
        if r[0] and str(r[0]).strip().startswith("Modeling and analysis completion"): info["completed"] = r[3].date()
        if r[1] == "AAAA_DATEX or AAAA_YEAR": ym = [int(c) for c in r[2:] if c is not None]
        if r[1] == "Historical": hist = [int(c) for c in r[2:] if c is not None]
    assert len(ym) == len(hist) == 72 and ym == sorted(ym), (len(ym), len(hist))
    assert [m for m, h in zip(ym, hist) if h] == [m for m in ym if m <= info["last_hist"]], "STEO: historical flags are not contiguous up to the last historical month"
    def tab(name, headers_ok):
        rows = list(wb[name].iter_rows(values_only=True)); keyed, section = {}, None
        hdr_years, hdr_months = rows[2], rows[3]
        year, got = None, []
        for j in range(2, 74):
            if hdr_years[j] is not None: year = int(hdr_years[j])
            got.append(year * 100 + list(calendar.month_abbr).index(hdr_months[j]))
        assert got == ym, f"STEO {name}: column dates do not match the Dates tab"
        for r in rows[4:]:
            if r[0] is None and r[1] is not None: section = str(r[1])
            if r[0] is not None and r[1] is not None and r[2] is not None:
                vals = [float(x) if isinstance(x, (int, float)) else None for x in r[2:74]]   # '-' = no data
                if r[0] in keyed: assert keyed[r[0]][1] == vals, f"STEO {name}: {r[0]} appears twice with different values"
                else: keyed[r[0]] = (section, vals, str(r[1]))
        for k, want in headers_ok.items():
            assert k in keyed and want in keyed[k][0].lower(), f"STEO {name}: {k} section {keyed.get(k, ('?',))[0]!r} lacks {want!r}"
        return keyed
    t3 = tab("3atab", {"papr_world": "million barrels per day", "copr_world": "million barrels per day", "papr_opec": "million barrels per day",
                       "patc_world": "million barrels per day", "t3_stchange_world": "million barrels per day",
                       "t3_stchange_us": "million barrels per day", "t3_stchange_ooecd": "million barrels per day",
                       "t3_stchange_noecd": "million barrels per day", "pasc_oecd_t3": "(million barrels)", "pasc_us": "(million barrels)",
                       "pasc_ooecd_t3": "(million barrels)"})
    t2 = tab("2tab", {"BREPUUS": "dollars per barrel", "WTIPUUS": "dollars per barrel"})
    return info, ym, hist, t3, t2

def parse_table1(path):
    """WPSR Table 1 as published: block 1 stocks (million barrels), block 2 the supply lines by number (kb/d)."""
    txt = rd(path).decode("utf-8", "replace"); rows = list(csv.reader(txt.splitlines()))
    def d(s):
        m, dd, yy = [int(x) for x in s.split("/")]; return Dt(2000 + yy, m, dd)
    h1 = rows[0]; assert h1[0] == "STUB_1"
    dates1 = (d(h1[1]), d(h1[2]), d(h1[5]))
    stocks = {}
    for r in rows[1:]:
        if r[0] == "STUB_1": break
        stocks[r[0]] = tuple(float(r[i].replace(",", "")) for i in (1, 2, 5))
    i2 = [i for i, r in enumerate(rows) if r[0] == "STUB_1" and i > 0][0]; h2 = rows[i2]
    dates2 = (d(h2[2]), d(h2[3]), d(h2[5])); assert dates1 == dates2, (dates1, dates2)
    lines = {}
    for r in rows[i2 + 1:]:
        if len(r) < 6: continue
        m = re.match(r"\((\d+)\)", r[1].strip())
        if m and r[0].strip() in ("Crude Oil Supply", "Products Supplied"):
            try: lines[int(m.group(1))] = tuple(float(r[i].replace(",", "")) for i in (2, 3, 5))
            except ValueError: pass
    return dates1, stocks, lines

def load_wb():
    """Local-only: {YYYYMM: (brent, wti)} from the World Bank Pink Sheet, or None if the file is absent."""
    if not os.path.exists(WB_X): return None
    import openpyxl
    wb = openpyxl.load_workbook(WB_X, read_only=True, data_only=True); rows = list(wb["Monthly Prices"].iter_rows(values_only=True))
    hdr, units = rows[4], rows[5]; ix = {}
    for k, name in (("brent", "Crude oil, Brent"), ("wti", "Crude oil, WTI")):
        j = hdr.index(name); assert units[j] == "($/bbl)", (name, units[j]); ix[k] = j
    out = {}
    for r in rows[6:]:
        if r[0] and re.fullmatch(r"\d{4}M\d\d", str(r[0])):
            out[int(r[0][:4]) * 100 + int(r[0][5:])] = tuple(float(r[ix[k]]) if isinstance(r[ix[k]], (int, float)) else None for k in ("brent", "wti"))
    return out


# ------------------------------------------------------------------------------------------------ calendar helpers
def shift_year(d, k):
    try: return d.replace(year=d.year - k)
    except ValueError: return d.replace(year=d.year - k, day=28)

def nearest(dsorted, target, tol=3):
    i = bisect.bisect_left(dsorted, target)
    c = [dsorted[j] for j in (i - 1, i) if 0 <= j < len(dsorted)]
    if not c: return None
    b = min(c, key=lambda x: abs((x - target).days))
    return b if abs((b - target).days) <= tol else None

def last_on_or_before(dsorted, d):
    i = bisect.bisect_right(dsorted, d); return dsorted[i - 1] if i else None

def week_avg(daily, w):
    """Mean of the daily prints from the Saturday before to the Friday w; (mean, n days)."""
    v = [daily[w - datetime.timedelta(days=k)] for k in range(7) if (w - datetime.timedelta(days=k)) in daily]
    return (sum(v) / len(v), len(v)) if v else (None, 0)

def avg4(ser, w):
    v = [ser.get(w - datetime.timedelta(days=7 * k)) for k in range(4)]
    return None if any(x is None for x in v) else sum(v) / 4.0

def five_year(ser, dsorted, w, change=False):
    vals = []
    for k in range(1, 6):
        t = nearest(dsorted, shift_year(w, k))
        if t is None or t not in ser: return None
        if change:
            if (t - ONE_WEEK) not in ser: return None
            vals.append(ser[t] - ser[t - ONE_WEEK])
        else: vals.append(ser[t])
    return min(vals), max(vals), sum(vals) / 5.0

def iso_prior(ser, w, k, change=False):
    """The Friday with the same ISO week number k years earlier (week 53 maps to the last week of a 52-week year)."""
    y, wk, _ = w.isocalendar(); ty = y - k
    t = Dt.fromisocalendar(ty, min(wk, Dt(ty, 12, 28).isocalendar()[1]), 5)
    if t not in ser: return None
    if change:
        return ser[t] - ser[t - ONE_WEEK] if (t - ONE_WEEK) in ser else None
    return ser[t]

def five_year_iso(ser, w, change=False):
    v = [iso_prior(ser, w, k, change) for k in range(1, 6)]
    return None if any(x is None for x in v) else (min(v), max(v), sum(v) / 5.0)

def vs_range(x, rng):
    if x is None or rng is None: return None
    return "above" if x > rng[1] else ("below" if x < rng[0] else "within")


# ------------------------------------------------------------------------------------------------ the weekly ledger
def build_ledger(H, PS, WTI, BRENT):
    weeks = [d for d in sorted(H["WCRSTUS1"]) if d >= LEDGER_START]
    assert weeks[0] == LEDGER_START and all((weeks[i] - weeks[i - 1]).days == 7 for i in range(1, len(weeks))), "ledger weeks are not consecutive"
    latest = {k: max(v) for k, v in H.items()}; assert len(set(latest.values())) == 1, f"series end on different weeks: {latest}"
    comm_d = sorted(H["WCESTUS1"])
    S = lambda k, w: (H[k] if k in H else PS[k]).get(w)
    rows = []
    for w in weeks:
        p = w - ONE_WEEK; r = {"week_end": w.isoformat()}
        r["comm_stocks_kb"], r["spr_kb"], r["crude_total_kb"] = S("WCESTUS1", w), S("WCSSTUS1", w), S("WCRSTUS1", w)
        r["cushing_kb"], r["gasoline_kb"], r["distillate_kb"], r["petroleum_total_kb"] = S("W_EPC0_SAX_YCUOK_MBBL", w), S("WGTSTUS1", w), S("WDISTUS1", w), S("WTTSTUS1", w)
        for nm, k in (("comm", "WCESTUS1"), ("spr", "WCSSTUS1"), ("crude_total", "WCRSTUS1"), ("cushing", "W_EPC0_SAX_YCUOK_MBBL"),
                      ("gasoline", "WGTSTUS1"), ("distillate", "WDISTUS1"), ("petroleum_total", "WTTSTUS1")):
            r[f"d_{nm}_kb"] = sub(S(k, w), S(k, p))
        assert r["crude_total_kb"] == r["comm_stocks_kb"] + r["spr_kb"], f"WCRSTUS1 != WCESTUS1 + WCSSTUS1 in {w}"
        r["production_kbd"], r["imports_kbd"], r["exports_kbd"], r["net_imports_kbd"] = S("WCRFPUS2", w), S("WCRIMUS2", w), S("WCREXUS2", w), S("WCRNTUS2", w)
        r["refinery_input_kbd"], r["product_supplied_kbd"] = S("WCRRIUS2", w), S("WRPUPUS2", w)
        r["unaccounted_kbd"], r["transfers_kbd"], r["stock_change_published_kbd"] = S("WCRAUUS2", w), S("W_EPC0_TVP_NUS_MBBLD", w), S("W_EPC0_SCG_NUS_MBBLD", w)
        for k in ("production_kbd", "net_imports_kbd", "refinery_input_kbd", "product_supplied_kbd", "stock_change_published_kbd"):
            assert r[k] is not None, f"{k} missing in the week ending {w}"
        both = r["imports_kbd"] is not None and r["exports_kbd"] is not None     # weekly imports / exports detail is missing before 2010-10 / 2007-06
        if both: assert abs(r["net_imports_kbd"] - (r["imports_kbd"] - r["exports_kbd"])) <= 1.0, f"WCRNTUS2 != imports - exports (+-1 kb/d) in {w}"
        netimp = r["imports_kbd"] - r["exports_kbd"] if both else r["net_imports_kbd"]
        r["imports_basis"] = "imports_minus_exports" if both else "published_net_imports_wcrntus2"
        r["transfers_in_identity"] = "published" if r["transfers_kbd"] is not None else "zero_line_did_not_exist"
        tr = r["transfers_kbd"] if r["transfers_kbd"] is not None else 0.0
        r["levels_vs_published_kb"] = r["d_crude_total_kb"] - 7.0 * r["stock_change_published_kbd"]
        rev = abs(r["levels_vs_published_kb"]) > REV_FLAG
        r["revision_week"] = 1 if rev else 0
        if r["unaccounted_kbd"] is None:
            r["bal_status"] = "unaccounted_for_not_published"       # nine weeks in 2005-2007 where EIA's file has no value for the balancing item
        else:
            r["bal_status"] = ("revision_week" if rev else "asserted") if w >= ASSERT_FROM else "before_2017_not_asserted"
            rhs = 7.0 * (r["production_kbd"] + tr + netimp + r["unaccounted_kbd"] - r["refinery_input_kbd"])
            r["bal_supply_minus_disposition_kb"] = rhs
            r["bal_residual_published_kb"] = 7.0 * r["stock_change_published_kbd"] - rhs
            r["bal_residual_kb"] = r["d_crude_total_kb"] - rhs
            r["bal_without_transfers_residual_kb"] = r["d_crude_total_kb"] - 7.0 * (r["production_kbd"] + netimp + r["unaccounted_kbd"] - r["refinery_input_kbd"])
            # total petroleum (transfers cancel between crude and products)
            oth, adjo, pe = S("W_EPP1_YPT_NUS_MBBLD", w), S("W_EPP0_VUA_NUS_MBBLD", w), S("WPESCUS2", w)
            r["other_production_kbd"], r["net_imports_all_kbd"], r["products_adjustment_kbd"] = oth, S("WTTNTUS2", w), adjo
            if oth is not None and adjo is not None and pe is not None and r["net_imports_all_kbd"] is not None:
                trhs = 7.0 * (r["production_kbd"] + oth + r["net_imports_all_kbd"] + r["unaccounted_kbd"] + adjo - r["product_supplied_kbd"])
                r["tp_bal_residual_kb"] = r["d_petroleum_total_kb"] - trhs
                r["tp_levels_vs_published_kb"] = r["d_petroleum_total_kb"] - 7.0 * (r["stock_change_published_kbd"] + pe)
                r["tp_revision_week"] = 1 if abs(r["tp_levels_vs_published_kb"]) > REV_FLAG else 0
        # five-year range, commercial stocks
        ly, lc = five_year(H["WCESTUS1"], comm_d, w), five_year(H["WCESTUS1"], comm_d, w, change=True)
        if ly:
            r["comm_5y_min_kb"], r["comm_5y_max_kb"], r["comm_5y_mean_kb"] = ly; r["comm_vs_5y"] = vs_range(r["comm_stocks_kb"], ly)
        if lc:
            r["d_comm_5y_min_kb"], r["d_comm_5y_max_kb"], r["d_comm_5y_mean_kb"] = lc; r["d_comm_vs_5y"] = vs_range(r["d_comm_kb"], lc)
        ii, ic = five_year_iso(H["WCESTUS1"], w), five_year_iso(H["WCESTUS1"], w, change=True)
        if ii:
            r["comm_5yiso_min_kb"], r["comm_5yiso_max_kb"], r["comm_5yiso_mean_kb"] = ii; r["comm_vs_5yiso"] = vs_range(r["comm_stocks_kb"], ii)
        if ic:
            r["d_comm_5yiso_min_kb"], r["d_comm_5yiso_max_kb"], r["d_comm_5yiso_mean_kb"] = ic; r["d_comm_vs_5yiso"] = vs_range(r["d_comm_kb"], ic)
        # 4-week averages
        for nm, ser in (("production", H["WCRFPUS2"]), ("net_imports", H["WCRNTUS2"]), ("imports", H["WCRIMUS2"]), ("exports", H["WCREXUS2"]),
                        ("refinery_input", H["WCRRIUS2"]), ("product_supplied", H["WRPUPUS2"])):
            r[f"{nm}_4wk_kbd"] = avg4(ser, w)
        # prices and dollars
        wt, nw = week_avg(WTI, w); br, nb = week_avg(BRENT, w)
        r["wti_avg_usd_bbl"], r["wti_n_days"], r["brent_avg_usd_bbl"], r["brent_n_days"] = wt, nw, br, nb
        r["comm_change_usd_mn"], r["spr_change_usd_mn"] = usd_mn(r["d_comm_kb"], wt), usd_mn(r["d_spr_kb"], wt)
        r["crude_total_change_usd_mn"] = usd_mn(r["d_crude_total_kb"], wt)
        r["comm_change_usd_mn_brent"], r["spr_change_usd_mn_brent"] = usd_mn(r["d_comm_kb"], br), usd_mn(r["d_spr_kb"], br)
        for nm, k in (("production", "production_kbd"), ("imports", "imports_kbd"), ("exports", "exports_kbd"), ("net_imports", "net_imports_kbd"),
                      ("refinery_input", "refinery_input_kbd"), ("unaccounted", "unaccounted_kbd")):
            r[f"{nm}_usd_mn"] = usd_mn(mul(r[k], 7.0), wt)
        r["comm_value_usd_bn"] = None if wt is None else r["comm_stocks_kb"] * wt / 1e6
        r["spr_value_usd_bn"] = None if wt is None else r["spr_kb"] * wt / 1e6
        rows.append(r)
    # cumulative since end-2025
    anchor = last_on_or_before(weeks, Dt(2025, 12, 31)); cum = dict(comm=0.0, spr=0.0, cu=0.0, su=0.0, cub=0.0, sub=0.0, una=0.0); ok = True
    for r in rows:
        if Dt.fromisoformat(r["week_end"]) <= anchor: continue
        for a, b in (("comm", "d_comm_kb"), ("spr", "d_spr_kb"), ("cu", "comm_change_usd_mn"), ("su", "spr_change_usd_mn"),
                     ("cub", "comm_change_usd_mn_brent"), ("sub", "spr_change_usd_mn_brent")):
            cum[a] = None if cum[a] is None or r[b] is None else cum[a] + r[b]
        cum["una"] += 7.0 * r["unaccounted_kbd"]
        r["since_end2025_comm_change_kb"], r["since_end2025_spr_change_kb"] = cum["comm"], cum["spr"]
        r["since_end2025_comm_change_usd_mn"], r["since_end2025_spr_change_usd_mn"] = cum["cu"], cum["su"]
        r["since_end2025_comm_change_usd_mn_brent"], r["since_end2025_spr_change_usd_mn_brent"] = cum["cub"], cum["sub"]
        r["since_end2025_unaccounted_kb"] = cum["una"]
    return rows, anchor

LEDGER_FIELDS = (["week_end", "comm_stocks_kb", "spr_kb", "crude_total_kb", "cushing_kb", "gasoline_kb", "distillate_kb", "petroleum_total_kb",
                  "d_comm_kb", "d_spr_kb", "d_crude_total_kb", "d_cushing_kb", "d_gasoline_kb", "d_distillate_kb", "d_petroleum_total_kb",
                  "production_kbd", "transfers_kbd", "imports_kbd", "exports_kbd", "net_imports_kbd", "unaccounted_kbd", "refinery_input_kbd",
                  "product_supplied_kbd", "stock_change_published_kbd", "other_production_kbd", "net_imports_all_kbd", "products_adjustment_kbd",
                  "transfers_in_identity", "imports_basis", "bal_supply_minus_disposition_kb", "bal_residual_kb", "bal_residual_published_kb", "levels_vs_published_kb",
                  "revision_week", "bal_status", "bal_without_transfers_residual_kb", "tp_bal_residual_kb", "tp_levels_vs_published_kb", "tp_revision_week",
                  "comm_5y_min_kb", "comm_5y_max_kb", "comm_5y_mean_kb", "comm_vs_5y", "d_comm_5y_min_kb", "d_comm_5y_max_kb", "d_comm_5y_mean_kb", "d_comm_vs_5y",
                  "comm_5yiso_min_kb", "comm_5yiso_max_kb", "comm_5yiso_mean_kb", "comm_vs_5yiso", "d_comm_5yiso_min_kb", "d_comm_5yiso_max_kb", "d_comm_5yiso_mean_kb", "d_comm_vs_5yiso",
                  "production_4wk_kbd", "net_imports_4wk_kbd", "imports_4wk_kbd", "exports_4wk_kbd", "refinery_input_4wk_kbd", "product_supplied_4wk_kbd",
                  "wti_avg_usd_bbl", "wti_n_days", "brent_avg_usd_bbl", "brent_n_days",
                  "comm_change_usd_mn", "spr_change_usd_mn", "crude_total_change_usd_mn", "comm_change_usd_mn_brent", "spr_change_usd_mn_brent",
                  "production_usd_mn", "imports_usd_mn", "exports_usd_mn", "net_imports_usd_mn", "refinery_input_usd_mn", "unaccounted_usd_mn",
                  "comm_value_usd_bn", "spr_value_usd_bn",
                  "since_end2025_comm_change_kb", "since_end2025_spr_change_kb", "since_end2025_unaccounted_kb",
                  "since_end2025_comm_change_usd_mn", "since_end2025_spr_change_usd_mn",
                  "since_end2025_comm_change_usd_mn_brent", "since_end2025_spr_change_usd_mn_brent"])


def dist(xs):
    a = [abs(x) for x in xs]; n = len(a)
    return dict(n=n, mean_signed=sum(xs) / n, max=max(a), p50=pctl(a, 50), p90=pctl(a, 90), p95=pctl(a, 95), p99=pctl(a, 99),
                le7=sum(1 for x in a if x <= 7) / n, le14=sum(1 for x in a if x <= 14) / n, le21=sum(1 for x in a if x <= 21) / n)

def assert_identities(rows, H, PS):
    """Hard stops. Returns the statistics the report prints."""
    st = {}
    st["n_missing_ua"] = sum(1 for r in rows if r["bal_status"] == "unaccounted_for_not_published")
    rows = [r for r in rows if r["bal_status"] != "unaccounted_for_not_published"]
    # (0) two routes agree; stocks add up; line 18; transfers sign and parts
    mism = []
    for sid, _, _ in WEEKLY:
        if sid in PS:
            a, b = H[sid], PS[sid]; common = set(a) & set(b)
            bad = [d for d in common if a[d] != b[d]]
            mism.append((sid, len(a), len(b), len(common), len(bad)))
            assert not bad and len(a) == len(b) == len(common), f"{sid}: the dnav file and psw01.xls differ ({len(bad)} values, {len(a)} vs {len(b)} weeks)"
    st["routes"] = mism
    ds = [d for d in PS["W_EPP1_YPT_NUS_MBBLD"] if all(d in PS[k] for k in ("W_EPL0_YPG_NUS_MBBLD", "W_EPOOXR_YPT_NUS_MBBLD", "WPGNPUS2"))]
    worst18 = max(abs(PS["W_EPP1_YPT_NUS_MBBLD"][d] - PS["W_EPL0_YPG_NUS_MBBLD"][d] - PS["W_EPOOXR_YPT_NUS_MBBLD"][d] - PS["WPGNPUS2"][d]) for d in ds)
    assert worst18 <= 1.0, f"line 18 != NGPL + renewables + processing gain (max {worst18})"
    st["line18"] = (len(ds), worst18)
    tv = PS["W_EPC0_TVP_NUS_MBBLD"]
    assert min(tv) == TRANSFERS_START, f"transfers series starts {min(tv)}, expected {TRANSFERS_START}"
    assert all(PS["W_EPC0NG_TVP_NUS_MBBLD"][d] == -tv[d] for d in tv) and set(tv) == set(PS["W_EPC0NG_TVP_NUS_MBBLD"]), "products-side transfers != -crude-side transfers"
    worst_t = max(abs(tv[d] - PS["W_EPLLCL2A_TVP_NUS_MBBLD"][d] - PS["W_EPLLNGPU_TVP_NUS_MBBLD"][d]) for d in tv)
    assert worst_t <= 1.0, f"transfers != its two parts (max {worst_t})"
    st["transfers"] = (len(tv), worst_t)
    sc = PS["W_EPC0_SCG_NUS_MBBLD"]
    worst_sc = max(abs(sc[d] - PS["WCESCUS2"][d] - PS["WCSSCUS2"][d]) for d in sc if d in PS["WCESCUS2"] and d in PS["WCSSCUS2"])
    assert worst_sc <= 1.0, f"stock change != commercial + SPR (max {worst_sc})"
    st["scg_parts"] = worst_sc
    # (1) the balance sheet closes as published, every week
    r1 = [r["bal_residual_published_kb"] for r in rows]
    assert max(abs(x) for x in r1) <= TOL_PUBLISHED, f"published balance sheet does not close: max {max(abs(x) for x in r1)} kb"
    st["R1"] = dist(r1)
    # (2) the brief's identity on the stock levels
    win = [r for r in rows if Dt.fromisoformat(r["week_end"]) >= ASSERT_FROM]
    rev = [r for r in win if r["revision_week"]]
    assert len(rev) <= MAX_REV_SHARE * len(win), f"{len(rev)} revision weeks in {len(win)}"
    ok = [r["bal_residual_kb"] for r in win if not r["revision_week"]]
    assert max(abs(x) for x in ok) <= TOL_LEVELS, f"level identity fails: max {max(abs(x) for x in ok)} kb"
    st["R2"] = dist(ok); st["R2_all"] = dist([r["bal_residual_kb"] for r in win]); st["rev"] = [(r["week_end"], r["levels_vs_published_kb"], r["bal_residual_kb"]) for r in rev]
    st["R2_nrev"], st["R2_nwin"] = len(rev), len(win)
    st["R2_post_nov23"] = dist([r["bal_residual_kb"] for r in win if not r["revision_week"] and Dt.fromisoformat(r["week_end"]) >= TRANSFERS_START])
    st["R2_2026"] = dist([r["bal_residual_kb"] for r in win if not r["revision_week"] and r["week_end"] >= "2026-01-01"])
    pre = [r["bal_residual_kb"] for r in rows if Dt.fromisoformat(r["week_end"]) < ASSERT_FROM]
    st["R2_pre"] = dist(pre); st["R2_pre_n_gt22"] = sum(1 for x in pre if abs(x) > TOL_LEVELS)
    # first-week-of-month pattern before 2017 (inference check)
    big = [r for r in rows if Dt.fromisoformat(r["week_end"]) < Dt(2016, 1, 1) and r["week_end"] >= "2010-01-01" and abs(r["levels_vs_published_kb"]) > REV_FLAG]
    st["pre_big"] = (len(big), sum(1 for r in big if Dt.fromisoformat(r["week_end"]).day <= 7))
    # check on the check: without the transfers term the identity fails after Nov 2023
    miss = [abs(r["bal_without_transfers_residual_kb"]) for r in win if r["transfers_kbd"] is not None]
    assert min(miss) > 1000, "dropping the transfers line should break the identity by thousands of kb"
    st["no_transfers"] = (len(miss), statistics.median(miss), max(miss))
    # (3) total petroleum
    tw = [r for r in win if r.get("tp_bal_residual_kb") is not None]
    trev = [r for r in tw if r["tp_revision_week"]]
    assert len(trev) <= MAX_REV_SHARE * len(tw), f"total petroleum: {len(trev)} revision weeks in {len(tw)}"
    tok = [r["tp_bal_residual_kb"] for r in tw if not r["tp_revision_week"]]
    assert max(abs(x) for x in tok) <= TOL_LEVELS, f"total-petroleum identity fails: max {max(abs(x) for x in tok)} kb"
    st["TP"] = dist(tok); st["TP_rev"] = [(r["week_end"], r["tp_levels_vs_published_kb"], r["tp_bal_residual_kb"]) for r in trev]; st["TP_n"] = len(tw)
    # yearly table of the balancing item
    by, bt = {}, {}
    for r in rows:
        y = int(r["week_end"][:4]); by.setdefault(y, []).append(r["unaccounted_kbd"])
        if r["transfers_kbd"] is not None: bt.setdefault(y, []).append(r["transfers_kbd"])
    st["una_year"] = {y: (statistics.mean(v), statistics.mean(bt[y]) if y in bt else None, len(v), len(bt.get(y, []))) for y, v in by.items() if y >= 2015}
    pf = [r["week_end"] for r in rows if r["week_end"] < ASSERT_FROM.isoformat() and r["revision_week"]]
    st["last_pre_flag"] = pf[-1] if pf else None
    return st


# ------------------------------------------------------------------------------------------------ the world, monthly
def build_world(steo, BRENT, WTI):
    info, ym, hist, t3, t2 = steo
    G = lambda k: t3[k][1]
    prod, cons, wd = G("papr_world"), G("patc_world"), G("t3_stchange_world")
    # STEO's own arithmetic, asserted (it carries many decimals, so the tolerance is tight)
    for i in range(72):
        assert abs(prod[i] - G("copr_world")[i] - G("world_nc")[i]) < 1e-4
        assert abs(wd[i] - (cons[i] - prod[i])) < 1e-4, f"STEO: withdrawal != consumption - production in {ym[i]}"
        assert abs(wd[i] - G("t3_stchange_us")[i] - G("t3_stchange_ooecd")[i] - G("t3_stchange_noecd")[i]) < 1e-4
        assert abs(G("pasc_oecd_t3")[i] - G("pasc_us")[i] - G("pasc_ooecd_t3")[i]) < 1e-3
    daily_last = max(BRENT)
    byb, byw = {}, {}
    for d, v in BRENT.items(): byb.setdefault(d.year * 100 + d.month, []).append(v)
    for d, v in WTI.items(): byw.setdefault(d.year * 100 + d.month, []).append(v)
    rows, chk = [], []
    for i, m in enumerate(ym):
        y, mo = divmod(m, 100); days = calendar.monthrange(y, mo)[1]; end = Dt(y, mo, days)
        complete = (m in byb) and daily_last >= end - datetime.timedelta(days=3)
        r = {"month": mon(m), "steo_basis": "historical" if hist[i] else "forecast", "days": days,
             "world_production_mbd": prod[i], "world_crude_production_mbd": G("copr_world")[i], "opec_production_mbd": G("papr_opec")[i],
             "world_consumption_mbd": cons[i], "implied_net_withdrawal_mbd": cons[i] - prod[i], "steo_net_withdrawal_mbd": wd[i],
             "us_withdrawal_mbd": G("t3_stchange_us")[i], "other_oecd_withdrawal_mbd": G("t3_stchange_ooecd")[i],
             "other_draws_and_balance_mbd": G("t3_stchange_noecd")[i], "implied_net_withdrawal_mb": (cons[i] - prod[i]) * days,
             "oecd_commercial_stocks_mb": G("pasc_oecd_t3")[i], "us_commercial_stocks_mb": G("pasc_us")[i], "other_oecd_commercial_stocks_mb": G("pasc_ooecd_t3")[i]}
        r["oecd_commercial_stock_change_mb"] = None if i == 0 else G("pasc_oecd_t3")[i] - G("pasc_oecd_t3")[i - 1]
        b = sum(byb[m]) / len(byb[m]) if m in byb else None
        r["brent_avg_usd_bbl"] = b if (complete and hist[i]) else None
        r["brent_n_days"] = len(byb.get(m, [])) or None
        r["steo_brent_usd_bbl"] = t2["BREPUUS"][1][i]
        r["implied_net_withdrawal_usd_bn"] = None if r["brent_avg_usd_bbl"] is None else r["implied_net_withdrawal_mb"] * r["brent_avg_usd_bbl"] / 1000.0
        if hist[i] and complete and m in byb:
            chk.append((m, abs(t2["BREPUUS"][1][i] - b), abs(t2["WTIPUUS"][1][i] - sum(byw[m]) / len(byw[m]))))
        rows.append(r)
    worst = max(chk, key=lambda x: x[1]), max(chk, key=lambda x: x[2])
    assert worst[0][1] < 0.011 and worst[1][2] < 0.011, f"STEO spot averages differ from the daily files: {worst}"
    return rows, dict(chk_n=len(chk), chk_brent=worst[0], chk_wti=worst[1])

WORLD_FIELDS = ["month", "steo_basis", "days", "world_production_mbd", "world_crude_production_mbd", "opec_production_mbd", "world_consumption_mbd",
                "implied_net_withdrawal_mbd", "steo_net_withdrawal_mbd", "us_withdrawal_mbd", "other_oecd_withdrawal_mbd", "other_draws_and_balance_mbd",
                "implied_net_withdrawal_mb", "oecd_commercial_stocks_mb", "us_commercial_stocks_mb", "other_oecd_commercial_stocks_mb",
                "oecd_commercial_stock_change_mb", "brent_avg_usd_bbl", "brent_n_days", "steo_brent_usd_bbl", "implied_net_withdrawal_usd_bn"]

def wnd(name):
    if name.endswith("_mbd"): return 3
    if name.endswith("_mb"): return 2
    if name.endswith("_usd_bbl"): return 2
    if name.endswith("_usd_bn"): return 3
    return 3

def world_window(world, a, b):
    """Sum over months a..b (YYYY-MM inclusive)."""
    rs = [r for r in world if a <= r["month"] <= b]
    n_h = sum(1 for r in rs if r["steo_basis"] == "historical"); n_f = len(rs) - n_h
    wd = sum(r["implied_net_withdrawal_mb"] for r in rs); days = sum(r["days"] for r in rs)
    usd = [r["implied_net_withdrawal_usd_bn"] for r in rs]
    us = sum(r["us_withdrawal_mbd"] * r["days"] for r in rs); oo = sum(r["other_oecd_withdrawal_mbd"] * r["days"] for r in rs); nn = sum(r["other_draws_and_balance_mbd"] * r["days"] for r in rs)
    first = world.index(rs[0]); oecd0 = world[first - 1]["oecd_commercial_stocks_mb"] if first > 0 else None
    return dict(a=a, b=b, n_h=n_h, n_f=n_f, wd_mb=wd, wd_mbd=wd / days, usd_bn=None if any(u is None for u in usd) else sum(usd),
                us_mb=us, oo_mb=oo, nn_mb=nn, oecd0=oecd0, oecd1=rs[-1]["oecd_commercial_stocks_mb"],
                prod_mbd=sum(r["world_production_mbd"] * r["days"] for r in rs) / days, cons_mbd=sum(r["world_consumption_mbd"] * r["days"] for r in rs) / days)


# ------------------------------------------------------------------------------------------------ Brent check
def brent_check(WTI, BRENT):
    ds = sorted(set(WTI) & set(BRENT)); sp = {d: BRENT[d] - WTI[d] for d in ds}
    L = dict(); spike = Dt(2026, 10, 2)
    bd = sorted(BRENT); i = bd.index(spike); prev, nxt = bd[i - 1], bd[i + 1:]
    L["spike_date"], L["prev_date"], L["prev"], L["spike"] = spike, prev, BRENT[prev], BRENT[spike]
    L["jump_pct"] = 100.0 * (BRENT[spike] / BRENT[prev] - 1.0); L["wti_prev"], L["wti_spike"] = WTI[prev], WTI[spike]
    L["after"] = [(d, BRENT[d], WTI.get(d), sp.get(d)) for d in nxt]
    L["spread_spike"] = sp[spike]; L["spread_prev"] = sp[prev]
    L["rank_spread"] = 1 + sum(1 for v in sp.values() if v > sp[spike]); L["n_spread"] = len(sp)
    L["top_spreads"] = sorted(((v, d) for d, v in sp.items()), reverse=True)[:6]
    L["prior_max_ex2020"] = max((v, d) for d, v in sp.items() if not (Dt(2020, 3, 1) <= d <= Dt(2020, 6, 30)) and d < Dt(2026, 1, 1))
    chg = [(math.log(BRENT[bd[k]] / BRENT[bd[k - 1]]), bd[k]) for k in range(1, len(bd)) if BRENT[bd[k]] > 0 and BRENT[bd[k - 1]] > 0]
    ups = sorted(chg, reverse=True); L["rank_up"] = 1 + sum(1 for c, d in chg if c > math.log(BRENT[spike] / BRENT[prev])); L["n_chg"] = len(chg); L["top_up"] = ups[:6]
    # spread level by month, 2026
    msp = {}
    for d, v in sp.items():
        if d >= Dt(2026, 1, 1): msp.setdefault(d.month, []).append(v)
    L["spread_2026"] = {m: (sum(v) / len(v), min(v), max(v), len(v)) for m, v in sorted(msp.items())}
    # sensitivity of the week ending 2 Oct
    wk = [spike - datetime.timedelta(days=k) for k in range(7)]
    v = [BRENT[d] for d in wk if d in BRENT]; v2 = [BRENT[prev] if d == spike else BRENT[d] for d in wk if d in BRENT]
    L["wk_avg"], L["wk_avg_alt"], L["wk_n"] = sum(v) / len(v), sum(v2) / len(v2), len(v)
    # WTI outlier of 25 Sep
    w925 = Dt(2026, 9, 25); wd_ = sorted(WTI); j = wd_.index(w925); L["wti_925"] = (WTI[wd_[j - 1]], WTI[w925], WTI[wd_[j + 1]], WTI[wd_[j + 2]])
    L["last"] = [(d, WTI.get(d), BRENT.get(d), sp.get(d), (100 * (BRENT[d] / BRENT[bd[bd.index(d) - 1]] - 1)) if d in BRENT and bd.index(d) > 0 else None) for d in ds[-60:]]
    return L


# ------------------------------------------------------------------------------------------------ report
def table(headers, rows, left=(0,)):
    w = [max(len(str(h)), *(len(str(r[i])) for r in rows)) for i, h in enumerate(headers)]
    def line(r): return "  ".join(str(c).ljust(w[i]) if i in left else str(c).rjust(w[i]) for i, c in enumerate(r)).rstrip()
    return [line(headers), "  ".join("-" * x for x in w)] + [line(r) for r in rows]

def n0(x): return "n/a" if x is None else f"{x:,.0f}"
def n1(x): return "n/a" if x is None else f"{x:,.1f}"
def n2(x): return "n/a" if x is None else f"{x:,.2f}"
def sg0(x): return "n/a" if x is None else f"{(round(x, 0) + 0.0):+,.0f}"
def sg1(x): return "n/a" if x is None else f"{(round(x, 1) + 0.0):+,.1f}"

def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--offline", action="store_true", help="rebuild from the saved raw files; no network")
    a = ap.parse_args()
    print("Crude oil physical ledger" + (" (offline, from the saved raw files)" if a.offline else ""))
    if not a.offline: pull()
    verify_raw()

    # ---- load everything from the saved raw files (the only path to the data, online or offline)
    H, HT = {}, {}
    for sid, unit, _ in WEEKLY: H[sid], HT[sid] = parse_hist(os.path.join(RAW, f"eia_{sid}w.xls"), sid, unit)
    PS, PT, psw_release = parse_psw01(os.path.join(RAW, "eia_wpsr_psw01.xls"))
    for k, unit in PSW_KEYS.items():
        assert k in PS, f"psw01.xls has no {k}"; assert unit_of(PT[k]) == unit, f"{k}: header unit {unit_of(PT[k])!r} != {unit!r}"
    WTI, wti_t = parse_hist(os.path.join(RAW, "eia_RWTCd.xls"), "RWTC", "Dollars per Barrel")
    BRENT, br_t = parse_hist(os.path.join(RAW, "eia_RBRTEd.xls"), "RBRTE", "Dollars per Barrel")
    FUT1, fut_t = parse_hist(os.path.join(RAW, "eia_RCLC1d.xls"), "RCLC1", "Dollars per Barrel")
    steo = parse_steo(os.path.join(RAW, "eia_STEO_m.xlsx"))
    t1_dates, t1_stocks, t1_lines = parse_table1(os.path.join(RAW, "eia_wpsr_table1.csv"))
    fut_html = rd(os.path.join(RAW, "eia_futures_page.htm")).decode("utf-8", "replace")
    fut_title = re.search(r"<title>(.*?)</title>", fut_html, re.S).group(1).strip()

    rows, anchor = build_ledger(H, PS, WTI, BRENT)
    by = {Dt.fromisoformat(r["week_end"]): r for r in rows}; weeks = sorted(by)
    st = assert_identities(rows, H, PS)
    world, wchk = build_world(steo, BRENT, WTI)
    info = steo[0]; latest = weeks[-1]
    bc = brent_check(WTI, BRENT)

    # ---- outputs: ledger, world, history files
    write_csv(os.path.join(OUT, "us_weekly_ledger.csv"), LEDGER_FIELDS, rows)
    write_csv(os.path.join(OUT, "world_monthly_balance.csv"), WORLD_FIELDS, world, nd=wnd)
    write_csv(os.path.join(OUT, "brent_spread_check.csv"), ["date", "wti_usd_bbl", "brent_usd_bbl", "brent_minus_wti_usd_bbl", "brent_1day_change_pct"],
              [{"date": d.isoformat(), "wti_usd_bbl": w, "brent_usd_bbl": b, "brent_minus_wti_usd_bbl": s, "brent_1day_change_pct": c} for d, w, b, s, c in bc["last"]],
              nd=lambda c: 2)
    for nm, sid, nd in (("us_commercial_stocks_kb", "WCESTUS1", 0), ("us_spr_stocks_kb", "WCSSTUS1", 0), ("us_cushing_stocks_kb", "W_EPC0_SAX_YCUOK_MBBL", 0),
                        ("us_field_production_kbd", "WCRFPUS2", 0), ("us_exports_kbd", "WCREXUS2", 0)):
        write_hist(nm, sorted(H[sid].items()), nd)
    write_hist("us_balance_residual_kb", [(Dt.fromisoformat(r["week_end"]), r["bal_residual_kb"]) for r in rows if r.get("bal_residual_kb") is not None], 0)

    # ---- checkpoints
    chk_week = {}
    for lab, d in CHECKPOINTS: chk_week[lab] = latest if d is None else last_on_or_before(weeks, d)
    ck = []   # tidy rows
    def C(block, label, a_, b_, measure, value, unit, note=""):
        nd = {"kb": 0, "kb/d": 1, "$mn": 1, "mb": 1, "$bn": 3, "$/bbl": 2, "mb/d": 3, "share of weeks": 3}.get(unit, 3 if unit not in ("weeks", "months") else 0)
        if unit == "kb" and measure in ("n", "mean_signed", "p50", "p90", "p95", "p99", "max"): nd = 1 if measure == "mean_signed" else 0
        ck.append(dict(block=block, label=label, start=a_.isoformat() if a_ else "", end=b_.isoformat() if b_ else "", measure=measure,
                       value=fmt(int(value) if unit in ("weeks", "months") and value is not None else value, nd), unit=unit, note=note))
    for lab, _ in CHECKPOINTS:
        w = chk_week[lab]; r = by[w]
        for m, k, u in (("commercial_crude_stocks", "comm_stocks_kb", "kb"), ("spr_crude_stocks", "spr_kb", "kb"), ("cushing_crude_stocks", "cushing_kb", "kb"),
                        ("gasoline_stocks", "gasoline_kb", "kb"), ("distillate_stocks", "distillate_kb", "kb"), ("crude_total_incl_spr", "crude_total_kb", "kb"),
                        ("wti_week_avg", "wti_avg_usd_bbl", "$/bbl"), ("brent_week_avg", "brent_avg_usd_bbl", "$/bbl"),
                        ("commercial_stocks_value_at_wti", "comm_value_usd_bn", "$bn"), ("spr_value_at_wti", "spr_value_usd_bn", "$bn"),
                        ("commercial_stocks_5y_min", "comm_5y_min_kb", "kb"), ("commercial_stocks_5y_max", "comm_5y_max_kb", "kb")):
            C("level", lab, None, w, m, r[k], u)
        for m, k in (("production", "production_4wk_kbd"), ("net_imports", "net_imports_4wk_kbd"), ("imports", "imports_4wk_kbd"), ("exports", "exports_4wk_kbd"),
                     ("refinery_input", "refinery_input_4wk_kbd"), ("product_supplied", "product_supplied_4wk_kbd")):
            C("flow_4wk_avg", lab, w - 3 * ONE_WEEK - datetime.timedelta(days=6), w, m, r[k], "kb/d", "mean of the four weekly values ending at week_end")
    e23, e24, e25, j26 = (chk_week[k] for k in ("end-2023", "end-2024", "end-2025", "30 Jun 2026"))
    WINDOWS = [("2024", e23, e24), ("2025", e24, e25), ("2026 H1 (to 30 Jun)", e25, j26), ("2026 Q3 to date", j26, latest), ("2026 to date", e25, latest)]
    def wsum(k, a_, b_):
        v = [by[w][k] for w in weeks if a_ < w <= b_]
        return None if any(x is None for x in v) else sum(v)
    WIN = {}
    for lab, a_, b_ in WINDOWS:
        n_w = sum(1 for w in weeks if a_ < w <= b_)
        W = dict(n=n_w, d_comm=by[b_]["comm_stocks_kb"] - by[a_]["comm_stocks_kb"], d_spr=by[b_]["spr_kb"] - by[a_]["spr_kb"],
                 d_cush=by[b_]["cushing_kb"] - by[a_]["cushing_kb"], d_gas=by[b_]["gasoline_kb"] - by[a_]["gasoline_kb"], d_dis=by[b_]["distillate_kb"] - by[a_]["distillate_kb"],
                 usd_comm=wsum("comm_change_usd_mn", a_, b_), usd_spr=wsum("spr_change_usd_mn", a_, b_), usd_comm_b=wsum("comm_change_usd_mn_brent", a_, b_),
                 usd_spr_b=wsum("spr_change_usd_mn_brent", a_, b_))
        for nm, k in (("prod", "production_kbd"), ("trans", "transfers_kbd"), ("imp", "imports_kbd"), ("exp", "exports_kbd"), ("ref", "refinery_input_kbd"),
                      ("unacc", "unaccounted_kbd"), ("ps", "product_supplied_kbd")):
            v = [by[w][k] for w in weeks if a_ < w <= b_]; W[nm] = None if any(x is None for x in v) else 7.0 * sum(v)
        W["usd_unacc"] = wsum("unaccounted_usd_mn", a_, b_); W["usd_prod"] = wsum("production_usd_mn", a_, b_); W["usd_imp"] = wsum("imports_usd_mn", a_, b_)
        W["usd_exp"] = wsum("exports_usd_mn", a_, b_); W["usd_ref"] = wsum("refinery_input_usd_mn", a_, b_)
        W["resid"] = wsum("bal_residual_kb", a_, b_)
        WIN[lab] = W
        for m, k, u in (("d_commercial_stocks", "d_comm", "kb"), ("d_spr", "d_spr", "kb"), ("d_cushing", "d_cush", "kb"), ("d_gasoline", "d_gas", "kb"), ("d_distillate", "d_dis", "kb"),
                        ("production", "prod", "kb"), ("transfers_to_crude_supply", "trans", "kb"), ("imports", "imp", "kb"), ("exports", "exp", "kb"),
                        ("refinery_input", "ref", "kb"), ("unaccounted_for", "unacc", "kb"), ("product_supplied_total_petroleum", "ps", "kb"),
                        ("balance_residual_sum", "resid", "kb"), ("commercial_change_at_wti", "usd_comm", "$mn"), ("spr_change_at_wti", "usd_spr", "$mn"),
                        ("commercial_change_at_brent", "usd_comm_b", "$mn"), ("spr_change_at_brent", "usd_spr_b", "$mn"), ("unaccounted_for_at_wti", "usd_unacc", "$mn"),
                        ("production_at_wti", "usd_prod", "$mn"), ("imports_at_wti", "usd_imp", "$mn"), ("exports_at_wti", "usd_exp", "$mn"), ("refinery_input_at_wti", "usd_ref", "$mn")):
            C("window_change", lab, a_, b_, m, W[k], u, "sum of 7 x weekly kb/d" if u == "kb" and k in ("prod", "trans", "imp", "exp", "ref", "unacc", "ps") else "")
    WW = [("2024", "2024-01", "2024-12"), ("2025", "2025-01", "2025-12"), ("2026 to date (Jan-Sep, historical)", "2026-01", mon(info["last_hist"])),
          ("2026 Mar-Sep (historical)", "2026-03", mon(info["last_hist"])), ("2026 full year (Oct-Dec forecast)", "2026-01", "2026-12")]
    WINW = {}
    for lab, a_, b_ in WW:
        x = world_window(world, a_, b_); WINW[lab] = x
        for m, k, u in (("implied_net_withdrawal", "wd_mb", "mb"), ("implied_net_withdrawal_rate", "wd_mbd", "mb/d"), ("implied_net_withdrawal_at_brent", "usd_bn", "$bn"),
                        ("of_which_us_scope_incl_spr", "us_mb", "mb"), ("of_which_other_oecd", "oo_mb", "mb"), ("of_which_other_draws_and_balance", "nn_mb", "mb"),
                        ("oecd_commercial_stocks_start", "oecd0", "mb"), ("oecd_commercial_stocks_end", "oecd1", "mb"),
                        ("months_historical", "n_h", "months"), ("months_forecast", "n_f", "months")):
            C("world_window", lab, Dt.fromisoformat(a_ + "-01"), Dt.fromisoformat(b_ + "-01"), m, x[k], u)
    for k_, v_ in (("R1_published_balance", st["R1"]), ("R2_levels_identity_asserted", st["R2"]), ("R2_levels_identity_incl_revision_weeks", st["R2_all"]),
                   ("R2_since_2023_11_10", st["R2_post_nov23"]), ("R2_2026", st["R2_2026"]), ("R2_before_2017_not_asserted", st["R2_pre"]), ("total_petroleum_identity", st["TP"])):
        for m in ("n", "mean_signed", "max", "p50", "p90", "p95", "p99", "le7", "le14", "le21"):
            C("identity", k_, None, None, m, v_[m], "weeks" if m == "n" else ("share of weeks" if m.startswith("le") else "kb"), "|residual|, kb (n in weeks)" if m != "mean_signed" else "signed mean, kb")
    write_csv(os.path.join(OUT, "checkpoints.csv"), ["block", "label", "start", "end", "measure", "value", "unit", "note"], ck,
              nd=lambda c: 3)

    # ---- spot check against the latest WPSR Table 1 (three weeks: latest, previous, a year earlier)
    spot = []; spot_bad = 0
    for idx, w in enumerate(t1_dates):
        r = by.get(w)
        for nm, lab, key in (("commercial crude stocks excl. SPR", "Commercial (Excluding SPR)", "comm_stocks_kb"), ("SPR crude stocks", "Strategic Petroleum Reserve (SPR)", "spr_kb"),
                             ("crude incl. SPR", "Crude Oil", "crude_total_kb"), ("total motor gasoline", "Total Motor Gasoline", "gasoline_kb"),
                             ("distillate fuel oil", "Distillate Fuel Oil", "distillate_kb"), ("total stocks incl. SPR", "Total Stocks (Including SPR)", "petroleum_total_kb")):
            t1 = round(t1_stocks[lab][idx] * 1000.0); mine = r[key] if r else None
            ok = mine is not None and abs(mine - t1) < 0.5; spot_bad += not ok; spot.append((w.isoformat(), nm, "kb", t1, mine, "match" if ok else "DIFFERENT"))
        for nm, ln, key in (("field production", 1, "production_kbd"), ("transfers to crude supply", 4, "transfers_kbd"), ("imports", 8, "imports_kbd"), ("exports", 12, "exports_kbd"),
                            ("unaccounted for (adjustment)", 16, "unaccounted_kbd"), ("refinery crude input", 17, "refinery_input_kbd"),
                            ("stock change, published", 13, "stock_change_published_kbd"), ("product supplied, total", 30, "product_supplied_kbd")):
            t1 = t1_lines[ln][idx]; mine = r[key] if r else None
            ok = mine is not None and abs(mine - t1) < 0.5; spot_bad += not ok; spot.append((w.isoformat(), nm, "kb/d", t1, mine, "match" if ok else "DIFFERENT"))

    # ================================================================ the report
    R = []; P = R.append
    P("CRUDE OIL PHYSICAL LEDGER  (E-021 follow-up: where the barrels came from and went)")
    P(f"Sources: EIA dnav weekly history files, WPSR Table 1 history (psw01.xls, release {psw_release}), EIA Short-Term Energy Outlook {info['release']} (modelling completed {info['completed']}),")
    P(f"EIA daily spot WTI (to {max(WTI)}) and Brent (to {max(BRENT)}). Latest US week: ending {latest}. STEO last historical month: {mon(info['last_hist'])}.")
    P("Units: kb = thousand barrels; kb/d = thousand barrels per day; mb = million barrels (1 mb = 1,000 kb); mb/d = million barrels per day; $mn / $bn = US$ millions / billions.")
    P("A week is Saturday to Friday ending on the week_end date (7 days: a flow of F kb/d is 7 x F kb). Dollars = barrels x the week's mean EIA spot price (WTI unless stated).")
    P("")
    W26k = WIN["2026 to date"]; r2k, r1k = st["R2"], st["R1"]; lwk = by[latest]; wm = WINW["2026 Mar-Sep (historical)"]; wy = WINW["2026 to date (Jan-Sep, historical)"]
    nt26 = (W26k["imp"] - W26k["exp"]) / 1000
    P("KEY RESULTS")
    P(f"  Weekly identity (with the transfers term the brief omitted): |residual| <= {r2k['max']:.0f} kb in {r2k['n']} of {st['R2_nwin']} weeks since {ASSERT_FROM} (median {r2k['p50']:.0f} kb, p99 {r2k['p99']:.0f} kb); the exception is a stock-revision week")
    P(f"    ({', '.join(f'{w_} {res:+.0f} kb' for w_, lv, res in st['rev'])}). The published Table 1 closes within {r1k['max']:.0f} kb in all {r1k['n']:,} weeks since {rows[0]['week_end']}.")
    P(f"  Latest week {latest}: commercial crude {lwk['comm_stocks_kb'] / 1000:,.1f} mb, SPR {lwk['spr_kb'] / 1000:,.1f} mb, Cushing {lwk['cushing_kb'] / 1000:,.1f} mb. Since the week ending {e25}: SPR {W26k['d_spr'] / 1000:+,.1f} mb")
    P(f"    ({W26k['usd_spr'] / 1000:+,.2f} $bn at weekly WTI), commercial {W26k['d_comm'] / 1000:+,.1f} mb ({W26k['usd_comm'] / 1000:+,.2f} $bn); WTI week mean ${by[e25]['wti_avg_usd_bbl']:.2f} -> ${lwk['wti_avg_usd_bbl']:.2f}.")
    P(f"  Where 2026's barrels went (mb, {W26k['n']} weeks): refinery input {W26k['ref'] / 1000:,.0f} was supplied by production {W26k['prod'] / 1000:,.0f} + transfers {W26k['trans'] / 1000:,.0f} + net imports {nt26:,.0f}; the stock draw")
    P(f"    {(W26k['d_comm'] + W26k['d_spr']) / 1000:+,.1f} = those supplies - refinery input + unaccounted-for ({W26k['unacc'] / 1000:+,.1f}); the plug is small beside the SPR draw.")
    P(f"  World (STEO {info['release']}, last historical month {mon(info['last_hist'])}): implied net stock withdrawal {wm['wd_mb']:+,.0f} mb ({wm['wd_mbd']:+.2f} mb/d) over Mar-Sep 2026, {wy['wd_mb']:+,.0f} mb Jan-Sep; 2025 {WINW['2025']['wd_mb']:+,.0f} mb, 2024 {WINW['2024']['wd_mb']:+,.0f} mb;")
    P(f"    {wm['usd_bn']:+,.1f} $bn for Mar-Sep at monthly Brent. All 2026 months to September are flagged historical (EIA estimates); October-December are forecast.")
    P(f"  Brent {bc['spike_date']}: ${bc['prev']:.2f} -> ${bc['spike']:.2f}; it persists on the next {len(bc['after'])} prints (${', $'.join(f'{b:.2f}' for _, b, _, _ in bc['after'])}); spread rank {bc['rank_spread']} of {bc['n_spread']:,}. Reported, not corrected; unverified.")
    P(f"  Term structure: EIA's NYMEX series ends {max(FUT1)}; no other free key-free official curve found; leg stopped.")
    P(f"  Spot check against EIA's current Table 1: {len(spot) - spot_bad} of {len(spot)} figures match exactly (3 weeks).")
    P("")
    P("1. THE WEEKLY BALANCE (US crude oil, commercial + SPR)")
    P("   change in stocks = 7 x ( production + TRANSFERS + imports - exports + unaccounted-for - refinery input ),  each term in kb/d, stocks in kb.")
    P("   The brief's identity has no transfers term. EIA's Table 1 line (4) 'Transfers to crude oil supply' begins with the week ending 2023-11-10 (no earlier value in EIA's file); it is in the identity")
    P("   from then and is zero before (the balancing item carried those barrels). 'Unaccounted for' is EIA's own balancing item, so the published balance sheet")
    P("   closes by construction; the tests below are rounding closure, the tie to the stock LEVELS, and the size of the plug.")
    P("")
    r1, r2, r2a = st["R1"], st["R2"], st["R2_all"]
    P(f"   ASSERTED (the script stops otherwise):")
    P(f"   (a) published balance sheet, every week {rows[0]['week_end']} to {latest} ({r1['n']} weeks; {st['n_missing_ua']} weeks in 2005-2007 where EIA publishes no balancing item are excluded):")
    P(f"       |7 x published stock change - 7 x terms| <= {TOL_PUBLISHED} kb.  Observed max {r1['max']:.0f} kb.")
    P(f"   (b) THE BRIEF'S IDENTITY on the stock levels, weeks from {ASSERT_FROM} to {latest} ({st['R2_nwin']} weeks, of which {st['R2_nrev']} revision week(s) excluded: {r2['n']} asserted):")
    P(f"       |change in (WCESTUS1 + WCSSTUS1) - 7 x terms| <= {TOL_LEVELS} kb (the rounding bound: six published terms x 0.5 kb/d x 7 days = 21 kb, +1 kb for whole-barrel levels).")
    wmean = statistics.mean(abs(r["d_crude_total_kb"]) for r in rows if Dt.fromisoformat(r["week_end"]) >= ASSERT_FROM)
    P(f"       Observed max {r2['max']:.0f} kb, against a mean absolute weekly change in total crude stocks of {wmean:,.0f} kb (denominator: the {r2['n']} asserted weeks).")
    P("")
    P("   Residual = (change in stocks) - 7 x (supply minus disposition), kb, absolute value, weeks from 2017-01-06, revision weeks excluded:")
    rows_ = []
    for nm, s in (("all asserted weeks", r2), ("since 2023-11-10 (transfers in)", st["R2_post_nov23"]), ("2026 to date", st["R2_2026"]), ("published balance sheet, 2001-11-09 on", r1),
                  ("total-petroleum identity", st["TP"])):
        rows_.append((nm, s["n"], f"{s['mean_signed']:+.1f}", n0(s["p50"]), n0(s["p90"]), n0(s["p95"]), n0(s["p99"]), n0(s["max"]), f"{100 * s['le7']:.1f}%", f"{100 * s['le14']:.1f}%", f"{100 * s['le21']:.1f}%"))
    R += ["   " + x for x in table(["residual set", "weeks", "mean(signed)", "median", "p90", "p95", "p99", "max", "<=7 kb", "<=14 kb", "<=21 kb"], rows_)]
    P("   (signed mean: kb; the other columns: kb; the last three: share of weeks with |residual| at or under 7, 14 and 21 kb. 7 kb = 1 kb/d.)")
    P("")
    P(f"   Revision weeks (level change differs from 7 x the published stock change by more than {REV_FLAG:.0f} kb; the stock history was revised after the line was published), from {ASSERT_FROM}:")
    for w_, lv, res in st["rev"]: P(f"       {w_}: levels minus published {lv:+.0f} kb; the brief's residual that week {res:+.0f} kb  (not asserted; listed)")
    for w_, lv, res in st["TP_rev"]: P(f"       {w_}: total petroleum, levels minus published {lv:+.0f} kb; residual {res:+.0f} kb")
    P(f"   Before {ASSERT_FROM} the level change and the published stock change disagree by up to {st['R2_pre']['max']:,.0f} kb; {st['R2_pre_n_gt22']} of {st['R2_pre']['n']} weeks exceed {TOL_LEVELS:.0f} kb; the last revision week before then is {st['last_pre_flag']}. In 2010-2015, {st['pre_big'][1]} of the {st['pre_big'][0]} weeks that")
    P("   disagree by more than 14 kb are the first week of a month. Reading (INFERENCE, not stated by EIA in these files): the level history was re-benchmarked monthly and the flow lines were not restated.")
    P(f"   Check on the check: leaving the transfers term out, the identity misses by a median {st['no_transfers'][1]:,.0f} kb (max {st['no_transfers'][2]:,.0f} kb) in each of the {st['no_transfers'][0]} weeks since 2023-11-10.")
    P("")
    P("   The balancing item ('unaccounted for', kb/d, annual mean of the weekly values) and the transfers line:")
    rows_ = [(y, n0(m), (f"{n0(t)} ({nt} wks)" if t is not None else "line did not exist"), n) for y, (m, t, n, nt) in sorted(st["una_year"].items())]
    R += ["   " + x for x in table(["year", "unaccounted for, mean kb/d", "transfers, mean kb/d", "weeks"], rows_)]
    P("   The balancing item fell from +600 to +860 kb/d (2022 to Oct 2023) to about zero from 2024, in the same weeks that the ~500-700 kb/d transfers line began. READING (inferred from timing and size,")
    P("   not stated by EIA in these files): the 2023 change was a redefinition that moved those barrels out of the plug, not a change in the oil. From 2024 the plug averages -8 to +96 kb/d (annual means).")
    P("")
    P("2. CHECKPOINTS (US, weekly; stocks are end-of-week levels, flows are 4-week averages of the weekly kb/d)")
    labs = [l for l, _ in CHECKPOINTS]
    hdr = ["", *[f"{l} ({chk_week[l]})" for l in labs]]
    def rowf(name, key, f=n0): return [name, *[f(by[chk_week[l]][key]) for l in labs]]
    R += ["   " + x for x in table(hdr, [
        rowf("commercial crude stocks, kb", "comm_stocks_kb"), rowf("SPR crude stocks, kb", "spr_kb"), rowf("Cushing crude stocks, kb", "cushing_kb"),
        rowf("gasoline stocks, kb", "gasoline_kb"), rowf("distillate stocks, kb", "distillate_kb"),
        rowf("production, 4wk avg kb/d", "production_4wk_kbd", n1), rowf("imports, 4wk avg kb/d", "imports_4wk_kbd", n1), rowf("exports, 4wk avg kb/d", "exports_4wk_kbd", n1),
        rowf("net imports (imports-exports), 4wk kb/d", "net_imports_4wk_kbd", n1), rowf("refinery crude input, 4wk avg kb/d", "refinery_input_4wk_kbd", n1),
        rowf("product supplied, total, 4wk kb/d", "product_supplied_4wk_kbd", n1), rowf("WTI, week mean $/bbl", "wti_avg_usd_bbl", n2),
        rowf("commercial stocks at WTI, $bn", "comm_value_usd_bn", n1), rowf("SPR at WTI, $bn", "spr_value_usd_bn", n1)])]
    P("   (net imports = imports - exports, including SPR flows, as in WCRNTUS2; the SPR value is barrels at spot, not cost basis.)")
    P("")
    P("3. CHANGES IN STOCKS, BARRELS AND DOLLARS (sum of weekly changes; dollars = each week's change x that week's mean WTI; negative = draw)")
    hdr = ["window (weeks)", "d commercial kb", "d SPR kb", "d Cushing kb", "d gasoline kb", "d distillate kb", "commercial $mn", "SPR $mn", "SPR $mn at Brent"]
    rows_ = [(f"{lab} ({WIN[lab]['n']})", sg0(W["d_comm"]), sg0(W["d_spr"]), sg0(W["d_cush"]), sg0(W["d_gas"]), sg0(W["d_dis"]), sg0(W["usd_comm"]), sg0(W["usd_spr"]), sg0(W["usd_spr_b"])) for lab, W in WIN.items()]
    R += ["   " + x for x in table(hdr, rows_)]
    W26 = WIN["2026 to date"]
    P(f"   2026 to date (after the week ending {e25}, to {latest}): SPR {W26['d_spr'] / 1000:+,.1f} mb = {W26['usd_spr'] / 1000:+,.2f} $bn at weekly WTI ({W26['usd_spr_b'] / 1000:+,.2f} $bn at weekly Brent);")
    P(f"   commercial crude {W26['d_comm'] / 1000:+,.1f} mb = {W26['usd_comm'] / 1000:+,.2f} $bn; SPR + commercial {(W26['d_spr'] + W26['d_comm']) / 1000:+,.1f} mb = {(W26['usd_spr'] + W26['usd_comm']) / 1000:+,.2f} $bn.")
    P("   The dollar figures are barrels at the week's spot: not proceeds to the Treasury, not the DOE's cost basis (the SPR level change nets sales, exchanges and returns).")
    P("")
    P("   Where the barrels came from and went (sum of 7 x weekly kb/d, mb; dollars at weekly WTI, $bn):")
    hdr = ["window", "production", "transfers", "imports", "exports", "refinery input", "unaccounted-for", "d stocks (comm+SPR)", "residual", "product supplied"]
    rows_ = [(lab, n1(W["prod"] / 1000), n1(W["trans"] / 1000), n1(W["imp"] / 1000), n1(W["exp"] / 1000), n1(W["ref"] / 1000), sg1(W["unacc"] / 1000),
              sg1((W["d_comm"] + W["d_spr"]) / 1000), sg1(W["resid"] / 1000), n1(W["ps"] / 1000)) for lab, W in WIN.items()]
    R += ["   " + x for x in table(hdr, rows_)]
    P("   (all in mb; 'residual' = change in stocks minus supply-minus-disposition, summed over the weeks; 'unaccounted-for' is EIA's plug and enters the identity; product supplied is")
    P("    total petroleum products, a different quantity from crude, shown for scale.)")
    rows_ = [(lab, n1(W["usd_prod"] / 1000), n1(W["usd_imp"] / 1000), n1(W["usd_exp"] / 1000), n1(W["usd_ref"] / 1000), sg1(W["usd_unacc"] / 1000)) for lab, W in WIN.items()]
    R += ["   " + x for x in table(["window ($bn at WTI)", "production", "imports", "exports", "refinery input", "unaccounted-for"], rows_)]
    P("")
    P("4. COMMERCIAL STOCKS AGAINST THE SAME WEEK'S FIVE-YEAR RANGE (the five preceding calendar years; 2021-2025 for 2026; kb)")
    hdr = ["week ending", "level", "5y min", "5y max", "5y mean", "position", "weekly chg", "chg 5y min", "chg 5y max", "chg 5y mean", "position"]
    for tag, ttl in (("5y", "date-aligned (primary): the weekly observation within 3 days of the same calendar date"), ("5yiso", "ISO-week-aligned (the first-pass map's alignment)")):
        P(f"   {ttl}")
        rows_ = []
        for lab in labs:
            r = by[chk_week[lab]]
            rows_.append((f"{r['week_end']} ({lab})", n0(r["comm_stocks_kb"]), n0(r.get(f"comm_{tag}_min_kb")), n0(r.get(f"comm_{tag}_max_kb")), n0(r.get(f"comm_{tag}_mean_kb")),
                          r.get("comm_vs_" + tag) or "n/a", sg0(r["d_comm_kb"]), sg0(r.get(f"d_comm_{tag}_min_kb")), sg0(r.get(f"d_comm_{tag}_max_kb")),
                          sg0(r.get(f"d_comm_{tag}_mean_kb")), r.get("d_comm_vs_" + tag) or "n/a"))
        R += ["   " + x for x in table(hdr, rows_)]
    wk = [w for w in weeks if w > e25]
    for tag in ("5y", "5yiso"):
        na = sum(1 for w in wk if by[w].get("comm_vs_" + tag) == "above"); nb = sum(1 for w in wk if by[w].get("comm_vs_" + tag) == "below")
        P(f"   Weeks since {e25} ({tag}): {len(wk)}; level above the five-year range in {na}, below in {nb}, within in {len(wk) - na - nb}.")
    lw = by[latest]
    P(f"   Latest week {latest}: the commercial level is {lw['comm_vs_5y']} the date-aligned range and {lw['comm_vs_5yiso']} the ISO-aligned range; the week's change ({lw['d_comm_kb']:+,.0f} kb) is "
      f"{lw['d_comm_vs_5y']} / {lw['d_comm_vs_5yiso']} the corresponding ranges of weekly changes.")
    P("")
    P("5. WORLD, MONTHLY (EIA Short-Term Energy Outlook; implied net stock WITHDRAWAL = consumption - production; positive = stocks drawn; mb/d and mb)")
    P(f"   STEO {info['release']}: last historical month {mon(info['last_hist'])}. 'Historical' months are EIA estimates of world data, not observations: consumption outside the OECD is apparent consumption and")
    P("   'other draws and balance' is a balancing item. The workbook was completed on 1 Oct 2026 and flags September 2026 'historical': for the world that month can only be an estimate.")
    P("   STEO's own arithmetic is asserted for all 72 months (production = crude + other liquids; withdrawal = consumption - production = US + other OECD + other).")
    hdr = ["window", "months hist/fcst", "production mb/d", "consumption mb/d", "net withdrawal mb", "mb/d", "$bn at Brent", "US scope mb", "other OECD mb", "other+balance mb", "OECD comm. stocks start->end mb"]
    rows_ = []
    for lab, x in WINW.items():
        rows_.append((lab, f"{x['n_h']}/{x['n_f']}", n2(x["prod_mbd"]), n2(x["cons_mbd"]), sg0(x["wd_mb"]), f"{x['wd_mbd']:+.2f}",
                      "n/a (forecast months)" if x["usd_bn"] is None else sg1(x["usd_bn"]), sg0(x["us_mb"]), sg0(x["oo_mb"]), sg0(x["nn_mb"]),
                      f"{x['oecd0']:,.0f} -> {x['oecd1']:,.0f} ({x['oecd1'] - x['oecd0']:+,.0f})"))
    R += ["   " + x for x in table(hdr, rows_)]
    P("   (net withdrawal in mb = sum of monthly mb/d x days; $bn = each month's mb x the month's mean EIA Brent spot, complete historical months only; 'US scope' is STEO's United States")
    P("    line, which tracks WPSR total petroleum stocks INCLUDING the SPR, not commercial stocks; OECD commercial stocks exclude government stocks, so their change is smaller than the withdrawal.)")
    P("   Monthly detail, 2026 (mb/d unless stated; basis = STEO flag):")
    hdr = ["month", "basis", "production", "consumption", "net withdrawal", "OECD comm. stocks mb", "d OECD mb", "Brent $/bbl"]
    rows_ = [(r["month"], r["steo_basis"], n2(r["world_production_mbd"]), n2(r["world_consumption_mbd"]), f"{r['implied_net_withdrawal_mbd']:+.2f}", n0(r["oecd_commercial_stocks_mb"]),
              sg0(r["oecd_commercial_stock_change_mb"]), n2(r["brent_avg_usd_bbl"]) if r["brent_avg_usd_bbl"] is not None else f"({n2(r['steo_brent_usd_bbl'])} STEO)")
             for r in world if r["month"].startswith("2026") or r["month"] == "2025-12"]
    R += ["   " + x for x in table(hdr, rows_)]
    P("   Scope of STEO's 'United States' withdrawal (mb; negative = stocks fell). Sum of the WPSR weekly changes in total crude + petroleum products for weeks ending in each month (weeks straddle month ends):")
    sc_rows, tot = [], [0.0, 0.0, 0.0]
    for r in world:
        if "2026-03" <= r["month"] <= mon(info["last_hist"]):
            ws = [x for x in rows if x["week_end"][:7] == r["month"]]
            incl = sum(x["d_petroleum_total_kb"] for x in ws) / 1000.0; excl = incl - sum(x["d_spr_kb"] for x in ws) / 1000.0
            steo_us = -r["us_withdrawal_mbd"] * r["days"]; tot = [tot[0] + steo_us, tot[1] + incl, tot[2] + excl]
            sc_rows.append((r["month"], sg1(steo_us), sg1(incl), sg1(excl), len(ws)))
    sc_rows.append(("Mar-Sep", sg1(tot[0]), sg1(tot[1]), sg1(tot[2]), ""))
    R += ["   " + x for x in table(["month", "STEO US, implied stock change", "WPSR incl. SPR", "WPSR excl. SPR", "weeks"], sc_rows)]
    P("   STEO's US line follows the stocks INCLUDING the SPR, so the world withdrawal includes government stock releases; the OECD 'commercial' stock series excludes them.")
    P(f"   Check: STEO's own Brent and WTI monthly averages equal the monthly means of the EIA daily files in all {wchk['chk_n']} historical months (largest difference ${wchk['chk_brent'][1]:.3f} Brent, ${wchk['chk_wti'][2]:.3f} WTI).")
    P("")
    P("6. THE BRENT PRINT OF 2026-10-02 (reported, not corrected)")
    P(f"   EIA Europe Brent spot: {bc['prev_date']} ${bc['prev']:.2f} -> {bc['spike_date']} ${bc['spike']:.2f} ({bc['jump_pct']:+.1f}%), while Cushing WTI went ${bc['wti_prev']:.2f} -> ${bc['wti_spike']:.2f}.")
    P("   It PERSISTS in the data: " + "; ".join(f"{d} Brent ${b:.2f} (WTI ${w:.2f}, spread ${s:.2f})" for d, b, w, s in bc["after"]) + f". The file ends {max(BRENT)}; there is no later print.")
    P(f"   Brent minus WTI: ${bc['spread_prev']:.2f} on {bc['prev_date']}, ${bc['spread_spike']:.2f} on {bc['spike_date']}: rank {bc['rank_spread']} of {bc['n_spread']:,} days with both prices since {min(set(WTI) & set(BRENT))}.")
    P("   Largest spreads on record: " + "; ".join(f"${v:.2f} {d}" for v, d in bc["top_spreads"]) + ". (2020-04-20 is the day WTI settled at -$36.98.)")
    pm = bc["prior_max_ex2020"]
    P(f"   Before 2026 the largest spread outside Mar-Jun 2020 was ${pm[0]:.2f} ({pm[1]}). Monthly Brent minus WTI in 2026 (mean, min, max, days): " +
      "; ".join(f"{calendar.month_abbr[m]} ${v[0]:.1f} ({v[1]:.1f}-{v[2]:.1f}, {v[3]}d)" for m, v in bc["spread_2026"].items()) + ".")
    P(f"   One-day change: +{bc['jump_pct']:.1f}% is rank {bc['rank_up']} of {bc['n_chg']:,} daily Brent moves by log change; the largest are " +
      "; ".join(f"{d} {100 * (math.exp(c) - 1):+.1f}%" for c, d in bc["top_up"]) + ".")
    P("   So it is not a one-tick error in the file: the later prints stay at the new level and the spread had been widening for weeks (monthly line above). It is a record outside 2020 and no second free daily")
    P("   Brent series was fetched to corroborate it (STEO was completed on 1 Oct, before the print). Status: UNVERIFIED as a market price, PERSISTENT as a data point.")
    P(f"   Effect on this ledger: the week ending {bc['spike_date']} Brent mean is ${bc['wk_avg']:.2f} ({bc['wk_n']} days) with the print, ${bc['wk_avg_alt']:.2f} if that one day were set to the day before. Brent-priced columns")
    P("   (comm_change_usd_mn_brent, spr_change_usd_mn_brent, since_end2025_*_brent) move by that; every WTI-priced column and the world monthly dollars (complete months to September) do not.")
    a_, b_, c_, d_ = bc["wti_925"]
    P(f"   Also in the WTI file: 2026-09-25 ${b_:.2f} between ${a_:.2f} (09-24) and ${c_:.2f} (09-28), then ${d_:.2f}: a one-day dip that reverses. It is in that week's WTI mean and was not adjusted.")
    P("   A World Bank Pink Sheet cross-check of monthly Brent is run when its local file is present (data/vintages/crude/derived/brent_worldbank_check.csv); its figures are not in this public file.")
    P("")
    P("7. TERM STRUCTURE (looked for briefly in EIA's own pages, then stopped)")
    fl = max(FUT1)
    P(f"   EIA's NYMEX WTI futures series (RCLC1, Cushing contract 1) ends {fl} (last print ${FUT1[fl]:.2f}); EIA's page is titled '{fut_title}'.")
    P(f"   WTI spot prints since then: {sum(1 for d in WTI if d > fl)}. CME refused scripted access in the first pass (HTTP 403) and was not retried here; CFTC data carry no prices.")
    P("   No other free, key-free, official source of a WTI curve or of front- and second-month settlements was found in EIA's futures page, its STEO page or its 'What drives crude oil prices' page")
    P("   (the last two were opened and not saved; the last says updates to its figures were discontinued). That leg is stopped; no workaround was tried.")
    P("")
    P("8. SPOT CHECK AGAINST EIA'S CURRENT WPSR TABLE 1 (https://ir.eia.gov/wpsr/table1.csv, saved raw; three columns: latest week, previous week, a year earlier)")
    hdr = ["week ending", "line", "unit", "Table 1", "this ledger", "result"]
    R += ["   " + x for x in table(hdr, [(w_, nm, u, n0(a_), n0(b_), res) for w_, nm, u, a_, b_, res in spot], left=(0, 1, 2, 5))]
    P(f"   {len(spot) - spot_bad} of {len(spot)} comparisons match exactly (stocks converted from million barrels to kb); {spot_bad} differ.")
    P("")
    P("9. TWO ROUTES, ONE ANSWER: every series available both as a dnav history file and in psw01.xls is identical date by date:")
    P("   " + ", ".join(f"{s} ({c} weeks)" for s, _, _, c, _ in st["routes"]) + ".")
    P(f"   Line 18 (other production) = NGPL + renewables/oxygenates + processing gain: max difference {st['line18'][1]:.0f} kb/d over {st['line18'][0]} weeks. Transfers = Alaskan NGL/condensate + natural gasoline/unfinished oils: max difference {st['transfers'][1]:.0f} kb/d over {st['transfers'][0]} weeks.")
    P("")
    P("10. SERIES USED (id; unit read from the file's own header and asserted; first and last observation; observations)")
    srows = []
    for sid, unit, role in WEEKLY:
        ds_ = sorted(H[sid]); srows.append((sid, "dnav hist_xls", unit_of(HT[sid]), ds_[0], ds_[-1], len(ds_), role))
    for k in PSW_KEYS:
        ds_ = sorted(PS[k]); srows.append((k, "psw01.xls (Table 1 history)", unit_of(PT[k]), ds_[0], ds_[-1], len(ds_), re.sub(r"\s+", " ", PT[k]).strip()))
    for sid, ser, t_ in (("RWTC", WTI, wti_t), ("RBRTE", BRENT, br_t), ("RCLC1", FUT1, fut_t)):
        ds_ = sorted(ser); srows.append((sid, "dnav hist_xls (daily)", unit_of(t_), ds_[0], ds_[-1], len(ds_), re.sub(r"\s+", " ", t_).strip()))
    R += ["   " + x for x in table(["series id", "file", "unit (header)", "first", "last", "obs", "description"], [(a_, b_, c_, str(d_), str(e_), f_, g_) for a_, b_, c_, d_, e_, f_, g_ in srows], left=(0, 1, 2, 3, 4, 6))]
    P("   (W_EPC0NG_TVP_NUS_MBBLD carries the same EIA title as W_EPC0_TVP_NUS_MBBLD; it is the products-side entry, equal to the negative of the crude-side line: asserted.)")
    P("   STEO_m.xlsx tab 3atab (million barrels per day unless stated): papr_world production, copr_world crude production, papr_opec OPEC production, patc_world consumption, t3_stchange_world / _us / _ooecd / _noecd")
    P("   net withdrawals; pasc_oecd_t3 / pasc_us / pasc_ooecd_t3 end-of-month commercial inventories (million barrels). Tab 2tab: BREPUUS Brent and WTIPUUS WTI monthly spot averages ($/bbl). Tab Dates: last historical month, per-month flag.")
    P("")
    P("LICENSING. EIA data are US-government work: public domain in the United States (17 U.S.C. 105). All raw files, outputs and history files here may be public; credit 'Source: U.S. Energy Information Administration'.")
    P("")
    P("WHAT IS UNCERTAIN: (1) the unaccounted-for plug is EIA's residual and absorbs every error in the production, import, export and refinery lines; (2) weekly production is a model-based estimate; (3) the SPR dollars are barrels at spot, not proceeds;")
    P("(4) the world withdrawal is an EIA model output whose last 'historical' months are estimates and which contains a balancing item; (5) the 2026-10-02 Brent print is unverified; (6) the pre-2017 stock-level revisions are an inference from a pattern.")
    report = "\n".join(R) + "\n"
    with open(os.path.join(OUT, "report.txt"), "w", encoding="utf-8", newline="") as f: f.write(report)
    print(report)

    # ---- LOCAL: World Bank Pink Sheet cross-check of monthly Brent (not part of any public file)
    wb = load_wb()
    if wb is None:
        print("(no local World Bank Pink Sheet file: Brent cross-check skipped; public outputs are unaffected)")
    else:
        byb, byw = {}, {}
        for d, v in BRENT.items(): byb.setdefault(d.year * 100 + d.month, []).append(v)
        for d, v in WTI.items(): byw.setdefault(d.year * 100 + d.month, []).append(v)
        out = []
        for m in sorted(wb):
            if m >= 202201 and m in byb and wb[m][0] is not None:
                eb = sum(byb[m]) / len(byb[m]); ew = sum(byw[m]) / len(byw[m]) if m in byw else None
                out.append({"month": mon(m), "eia_brent_days": len(byb[m]), "eia_brent_mean": eb, "worldbank_brent": wb[m][0], "brent_pct_diff": 100 * (eb / wb[m][0] - 1),
                            "eia_wti_mean": ew, "worldbank_wti": wb[m][1], "wti_pct_diff": None if ew is None or wb[m][1] is None else 100 * (ew / wb[m][1] - 1)})
        write_csv(os.path.join(DER, "brent_worldbank_check.csv"), list(out[0].keys()), out, nd=lambda c: 3)
        print("LOCAL World Bank Pink Sheet check of monthly Brent (EIA daily mean vs World Bank), months since 2022-01 present in both:")
        for o in out[-9:]:
            print(f"  {o['month']}  EIA {o['eia_brent_mean']:.2f} ({o['eia_brent_days']} days)  World Bank {o['worldbank_brent']:.2f}  diff {o['brent_pct_diff']:+.2f}%")
        print(f"  mean absolute difference {sum(abs(o['brent_pct_diff']) for o in out) / len(out):.2f}%, max {max(abs(o['brent_pct_diff']) for o in out):.2f}% over {len(out)} months")
    print("done. Outputs: data/crude/{us_weekly_ledger,world_monthly_balance,checkpoints,brent_spread_check}.csv, report.txt; data/history/crude_*.csv")


if __name__ == "__main__":
    main()
