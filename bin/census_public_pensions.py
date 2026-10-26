#!/usr/bin/env python3
"""US state and local government pension systems, national contributions, benefit payments and assets, survey years 2013-2025 (Census Bureau, Annual Survey of
Public Pensions). Python 3 standard library plus openpyxl.

  python3 bin/census_public_pensions.py fetch  [--cache DIR] [--out DIR] [--sleep 3.05] [--max-files N]   # the only network command; at most 30 requests a run
  python3 bin/census_public_pensions.py build  [--cache DIR] [--out DIR]                                  # offline; byte-identical on re-run
  python3 bin/census_public_pensions.py selftest [--no-mutants] [--mutant NAME]                           # synthetic data only; no network

Specification: _research/2026-10-26-Census-Public-Pensions-Spec.md (committed at 104218d, with its supervisor rulings of rounds 1 to 4; where rounds differ the
latest round is applied). Outputs: data/census_public_pensions/ (seven CSVs, README.md, SOURCES.md).

WHAT TOUCHES THE NETWORK. Only `fetch`, and only www2.census.gov: one directory listing (read only to print a notice when a survey-year folder above 2025 appears)
and the 20 files of the spec's table. Any other host, and any path that is not in that table, is refused before a request is made. GET only, one request at a
time, at least 3 seconds apart, at most 30 requests a run, no redirect followed, no retry of any kind, the neutral identification string (no address, no name).
Any status other than 200, or a body that is empty or is a web page, writes STOP_census in data/vintages/census_public_pensions/ and ends the fetch. A timeout is
not a refusal: it is logged, the invocation ends with exit 3, nothing is latched and nothing is retried. This script never deletes, renames over or truncates a
STOP file; only a person deletes one.

WHAT IT RECORDS. The Census Bureau's own national totals of what US state- and locally-administered defined-benefit pension systems took in and paid out, and the
cash and investments they reported, for each survey year. It records the systems' own accounts only; it does not show what they bought or sold in markets.
"""
import argparse
import ast
import codecs
import collections
import contextlib
import csv
import datetime as dt
import decimal
import email.utils
import gzip
import hashlib
import http.client
import io
import os
import re
import shutil
import sys
import tempfile
import time
import warnings
from decimal import Decimal
from urllib.parse import urlsplit

warnings.filterwarnings("ignore", message="Cannot parse header or footer")
decimal.getcontext().prec = 60

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# ------------------------------------------------------------------------------------------------------------------------
# mutants: deliberately broken variants, switched on only by the selftest, in this process
# ------------------------------------------------------------------------------------------------------------------------
_ACTIVE = set()


def M(name):
    return name in _ACTIVE


@contextlib.contextmanager
def mutant(name):
    added = name not in _ACTIVE
    _ACTIVE.add(name)
    try:
        yield
    finally:
        if added:
            _ACTIVE.discard(name)


class Refused(Exception):
    """A request, path or setting this script will not honour."""


class MaxRequests(Refused):
    """The request cap for this run was reached."""


class Latched(Exception):
    """A technical refusal: STOP was written (or already exists); the fetch ends."""


class SocketFault(Exception):
    """Timeout, reset or TLS error: the invocation ends with exit 3, without latching."""


class ContentError(Exception):
    """A body could not be read (decoding, workbook)."""


class DataError(Exception):
    """A hard check on the data failed."""


# ------------------------------------------------------------------------------------------------------------------------
# constants
# ------------------------------------------------------------------------------------------------------------------------
OUT_DEFAULT = os.path.join(ROOT, "data", "census_public_pensions")
CACHE_DEFAULT = os.path.join(ROOT, "data", "vintages", "census_public_pensions")
STOP_DIR = CACHE_DEFAULT                      # fixed, whatever --cache says: a latch cannot be dodged by pointing the cache elsewhere
STOP_NAME = "STOP_census"

UA = "Mozilla/5.0 (compatible; ThirdDerivativeResearch/1.0; non-commercial research)"
HEADERS = {"User-Agent": UA, "Accept-Encoding": "identity", "Accept": "*/*"}
HOST = "www2.census.gov"
HOST_ALLOW = frozenset({HOST})
MIN_SLEEP = 3.0
DEFAULT_SLEEP = 3.05
MAX_REQUESTS = 30
TIMEOUT = 120
BODY_CAP = 100 * 1024 * 1024
ERR_CAP = 4096
BASE = "/programs-surveys/aspp/"
LISTING_PATH = BASE + "datasets/"
LISTING_MARK = b"Index of /programs-surveys/aspp/datasets"

SURVEY_YEARS = tuple(range(2013, 2026))
OLD_YEARS = (2013, 2014, 2015, 2016)
CENSUS_YEARS = (2017, 2022)                   # years ending in 2 and 7: every system is counted, every weight is 1
SUMMARY_SOURCED = (2013, 2014, 2015, 2016, 2020)
NEWEST_PIN_DATE = "2026-05-14"
UNIT_2020_PATH = BASE + "datasets/2020/aspp-historical-datasets/ASPP_Unit_File_2020.dsv"   # a byte copy of 2019 (R2-1): never requested
STALE_2024_ROOT_PATHS = (BASE + "datasets/2024/ASPP_Unit_File_2024.dsv", BASE + "datasets/2024/ASPP_Unit_File_2024.xlsx")   # dated 23 May 2025, before the revision (R3-6): never requested
UNIT_2024_PATH = BASE + "datasets/2024/aspp-historical-datasets/ASPP_Unit_File_2024.csv"


def _file(year, kind, path):
    return {"year": year, "kind": kind, "path": path, "name": path.rsplit("/", 1)[1]}


def build_files():
    out = []
    for y in (2013, 2014, 2015, 2016):
        out.append(_file(y, "summary", BASE + "tables/%d/aspp-historical-tables/%dasppsummary.xlsx" % (y, y)))
    for y in (2017, 2018, 2019, 2021, 2022, 2023):
        out.append(_file(y, "unit", BASE + "datasets/%d/aspp-historical-datasets/ASPP_Unit_File_%d.dsv" % (y, y)))
    out.append(_file(2024, "unit", UNIT_2024_PATH))
    if M("g19_root_2024"):
        out[-1] = _file(2024, "unit", STALE_2024_ROOT_PATHS[0])
    out.append(_file(2025, "unit", BASE + "datasets/2025/aspp-historical-datasets/ASPP_Unit_File_2025.csv"))
    for y in (2017, 2018, 2019, 2020, 2021, 2022, 2023):
        out.append(_file(y, "summary", BASE + "datasets/%d/aspp-historical-datasets/ASPP_%d_Summary.xlsx" % (y, y)))
    out.append(_file(2024, "summary", BASE + "datasets/2024/ASPP_2024_Summary.xlsx"))
    return out


FILES = build_files()
N_FILES = len(FILES)                          # 20

# the 53 item codes of the 2025 file and the layouts (spec R3b)
CODES = frozenset(["RJ01", "RJ10", "RJ11", "RJ12", "RJ20", "RJ30", "RJ31", "RJ32"] + ["RP%02d" % i for i in range(1, 22)] +
                  ["RX01", "RX10", "RX11", "RX12", "RX20", "RX30", "RX40", "RX50", "RX51", "RX52", "RX60", "RX61", "RX62", "RX70", "RX71", "RX72", "RX73",
                   "RX80", "RX90", "RX91", "RX92"] + ["RZ01", "RZ10", "RZ90"])

# the national columns and the unit-file codes that make them (spec 2.3)
UNIT_MAP = collections.OrderedDict([
    ("contrib_employee", ("RX51", "RX52")),
    ("contrib_government", ("RX60", "RX70")),
    ("contrib_total", ("RX51", "RX52", "RX60", "RX70")),
    ("benefits", ("RX40",)),
    ("withdrawals", ("RX80",)),
    ("earnings_ex_gains", ("RX01",)),
    ("gains_losses", ("RX30",)),
    ("expenses", ("RX90",)),
    ("assets_end", ("RZ01",)),
])
ITEMS = ("contrib_employee", "contrib_government", "contrib_total", "benefits", "withdrawals", "earnings_ex_gains", "gains_losses", "assets_end")
# summary labels (spec 2.3), matched by text and never by position
LABELS_CUR = collections.OrderedDict([
    ("contrib_total", "Total Contributions"), ("contrib_employee", "Employee contributions"), ("contrib_government", "Government contributions"),
    ("earnings", "Earnings on investments"), ("gains", "Gains/Losses on Investments"), ("benefits", "Total Benefits Paid"),
    ("expenses", "Administrative Costs"), ("withdrawals", "Withdrawals"), ("assets", "Total cash and investments"), ("systems", "Number of systems")])
LABELS_OLD = collections.OrderedDict([
    ("contrib_total", "Total Contributions"), ("contrib_employee", "Employee contributions"), ("contrib_government", "Government contributions"),
    ("benefits", "Benefits"), ("withdrawals", "Withdrawals"), ("earnings", "Earnings on investments"),
    ("assets", "Total cash and investment holdings"), ("systems", "Number of systems")])
CV_KEYS = ("contrib_total", "benefits", "assets")
CV_MAX = Decimal("0.05")

# the frozen revision band (R2-3), measured by the first build on vintage 1e8f842133376a6b (newest unit file 2026-05-14) and hard-coded here; every later build uses these
# constants (SOURCES.md, Builder decisions). The first build had them as None and used what it measured, rounded to three decimals.
FROZEN_BAND_PCT = Decimal("1.898")
FROZEN_SMP_BAND_PCT = Decimal("0.176")

# canonical column allowlists: the spec's seven lines, byte for byte (section 5)
SPEC_ALLOW_TEXT = (
    "national.csv:survey_year,schema,year_kind,contrib_employee_usd_mn,contrib_government_usd_mn,contrib_total_usd_mn,benefits_usd_mn,withdrawals_usd_mn,"
    "net_cash_usd_mn,earnings_ex_gains_usd_mn,earnings_old_schema_usd_mn,gains_losses_usd_mn,expenses_admin_invest_usd_mn,assets_end_usd_mn,systems_published,"
    "cv_contrib_total,cv_benefits,cv_assets,summary_file_date,source_kind,source_file_date,vintage\n"
    "ratios.csv:survey_year,schema,benefits_per_contrib_usd,benefits_withdrawals_per_contrib_usd,net_cash_pct_of_assets_start,net_cash_pct_of_assets_end,"
    "income_per_net_cash_gap_usd,income_per_gap_after_expenses_usd,ratio_cv_bound_pct,vintage\n"
    "summary.csv:question,segment,from_year,to_year,statistic,value,unit,pattern,primary,vintage\n"
    "revisions.csv:survey_year,item,source_prev,date_prev,value_prev_usd_mn,source_new,date_new,value_new_usd_mn,change_usd_mn,change_pct\n"
    "checks.csv:check_id,description,expected,observed,status\n"
    "vintages.csv:vintage,newest_unit_file,newest_unit_file_date,files,requests,manifest_sha256\n"
    "fetch_log.csv:utc_time,host,path,status,bytes,sha256,last_modified\n")
SPEC_ALLOW_SHA256 = "652a264d517b39f333a2a1270ca4abb4e140d40a1235c64ef98905c40a2bbdb9"
ALLOW_ORDER = tuple(l.split(":", 1)[0] for l in SPEC_ALLOW_TEXT.splitlines())
ALLOW = {l.split(":", 1)[0]: l.split(":", 1)[1].split(",") for l in SPEC_ALLOW_TEXT.splitlines()}
COMMITTED = tuple(ALLOW_ORDER) + ("README.md", "SOURCES.md")
RAW_EXT = (".dsv", ".xlsx", ".xls", ".txt", ".pdf", ".html", ".htm", ".zip", ".gz")

ALLOWED_VALUES = {
    "schema": ("old", "current"), "segment": ("old", "current"), "year_kind": ("census", "sample"), "source_kind": ("unit_file", "summary"),
    "source_prev": ("unit_file", "summary"), "source_new": ("unit_file", "summary"),
    "item": ITEMS, "question": ("Q2", "Q3", "Q4", "Q5", "Q6", "S1", "Q1"),
    "statistic": ("benefits_per_contrib_usd", "change_benefits_per_contrib_usd", "change_benefits_per_contrib_pct", "net_cash_pct_of_assets_start",
                  "net_cash_pct_of_assets_end", "change_net_cash_start_pp", "change_net_cash_end_pp", "n_years_net_cash_negative", "income_per_net_cash_gap_usd",
                  "income_per_gap_after_expenses_usd") + tuple("max_abs_revision_pct__" + i for i in ITEMS) +
                 ("ratio_revision_pct", "revision_band_pct", "revision_band_sample_years_pct", "ratio_cv_bound_pct", "s1_revision_2025_pct"),
    "unit": ("usd_per_usd", "pct", "pp", "count"), "pattern": ("above_band", "below_band", "within_band", "n_a"), "primary": ("yes", "no"),
}
SUMMARY_KEY = ("question", "segment", "from_year", "to_year", "statistic")

# wording guard patterns (G14), pasted from the spec
CAUSAL_RX = re.compile(r"\b(caus(?:e|ed|es|ing)|drove|driv(?:e|es|en|ing)|leaned?\s+against|because\s+of\s+(?:the\s+)?intervention|led\s+to|prompted|pushed|in\s+response\s+to|because\s+of)\b", re.I)
EXTRA_RX = re.compile(r"\b(?:in|out)flows?\b|\bflows?\s+(?:into|out\s+of|from|to)\b|\bnet\s+(?:sell|buy)(?:ers?|ing)\b|\b(?:sell|sells|sold|selling|sale|sales|buy|buys|bought|buying|purchas(?:e|es|ed|ing))\s+(?:of\s+)?(?:equit(?:y|ies)|stocks?|shares|bonds|assets)\b|\bforced\s+(?:sell|sale)\w*|\bliquidat\w*|\bsell[- ]?off\b|\bbid\b|\bdrain\w*|\bdeplet\w*|\bunderfunded\b|\binsolven\w*|\bcrisis\b|\brunning\s+out\b|\bretirement\s+flows?\b", re.I)
ROSE_FELL_RX = re.compile(r"\b(rose|fell)\b", re.I)          # R3-1: the pattern labels are descriptive; no direction verbs in the README
DENIAL_SENTENCE = ("It records the pension systems' own accounts only: a benefit payment is not a sale of equities, a contribution is not a purchase of them, "
                   "and nothing here shows what the systems bought or sold in markets.")

# the raw-row pattern of leak rule 6.1 (quote-aware, no start-of-line anchor)
RAW_ROW_RX = re.compile(r'(?<!\d)\d{4}(?:\|[^|\n]*\||,(?:"[^"]*"|[^,"\n]*),)\d{6}[|,][A-Z]{2}[|,]\d[|,]R[XZJP]\d\d[|,]')
RAW_ROW_RX_NAIVE = re.compile(r"^\d{4}(?:\|[^|,\n]*\||,[^,\n]*,)\d{6}[|,][A-Z]{2}[|,]\d[|,]R[XZJP]\d\d[|,]", re.M)    # the reviewer's first version, kept for the mutant

# reference constants typed in from other sources (C15-C17); they never enter any table except checks.csv (R10)
PRESS = {"contrib_2025_bn": Decimal("315.02"), "benefits_2025_bn": Decimal("418.25"), "assets_2025_trn": Decimal("6.49"), "benefits_2024_bn": Decimal("404.46")}
NASRA_FY2024_TRN = Decimal("5.13")
Z1_2025Q2_BN = Decimal("6541.4")


# ------------------------------------------------------------------------------------------------------------------------
# helpers
# ------------------------------------------------------------------------------------------------------------------------
def tilde(p):
    home = os.path.expanduser("~")
    return str(p).replace(home, "~") if home and home != "~" else str(p)


def sha256_of(b):
    return hashlib.sha256(b).hexdigest()


def D(x):
    """Decimal from an int, a decimal string or a float read from a workbook (via its shortest text, never the binary value)."""
    if x is None:
        return None
    if isinstance(x, Decimal):
        return x
    if isinstance(x, float):
        return Decimal(repr(x))
    return Decimal(str(x).strip())


def fdec(v, nd):
    """Fixed-decimal text of a Decimal, '' for None, no negative zero. Rounding is half-even, and happens only here, at write time."""
    if v is None:
        return ""
    q = Decimal(1).scaleb(-nd)
    s = format(Decimal(v).quantize(q, rounding=decimal.ROUND_HALF_EVEN), "f")
    if s.startswith("-") and Decimal(s) == 0:
        s = s[1:]
    return s


def assert_ua(ua=UA):
    if not M("g16_at_in_ua_ok"):
        if "@" in ua:
            raise Refused("the identification string contains an address")
        for var in ("SEC_UA", "CLAUDE_CODE_USER_EMAIL"):
            v = os.environ.get(var, "").strip()
            if v and (v in ua or ua in v):
                raise Refused("the identification string contains a configured contact value (not printed)")


assert_ua()


def write_text(path, text):
    """Atomic write via a temporary sibling. Refuses any path whose name starts with STOP: this script never writes over a latch."""
    if os.path.basename(path).upper().startswith("STOP"):
        raise Refused("write_text never touches a STOP file")
    os.makedirs(os.path.dirname(path), exist_ok=True)
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8", newline="") as fh:
        fh.write(text)
    os.replace(tmp, path)


def write_bytes(path, data):
    if os.path.basename(path).upper().startswith("STOP"):
        raise Refused("write_bytes never touches a STOP file")
    os.makedirs(os.path.dirname(path), exist_ok=True)
    tmp = path + ".tmp"
    with open(tmp, "wb") as fh:
        fh.write(data)
    os.replace(tmp, path)


def read_file(path, mode="rb"):
    with open(path, mode) as fh:
        return fh.read()


def read_csv_rows(path):
    """(header, rows) of a csv with comment lines (# ...) skipped; ([], []) if absent."""
    if not os.path.exists(path):
        return [], []
    with open(path, encoding="utf-8", newline="") as fh:
        lines = [ln for ln in fh.read().splitlines(True) if not ln.startswith("#")]
    rd = list(csv.reader(io.StringIO("".join(lines))))
    return (rd[0], rd[1:]) if rd else ([], [])


def http_date_iso(s):
    """An HTTP date (always GMT) as a UTC ISO text. Never local time (G9)."""
    if M("g9_lm_local"):
        ts = email.utils.mktime_tz(email.utils.parsedate_tz(s))
        return dt.datetime.fromtimestamp(ts).strftime("%Y-%m-%dT%H:%M:%SZ")          # the mutant: the machine's own time zone
    d = email.utils.parsedate_to_datetime(s)
    if d.tzinfo is None:
        d = d.replace(tzinfo=dt.timezone.utc)
    return d.astimezone(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def is_unit_path(p):
    return "Unit_File" in p.rsplit("/", 1)[-1]


def vintage_key(rows):
    """R5 / R3-4: 16 hex characters of the sha256 of the lines `url_path|last_modified_utc_iso|bytes\\n` for all 20 files, sorted by url_path. From the server's
    headers, never the wall clock. rows: (url_path, last_modified_iso, bytes)."""
    rows = sorted(rows)
    if M("g9_key_newest_only"):
        unit = [r for r in rows if is_unit_path(r[0])]
        rows = [max(unit, key=lambda r: (r[1], r[0]))] if unit else rows
    text = "".join("%s|%s|%s\n" % ((p, lm, b) if not M("g9_key_no_size") else (p, lm, "")) for p, lm, b in rows)
    if M("g9_key_wallclock"):
        text += repr(time.time())
    return sha256_of(text.encode("utf-8"))[:16]


def newest_unit(rows):
    """(basename, date) of the newest `Last-Modified` among unit files."""
    unit = [r for r in rows if is_unit_path(r[0])]
    if not unit:
        return "", ""
    p, lm, _ = max(unit, key=lambda r: (r[1], r[0]))
    return p.rsplit("/", 1)[1], lm[:10]


# ------------------------------------------------------------------------------------------------------------------------
# requests: the URL check (R1, R3), the content tests (R3a) and the fetcher
# ------------------------------------------------------------------------------------------------------------------------
def allowed_targets():
    t = {f["path"] for f in FILES}
    t.add(LISTING_PATH)
    if M("g16_2020_ok"):
        t.add(UNIT_2020_PATH)
    return t


def check_url(url):
    """Only plain https GETs to www2.census.gov and only the 21 request targets of the spec. Returns (host, target)."""
    p = urlsplit(url)
    host = (p.hostname or "").lower()
    target = p.path + ("?" + p.query if p.query else "")
    if M("g16_allow_other_host"):
        return host, target
    if p.scheme != "https" or p.username or p.password or p.port not in (None, 443) or p.fragment:
        raise Refused("only plain https URLs are allowed")
    if host not in HOST_ALLOW:
        raise Refused("host %s is not on the allowlist" % host)
    if target in (UNIT_2020_PATH,) + STALE_2024_ROOT_PATHS and not M("g16_2020_ok"):
        raise Refused("that file is never requested (a byte copy of 2019, or the stale 2024 root file)")
    if target not in allowed_targets() and not M("g16_url_any_ok"):
        raise Refused("that path is not one of the 21 requests of the spec")
    return host, target


def check_cache(cache, root=ROOT, out_dir=None):
    """The cache may be outside the repository, or inside it only where git ignores it (never in the committed output folder)."""
    real = os.path.realpath(os.path.expanduser(cache))
    rroot = os.path.realpath(root)
    if real == rroot or real.startswith(rroot + os.sep):
        outd = os.path.realpath(out_dir or os.path.join(root, "data", "census_public_pensions"))
        if real == outd or real.startswith(outd + os.sep):
            raise Refused("the cache may not live in the committed output folder")
        rel = os.path.relpath(real, rroot).replace(os.sep, "/") + "/"
        if not rel.startswith("data/vintages/"):
            raise Refused("a cache inside the repository must be under data/vintages/ (git-ignored)")
    return real


def real_opener(host, target, headers, timeout, cap):
    """GET https://host/target. Returns (status, body, response headers). No redirect is followed (a 3xx comes back as a status and latches), no cookies, the only
    verb is GET. A non-200 body is read up to ERR_CAP bytes (its head goes into the STOP file). Socket faults raise SocketFault."""
    conn = http.client.HTTPSConnection(host, timeout=timeout)
    try:
        conn.request("GET", target, headers=dict(headers))
        resp = conn.getresponse()
        status = resp.status
        hdrs = {k.lower(): v for k, v in resp.getheaders()}
        if status != 200:
            return status, resp.read(ERR_CAP), hdrs
        chunks, total = [], 0
        while True:
            b = resp.read(1 << 16)
            if not b:
                break
            total += len(b)
            if total > cap:
                return status, None, hdrs
            chunks.append(b)
        body = b"".join(chunks)
        if hdrs.get("content-encoding", "").lower() == "gzip":      # asked for identity; a server that compresses anyway is still read
            body = gzip.decompress(body)
        return status, body, hdrs
    except (OSError, EOFError, http.client.HTTPException) as e:
        raise SocketFault(type(e).__name__)
    finally:
        conn.close()


def esc_head(body, n=300):
    return (body or b"")[:n].decode("latin-1").encode("unicode_escape").decode("ascii")


STOP_FIELDS = ("status", "utc_time", "url_path", "retry_after", "body_head")


def stop_fields(status, utc_time, url_path, retry_after, body_head):
    lines = ["status: %s" % status, "utc_time: %s" % utc_time, "url_path: %s" % url_path]
    if not M("g15_stop_short"):
        lines.append("retry_after: %s" % retry_after)
    lines.append("body_head: %s" % body_head)
    return "\n".join(lines) + "\n"


def stop_text_ok(text):
    return [l.split(":", 1)[0] for l in text.splitlines()] == list(STOP_FIELDS)


def create_stop(path, text):
    """The one place a STOP file is created: exclusive create ("x"), so an existing latch is never overwritten. The AST scan (G15) allows nothing else."""
    os.makedirs(os.path.dirname(path), exist_ok=True)
    try:
        with open(path, "x", encoding="utf-8") as fh:
            fh.write(text)
    except FileExistsError:
        if M("g15_overwrite"):
            _clobber(path, text)                  # the mutant: what create_stop must never do; run-time test only


def _clobber(p, text):
    with open(p, "w", encoding="utf-8") as fh:
        fh.write(text)


def looks_like_page(body):
    return body.lstrip()[:15].lower().startswith((b"<!", b"<html"))


def file_body_problem(body):
    """R3a: a body that is empty or is HTML latches the STOP file."""
    if not body and not M("g15_empty_ok"):
        return "empty body"
    if body and looks_like_page(body) and not M("g15_html_ok"):
        return "the body is a web page"
    return ""


def listing_body_problem(body):
    """The one request whose reply is meant to be a web page: the Apache directory index. It must still be that index (its title), or it is a refusal page."""
    if not body and not M("g15_empty_ok"):
        return "empty body"
    if LISTING_MARK not in (body or b"") and not M("g15_html_ok"):
        return "the page is not the directory index"
    return ""


def listing_years(body):
    return sorted({int(m) for m in re.findall(rb'href="(\d{4})/"', body or b"")})


MANIFEST_COLS = ["url_path", "sha256", "bytes", "last_modified"]


def read_manifest(folder):
    p = os.path.join(folder, "manifest.csv")
    if not os.path.exists(p):
        return {}
    with open(p, encoding="utf-8", newline="") as fh:
        return {r["url_path"]: r for r in csv.DictReader(fh)}


def write_manifest(folder, man):
    out = io.StringIO()
    w = csv.writer(out, lineterminator="\n")
    w.writerow(MANIFEST_COLS)
    for k in sorted(man):
        w.writerow([man[k][c] for c in MANIFEST_COLS])
    write_text(os.path.join(folder, "manifest.csv"), out.getvalue())


def fetch_source_line(rows):
    ds = sorted({r[0][:10] for r in rows if len(r) > 4 and r[3] == "200" and r[4]})
    return ("# source: request log of bin/census_public_pensions.py fetch against %s, retrieved %s; one row per request" %
            (HOST, ", ".join(ds) if ds else "(no logged retrieval)"))


def remove_partial(partial):
    """The only removal in the script: the `partial/` folder by its exact path, when the vintage it would have become already exists complete (R4)."""
    if os.path.basename(partial) != "partial":
        raise Refused("remove_partial removes only a folder named partial")
    shutil.rmtree(partial)


class Fetcher:
    """One host. Opener, clock, sleeper and the STOP directory are injectable so every rule is testable without a network."""

    def __init__(self, cache, out_dir, stop_dir=STOP_DIR, sleep=None, max_requests=None, max_files=None, opener=real_opener, wall=time.time, mono=time.monotonic,
                 sleeper=time.sleep, now_utc=None, root=ROOT):
        self.sleep = float(DEFAULT_SLEEP if sleep is None else sleep)
        if self.sleep < MIN_SLEEP and not M("g16_pacing_off"):
            raise Refused("--sleep must be at least %g seconds" % MIN_SLEEP)
        self.max_requests = MAX_REQUESTS if max_requests is None else max_requests
        if self.max_requests > MAX_REQUESTS or self.max_requests < 1:
            raise Refused("the request cap must be between 1 and %d" % MAX_REQUESTS)
        self.max_files = max_files
        self.out_dir = out_dir
        self.cache = check_cache(cache, root=root, out_dir=out_dir)
        self.stop_dir = stop_dir
        self.opener, self.wall, self.mono, self.sleeper = opener, wall, mono, sleeper
        self.now_utc = now_utc or (lambda: dt.datetime.now(dt.timezone.utc))
        self.requests = 0
        self._last_mono = None
        self._last_file = os.path.join(self.cache, ".last_request")
        self.log_path = os.path.join(out_dir, "fetch_log.csv")
        self.partial = os.path.join(self.cache, "partial")

    # ---- state ----
    def utc_str(self):
        return self.now_utc().strftime("%Y-%m-%dT%H:%M:%SZ")

    def stop_path(self):
        return os.path.join(self.stop_dir, STOP_NAME)

    def check_stop(self):
        if os.path.exists(self.stop_path()) and not M("g15_stop_ignored"):
            raise Latched("STOP present: %s" % tilde(self.stop_path()))

    def run_requests(self):
        """Requests already made in this run (a run may take several invocations, each resuming the same partial folder)."""
        p = os.path.join(self.partial, "requests.txt")
        try:
            return int(read_file(p, "r").strip())
        except (OSError, ValueError):
            return 0

    def _pace(self):
        wait = 0.0
        try:
            with open(self._last_file, encoding="utf-8") as fh:
                last = float(fh.read().strip())
            wait = max(wait, self.sleep - (self.wall() - last))
        except (OSError, ValueError):
            pass
        if self._last_mono is not None:
            wait = max(wait, self.sleep - (self.mono() - self._last_mono))
        if wait > 0 and not M("g16_pacing_off"):
            self.sleeper(min(wait, self.sleep))

    @staticmethod
    def _stamp(ts):
        return dt.datetime.fromtimestamp(ts, dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

    def _log(self, ts, target, status, nbytes, sha, lm):
        _, rows = read_csv_rows(self.log_path)
        rows.append([self._stamp(ts), HOST, target, str(status), str(nbytes), sha, lm])
        write_text(self.log_path, render_csv("fetch_log.csv", rows, fetch_source_line(rows)))

    def _latch(self, status, target, retry_after, body, why):
        """Write the STOP file (exclusive create) and end the fetch."""
        text = stop_fields(status, self.utc_str(), target, retry_after or "", esc_head(body))
        create_stop(self.stop_path(), text)
        raise Latched("STOP written: %s (%s)" % (tilde(self.stop_path()), why))

    def _get(self, target):
        check_url("https://" + HOST + target)
        if self.run_requests() >= self.max_requests and not M("g16_cap_off"):
            raise MaxRequests("request cap reached (%d for this run)" % self.max_requests)
        self._pace()
        started = self.wall()
        self._last_mono = self.mono()
        os.makedirs(self.cache, exist_ok=True)
        write_text(self._last_file, "%r\n" % started)
        os.makedirs(self.partial, exist_ok=True)
        write_text(os.path.join(self.partial, "requests.txt"), "%d\n" % (self.run_requests() + 1))
        self.requests += 1
        try:
            status, body, rh = self.opener(HOST, target, dict(HEADERS), TIMEOUT, BODY_CAP)
        except SocketFault:
            self._log(started, target, "socket_fault", "", "", "")
            if M("g15_stop_on_timeout"):
                create_stop(self.stop_path(), stop_fields("timeout", self.utc_str(), target, "", ""))
            raise
        return status, body, rh, started

    def get_checked(self, target, problem_fn):
        """GET one URL. Anything but a 200 with a usable body latches the STOP file. Returns (body, headers)."""
        self.check_stop()
        status, body, rh, started = self._get(target)
        if status in (429, 500) and M("g16_retry"):
            status, body, rh, started = self._get(target)
        if status != 200:
            self._log(started, target, status, "", "", "")
            if status // 100 == 3 and M("g15_3xx_ignored"):
                return None, rh
            if status // 100 == 3 and M("g16_redirect_follow"):
                status, body, rh, started = self._get(rh.get("location", target))
            else:
                self._latch(status, target, rh.get("retry-after", ""), body, "HTTP %s" % status)
        if body is None:
            self._log(started, target, status, "", "", "")
            raise SocketFault("body larger than the cap")
        prob = problem_fn(body)
        lm = ""
        if rh.get("last-modified"):
            lm = http_date_iso(rh["last-modified"])
        self._log(started, target, status, len(body), sha256_of(body), lm)
        if prob:
            self._latch(status, target, "", body, "content test: %s" % prob)
        return body, rh

    # ---- the run ----
    def month_gate(self):
        """At most one run per calendar month (R3): a new run is refused when the request log already holds a retrieval this month."""
        if M("g16_month_gate_off"):
            return
        _, rows = read_csv_rows(self.log_path)
        now = self.utc_str()[:7]
        if any(len(r) > 3 and r[0][:7] == now and r[3] == "200" for r in rows):
            raise Refused("at most one run per calendar month: the request log already holds a retrieval in %s" % now)

    def run(self, say=print):
        """Returns the exit code: 0 done (or a slice done), 2 STOP present or written, 3 network fault."""
        try:
            self.check_stop()
        except Latched as e:
            say(str(e))
            return 2
        new_run = not os.path.isdir(self.partial)
        if new_run:
            self.month_gate()
        os.makedirs(self.partial, exist_ok=True)
        man = read_manifest(self.partial)
        try:
            if not os.path.exists(os.path.join(self.partial, "notice.txt")):
                body, _ = self.get_checked(LISTING_PATH, listing_body_problem)
                yrs = listing_years(body or b"")
                new = [y for y in yrs if y > SURVEY_YEARS[-1]]
                for y in new:
                    say("NOTICE: new survey year folder %d" % y)
                write_text(os.path.join(self.partial, "notice.txt"), "listing years above %d: %s\n" % (SURVEY_YEARS[-1], ",".join(str(y) for y in new) or "none"))
            n_new = 0
            for f in FILES:
                row = man.get(f["path"])
                if row and os.path.exists(os.path.join(self.partial, f["name"])):
                    continue
                if self.max_files is not None and n_new >= self.max_files:
                    say("slice done: %d of %d files in %s; run fetch again to resume" % (len(man), N_FILES, tilde(self.partial)))
                    return 0
                body, rh = self.get_checked(f["path"], file_body_problem)
                if body is None:
                    continue
                lm = http_date_iso(rh["last-modified"]) if rh.get("last-modified") else ""
                if not lm:
                    raise DataError("no Last-Modified header on %s: the vintage key cannot be formed" % f["path"])
                write_bytes(os.path.join(self.partial, f["name"]), body)
                man[f["path"]] = {"url_path": f["path"], "sha256": sha256_of(body), "bytes": str(len(body)), "last_modified": lm}
                write_manifest(self.partial, man)
                n_new += 1
                say("saved %s (%d bytes)" % (f["name"], len(body)))
        except Latched as e:
            say(str(e))
            return 2
        except SocketFault as e:
            say("FAIL network (%s): the run ends; nothing was latched and nothing is retried" % e)
            return 3
        except MaxRequests as e:
            say("FAIL %s" % e)
            return 3
        if len(man) < N_FILES:
            say("partial: %d of %d files; run fetch again to resume" % (len(man), N_FILES))
            return 0
        return self.complete(man, say)

    def complete(self, man, say):
        rows = [(p, r["last_modified"], int(r["bytes"])) for p, r in man.items()]
        key = vintage_key(rows)
        final = os.path.join(self.cache, key)
        if os.path.isdir(final):
            remove_partial(self.partial)
            say("no change: vintage %s already complete" % key)
            return 0
        os.rename(self.partial, final)
        name, date = newest_unit(rows)
        say("complete: vintage %s (newest unit file %s, %s); %d requests" % (key, name, date, self.run_requests_in(final)))
        return 0

    @staticmethod
    def run_requests_in(folder):
        try:
            return int(read_file(os.path.join(folder, "requests.txt"), "r").strip())
        except (OSError, ValueError):
            return 0


# ------------------------------------------------------------------------------------------------------------------------
# parsing the unit files (R8, method 2-4)
# ------------------------------------------------------------------------------------------------------------------------
def decode_text(raw):
    """UTF-8 with or without a BOM, and if that fails cp1252 (R8). Strict: a byte that neither can read is an error, never dropped."""
    bom = raw.startswith(codecs.BOM_UTF8)
    try:
        return raw.decode("utf-8-sig", errors="ignore" if M("g1_decode_ignore") else "strict"), ("utf-8-sig" if bom else "utf-8")
    except UnicodeDecodeError:
        pass
    try:
        return raw.decode("cp1252"), "cp1252"
    except UnicodeDecodeError:
        raise ContentError("the file decodes neither as UTF-8 nor as cp1252")


def delim_for(ext):
    if M("g1_dsv_comma"):
        return ","
    return "|" if ext == ".dsv" else ","


def _dec_or_none(s):
    s = s.strip()
    if s == "":
        return None
    return Decimal(s)


class UnitParse:
    """What one unit file yields. No unit identifier, name or value is kept except the sets used by the weight counts and the name guard (never written)."""

    def __init__(self, year):
        self.year = year
        self.enc = ""
        self.header = False
        self.rows = 0
        self.bad_fields = 0
        self.bad_years = 0
        self.bad_values = 0
        self.blank_weight = 0
        self.unknown = set()
        self.sums = collections.defaultdict(Decimal)
        self.all_one = True
        self.max_weight = Decimal(0)
        self.units = set()
        self.names = set()
        self.neg_rx30 = 0
        self.units_second = 0
        self.rows_second = 0
        self.rows_second_nonblank = 0
        self.example_id = ""
        self.problems = []


def parse_unit(raw, year, ext):
    up = UnitParse(year)
    try:
        text, up.enc = decode_text(raw)
    except ContentError as e:
        up.problems.append(str(e))
        return up
    delim = delim_for(ext)
    if M("g1_split_comma"):
        reader = (ln.split(delim) for ln in text.splitlines())
    else:
        reader = csv.reader(io.StringIO(text, newline=""), delimiter=delim)
    buf = []
    first = True
    for row in reader:
        if not row:
            continue
        if first:
            first = False
            if row[0].strip().upper().startswith("SAMPLE_YEAR") and not M("g1_header_as_data"):
                up.header = True
                continue
            if M("g1_skip_first_row"):
                continue
        if len(row) != 10:
            if M("g1_nine_eleven_ok") and len(row) in (9, 11):
                row = (list(row) + [""] * 10)[:10]
            else:
                up.bad_fields += 1
                continue
        y, name, uid, _state, _utype, code, val, _eic, _isc, w = [c.strip() for c in row]
        if y != str(year):
            up.bad_years += 1
            continue
        if code not in CODES and not (M("g4_54th_ok") and code == "RX99"):
            up.unknown.add(code)
            if M("g4_unknown_one_at_a_time"):
                break
            continue
        up.rows += 1
        try:
            wt = _dec_or_none(w)
            v = _dec_or_none(val)
        except decimal.InvalidOperation:
            up.bad_values += 1
            continue
        buf.append((uid, name, code, v, wt))
    # pass 1: weights per unit
    per_unit = collections.defaultdict(collections.Counter)
    for uid, name, code, v, wt in buf:
        up.units.add(uid)
        if len(name) >= 6:
            up.names.add(name)
        if wt is not None:
            per_unit[uid][wt] += 1
            if wt != 1:
                up.all_one = False
            if wt > up.max_weight:
                up.max_weight = wt
    modal = {}
    for uid, cnt in per_unit.items():
        modal[uid] = sorted(cnt.items(), key=lambda kv: (-kv[1], kv[0]))[0][0]
    second_units = set()
    for uid, name, code, v, wt in buf:
        if wt is not None and wt != modal.get(uid):
            second_units.add(uid)
            up.rows_second += 1
            if v is not None:
                up.rows_second_nonblank += 1
    up.units_second = len(second_units)
    if second_units:
        up.example_id = sorted(second_units)[0]
    # pass 2: the weighted sums
    for uid, name, code, v, wt in buf:
        if v is None:
            continue                                   # a blank value adds nothing (Census's convention)
        if wt is None:
            up.blank_weight += 1
            continue
        if code == "RX30" and v < 0:
            up.neg_rx30 += 1
        if code == "RX30" and M("g5_abs_rx30"):
            v = abs(v)
        if M("g2_unweighted"):
            f = v
        elif M("g2_modal_weight"):
            f = v * modal.get(uid, wt)
        elif M("g2_weight_twice"):
            f = v * wt * wt
        elif M("g8_float_sum"):
            f = Decimal(float(v) * float(wt))
        else:
            f = v * wt
        if M("g8_float_sum"):
            up.sums[code] = Decimal(float(up.sums[code]) + float(f))
        else:
            up.sums[code] += f
    if up.bad_fields:
        up.problems.append("%d row(s) with a field count other than 10" % up.bad_fields)
    if up.bad_years:
        up.problems.append("%d row(s) whose first field is not the survey year %d" % (up.bad_years, year))
    if up.unknown:
        up.problems.append("%d item code(s) outside the 53 allowed: %s" % (len(up.unknown), ", ".join(sorted(up.unknown))))
    if up.bad_values:
        up.problems.append("%d row(s) with a value or weight that is not a number" % up.bad_values)
    if up.blank_weight:
        up.problems.append("%d non-blank row(s) with a blank weight" % up.blank_weight)
    return up


def codes_for(col):
    c = list(UNIT_MAP[col])
    if M("g3_rx50_added") and col in ("contrib_employee", "contrib_total"):
        c.append("RX50")
    if M("g3_sublines_added") and col in ("contrib_government", "contrib_total"):
        c += ["RX61", "RX71"]
    if M("g3_rx30_in_earnings") and col == "earnings_ex_gains":
        c.append("RX30")
    if M("g3_rx91_added") and col == "expenses":
        c.append("RX91")
    return c


def pick(sums, codes):
    if M("g3_substring_codes"):
        return sum((v for k, v in sums.items() if any(k.startswith(c[:3]) for c in codes)), Decimal(0))
    return sum((sums.get(c, Decimal(0)) for c in codes), Decimal(0))


def unit_to_mn(x):
    return x / (Decimal(1000) if M("g6_unit_as_thousands") else Decimal(1000000))


def summary_to_mn(x):
    return x / (Decimal(1000000) if M("g6_summary_as_dollars") else Decimal(1000))


def unit_cols(up):
    """All nine national columns of a unit file, in $ million (exact)."""
    return {col: unit_to_mn(pick(up.sums, codes_for(col))) for col in UNIT_MAP}


def norm_label(s):
    return " ".join(str(s).split()).casefold()


def _num(c):
    return isinstance(c, (int, float, Decimal)) and not isinstance(c, bool)


class SummaryParse:
    def __init__(self, year):
        self.year = year
        self.values = {}                  # key -> Decimal, as printed (thousands of dollars; counts for `systems`)
        self.cv = {}                      # key -> Decimal fraction or None
        self.problems = []


def parse_summary(raw, year):
    import openpyxl
    sp = SummaryParse(year)
    labels = LABELS_OLD if year <= 2016 else LABELS_CUR
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        grid = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        wb.close()
    except Exception as e:                             # the workbook is not readable: a content failure, not a refusal
        sp.problems.append("the workbook does not open (%s)" % type(e).__name__)
        return sp
    item_i = next((i for i, r in enumerate(grid) if r and r[0] is not None and norm_label(r[0]) == "item"), None)
    if item_i is None:
        sp.problems.append("no row whose first cell is Item")
        return sp
    geo = None
    for r in grid[:item_i]:
        for j, c in enumerate(r):
            if c is not None and str(c).strip() == "United States":
                geo = j
    if M("g4_col_by_position"):
        geo = 1
    header = grid[item_i]
    if geo is None or geo >= len(header) or header[geo] is None or str(header[geo]).strip() != "State & Local":
        sp.problems.append("no column headed State & Local under United States")
        return sp
    cvc = geo + 1 if geo + 1 < len(header) and header[geo + 1] is not None and str(header[geo + 1]).strip().upper().startswith("CV") else None
    byl = collections.defaultdict(list)
    for r in grid[item_i + 1:]:
        if r and r[0] is not None and str(r[0]).strip():
            byl[norm_label(r[0])].append(r)
    missing, dup, nonnum = [], [], []
    for key, label in labels.items():
        rows = byl.get(norm_label(label), [])
        if not rows:
            if M("g4_missing_zero"):
                sp.values[key] = Decimal(0)
            else:
                missing.append(label)
            continue
        if len(rows) > 1:
            dup.append(label)
            continue
        r = rows[0]
        cell = r[geo] if geo < len(r) else None
        if not _num(cell):
            nonnum.append(label)
            continue
        sp.values[key] = D(cell)
        if key in CV_KEYS and cvc is not None:
            c2 = r[cvc] if cvc < len(r) else None
            sp.cv[key] = D(c2) if _num(c2) else None
    if year > 2016 and M("g5_losses_gain"):
        g = byl.get("investment gains", [])
        l = byl.get("investment losses", [])
        if g and l and _num(g[0][geo]) and _num(l[0][geo]):
            sp.values["gains"] = D(g[0][geo]) + D(l[0][geo])
    if missing:
        sp.problems.append("missing label(s): %s" % "; ".join(missing))
    if dup:
        sp.problems.append("label(s) found more than once: %s" % "; ".join(dup))
    if nonnum:
        sp.problems.append("label(s) with no number in the State & Local column: %s" % "; ".join(nonnum))
    return sp


def sum_cols(sp, schema):
    """The national columns a summary workbook gives, in $ million."""
    out = {}
    v = sp.values
    for src, col in (("contrib_employee", "contrib_employee"), ("contrib_government", "contrib_government"), ("contrib_total", "contrib_total"), ("benefits", "benefits"),
                     ("withdrawals", "withdrawals"), ("assets", "assets_end")):
        out[col] = summary_to_mn(v[src]) if src in v else None
    if schema == "current":
        for src, col in (("earnings", "earnings_ex_gains"), ("gains", "gains_losses"), ("expenses", "expenses")):
            out[col] = summary_to_mn(v[src]) if src in v else None
    else:
        out["earnings_old"] = summary_to_mn(v["earnings"]) if "earnings" in v else None
    return out


# ------------------------------------------------------------------------------------------------------------------------
# vintages: a complete cache folder (R4, R5), real or in memory
# ------------------------------------------------------------------------------------------------------------------------
class Vintage:
    def __init__(self, key, meta, reader, requests=0, notice_years=(), manifest_sha="", retrieved=""):
        self.key = key
        self.meta = meta                                # url_path -> {"sha256", "bytes", "last_modified"}
        self._read = reader
        rows = [(p, m["last_modified"], int(m["bytes"])) for p, m in meta.items()]
        self.rows = rows
        self.newest_unit_file, self.newest_unit_file_date = newest_unit(rows)
        self.requests = requests
        self.notice_years = tuple(notice_years)
        self.manifest_sha = manifest_sha
        self.retrieved = retrieved

    def date(self, path):
        return self.meta[path]["last_modified"][:10]

    def read(self, path):
        return self._read(path)

    def order(self):
        return (self.newest_unit_file_date, self.retrieved, self.key)


KEY_RX = re.compile(r"^[0-9a-f]{16}$")


def list_vintage_names(cache):
    names = []
    if os.path.isdir(cache):
        for n in sorted(os.listdir(cache)):
            if KEY_RX.match(n) and os.path.isdir(os.path.join(cache, n)):
                names.append(n)
        if M("g10_partial_as_vintage") and os.path.isdir(os.path.join(cache, "partial")):
            names.append("partial")
    return names


def load_vintage(folder, fetch_rows=()):
    """A complete cache folder as a Vintage; every file is checked against the manifest and the folder name against the recomputed key."""
    man = read_manifest(folder)
    missing = [f["path"] for f in FILES if f["path"] not in man]
    if missing:
        raise DataError("%s: the manifest lacks %d of the %d files" % (os.path.basename(folder), len(missing), N_FILES))
    names = {f["path"]: f["name"] for f in FILES}
    for p in names:
        fp = os.path.join(folder, names[p])
        if not os.path.exists(fp) or sha256_of(read_file(fp)) != man[p]["sha256"]:
            raise DataError("%s: %s is absent or does not match its manifest sha256" % (os.path.basename(folder), names[p]))
    rows = [(p, man[p]["last_modified"], int(man[p]["bytes"])) for p in names]
    key = vintage_key(rows)
    if key != os.path.basename(folder):
        raise DataError("%s: the folder name is not the vintage key of its manifest (%s)" % (os.path.basename(folder), key))
    req = Fetcher.run_requests_in(folder)
    notice = []
    try:
        t = read_file(os.path.join(folder, "notice.txt"), "r")
        tail = t.split(":", 1)[1].strip()
        notice = [int(x) for x in tail.split(",") if x.strip().isdigit()]
    except (OSError, IndexError):
        pass
    shas = {man[p]["sha256"] for p in names}
    dates = [r[0] for r in fetch_rows if len(r) > 5 and r[3] == "200" and r[5] in shas]
    return Vintage(key, {p: man[p] for p in names}, lambda p: read_file(os.path.join(folder, names[p])), requests=req, notice_years=notice,
                   manifest_sha=sha256_of(read_file(os.path.join(folder, "manifest.csv"))), retrieved=max(dates) if dates else "")


# ------------------------------------------------------------------------------------------------------------------------
# the national rows, the ratios and the revisions
# ------------------------------------------------------------------------------------------------------------------------
def parse_vintage(V):
    """{year: {"unit": UnitParse, "unit_date", "sum": SummaryParse, "sum_date"}} for every survey year."""
    P = {y: {"year": y} for y in SURVEY_YEARS}
    for f in FILES:
        y = f["year"]
        raw = V.read(f["path"])
        if f["kind"] == "unit":
            P[y]["unit"] = parse_unit(raw, y, os.path.splitext(f["name"])[1].lower())
            P[y]["unit_date"] = V.date(f["path"])
            P[y]["unit_path"] = f["path"]
        else:
            P[y]["sum"] = parse_summary(raw, y)
            P[y]["sum_date"] = V.date(f["path"])
    return P


def net_of(tot, ben, wd, exp):
    if M("g5_net_benefits_minus_contrib"):
        return ben - tot
    if M("g5_withdrawals_added"):
        return tot - ben + wd
    if M("g3_rx90_in_net"):
        return tot - ben - wd - (exp or Decimal(0))
    return tot - ben - wd


NAT_COLS = ("contrib_employee", "contrib_government", "contrib_total", "benefits", "withdrawals", "net_cash", "earnings_ex_gains", "earnings_old", "gains_losses",
            "expenses", "assets_end")


def national_rows(P, vkey):
    """One dict per survey year, values as exact Decimals in $ million."""
    N = {}
    for y in SURVEY_YEARS:
        e = P[y]
        up, sp = e.get("unit"), e.get("sum")
        schema = "old" if y <= 2016 else "current"
        r = {"survey_year": y, "schema": schema, "year_kind": "census" if y in CENSUS_YEARS else "sample", "vintage": vkey}
        for c in NAT_COLS:
            r[c] = None
        if up is not None and y not in SUMMARY_SOURCED:
            r.update(unit_cols(up))
            r["source_kind"], r["source_file_date"] = "unit_file", e["unit_date"]
        elif sp is not None:
            r.update({k: v for k, v in sum_cols(sp, schema).items()})
            r["source_kind"], r["source_file_date"] = "summary", e["sum_date"]
        else:
            raise DataError("survey year %d has neither a unit file nor a summary workbook" % y)
        if M("g7_old_earnings_in_exgains") and schema == "old":
            r["earnings_ex_gains"] = r["earnings_old"]
        if r["contrib_total"] is not None and r["benefits"] is not None and r["withdrawals"] is not None:
            r["net_cash"] = net_of(r["contrib_total"], r["benefits"], r["withdrawals"], r["expenses"])
        r["systems"] = int(sp.values["systems"]) if sp is not None and "systems" in sp.values else None
        r["summary_file_date"] = e.get("sum_date", "") if sp is not None else ""
        for k in CV_KEYS:
            r["cv_" + k] = None
        if sp is not None and (y not in CENSUS_YEARS or M("g18_census_cv_nonblank")):
            for k in CV_KEYS:
                r["cv_" + k] = sp.cv.get(k)
        N[y] = r
    return N


def band_pattern(change_pct, band):
    """R2-3 / R3-1: a percent change is compared with the percent band; descriptive, not a test."""
    if change_pct is None or band is None:
        return "n_a"
    if M("g11_rose_fell"):
        return "rose" if change_pct > band else ("fell" if change_pct < -band else "within_band")
    if change_pct > band:
        return "above_band"
    if change_pct < -band:
        return "below_band"
    return "within_band"


def ratio_rows(N):
    """Per year: the ratios of method 6, exact Decimals (rounded only at write time)."""
    R = {}
    for y in SURVEY_YEARS:
        n = N[y]
        r = {"survey_year": y, "schema": n["schema"]}
        tot, ben, wd, exp, earn, net, ast_ = n["contrib_total"], n["benefits"], n["withdrawals"], n["expenses"], n["earnings_ex_gains"], n["net_cash"], n["assets_end"]
        if M("g8_round_before_divide"):
            tot, ben, wd = (Decimal(fdec(x, 3)) if x is not None else None for x in (tot, ben, wd))
        r["bpc"] = ben / tot if tot else None
        r["bwpc"] = (ben + wd) / tot if tot else None
        hundred = Decimal(1) if M("g11_pct_no_100") else Decimal(100)
        prev = N.get(y - 1)
        r["start"] = None
        if prev is not None and prev["assets_end"] and net is not None and y not in (2015,) and (prev["schema"] == n["schema"] or M("g7_start_across_break")):
            den = ast_ if M("g11_start_same_year") else prev["assets_end"]
            r["start"] = hundred * net / den
        r["end"] = None
        if ast_ and net is not None:
            den = (prev["assets_end"] if (M("g11_end_prev_year") and prev is not None) else ast_)
            r["end"] = hundred * net / den
        r["income_gap"] = r["income_gap_exp"] = None
        use = earn if n["schema"] == "current" else (n["earnings_old"] if M("g7_income_for_old") else None)
        if use is not None and ben is not None and tot is not None:
            den = ben + wd - tot
            if den > 0:
                r["income_gap"] = use / den
            den2 = ben + wd - tot if M("g11_gap_no_rx90") else ben + wd + (exp or Decimal(0)) - tot
            if den2 > 0:
                r["income_gap_exp"] = use / den2
        r["cv_bound"] = None
        if y in CENSUS_YEARS:
            r["cv_bound"] = Decimal(0)
        elif n["cv_benefits"] is not None and n["cv_contrib_total"] is not None:
            r["cv_bound"] = (Decimal(1) if M("g18_bound_no_100") else Decimal(100)) * (n["cv_benefits"] + n["cv_contrib_total"])
        R[y] = r
    return R


def pct_change(new, prev):
    """100 x (new - prev) / |prev|: the usual percent change for a positive base, and the sign that reads naturally for a negative base (gains)."""
    if prev is None or new is None or prev == 0:
        return None
    return Decimal(100) * (new - prev) / abs(prev)


def revision_row(y, item, sk_prev, d_prev, v_prev, sk_new, d_new, v_new):
    return {"survey_year": y, "item": item, "source_prev": sk_prev, "date_prev": d_prev, "value_prev": v_prev, "source_new": sk_new, "date_new": d_new,
            "value_new": v_new, "change": v_new - v_prev, "change_pct": pct_change(v_new, v_prev)}


def revisions_a(P):
    """(a) each year's summary against its unit file where their `Last-Modified` dates differ (2.8). Returns (rows, {year: ratio revision percent})."""
    rows, ratio = [], {}
    for y in SURVEY_YEARS:
        e = P[y]
        up, sp = e.get("unit"), e.get("sum")
        if up is None or sp is None or y in SUMMARY_SOURCED:
            continue
        if e["unit_date"] == e["sum_date"] and not M("g10_all_years_revised"):
            continue
        uc, sc = unit_cols(up), sum_cols(sp, "current")
        for item in ITEMS:
            if sc.get(item) is None:
                continue
            rows.append(revision_row(y, item, "summary", e["sum_date"], sc[item], "unit_file", e["unit_date"], uc[item]))
        if sc["benefits"] and sc["contrib_total"] and uc["contrib_total"]:
            rs, ru = sc["benefits"] / sc["contrib_total"], uc["benefits"] / uc["contrib_total"]
            ratio[y] = Decimal(100) * (ru / rs - 1)
    return rows, ratio


def revisions_b(Nprev, Nnew):
    """(b) an adjacent pair of complete cached vintages: every cell whose value differs."""
    rows = []
    for y in SURVEY_YEARS:
        a, b = Nprev.get(y), Nnew.get(y)
        if a is None or b is None:
            continue
        for item in ITEMS:
            va, vb = a.get(item), b.get(item)
            if va is None or vb is None or va == vb:
                continue
            rows.append(revision_row(y, item, a["source_kind"], a["source_file_date"], va, b["source_kind"], b["source_file_date"], vb))
    return rows


def sort_revisions(rows):
    return sorted(rows, key=lambda r: (r["survey_year"], ITEMS.index(r["item"]), r["date_prev"], r["date_new"]))


def measured_bands(ratio_revs):
    allv = [abs(p) for p in ratio_revs.values()]
    samp = [abs(p) for y, p in ratio_revs.items() if (y not in CENSUS_YEARS or M("g11_sample_band_incl_census"))]
    return (max(allv) if allv else None), (max(samp) if samp else None)


def q3(x):
    return None if x is None else x.quantize(Decimal("0.001"), rounding=decimal.ROUND_HALF_EVEN)


def band_values(ratio_revs, frozen=None, frozen_sample=None):
    """(band, sample-years band): the frozen constants once they are set; the first build uses what it measured, rounded to three decimals."""
    fb = FROZEN_BAND_PCT if frozen is None else frozen
    fs = FROZEN_SMP_BAND_PCT if frozen_sample is None else frozen_sample
    mb, ms = measured_bands(ratio_revs)
    if M("g11_band_recomputed"):
        return q3(mb), q3(ms)
    return (fb if fb is not None else q3(mb)), (fs if fs is not None else q3(ms))


# ------------------------------------------------------------------------------------------------------------------------
# summary.csv rows (section 4)
# ------------------------------------------------------------------------------------------------------------------------
SPANS_Q2 = ((2017, 2023), (2023, 2024), (2024, 2025), (2023, 2025))
SPANS_Q3 = ((2018, 2023), (2023, 2025))
PRIMARY_SPAN = (2023, 2025)
UNIT_DECIMALS = {"usd_per_usd": 4, "pct": 3, "pp": 3, "count": 0}


def span_ok(a, b):
    """G7: no change within the old schema (2013-2016) and none that spans 2016/2017; a change is between current-schema years only."""
    return a >= 2017 and b > a


def summary_rows(N, R, ratio_revs, revs, band, band_s, vkey, s1=None):
    S = []

    def add(q, seg, fy, ty, stat, val, unit, pat="n_a", prim="no"):
        if val is None:
            return
        S.append({"question": q, "segment": seg, "from_year": fy, "to_year": ty, "statistic": stat, "value": val, "unit": unit, "pattern": pat, "primary": prim,
                  "vintage": vkey})

    # Q2: benefit payments per dollar of contributions
    for y in SURVEY_YEARS:
        add("Q2", N[y]["schema"], None, y, "benefits_per_contrib_usd", R[y]["bpc"], "usd_per_usd")
    spans = list(SPANS_Q2)
    if M("g7_change_within_old"):
        spans.append((2014, 2015))
    if M("g7_change_span_break"):
        spans.append((2016, 2018))
    for a, b in spans:
        if not span_ok(a, b) and not (M("g7_change_within_old") or M("g7_change_span_break")):
            continue
        ra, rb = R[a]["bpc"], R[b]["bpc"]
        if ra is None or rb is None:
            continue
        seg = "old" if b <= 2016 else "current"
        usd = rb - ra
        pct = Decimal(100) * (rb / ra - 1)
        add("Q2", seg, a, b, "change_benefits_per_contrib_usd", usd, "usd_per_usd", band_pattern(usd, band) if M("g11_abs_vs_band") else "n_a")
        prim = "yes" if (a, b) == PRIMARY_SPAN or (M("g11_two_primary") and (a, b) == (2024, 2025)) else "no"
        add("Q2", seg, a, b, "change_benefits_per_contrib_pct", pct, "pct", band_pattern(pct, band), prim)
    # Q3: net cash flow as a percent of opening assets (primary measure) and of end-of-year assets (secondary)
    for y in SURVEY_YEARS:
        add("Q3", N[y]["schema"], None, y, "net_cash_pct_of_assets_start", R[y]["start"], "pct")
    for y in SURVEY_YEARS:
        add("Q3", N[y]["schema"], None, y, "net_cash_pct_of_assets_end", R[y]["end"], "pct")
    for a, b in SPANS_Q3:
        for key, stat in (("start", "change_net_cash_start_pp"), ("end", "change_net_cash_end_pp")):
            if R[a][key] is None or R[b][key] is None:
                continue
            pp = R[b][key] - R[a][key]
            add("Q3", "current", a, b, stat, pp, "pp", band_pattern(pp, band) if M("g11_pattern_on_pp") else "n_a")
    for seg, ys in (("old", OLD_YEARS), ("current", tuple(y for y in SURVEY_YEARS if y >= 2017))):
        vals = [N[y]["net_cash"] for y in ys if N[y]["net_cash"] is not None]
        add("Q3", seg, ys[0], ys[-1], "n_years_net_cash_negative", sum(1 for v in vals if v < 0), "count")
    # Q4: investment income per dollar by which payouts exceeded contributions (current schema only)
    for y in SURVEY_YEARS:
        if N[y]["schema"] != "current" and not M("g7_income_for_old"):
            continue
        add("Q4", N[y]["schema"], None, y, "income_per_net_cash_gap_usd", R[y]["income_gap"], "usd_per_usd")
        add("Q4", N[y]["schema"], None, y, "income_per_gap_after_expenses_usd", R[y]["income_gap_exp"], "usd_per_usd")
    # Q5: revisions and the frozen band
    for item in ITEMS:
        rs = [r for r in revs if r["item"] == item and r["change_pct"] is not None]
        if rs:
            add("Q5", "current", min(r["survey_year"] for r in rs), max(r["survey_year"] for r in rs), "max_abs_revision_pct__" + item,
                max(abs(r["change_pct"]) for r in rs), "pct")
    for y in sorted(ratio_revs):
        add("Q5", "current", None, y, "ratio_revision_pct", ratio_revs[y], "pct")
    add("Q5", "current", None, None, "revision_band_pct", band, "pct")
    add("Q5", "current", None, None, "revision_band_sample_years_pct", band_s, "pct")
    # Q6: the sampling bound
    for y in (2023, 2024):
        add("Q6", N[y]["schema"], None, y, "ratio_cv_bound_pct", R[y]["cv_bound"], "pct")
    if s1 is not None:
        add("S1", "current", None, 2025, "s1_revision_2025_pct", s1, "pct", band_pattern(s1, band))
    return S


def fmt_summary_value(row):
    v = row["value"]
    return str(int(v)) if row["unit"] == "count" else fdec(v, UNIT_DECIMALS[row["unit"]] - (1 if M("g8_wrong_decimals") and row["unit"] == "usd_per_usd" else 0))


def summary_problems(rows):
    """Rules on summary.csv rows given as lists of text (the committed form): allowed values, one primary row, a unique key, a pattern only where it is allowed, and no change
    row inside the old schema or across the 2016/2017 break (G7, G11, G12). Used by the build and by the selftest; check.sh holds its own copy."""
    h = ALLOW["summary.csv"]
    out, seen = [], set()
    n_primary = 0
    for r in rows:
        d = dict(zip(h, r))
        key = tuple(d[k] for k in SUMMARY_KEY)
        if key in seen and not M("g12_dup_key_ok"):
            out.append("duplicate key %s" % "|".join(key))
        seen.add(key)
        for col, allowed in (("question", ALLOWED_VALUES["question"]), ("segment", ALLOWED_VALUES["segment"]), ("statistic", ALLOWED_VALUES["statistic"]),
                             ("unit", ALLOWED_VALUES["unit"]), ("pattern", ALLOWED_VALUES["pattern"]), ("primary", ALLOWED_VALUES["primary"])):
            if d[col] not in allowed and not (M("g12_stat_ok") and col in ("statistic", "question")):
                out.append("%s %r is not an allowed value" % (col, d[col]))
        if d["primary"] == "yes":
            n_primary += 1
        if d["pattern"] != "n_a" and d["statistic"] not in ("change_benefits_per_contrib_pct", "s1_revision_2025_pct"):
            out.append("a pattern on %s" % d["statistic"])
        if d["statistic"].startswith("change_") and d["from_year"]:
            a, b = int(d["from_year"]), int(d["to_year"])
            if a <= 2016 and b <= 2016:
                out.append("a change row inside the old schema (%d to %d)" % (a, b))
            elif a <= 2016 < b:
                out.append("a change row spanning 2016/2017 (%d to %d)" % (a, b))
    if rows and n_primary != 1:
        out.append("%d primary rows (exactly one is allowed)" % n_primary)
    return out


# ------------------------------------------------------------------------------------------------------------------------
# the pinned checks (section 8)
# ------------------------------------------------------------------------------------------------------------------------
PINS = {
    2025: {"contrib_employee": "78232.942", "contrib_government": "236787.982", "contrib_total": "315020.924", "benefits": "418247.000", "withdrawals": "8662.260",
           "earnings_ex_gains": "123246.347", "gains_losses": "493615.768", "expenses": "32488.230", "assets_end": "6490183.499"},
    2024: {"contrib_total": "304507.314", "benefits": "404460.423", "withdrawals": "8485.325", "expenses": "30383.985", "assets_end": "5984012.295"},
    2023: {"contrib_total": "294746.197", "benefits": "386915.776", "withdrawals": "8073.515", "expenses": "26164.033", "assets_end": "5484493.565"},
    2022: {"contrib_total": "287958.518", "benefits": "368117.415", "gains_losses": "-115465.262", "expenses": "27858.658", "assets_end": "5422683.528"},
    2017: {"contrib_total": "200567.477", "benefits": "297948.620", "withdrawals": "7047.012", "expenses": "17254.773", "assets_end": "4077278.537"},
    2021: {"contrib_total": "247844.543", "benefits": "352030.838", "expenses": "25378.773", "assets_end": "5616036.664"},
    2013: {"contrib_total": "152814.152", "benefits": "242324.376", "withdrawals": "5808.003", "assets_end": "3289331.745"},
    2016: {"contrib_total": "190647.517", "benefits": "283350.356", "withdrawals": "6123.396", "assets_end": "3733413.704"},
    2018: {"contrib_total": "218847.610", "benefits": "311712.813", "withdrawals": "7141.871", "assets_end": "4345989.226"},
    2019: {"contrib_total": "226284.675", "benefits": "322701.634", "withdrawals": "6913.836", "assets_end": "4442571.521"},
    2020: {"contrib_total": "241972.089", "benefits": "336659.802", "withdrawals": "7060.712", "assets_end": "4597361.334"},
}
SYSTEMS_PINS = {2013: 3992, 2014: 3972, 2015: 6299, 2016: 6276}
UNITS_PINS = {2022: 4394, 2017: 5529}
SHA_2025 = ("ce8e85af", "b1533e")
REV_PINS = {(2022, "contrib_total"): ("282833.491", "287958.518", "5125.027"), (2024, "benefits"): ("405454.166", "404460.423", "-993.743")}
RATIO_REV_PINS = {2022: "-1.898", 2023: "0.102", 2024: "0.176"}
BAND_PIN, BAND_SAMPLE_PIN = "1.898", "0.176"
PIN_GROUP_NEW = (2022, 2023, 2024, 2025)


def chk(cid, desc, expected, observed, status):
    return [cid, desc, expected, observed, status]


def worst(statuses):
    for s in ("FAIL", "WARN"):
        if s in statuses:
            return s
    return "PASS"


class Cells:
    """A list of (label, expected text, observed text, status if they differ) cells for one check."""

    def __init__(self):
        self.items = []

    def add(self, label, expected, got, bad_status="FAIL"):
        self.items.append((label, str(expected), str(got), bad_status))

    def result(self):
        bad = [(l, e, g, s) for l, e, g, s in self.items if e != g]
        if not bad:
            return "all %d pinned values match" % len(self.items), "PASS"
        return "%d of %d differ: %s" % (len(bad), len(self.items), "; ".join("%s expected %s, got %s" % (l, e, g) for l, e, g, s in bad[:6])), worst([b[3] for b in bad])


def mismatch_status(group, V, revs, year, item):
    has_rev = any(r["survey_year"] == year and r["item"] == item and r["change"] != 0 for r in revs)
    if group == "new":
        return "FAIL" if V.newest_unit_file_date == NEWEST_PIN_DATE else ("WARN" if has_rev else "FAIL")
    return "WARN" if has_rev else "FAIL"


def pin_cells(years, N, V, revs, group):
    cells = Cells()
    for y in years:
        for col, want in PINS[y].items():
            item = {"expenses": None}.get(col, col)
            got = N[y].get(col)
            cells.add("%d %s" % (y, col), fdec(Decimal(want), 3), fdec(got, 3), mismatch_status(group, V, revs, y, item))
    return cells


def c04_compare(P):
    """Same-vintage comparison (R3-2): the weighted unit-file sum against the summary of the same Last-Modified date, in thousands of dollars."""
    tol = Decimal(10) if M("g10_same_date_5k_ok") else Decimal(1)
    gaps, years_same, years_diff, bad = [], [], [], []
    for y in SURVEY_YEARS:
        e = P[y]
        up, sp = e.get("unit"), e.get("sum")
        if up is None or sp is None or y in SUMMARY_SOURCED:
            continue
        same = e["unit_date"] == e["sum_date"]
        (years_same if same else years_diff).append(y)
        for col, key in (("contrib_total", "contrib_total"), ("benefits", "benefits"), ("withdrawals", "withdrawals"), ("assets_end", "assets")):
            if key not in sp.values:
                continue
            gap = abs(pick(up.sums, codes_for(col)) / Decimal(1000) - sp.values[key])
            if same:
                gaps.append((gap, y, col))
                if gap > tol:
                    bad.append("%d %s off by %s thousand" % (y, col, fdec(gap, 3)))
            elif gap > tol and M("g10_diffdate_fail"):
                bad.append("%d %s: a different-date gap of %s thousand written as a failure" % (y, col, fdec(gap, 3)))
    return years_same, years_diff, gaps, bad


def check_c09(N, P):
    n25 = N[2025]
    neg_ok = n25["net_cash"] is not None and (n25["net_cash"] < 0 or M("g5_c09_blind"))
    gap = n25["benefits"] - n25["contrib_total"]
    u25, u24 = P[2025].get("unit"), P[2024].get("unit")
    cells = Cells()
    cells.add("2025 net_cash_usd_mn is negative", "True", str(bool(neg_ok)))
    cells.add("2025 benefits minus contributions", "103226.076", fdec(gap, 3))
    cells.add("2025 negative RX30 rows", "8", str(u25.neg_rx30) if u25 else "none")
    cells.add("2024 negative RX30 rows", "24", str(u24.neg_rx30) if u24 else "none")
    obs, st = cells.result()
    return chk("C09", "Sign: survey year 2025 net cash is negative; benefits minus contributions is Census's gap that leaves out withdrawals; the negative gain rows stay negative", "as pinned", obs, st)


def check_c10(P):
    cen_bad = [y for y in CENSUS_YEARS if P[y].get("unit") is not None and not P[y]["unit"].all_one and not M("g2_census_weight_ok")]
    smp_bad = [y for y in SURVEY_YEARS if y not in CENSUS_YEARS and P[y].get("unit") is not None and not P[y]["unit"].max_weight > 1]
    obs = (("census-year weight other than 1 in %s; " % ", ".join(map(str, cen_bad))) if cen_bad else "") + (("no weight above 1 in sample year(s) %s" % ", ".join(map(str, smp_bad))) if smp_bad else "")
    return chk("C10", "Census years 2017 and 2022: every weight is 1; every sample year read from a unit file has at least one weight above 1", "no exception",
               obs or "census years all 1; every sample year has a weight above 1", "FAIL" if cen_bad else ("WARN" if smp_bad else "PASS"))


def check_c12(P):
    parts, nonzero = [], False
    for y in SURVEY_YEARS:
        up = P[y].get("unit")
        if up is None:
            continue
        d60 = unit_to_mn(up.sums.get("RX61", Decimal(0)) + up.sums.get("RX62", Decimal(0)) - up.sums.get("RX60", Decimal(0)))
        d70 = unit_to_mn(sum((up.sums.get(k, Decimal(0)) for k in ("RX71", "RX72", "RX73")), Decimal(0)) - up.sums.get("RX70", Decimal(0)))
        piece = []
        if d60 != 0 or d70 != 0:
            piece.append("sub-lines minus total: RX60 %s, RX70 %s" % (fdec(d60, 3), fdec(d70, 3)))
        if up.units_second:
            piece.append("%d unit(s) with a second weight, %d non-blank row(s) carrying it" % (up.units_second, up.rows_second_nonblank))
            if M("g2_c12_prints_ids"):
                piece.append("first unit %s" % up.example_id)
        if piece:
            nonzero = True
            parts.append("%d: %s" % (y, "; ".join(piece)))
    return chk("C12", "Sub-line identities (RX60 against RX61 + RX62, RX70 against RX71-RX73) and rows that carry a second weight; counts only, no unit identifier or value",
               "2023: -560.031 and -12.930; 515 units, 4 rows", "; ".join(parts) if parts else "every sub-line identity holds and no unit has a second weight", "WARN" if nonzero else "PASS")


def check_c13(N):
    bad = []
    for y in SURVEY_YEARS:
        n = N[y]
        if abs(n["contrib_employee"] + n["contrib_government"] - n["contrib_total"]) > Decimal("0.002"):
            bad.append("%d employee + government differs from total" % y)
        if abs(n["net_cash"] - (n["contrib_total"] - n["benefits"] - n["withdrawals"])) > Decimal("0.002"):
            bad.append("%d net cash differs from contributions minus benefits minus withdrawals" % y)
    return chk("C13", "Tamper check: on every national row employee + government = total contributions and net cash = contributions - benefits - withdrawals, within 0.002 $mn",
               "no row off", "; ".join(bad[:4]) if bad else "all %d rows hold" % len(N), "FAIL" if bad else "PASS")


def check_c18(N):
    cvs = [(n[k], y) for y, n in N.items() for k in ("cv_contrib_total", "cv_benefits", "cv_assets") if n[k] is not None]
    mxcv = max(cvs) if cvs else None
    return chk("C18", "CV scale: every national sample-year CV read for contributions, benefits and assets is at most 0.05 (a fraction; above that the column has changed to percent)",
               "at most 0.05", ("largest %s (survey year %d)" % (mxcv[0], mxcv[1])) if mxcv else "no CV read",
               "FAIL" if (mxcv and mxcv[0] > CV_MAX and not M("g18_cv_percent_ok")) else "PASS")


def check_c19(V):
    in_table = UNIT_2024_PATH in V.meta and not any(p in V.meta for p in STALE_2024_ROOT_PATHS)
    d24 = V.date(UNIT_2024_PATH) if UNIT_2024_PATH in V.meta else ""
    stale = bool(d24) and d24 < NEWEST_PIN_DATE and V.newest_unit_file_date >= NEWEST_PIN_DATE and not M("g19_stale_ok")
    return chk("C19", "The 2024 unit file came from datasets/2024/aspp-historical-datasets/ASPP_Unit_File_2024.csv and its Last-Modified is 2026-05-14 or later",
               "that path, dated 2026-05-14 or later", "path %s; Last-Modified %s" % ("as required" if in_table else "NOT as required", d24 or "n/a"), "PASS" if in_table and not stale else "FAIL")


def check_c20(P):
    probs = []
    for y in SURVEY_YEARS:
        for kind in ("unit", "sum"):
            o = P[y].get(kind)
            if o is not None and o.problems:
                probs.append("%d %s: %s" % (y, "unit file" if kind == "unit" else "summary", "; ".join(o.problems)))
    return chk("C20", "Content tests R3b: unit files decode, every row has 10 fields and the expected survey year, every item code is one of the 53; summary workbooks open and hold "
               "the labelled rows and the State & Local column", "no failure", " | ".join(probs[:6]) if probs else "no failure in %d files" % N_FILES, "FAIL" if probs else "PASS")


def build_checks(c):
    N, P, V, revs, ratio_revs = c["N"], c["P"], c["V"], c["revs"], c["ratio_revs"]
    out = []
    fresh = V.newest_unit_file_date == NEWEST_PIN_DATE

    def pin_row(cid, desc, years, group, extra=None):
        cells = pin_cells(years, N, V, revs, group)
        for lab, exp, got, st in (extra or []):
            cells.add(lab, exp, got, st)
        obs, st = cells.result()
        out.append(chk(cid, desc, "equal to the pinned values (survey years %s)" % ", ".join(str(y) for y in years), obs, st))

    # C01-C03
    p25 = next(f["path"] for f in FILES if f["year"] == 2025 and f["kind"] == "unit")
    sha = V.meta[p25]["sha256"]
    sha_st = "FAIL" if fresh else "WARN"
    sha_pin = c.get("sha_pin", SHA_2025)
    pin_row("C01", "Survey year 2025 (unit file dated 2026-05-14): contributions, employee and government, benefits, withdrawals, earnings excluding gains, gains, expenses, assets",
            [2025], "new", [("2025 file sha256 start", sha_pin[0], sha[:8], sha_st), ("2025 file sha256 end", sha_pin[1], sha[-6:], sha_st)])
    pin_row("C02", "Survey year 2024 (unit file dated 2026-05-14): contributions, benefits, withdrawals, expenses, assets", [2024], "new")
    pin_row("C03", "Survey year 2023 (unit file dated 2026-05-14): contributions, benefits, withdrawals, expenses, assets", [2023], "new")
    # C04: same-vintage comparison
    ys, yd, gaps, bad = c04_compare(P)
    mx = max(gaps) if gaps else None
    out.append(chk("C04", "Same-vintage comparison: each unit-file weighted sum (contributions, benefits, withdrawals, assets) against the summary of the same Last-Modified date, to the "
                   "thousand dollars; a different-date gap is a revision row, not a failure", "within 1 thousand dollars on every compared year",
                   ("compared %s; largest gap %s thousand; different dates (see revisions.csv): %s" % (", ".join(map(str, ys)) or "none", fdec(mx[0], 3) if mx else "n/a", ", ".join(map(str, yd)) or "none"))
                   if not bad else "; ".join(bad[:4]), "PASS" if not bad else "FAIL"))
    # C04p: the 2022 pin
    u22 = P[2022].get("unit")
    ex = [("2022 units", str(UNITS_PINS[2022]), str(len(u22.units)) if u22 else "none", "FAIL"), ("2022 every weight 1", "True", str(bool(u22 and u22.all_one)), "FAIL")]
    pin_row("C04p", "Survey year 2022, census (unit file dated 2026-05-14): contributions, benefits, gains (negative), expenses, assets, 4,394 units, every weight 1", [2022], "new", ex)
    # C05, C06
    u17 = P[2017].get("unit")
    ex = [("2017 units", str(UNITS_PINS[2017]), str(len(u17.units)) if u17 else "none", "FAIL"), ("2017 every weight 1", "True", str(bool(u17 and u17.all_one)), "FAIL"),
          ("2017 decodes only as cp1252", "cp1252", u17.enc if u17 else "none", "FAIL")]
    pin_row("C05", "Survey year 2017, census (unit file and summary dated 2025-05-23): contributions, benefits, withdrawals, expenses, assets, 5,529 units, every weight 1, "
            "decodes only as cp1252", [2017], "old", ex)
    pin_row("C06", "Survey year 2021, sample (both dated 2025-05-23): contributions, benefits, expenses, assets", [2021], "old")
    ex = [("%d systems" % y, str(n), str(N[y]["systems"]), "FAIL") for y, n in sorted(SYSTEMS_PINS.items())]
    pin_row("C07", "Old schema (summaries), survey years 2013 and 2016: contributions, benefits, withdrawals, assets; systems published 2013-2016", [2013, 2016], "old", ex)
    # C08: revisions
    cells = Cells()
    st_rev = "FAIL" if fresh else "WARN"
    for (y, item), (pv, nv, ch) in sorted(REV_PINS.items()):
        row = next((r for r in revs if r["survey_year"] == y and r["item"] == item), None)
        cells.add("%d %s prev" % (y, item), fdec(Decimal(pv), 3), fdec(row["value_prev"], 3) if row else "none", st_rev)
        cells.add("%d %s new" % (y, item), fdec(Decimal(nv), 3), fdec(row["value_new"], 3) if row else "none", st_rev)
        cells.add("%d %s change" % (y, item), fdec(Decimal(ch), 3), fdec(row["change"], 3) if row else "none", st_rev)
    for y, pv in sorted(RATIO_REV_PINS.items()):
        cells.add("%d ratio revision pct" % y, pv, fdec(ratio_revs.get(y), 3), st_rev)
    mb, ms = measured_bands(ratio_revs)
    cells.add("measured revision_band_pct", BAND_PIN, fdec(mb, 3), st_rev)
    cells.add("measured revision_band_sample_years_pct", BAND_SAMPLE_PIN, fdec(ms, 3), st_rev)
    if FROZEN_BAND_PCT is not None:
        cells.add("FROZEN_BAND_PCT in the script", BAND_PIN, fdec(FROZEN_BAND_PCT, 3), "FAIL")
        cells.add("FROZEN_SMP_BAND_PCT in the script", BAND_SAMPLE_PIN, fdec(FROZEN_SMP_BAND_PCT, 3), "FAIL")
    obs, st = cells.result()
    out.append(chk("C08", "Revisions, 2025-05-23 summary to 2026-05-14 unit file: 2022 contributions, 2024 benefits, the three ratio revisions and the two bands", "as pinned", obs, st))
    out.append(check_c09(N, P))
    out.append(check_c10(P))
    out.append(check_c12(P))
    out.append(check_c13(N))
    pin_row("C14", "Survey years 2018 and 2019 (unit files, equal to their summaries) and 2020 (summary): contributions, benefits, withdrawals, assets", [2018, 2019, 2020], "old")
    # C15-C17: reference constants typed from other sources; they appear only here
    a25, a24 = N[2025]["assets_end"], N[2024]["assets_end"]
    cells = Cells()
    cells.add("2025 contributions, $bn", fdec(PRESS["contrib_2025_bn"], 2), fdec(N[2025]["contrib_total"] / 1000, 2))
    cells.add("2025 benefits, $bn", fdec(PRESS["benefits_2025_bn"], 2), fdec(N[2025]["benefits"] / 1000, 2))
    cells.add("2025 assets, $ trillion", fdec(PRESS["assets_2025_trn"], 2), fdec(a25 / 1000000, 2))
    cells.add("2024 benefits, $bn", fdec(PRESS["benefits_2024_bn"], 2), fdec(N[2024]["benefits"] / 1000, 2))
    obs, st = cells.result()
    out.append(chk("C15", "Press release CB25-TPS.31 (14 May 2026): survey year 2025 contributions $315.02bn, benefits $418.25bn, assets $6.49 trillion; survey year 2024 benefits $404.46bn",
                   "round to the release", obs, st))
    share = Decimal(100) * NASRA_FY2024_TRN * 1000000 / a24
    out.append(chk("C16", "NASRA Public Fund Survey FY 2024 ($5.13 trillion, relayed, different universe) is 80 to 95 percent of the build's survey year 2024 assets", "between 80% and 95%",
                   "%s%% of the build's survey year 2024 assets" % fdec(share, 1), "PASS" if Decimal(80) <= share <= Decimal(95) else "WARN"))
    dz = Decimal(100) * (Z1_2025Q2_BN * 1000 / a25 - 1)
    out.append(chk("C17", "Fed Z.1 table S129.3.s (funded by assets), 2025:Q2, $6,541.4bn (relayed) is within 3 percent of the build's survey year 2025 assets (a level check only)",
                   "within 3%", ("%s%%" if dz < 0 else "+%s%%") % fdec(dz, 1) + " of the build's survey year 2025 assets", "PASS" if abs(dz) <= 3 else "WARN"))
    out.append(check_c18(N))
    out.append(check_c19(V))
    out.append(check_c20(P))
    return out


# ------------------------------------------------------------------------------------------------------------------------
# compute
# ------------------------------------------------------------------------------------------------------------------------
def source_line(c):
    V = c["V"]
    ret = V.retrieved[:10] if V.retrieved else "(no logged retrieval)"
    return ("# source: U.S. Census Bureau, Annual Survey of Public Pensions, survey years 2013-2025, unit files and summary tables, https://www.census.gov/programs-surveys/aspp.html, "
            "vintage %s (newest unit file %s), retrieved %s" % (V.key, V.newest_unit_file_date, ret))


PARSE_MUTANT_PREFIXES = ("g1_", "g2_unweighted", "g2_modal", "g2_weight_twice", "g4_", "g5_abs", "g5_losses", "g8_float")


def parse_is_clean():
    return not any(m.startswith(PARSE_MUTANT_PREFIXES) for m in _ACTIVE)


def compute(inp, frozen=None, frozen_sample=None, existing_revs=None):
    vints = sorted(inp["vints"], key=lambda v: v.order())
    if not vints:
        raise DataError("no complete vintage in the cache: run fetch first")
    V = vints[-1]
    cache = inp.get("_P")
    if cache is not None and V.key in cache and parse_is_clean():
        P = cache[V.key]
    else:
        P = parse_vintage(V)
        if cache is not None and parse_is_clean():
            cache[V.key] = P
    if M("g15_stop_on_content") and any(o.problems for y in SURVEY_YEARS for o in (P[y].get("unit"), P[y].get("sum")) if o is not None):
        create_stop(os.path.join(inp.get("stop_dir", STOP_DIR), STOP_NAME), stop_fields("content", "", "", "", ""))
    N = national_rows(P, V.key)
    R = ratio_rows(N)
    ra, ratio_revs = revisions_a(P)
    rb, s1 = [], None
    prevN = None
    for prev, new in zip(vints, vints[1:]):
        Np = national_rows(parse_vintage(prev), prev.key)
        Nn = N if new is V else national_rows(parse_vintage(new), new.key)
        rb += revisions_b(Np, Nn)
        if new is V:
            prevN = Np
    revs = sort_revisions(ra + rb)
    if M("g10_append") and existing_revs:
        revs = existing_revs + revs
    band, band_s = band_values(ratio_revs, frozen, frozen_sample)
    if prevN is not None and any(y >= 2026 for y in V.notice_years) and prevN[2025]["benefits"] and N[2025]["benefits"]:
        rp = prevN[2025]["benefits"] / prevN[2025]["contrib_total"]
        s1 = Decimal(100) * (N[2025]["benefits"] / N[2025]["contrib_total"] / rp - 1)
    c = {"inp": inp, "vints": vints, "V": V, "P": P, "N": N, "R": R, "revs": revs, "ratio_revs": ratio_revs, "band": band, "band_s": band_s, "s1": s1,
         "fetch_rows": inp.get("fetch_rows", []), "sha_pin": inp.get("sha_pin", SHA_2025)}
    c["summary"] = summary_rows(N, R, ratio_revs, revs, band, band_s, V.key, s1)
    c["checks"] = build_checks(c)
    c["names"] = set().union(*[P[y]["unit"].names for y in SURVEY_YEARS if P[y].get("unit") is not None])
    return c


# ------------------------------------------------------------------------------------------------------------------------
# rendering the csv files and the folder guards
# ------------------------------------------------------------------------------------------------------------------------
def allowlist(name):
    cols = list(ALLOW[name])
    if M("g12_allow_drift") and name == "ratios.csv":
        cols[cols.index("ratio_cv_bound_pct")] = "ratio_cv_bound"
    return cols


def allow_text():
    return "".join("%s:%s\n" % (n, ",".join(allowlist(n))) for n in ALLOW_ORDER)


def check_header(name, header):
    allow = allowlist(name)
    if M("g12_extra_col_ok"):
        ok = list(header[:len(allow)]) == allow
    elif M("g12_renamed_ok"):
        ok = len(header) == len(allow)
    else:
        ok = list(header) == allow
    if not ok:
        raise DataError("%s: header %r is not the exact column allowlist" % (name, list(header)))


def csv_body(text):
    """(header, rows) of a committed csv text, skipping every `# source:` line (mutant: the source line is read as the header)."""
    lines = text.splitlines()
    if not M("g12_source_as_header"):
        lines = [l for l in lines if not l.startswith("# source:")]
    rd = list(csv.reader(lines))
    return (rd[0], rd[1:]) if rd else ([], [])


def strip_exempt(text, readme=True):
    """Remove the one denial sentence (README only), once."""
    if not readme:
        return text
    if M("g14_denial_all"):
        return re.sub(r"[^.]*\b(?:sale|sold|bought|purchase)\b[^.]*\.", "", text)
    if M("g14_denial_loose"):
        return text.replace(DENIAL_SENTENCE, "")
    return text.replace(DENIAL_SENTENCE, "", 1)


def wording_problems(name, text):
    t = strip_exempt(text, readme=name == "README.md")
    out = []
    if not M("g14_causal_off"):
        out += ["causal wording %r" % re.sub(r"\s+", " ", m.group(0).lower()) for m in CAUSAL_RX.finditer(t)]
    if not M("g14_extra_off"):
        out += ["flow, sale or motive wording %r" % re.sub(r"\s+", " ", m.group(0).lower()) for m in EXTRA_RX.finditer(t)]
    if name == "README.md" and not M("g11_rose_fell_ok"):
        out += ["direction verb %r" % m.group(0).lower() for m in ROSE_FELL_RX.finditer(t)]
    return out


def guard_text(name, text):
    probs = wording_problems(name, text) if name != "SOURCES.md" else []
    if probs:
        raise DataError("%s: %s" % (name, "; ".join(sorted(set(probs))[:4])))
    home = os.path.expanduser("~")
    if home != "~" and home in text:
        raise DataError("%s: an absolute home path would be published" % name)


def render_csv(name, rows, source, header=None):
    header = allowlist(name) if header is None else header
    check_header(name, header)
    out = io.StringIO()
    out.write(source + "\n")
    w = csv.writer(out, lineterminator="\n")
    w.writerow(header)
    for r in rows:
        if len(r) != len(header):
            raise DataError("%s: a row has %d cells, the header %d" % (name, len(r), len(header)))
        w.writerow(r)
    text = out.getvalue()
    if not text.startswith("# source:"):
        raise DataError("%s: the first line is not a `# source:` line" % name)
    guard_text(name, text)
    return text


def money(v):
    if v is None:
        return ""
    if M("g6_bn_in_mn"):
        v = v / 1000
    return fdec(v, 2 if M("g8_wrong_decimals") else 3)


def national_csv_rows(N):
    out = []
    for y in SURVEY_YEARS:
        n = N[y]
        out.append([y, n["schema"], n["year_kind"], money(n["contrib_employee"]), money(n["contrib_government"]), money(n["contrib_total"]), money(n["benefits"]),
                    money(n["withdrawals"]), money(n["net_cash"]), money(n["earnings_ex_gains"]), money(n["earnings_old"]), money(n["gains_losses"]), money(n["expenses"]),
                    money(n["assets_end"]), "" if n["systems"] is None else n["systems"], fdec(n["cv_contrib_total"], 7), fdec(n["cv_benefits"], 7), fdec(n["cv_assets"], 7),
                    n["summary_file_date"], n["source_kind"], n["source_file_date"], n["vintage"]])
    return out


def ratio_csv_rows(R, vkey):
    out = []
    for y in SURVEY_YEARS:
        r = R[y]
        dn = 3 if M("g8_wrong_decimals") else 4
        out.append([y, r["schema"], fdec(r["bpc"], dn), fdec(r["bwpc"], dn), fdec(r["start"], 3), fdec(r["end"], 3), fdec(r["income_gap"], dn), fdec(r["income_gap_exp"], dn),
                    fdec(r["cv_bound"], 3), vkey])
    return out


def summary_csv_rows(S):
    return [[r["question"], r["segment"], "" if r["from_year"] is None else r["from_year"], "" if r["to_year"] is None else r["to_year"], r["statistic"],
             fmt_summary_value(r), r["unit"], r["pattern"], r["primary"], r["vintage"]] for r in S]


def revision_csv_rows(revs):
    return [[r["survey_year"], r["item"], r["source_prev"], r["date_prev"], money(r["value_prev"]), r["source_new"], r["date_new"], money(r["value_new"]), money(r["change"]),
             fdec(r["change_pct"], 3)] for r in revs]


def vintage_csv_rows(vints):
    return [[v.key, v.newest_unit_file, v.newest_unit_file_date, N_FILES, v.requests, v.manifest_sha] for v in vints]


def render_data_texts(c):
    src = source_line(c)
    t = {}
    t["national.csv"] = render_csv("national.csv", national_csv_rows(c["N"]), src)
    t["ratios.csv"] = render_csv("ratios.csv", ratio_csv_rows(c["R"], c["V"].key), src)
    t["summary.csv"] = render_csv("summary.csv", summary_csv_rows(c["summary"]), src)
    t["revisions.csv"] = render_csv("revisions.csv", revision_csv_rows(c["revs"]), src + "; regenerated on every build from the cache, never appended")
    t["vintages.csv"] = render_csv("vintages.csv", vintage_csv_rows(c["vints"]), src)
    rows = c["fetch_rows"]
    t["fetch_log.csv"] = render_csv("fetch_log.csv", rows, fetch_source_line(rows))
    return t


def leak_problems(files, names=()):
    """files: {name: text or bytes}. Section 6: only the nine committed names; no raw file, no pasted unit-file row, no system name, no local path, no at sign, no SAMPLE_YEAR."""
    out = []
    for n in sorted(files):
        raw = files[n] if isinstance(files[n], bytes) else files[n].encode("utf-8")
        text = raw.decode("utf-8", errors="replace")
        if (n.lower().endswith(RAW_EXT) and n not in COMMITTED) or raw[:2] == b"PK" or raw[:4] == b"%PDF":
            if not M("g13_raw_ok"):
                out.append("raw file in the committed folder: %s" % n)
            continue
        if n not in COMMITTED and n != ".DS_Store" and not M("g12_eighth_csv_ok"):
            out.append("file that is not one of the nine committed names: %s" % n)
        rx = RAW_ROW_RX_NAIVE if M("g13_rowpat_anchored") else RAW_ROW_RX
        if not M("g13_rowpat_off") and rx.search(text):
            out.append("a raw unit-file row in %s" % n)
        if not M("g13_samplehdr_off") and "SAMPLE_YEAR" in text:
            out.append("a unit-file header line in %s" % n)
        if not M("g13_names_off"):
            low = text.lower()
            for nm in names:
                if nm.lower() in low:
                    out.append("a system name in %s" % n)
                    break
        if not M("g13_path_ok") and any(p in text for p in ("/Users/", "/private/", "~")):
            out.append("a local path in %s" % n)
        if "@" in text:
            out.append("an at sign in %s" % n)
    for n in sorted(set(COMMITTED) - set(files)):
        out.append("missing: %s" % n)
    return out


def header_problems(files):
    out = []
    for n in ALLOW_ORDER:
        if n in files:
            try:
                check_header(n, csv_body(files[n])[0])
            except DataError as e:
                out.append(str(e))
    return out


def dup_key_problems(files):
    out = []
    keys = {"national.csv": ("survey_year",), "ratios.csv": ("survey_year",), "summary.csv": SUMMARY_KEY, "revisions.csv": ("survey_year", "item", "date_prev", "date_new"),
            "checks.csv": ("check_id",), "vintages.csv": ("vintage",), "fetch_log.csv": ("utc_time", "path")}
    for n, ks in keys.items():
        if n not in files:
            continue
        h, rows = csv_body(files[n])
        if list(h) != ALLOW[n]:
            continue
        idx = [h.index(k) for k in ks]
        seen = set()
        for r in rows:
            k = tuple(r[i] for i in idx)
            if k in seen and not M("g12_dup_key_ok"):
                out.append("duplicate key in %s: %s" % (n, "|".join(k)))
            seen.add(k)
    return out


def tamper_problems(files):
    """The C13 identities on the committed national.csv (M25 holds its own copy)."""
    out = []
    if "national.csv" not in files:
        return out
    h, rows = csv_body(files["national.csv"])
    if list(h) != ALLOW["national.csv"]:
        return out
    for r in rows:
        d = dict(zip(h, r))
        try:
            e, g, t = Decimal(d["contrib_employee_usd_mn"]), Decimal(d["contrib_government_usd_mn"]), Decimal(d["contrib_total_usd_mn"])
            b, w, nc = Decimal(d["benefits_usd_mn"]), Decimal(d["withdrawals_usd_mn"]), Decimal(d["net_cash_usd_mn"])
        except decimal.InvalidOperation:
            out.append("survey year %s: a money cell is not a number" % d["survey_year"])
            continue
        if abs(e + g - t) > Decimal("0.002"):
            out.append("survey year %s: employee + government is not total contributions" % d["survey_year"])
        if abs(nc - (t - b - w)) > Decimal("0.002"):
            out.append("survey year %s: net cash is not contributions minus benefits minus withdrawals" % d["survey_year"])
    return out


def pairing_problems(files, extra=()):
    """R10: the reference constants of C15-C17 appear in checks.csv and nowhere else among the committed files."""
    out = []
    needles = [r"(?<![\d.])5\.13(?!\d)", r"(?<![\d.])6,?541\.4(?!\d)"] + list(extra)
    for n, t in files.items():
        if n in ("checks.csv", "SOURCES.md") or not isinstance(t, str):
            continue
        if M("g17_pairing_ok"):
            continue
        for nd in needles:
            if re.search(nd, t):
                out.append("%s holds a figure from another source (%s)" % (n, nd))
    return out


def all_folder_problems(files, names=()):
    out = leak_problems(files, names) + header_problems(files) + dup_key_problems(files) + tamper_problems(files) + pairing_problems(files)
    if isinstance(files.get("summary.csv"), str):
        h, rows = csv_body(files["summary.csv"])
        if list(h) == ALLOW["summary.csv"]:
            out += summary_problems(rows)
    for n, t in files.items():
        if n.endswith((".csv", ".md")) and isinstance(t, str) and n != "SOURCES.md":
            out += ["%s: %s" % (n, p) for p in sorted(set(wording_problems(n, t)))]
    return out


# ------------------------------------------------------------------------------------------------------------------------
# README.md and SOURCES.md (generated, so the wording guards scan the README)
# ------------------------------------------------------------------------------------------------------------------------
def md_table(head, rows):
    out = ["| " + " | ".join(head) + " |", "|" + "|".join("---" for _ in head) + "|"]
    out += ["| " + " | ".join(str(x) for x in r) + " |" for r in rows]
    return "\n".join(out)


def bn(v, nd=1):
    if v is None:
        return "n/a"
    s = fdec(Decimal(v) / 1000, nd)
    ip, _, fp = s.partition(".")
    neg = ip.startswith("-")
    ip = "{:,}".format(int(ip.lstrip("-")))
    return ("-" if neg else "") + ip + ("." + fp if fp else "")


def pctx(v, nd=3, sign=False):
    if v is None:
        return "n/a"
    s = fdec(v, nd)
    return ("+" + s) if (sign and not s.startswith("-")) else s


def srow(c, stat, a=None, b=None, q=None):
    for r in c["summary"]:
        if r["statistic"] == stat and r["to_year"] == b and r["from_year"] == a and (q is None or r["question"] == q):
            return r
    return None


PATTERN_GLOSS = {
    "within_band": "smaller than the largest past revision of this ratio (%s%%); that is not the same as no change",
    "above_band": "an increase larger than the largest past revision of this ratio (%s%%)",
    "below_band": "a decrease larger than the largest past revision of this ratio (%s%%)",
}


def render_readme(c):
    N, R, V = c["N"], c["R"], c["V"]
    band, band_s = c["band"], c["band_s"]
    prim = srow(c, "change_benefits_per_contrib_pct", 2023, 2025, "Q2")
    a23, a25 = R[2023]["bpc"], R[2025]["bpc"]
    gloss = PATTERN_GLOSS[prim["pattern"]] % fdec(band, 3)
    rr = c["ratio_revs"]
    rr_txt = ", ".join("%s%% for survey year %d" % (pctx(rr[y], 3, True), y) for y in sorted(rr))
    gap25 = N[2025]["benefits"] - N[2025]["contrib_total"]
    gains25 = N[2025]["gains_losses"]
    L = []
    L.append("# US state and local government pension systems: money in, money out and holdings, survey years 2013 to 2025\n")
    L.append("This table shows how much US state and local government pension systems took in and paid out in each survey year from 2013 to 2025, and the cash and investments "
             "they reported holding, in the Census Bureau's Annual Survey of Public Pensions. "
             "Its one primary row compares survey year 2023 with survey year 2025: benefit payments per dollar of contributions went from %s to %s, a change of %s%%, "
             "labelled `%s`, which means %s.\n" % (fdec(a23, 4), fdec(a25, 4), pctx(prim["value"], 3, True), prim["pattern"], gloss))
    L.append(DENIAL_SENTENCE + " These systems promise retirees a set benefit. Money in is what employees and governments paid in; money out is benefits to retirees and refunds "
             "to members who left. Beside the table are the cash and investments the systems reported holding. The figures are the Census Bureau's national estimates from its Annual "
             "Survey of Public Pensions. Survey year 2025 covers the fiscal years that most systems ended on 30 June 2025. Census does not say on what date holdings are measured; "
             "we take it to be the end of each system's fiscal year.\n")
    L.append("Survey year 2025 is a first release (a release is one posting of the survey's files by the Census Bureau), and the next release is expected to revise it and the years back to 2022.\n")
    L.append("## Timing\n")
    L.append("A survey year covers the fiscal year that begins on 1 July of the previous year and ends on 30 June of the survey year. Systems whose fiscal year ends after 30 June "
             "report their previous fiscal year. The exceptions are state systems in Alabama and Michigan, which report October to September, and in Texas, which reports "
             "September to August. Survey year 2025 is therefore mostly July 2024 to June 2025. Census's publication, file layouts and methodology do not say on what date "
             "holdings are measured, so the column `assets_end_usd_mn` carries an assumption in its name: the end of each system's fiscal year. These are survey years, "
             "not calendar years, and nothing here is set beside a calendar-year or quarterly series.\n")
    L.append("## What is counted\n")
    L.append("Nine lines, nationally, in millions of US dollars: employee contributions; government contributions; total contributions (the two together); benefit payments; "
             "withdrawals (refunds to members who leave); investment earnings excluding gains and losses; gains and losses on investments (signed); administrative and "
             "investment expenses; and cash and investments held. Net cash flow is contributions minus benefits minus withdrawals. It is neither the press release's gap of "
             "benefits minus contributions ($%sbn in survey year 2025, which leaves out withdrawals) nor NASRA's external cash flow, which is a different measure on a different "
             "set of plans. Investment earnings here exclude gains. Administrative and investment expenses are their own line and are not part of net cash flow.\n" % bn(gap25, 2))
    L.append("## Full-count years, sample years and the old years\n")
    L.append("In a *full-count year* (a census year: survey years 2017 and 2022) Census collects every system and every weight is 1. In a *sample year* it collects a sample and "
             "weights each row so that the weighted sum estimates the national total, so the figure carries sampling error. The error is reported as a *CV*: the coefficient "
             "of variation, which is the sampling error as a share of the estimate. Census publishes CVs in its sample-year summary tables (a summary table is Census's workbook of national and state totals for one survey year); no 2025 CV file is posted on "
             "www2.census.gov, so survey year 2025 has no CV here. System counts are the published counts from each year's summary table (a unit file is Census's row-by-row file, one row per system and item; its row count is not a "
             "system count); survey year 2025 has no summary table, and its publication counts 311 state systems and 4,061 local systems.\n")
    L.append("Survey years 2013 to 2016 are a separate, labelled segment (`old`) built from Census's summary tables; survey years 2017 to 2025 "
             "are the `current` segment (2020 comes from its summary table only). The 2015 brief says that more than 2,200 local systems were added in 2015, that the data may include hybrid and other post-employment "
             "funds, and that 2012 and 2015 should be compared only with caution. In the old years \"earnings on investments\" include gains and losses, so they sit in their own "
             "column. From 2017 the schema, the earnings definition and the asset detail change. For both reasons no change is computed within the old years or between 2016 "
             "and 2017; old years appear as single-year rows only.\n")
    L.append("## Beside the project's other retirement measures\n")
    L.append("The project's figures for private defined-contribution plans stop at 2023 and cover different plans on different dates; nothing in this table is added to, subtracted "
             "from or compared in one number with them. The Federal Reserve's Z.1 accounts record state and local defined-benefit plans' equity transactions in calendar "
             "quarters; this table is a cash account on fiscal years. The two are never set side by side.\n")
    L.append("## Levels\n")
    L.append("Millions of dollars are shown here as billions of dollars (`$bn`) to one decimal; the CSV files hold millions to three decimals.\n")
    cur = [y for y in SURVEY_YEARS if y >= 2017]
    L.append("**Current segment, `$bn`**\n")
    L.append(md_table(["Survey year", "Kind", "Contributions", "Benefits", "Withdrawals", "Net cash flow", "Earnings excluding gains", "Gains and losses", "Expenses", "Cash and investments", "Systems published"],
                      [[y, N[y]["year_kind"], bn(N[y]["contrib_total"]), bn(N[y]["benefits"]), bn(N[y]["withdrawals"]), bn(N[y]["net_cash"]), bn(N[y]["earnings_ex_gains"]),
                        bn(N[y]["gains_losses"]), bn(N[y]["expenses"]), bn(N[y]["assets_end"]), "{:,}".format(N[y]["systems"]) if N[y]["systems"] else "not in a summary table"] for y in cur]) + "\n")
    L.append("**Old segment, `$bn` (earnings include gains and losses)**\n")
    L.append(md_table(["Survey year", "Contributions", "Benefits", "Withdrawals", "Net cash flow", "Earnings including gains", "Cash and investments", "Systems published"],
                      [[y, bn(N[y]["contrib_total"]), bn(N[y]["benefits"]), bn(N[y]["withdrawals"]), bn(N[y]["net_cash"]), bn(N[y]["earnings_old"]), bn(N[y]["assets_end"]),
                        "{:,}".format(N[y]["systems"])] for y in OLD_YEARS]) + "\n")
    L.append("## Ratios and the primary row\n")
    L.append("**The primary row** (the only one fixed in advance as primary): the change in benefit payments per dollar of contributions, segment `current`, from survey year 2023 to "
             "survey year 2025: %s to %s, %s%%, `%s`. Every other row below is exploratory. `within_band` means smaller than one past revision. It does not mean no change.\n"
             % (fdec(a23, 4), fdec(a25, 4), pctx(prim["value"], 3, True), prim["pattern"]))
    mx_year = max(rr, key=lambda y: abs(rr[y])) if rr else None
    if mx_year is None:
        origin = "No revision of the ratio could be measured in this vintage."
    elif mx_year in CENSUS_YEARS:
        origin = "It is the largest of %d past revisions of the ratio, set by one full-count-year re-post (survey year %d, %s%%)." % (len(rr), mx_year, pctx(rr[mx_year], 3, True))
    else:
        origin = "It is the largest of %d past revisions of the ratio, set by a sample-year revision (survey year %d, %s%%)." % (len(rr), mx_year, pctx(rr[mx_year], 3, True))
    samp = ", ".join("%s%% (survey year %d)" % (pctx(rr[y], 3, True), y) for y in sorted(rr) if y not in CENSUS_YEARS)
    L.append("**The band.** The band is %s%%. %s The sample-year revisions were %s, and the largest of them is %s%%. The band was fixed after the first build. The first build had already "
             "shown the survey year 2023 to 2025 change (about +1.1%% on the press-release figures), so both the band and that change were seen before the band was frozen. A label is "
             "a description set against one past revision. It is not a test.\n" % (fdec(band, 3), origin, samp or "none", fdec(band_s, 3)))
    chg = [r for r in c["summary"] if r["statistic"] in ("change_benefits_per_contrib_usd", "change_benefits_per_contrib_pct") and r["question"] == "Q2"]
    spans = sorted({(r["from_year"], r["to_year"]) for r in chg})
    rows = []
    for a, b in spans:
        u = srow(c, "change_benefits_per_contrib_usd", a, b, "Q2")
        p = srow(c, "change_benefits_per_contrib_pct", a, b, "Q2")
        rows.append(["survey year %d to survey year %d" % (a, b), fdec(R[a]["bpc"], 4), fdec(R[b]["bpc"], 4), pctx(u["value"], 4, True), pctx(p["value"], 3, True), p["pattern"],
                     "primary" if p["primary"] == "yes" else "exploratory"])
    L.append(md_table(["Span", "Benefits per dollar of contributions, first year", "second year", "Change in dollars", "Change in percent", "Label", "Status"], rows) + "\n")
    L.append("**Ratios by survey year.** Benefit payments per dollar of contributions is benefits divided by total contributions. Net cash flow as a share of opening assets is 100 "
             "times net cash flow divided by the previous survey year's cash and investments in the same release and the same segment; it is blank for 2013, 2015 (its opening "
             "assets cross the 2014-to-2015 jump in the number of systems) and 2017 (the segment changes). It is the main measure. Net cash flow as a share of end-of-year assets "
             "divides by the same survey year's cash and investments. End-of-year assets include that year's investment gains ($%sbn in survey year 2025), so net cash flow is a "
             "smaller share of end-of-year assets in years with large gains; the share of opening assets is the main measure for that reason.\n" % bn(gains25))
    L.append(md_table(["Survey year", "Segment", "Benefits per dollar of contributions", "Benefits and withdrawals per dollar of contributions", "Net cash flow, % of opening assets",
                       "Net cash flow, % of end-of-year assets", "Income per dollar of gap", "Income per dollar of gap, after expenses", "Sampling bound on the first ratio, %"],
                      [[y, R[y]["schema"], fdec(R[y]["bpc"], 4), fdec(R[y]["bwpc"], 4), fdec(R[y]["start"], 3) or "blank", fdec(R[y]["end"], 3), fdec(R[y]["income_gap"], 4) or "not defined",
                        fdec(R[y]["income_gap_exp"], 4) or "not defined", fdec(R[y]["cv_bound"], 3) or "not posted"] for y in SURVEY_YEARS]) + "\n")
    g23, g25 = R[2023], R[2025]
    L.append("**Investment income against the gap** (current segment only). The gap is the amount by which benefits plus withdrawals exceeded contributions. The first income ratio is "
             "interest, dividends and other investment income per dollar by which payouts exceeded contributions; it was %s in survey year 2023 and %s in survey year 2025. "
             "The second adds administrative and investment expenses to the payouts: %s and %s. The table does not show which money paid which benefit.\n"
             % (fdec(g23["income_gap"], 4), fdec(g25["income_gap"], 4), fdec(g23["income_gap_exp"], 4), fdec(g25["income_gap_exp"], 4)))
    L.append("**Sampling bound.** The last column is 100 times the sum of the published CVs of benefits and of contributions, a ceiling on the sampling CV of the first ratio whatever "
             "their correlation. It is %s%% for survey year 2023 and %s%% for survey year 2024, and 0 in full-count years. It is blank for survey year 2025, because no 2025 CV "
             "file is posted on www2.census.gov. Revisions have been larger than this sampling error.\n" % (fdec(R[2023]["cv_bound"], 3), fdec(R[2024]["cv_bound"], 3)))
    L.append("## Revisions and vintages\n")
    ud = collections.defaultdict(list)
    sd = collections.defaultdict(list)
    for y in SURVEY_YEARS:
        if c["P"][y].get("unit") is not None:
            ud[c["P"][y]["unit_date"]].append(y)
        if c["P"][y].get("sum") is not None:
            sd[c["P"][y]["sum_date"]].append(y)
    ud_txt = "; ".join("%s: %s" % (d, ", ".join(map(str, ys))) for d, ys in sorted(ud.items(), reverse=True))
    sd_txt = "; ".join("%s: %s" % (d, ", ".join(map(str, ys))) for d, ys in sorted(sd.items(), reverse=True))
    L.append("Each Census release revises earlier years. A vintage is the set of files fetched together, identified by their posting dates. The unit files in this vintage carry these posting dates (survey years after the colon): %s. The summary tables: %s. "
             "Where a summary table and a unit file carry different dates, the ratio of benefits to contributions changed by %s. `revisions.csv` is regenerated on every build from each year's "
             "summary table against its unit file where their dates differ, and from every adjacent pair of complete cached vintages. This file set is vintage `%s`; its newest unit file is "
             "%s, dated %s. The next release (survey year 2026) is not announced; May 2027 is expected on this pattern. One statement is registered in advance, and it is about revisions "
             "only: in the first release carrying survey year 2026, the re-posted survey year 2025 benefits per dollar of contributions differs from this build's value by less than %s%%. "
             "Nothing is registered about the cash flows themselves.\n" % (ud_txt, sd_txt, rr_txt or "n/a", V.key, V.newest_unit_file, V.newest_unit_file_date, fdec(band, 3)))
    L.append("## Column dictionary\n")
    L.append("Every file begins with a `# source:` line. `segment` and `schema` are `old` (survey years 2013 to 2016) or `current` (2017 to 2025). `year_kind` is `census` (full-count "
             "year) or `sample`. `pp` is a percentage point. `source_kind` is `unit_file` or `summary`. Money columns are millions of US dollars to three decimals.\n")
    L.append("- `national.csv`: one row per survey year, 2013 to 2025. `contrib_employee_usd_mn`, `contrib_government_usd_mn`, `contrib_total_usd_mn`, `benefits_usd_mn`, "
             "`withdrawals_usd_mn`, `earnings_ex_gains_usd_mn` (current segment), `earnings_old_schema_usd_mn` (old segment, includes gains), `gains_losses_usd_mn` (signed), "
             "`expenses_admin_invest_usd_mn`, and `assets_end_usd_mn` (cash and investments; an assumption that holdings are measured at the end of each system's fiscal year). "
             "`net_cash_usd_mn` = contributions minus benefits minus withdrawals; it is negative when payouts exceed contributions. `systems_published` is the published count. "
             "`cv_*` are published CVs as fractions, blank in full-count years and for 2025. `summary_file_date` and `source_file_date` are the posting dates of the summary "
             "table and of the file the row was built from.\n")
    L.append("- `ratios.csv`: `benefits_per_contrib_usd` = benefits / total contributions; `benefits_withdrawals_per_contrib_usd` = (benefits + withdrawals) / total contributions; "
             "`net_cash_pct_of_assets_start` = 100 x net cash / the previous survey year's `assets_end` (same release, same segment); `net_cash_pct_of_assets_end` = 100 x net cash / "
             "the same survey year's `assets_end`; `income_per_net_cash_gap_usd` = earnings excluding gains / (benefits + withdrawals - contributions), blank if that denominator is "
             "not positive; `income_per_gap_after_expenses_usd` = earnings excluding gains / (benefits + withdrawals + expenses - contributions); `ratio_cv_bound_pct` = 100 x (CV of "
             "benefits + CV of contributions).\n")
    L.append("- `summary.csv`: one row per statistic, keyed by (`question`, `segment`, `from_year`, `to_year`, `statistic`); `pattern` is `above_band`, `below_band`, `within_band` or `n_a`, "
             "and only the percent change of the ratio carries one; `primary` marks the one primary row. `revisions.csv`: the value in the earlier and the later source, the change in "
             "millions and in percent (against the size of the earlier value). `vintages.csv`, `checks.csv` and `fetch_log.csv` record the cache, the checks and the requests.\n")
    L.append("## Citation\n")
    L.append("U.S. Census Bureau, Annual Survey of Public Pensions, survey years 2013-2025, unit files and summary tables, https://www.census.gov/programs-surveys/aspp.html, vintage %s "
             "(newest unit file %s). Census asks users who make their own estimates to cite the Census Bureau as the source of the original data only. The raw files stay in the "
             "gitignored cache, and only national totals are committed. How the numbers were fetched, the checks and the builder's decisions are in [SOURCES.md](SOURCES.md).\n"
             % (V.key, V.newest_unit_file_date))
    text = "\n".join(L)
    if M("x_drop_gloss_open"):
        text = text.replace("; that is not the same as no change", "")
    if M("x_drop_gloss_dnm"):
        text = text.replace(" It does not mean no change.", "")
    return text


SPEC_AUTHOR_REQUESTS = """The spec author's 56 requests of 26 Oct 2026 (05:43 to 05:54 UTC, about 2 seconds apart, neutral User-Agent, no redirect followed, no retry), before any build:
`www.census.gov` (13): the three scoping URLs; `/programs-surveys/aspp/data/tables.html`, `/data/api.html`, `/data/datasets.html`, `/data/historical_data.html`;
`/programs-surveys/aspp/technical-documentation.html`; `/data/developers/data-sets/govspension.html`; `/data/tables/2013/econ/aspp/aspp-historical-tables.html`;
`/about/policies/citation.html`, all 200; `/about/policies/open-gov/open-data.html` **301, not followed**; `/robots.txt` **404**. `www2.census.gov` (43), all 200: the two layout PDFs,
the 2025 publication PDF, the 2024 methodology PDF, `/robots.txt`; directory listings of `datasets/`, `datasets/2013/`, `2016/`, `2017/`, `2021/`, `2022/`, `2024/`, `2025/`, the
`aspp-historical-datasets/` subfolders of 2013, 2016, 2017 (twice: an inadvertent repeat), 2018, 2019, 2020, 2021, 2022, 2023, 2024, 2025, and the file-layouts folder and its
`2017-2024/`; files: unit files 2013 (`.txt`), 2017, 2021, 2022, 2023 (`.dsv`), 2024, 2025 (`.csv`); summaries 2013, 2014, 2015, 2016, 2017, 2021, 2022, 2023, 2024. No request to
`api.census.gov` or `data.census.gov`. The spec author's copy was a temporary selftest-fixture seed, not a vintage."""

DISCLOSURE = """Before any build, the spec author summed the national components for 2017, 2021 and 2022-2025 and read the summaries for 2013-2017 and 2021-2024. The second reviewer and Grok then
computed the ratios themselves: benefits per dollar of contributions about 1.313 (2023) and 1.328 (2025), a change of about +1.1%, and the three ratio revisions that set the band.
Nothing in the spec is blind, and no output carries a `blind` column. The builder did not look at any headline value before the selftest passed in full."""


def decisions_text(c):
    """The Builder decisions section: fixed text plus what this build measured (encodings, the frozen band)."""
    P = c["P"]
    enc = ", ".join("%d %s" % (y, P[y]["unit"].enc) for y in SURVEY_YEARS if P[y].get("unit") is not None)
    mb, ms = measured_bands(c["ratio_revs"])
    L = []
    L.append("1. **Encodings per unit file (R8):** %s." % enc)
    L.append("2. **FROZEN_BAND_PCT (R2-3):** the first build measured the largest absolute revision of `benefits_per_contrib_usd` between a year's summary and its unit file as %s%% "
             "(`revision_band_pct`), and the largest over sample years as %s%% (`revision_band_sample_years_pct`). Both are hard-coded in the script as `FROZEN_BAND_PCT = Decimal(\"%s\")` and "
             "`FROZEN_SMP_BAND_PCT = Decimal(\"%s\")`, and every later build uses those constants. In this build the constants in the script are %s and %s."
             % (fdec(mb, 3), fdec(ms, 3), fdec(FROZEN_BAND_PCT, 3) if FROZEN_BAND_PCT is not None else "(not yet set)",
                fdec(FROZEN_SMP_BAND_PCT, 3) if FROZEN_SMP_BAND_PCT is not None else "(not yet set)",
                fdec(FROZEN_BAND_PCT, 3) if FROZEN_BAND_PCT is not None else "not set", fdec(FROZEN_SMP_BAND_PCT, 3) if FROZEN_SMP_BAND_PCT is not None else "not set"))
    L.append("3. **New item codes:** none outside the 53 listed in R3b were found (C20).")
    L += DECISIONS_FIXED
    k = len(L) + 1
    u23 = P[2023].get("unit")
    if u23 is not None:
        d60 = unit_to_mn(u23.sums.get("RX61", Decimal(0)) + u23.sums.get("RX62", Decimal(0)) - u23.sums.get("RX60", Decimal(0)))
        d70 = unit_to_mn(sum((u23.sums.get(x, Decimal(0)) for x in ("RX71", "RX72", "RX73")), Decimal(0)) - u23.sums.get("RX70", Decimal(0)))
        L.append("%d. **Spec prose against its own figures (C12).** The spec says that in 2023 RX60 exceeds RX61 + RX62 by 560.031 $mn, and lists the two differences as -560.031 and -12.930. This build finds, "
                 "in survey year 2023, (RX61 + RX62) - RX60 = %s and (RX71 + RX72 + RX73) - RX70 = %s: the magnitudes agree and the sub-lines are the larger side, so the spec's prose direction looks "
                 "reversed; the published total line is used either way. The counts (%d units with a second weight, %d non-blank rows carrying it) are read with the definition in item 11." %
                 (k, fdec(d60, 3), fdec(d70, 3), u23.units_second, u23.rows_second_nonblank))
        k += 1
    cvs = [(n[x], y) for y, n in c["N"].items() for x in ("cv_contrib_total", "cv_benefits", "cv_assets") if n[x] is not None]
    if cvs:
        mx = max(cvs)
        L.append("%d. **C18 reads three lines only.** The largest CV on contributions, benefits and assets is %s (survey year %d). The spec's 0.0334 (2013) is the largest on any line of the summaries, "
                 "which this build does not read." % (k, mx[0], mx[1]))
        k += 1
    L.append("%d. **Listing reply.** The directory listing's `Last-Modified` is the time of the request itself, so its logged date is the retrieval time and is not part of any vintage key." % k)
    L += DECISIONS_POST
    return "\n".join(L)


DECISIONS_FIXED = [
    "4. **The directory listing is a web page (R3a).** R3a latches STOP on any body that starts with `<!` or `<html`, but the one listing request (`/programs-surveys/aspp/datasets/`, an Apache index, "
    "observed to start with a DOCTYPE) is a web page by construction. The listing therefore gets its own content test: a non-200 status or an empty body latches STOP as for every request, and so does a "
    "page whose title is not `Index of /programs-surveys/aspp/datasets` (a refusal page). The 20 file requests use the R3a test as written. The listing is requested first, so a refusal stops the run before "
    "any file is downloaded.",
    "5. **Body over the cap, and slices.** A body above 100 MB is not a refusal; it ends the run with exit 3 and no STOP. A run may take several invocations (`fetch --max-files N`): the folder `partial/` "
    "keeps `manifest.csv`, `requests.txt` (the run's request count, cap 30 across the invocations) and `notice.txt`; a saved file is never requested again; at most one run per calendar month is enforced "
    "from `fetch_log.csv`.",
    "6. **O-10 (round 4), the 2025 system count.** The ruling asks for the count from the 2025 publication with `source_kind = publication`. The publication is not among the 20 files of the table and "
    "`source_kind` allows only `unit_file` and `summary`; taking it would need a request outside the table and a changed allowlist. The stricter reading is used: `systems_published` is blank for survey "
    "year 2025 and the README quotes the publication's 311 + 4,061 as text. Nothing is estimated.",
    "7. **O-2 against R3-3.** O-2 says the README states that the old and current definitions are assumed comparable from their labels; R3-3 (later) replaces that with the 2015 brief's facts. The later "
    "ruling is applied and the README does not call the segments comparable.",
    "8. **O-7 (round 4):** the README opening is two plain sentences; the second names the primary row and what its label means; the reviewer's remaining sentences follow the one denial sentence in the "
    "paragraph after the opening.",
    "9. **O-9:** `net_cash_pct_of_assets_start` for 2015 is blank (its opening assets are 2014 assets, across the jump in the number of systems); the reason is in the README, because `ratios.csv` has no note column.",
    "10. **C16 and C17** are PASS inside their stated ranges and WARN outside; C12 is WARN whenever a sub-line identity or a second weight is non-zero (as the spec observed for 2023), PASS if every "
    "identity holds. C12's sign is (sub-lines) minus (total line). C04 is split into C04 (same-vintage comparison) and C04p (the 2022 pin).",
    "11. **Unit counts and second weights.** A unit is a distinct value of the unit identifier; `second weight` means a weight other than the unit's most common weight (ties: the smaller). Counts only are printed.",
    "12. **C11 (fixed output)** compares a second, independent compute-and-render of every file except `checks.csv` with the first; the byte-identical re-run on disk is a separate step. "
    "**C21** is a check added by the builder: the folder guards (nine names, exact headers, no raw file, raw-row pattern, local path, at sign, banned wording, a real system name from the cache) run on the rendering.",
    "13. **Percent change on a negative base** (`revisions.csv`, gains): 100 x (new - prev) / |prev|, so an improvement reads positive; the maximum absolute value is unaffected.",
    "14. **Summary rows with no defined value are not written** (a ratio whose denominator is not positive). `n_years_net_cash_negative` carries the first and last year of its segment as `from_year` and `to_year`; "
    "`max_abs_revision_pct__<item>` carries the first and last year with a revision row.",
    "15. **`ratio_cv_bound_pct` is 0 in full-count years** and blank where no CV is published (2025, and any year whose workbook prints none).",
    "16. **Outputs of a build with a FAIL** are written to `data/census_public_pensions/_FAILED/` (gitignored), as in the house builds, and the committed files are left alone.",
    "17. **Wording:** the README is scanned for the two G14 patterns outside the one denial sentence, and also for the direction verbs `rose` and `fell` (R3-1); the two G14 patterns are applied to the CSV files as well, which is stricter than the spec.",
    "18. **S1** is written (`s1_revision_2025_pct`) only when two complete vintages are cached and the newest one's listing showed a survey-year folder above 2025; there is no such row yet.",
]
DECISIONS_POST = []                      # post-registration fixes (R9) are appended here with the old and the new value

REFERENCE_CITATIONS = """- C15 (press release): U.S. Census Bureau, CB25-TPS.31, 14 May 2026: survey year 2025 contributions $315.02bn, benefits $418.25bn, cash and investments $6.49 trillion; survey year 2024 benefits $404.46bn (typed in; the supervisor checked them against the release).
- C16 (NASRA): NASRA Public Fund Survey, FY 2024, $5.13 trillion; relayed via Grok and the supervisor, not fetched (nasra.org is never requested). A different universe; a level check only.
- C17 (Z.1): Federal Reserve Z.1, table S129.3.s, "Funded by assets", 2025:Q2, $6,541.4bn; relayed, not fetched (federalreserve.gov is never requested). A level check only; no Z.1 transaction series is read."""


def render_sources(c):
    V = c["V"]
    L = []
    L.append("# Sources, requests and builder decisions: Census Annual Survey of Public Pensions, national totals 2013-2025\n")
    L.append("Built by `bin/census_public_pensions.py` from the spec `_research/2026-10-26-Census-Public-Pensions-Spec.md` (committed at 104218d, with its rulings of rounds 1 to 4). "
             "Vintage `%s`; newest unit file %s (%s).\n" % (V.key, V.newest_unit_file, V.newest_unit_file_date))
    L.append("## What was seen of the answer before any build\n")
    L.append(DISCLOSURE + "\n")
    L.append("## Files requested by the build\n")
    L.append("Host `www2.census.gov` only; GET only; at least 3 seconds apart; no redirect followed; no retry; neutral User-Agent with no address; the 2020 unit file and the stale 2024 root files are never "
             "requested. One row per request, from `fetch_log.csv`:\n")
    rows = [[r[0], r[2], r[3], r[4], (r[5][:12] if len(r) > 5 and r[5] else ""), (r[6] if len(r) > 6 else "")] for r in c["fetch_rows"]]
    L.append(md_table(["UTC time", "Path", "Status", "Bytes", "sha256 (first 12)", "Last-Modified (UTC)"], rows) + "\n")
    L.append("## Requests made by the spec author (before the build)\n")
    L.append(SPEC_AUTHOR_REQUESTS + "\n")
    L.append("An unverified figure (R3-2): Grok relayed $286.21bn as 2022 contributions from indexed text of the data.census.gov table `GOVSPENSIONTIMESERIES.GS00PENSION`, which it could not render. "
             "It is recorded here as unverified; the build holds $%sbn from the 2022 unit file, and no data.census.gov request is made.\n" % bn(c["N"][2022]["contrib_total"], 2))
    L.append("## Layout mappings, labels and codes\n")
    L.append("Item codes come from the 2017-2024 and 2022-2025 layout PDFs (identical apart from the file-format sentence). Fields by position: year, name, identifier, state, government type, item code, "
             "value, value flag, source flag, weight.\n")
    L.append(md_table(["Column", "Codes", "Summary label (2017 on)"], [
        ["contrib_employee_usd_mn", "RX51 + RX52", "Employee contributions"], ["contrib_government_usd_mn", "RX60 + RX70", "Government contributions"],
        ["contrib_total_usd_mn", "RX51 + RX52 + RX60 + RX70", "Total Contributions"], ["benefits_usd_mn", "RX40", "Total Benefits Paid"], ["withdrawals_usd_mn", "RX80", "Withdrawals"],
        ["earnings_ex_gains_usd_mn", "RX01 (= RX10 + RX20; excludes gains)", "Earnings on investments"], ["gains_losses_usd_mn", "RX30, signed", "Gains/Losses on Investments"],
        ["expenses_admin_invest_usd_mn", "RX90", "Administrative Costs"], ["assets_end_usd_mn", "RZ01 (= RZ10 + RZ90)", "Total cash and investments"]]) + "\n")
    L.append("Never added on top: RX50 (equals RX51 + RX52), RX61 and RX62, RX71 to RX73, RX91 and RX92 (sub-lines), RJ (membership), RP (actuarial). Old-schema summary labels (2013-2016): "
             "`Total Contributions`, `Employee contributions`, `Government contributions`, `Benefits`, `Withdrawals`, `Earnings on investments`, `Total cash and investment holdings`, `Number of systems`. "
             "The summaries print `Investment Losses` as a positive number, so the build reads only the signed net line.\n")
    L.append("The 53 item codes: " + ", ".join(sorted(CODES)) + ".\n")
    L.append("## Timing, schema change and the 2015 brief\n")
    L.append("Timing (2024 methodology): \"The survey year encompasses the fiscal year beginning July 1 of the previous survey year and June 30 of the current survey year.\" "
             "Systems whose fiscal year ends after 30 June report the previous fiscal year; Alabama and Michigan state systems report October to September and Texas September to August. "
             "When holdings are measured is not stated in the publication, the layouts, the 2013 tables page or the methodology.\n")
    L.append("Schema change (2024 methodology, observed): Census streamlined the asset categories with the 2019 release and refitted 2017 and 2018 to the new schema, so the current files hold only "
             "cash and short-term versus long-term investments. Census points users who want granular asset detail to the Quarterly Survey of Public Pensions, which is out of scope here.\n")
    L.append("The 2015 brief (relayed via Grok, not fetched: `https://www.census.gov/content/dam/Census/library/publications/2016/econ/g15-aspp-sl.pdf`) says that more than 2,200 local systems were added, "
             "that the data may include hybrid and other post-employment funds, and that 2012 and 2015 should be compared only with caution.\n")
    L.append("The second reviewer observed that `ASPP_Unit_File_2020.dsv` on www2 has the same sha256 as the 2019 file and that every year field in it says 2019. It is never requested; 2020 comes from "
             "`ASPP_2020_Summary.xlsx`. The press release says benefits rose 3.40% from 2024; the unit files give 3.409%. That is a rounding note, not a check.\n")
    L.append("## Licence\n")
    L.append("Census's citation page asks users who make their own estimates to \"cite the Census Bureau as the source of the original data only\". The 2024 methodology treats the data as public record "
             "(relayed). Federal works carry no US copyright (17 U.S.C. 105; relayed; the open-data policy page answered 301 and was not followed). Raw files stay in the gitignored cache as a house "
             "choice (size, and they name individual systems), not as a licence requirement. `www2.census.gov/robots.txt` names a few crawlers and disallows nothing for others.\n")
    L.append("## Reference constants typed into the script (C15 to C17; never fetched)\n")
    L.append(REFERENCE_CITATIONS + "\n")
    L.append("## Builder decisions\n")
    L.append(decisions_text(c) + "\n")
    return "\n".join(L)


# ------------------------------------------------------------------------------------------------------------------------
# finish, inputs and the build
# ------------------------------------------------------------------------------------------------------------------------
CHECK_ORDER = ["C01", "C02", "C03", "C04", "C04p", "C05", "C06", "C07", "C08", "C09", "C10", "C11", "C12", "C13", "C14", "C15", "C16", "C17", "C18", "C19", "C20", "C21"]


def digests(texts):
    return {n: sha256_of(x.encode("utf-8")) for n, x in sorted(texts.items()) if n != "checks.csv"}


def render_all(c):
    t = render_data_texts(c)
    t["README.md"] = render_readme(c)
    t["SOURCES.md"] = render_sources(c)
    return t


def finish(c):
    """Add C11 (a second, independent compute and render leaves every file but checks.csv unchanged) and C21 (the folder guards run on the rendering)."""
    t1 = render_all(c)
    inp2 = dict(c["inp"])
    inp2.pop("_P", None)
    c2 = compute(inp2)
    t2 = render_all(c2)
    if M("x_render_unstable"):
        t2["national.csv"] += "x"
    same = digests(t1) == digests(t2)
    checks = c["checks"] + [chk("C11", "Fixed output: a second compute and render over the same cache leaves the sha256 of every file but checks.csv unchanged", "identical",
                                "identical (%d files)" % len(digests(t1)) if same else "DIFFERENT", "PASS" if same else "FAIL")]
    src = source_line(c)
    t1["checks.csv"] = render_csv("checks.csv", sorted(checks, key=lambda r: CHECK_ORDER.index(r[0])), src)
    probs = all_folder_problems(t1, c["names"])
    checks = sorted(checks + [chk("C21", "Folder guards on the rendering: nine names, exact headers, no raw file, no raw-row pattern, no system name from the cache, no local path, no at sign, banned wording, "
                                  "the summary rules (one primary row, no change in the old years or across 2016/2017)", "no problem",
                                  "; ".join(probs[:3]) if probs else "no problem", "FAIL" if probs else "PASS")], key=lambda r: CHECK_ORDER.index(r[0]))
    c["checks"] = checks
    t1["checks.csv"] = render_csv("checks.csv", checks, src)
    failed = [r[0] for r in checks if r[4] == "FAIL"]
    return t1, failed


def load_inputs(cache, out_dir):
    _, fetch_rows = read_csv_rows(os.path.join(out_dir, "fetch_log.csv"))
    vints = []
    for n in list_vintage_names(cache):
        vints.append(load_vintage(os.path.join(cache, n), fetch_rows))
    if not vints:
        raise DataError("no complete vintage in the cache %s: run fetch first" % tilde(cache))
    return {"vints": vints, "fetch_rows": fetch_rows}


def write_outputs(texts, out_dir, failed):
    target = os.path.join(out_dir, "_FAILED") if failed else out_dir
    for n, x in texts.items():
        if n == "fetch_log.csv" and target == out_dir and os.path.exists(os.path.join(out_dir, n)):
            continue                                              # the request log is written by fetch and only read here
        write_text(os.path.join(target, n), x)
    return target


def run_build(a):
    t0 = time.time()
    out_dir = a.out
    cache = check_cache(a.cache, out_dir=out_dir)
    try:
        inp = load_inputs(cache, out_dir)
        c = compute(inp)
        if not c["fetch_rows"]:
            raise DataError("fetch_log.csv holds no request: run fetch first")
        texts, failed = finish(c)
    except (DataError, ContentError) as e:
        print("STOP (data): %s" % e, file=sys.stderr)
        return 2
    target = write_outputs(texts, out_dir, failed)
    N, R, V = c["N"], c["R"], c["V"]
    prim = srow(c, "change_benefits_per_contrib_pct", 2023, 2025, "Q2")
    print("vintage %s (newest unit file %s, %s); %d vintage(s) cached" % (V.key, V.newest_unit_file, V.newest_unit_file_date, len(c["vints"])))
    print("== band: revision_band_pct %s, revision_band_sample_years_pct %s (measured %s, %s; frozen constants %s, %s)" %
          (fdec(c["band"], 3), fdec(c["band_s"], 3), fdec(measured_bands(c["ratio_revs"])[0], 3), fdec(measured_bands(c["ratio_revs"])[1], 3),
           fdec(FROZEN_BAND_PCT, 3) or "not set", fdec(FROZEN_SMP_BAND_PCT, 3) or "not set"))
    print("== primary row: change_benefits_per_contrib_pct, current, 2023 to 2025: %s%% (%s); ratio %s to %s" % (fdec(prim["value"], 3), prim["pattern"], fdec(R[2023]["bpc"], 4), fdec(R[2025]["bpc"], 4)))
    print("== checks")
    for r in c["checks"]:
        print("   %-5s %-5s %s" % (r[0], r[4], r[3][:230]))
    print("failed: %s; wrote %d files to %s (%.1f s)" % (", ".join(failed) or "none", len(texts), tilde(target), time.time() - t0))
    return 1 if failed else 0


# ------------------------------------------------------------------------------------------------------------------------
# static scans of this script's own source: the STOP latch (G15) and isolation (G17)
# ------------------------------------------------------------------------------------------------------------------------
REMOVERS = {"remove", "unlink", "rmtree", "rmdir"}
MUTATORS = {"replace", "rename", "move", "write_text", "write_bytes", "truncate", "copy", "copyfile", "copy2", "symlink", "link"}
LISTERS = {"glob", "iglob", "listdir", "walk", "scandir", "rglob", "iterdir"}


def _mentions_stop(node, tainted):
    for n in ast.walk(node):
        if isinstance(n, ast.Name) and ("stop" in n.id.lower() or n.id in tainted):
            return True
        if isinstance(n, ast.Attribute) and "stop" in n.attr.lower():
            return True
        if isinstance(n, ast.Constant) and isinstance(n.value, str) and "stop" in n.value.lower():
            return True
    return False


def _fname(call):
    f = call.func
    return f.attr if isinstance(f, ast.Attribute) else (f.id if isinstance(f, ast.Name) else "")


def _open_mode(call):
    """The mode of an open() call: the literal string, '' when absent (read), None when it is not a literal."""
    node = call.args[1] if len(call.args) > 1 else next((k.value for k in call.keywords if k.arg == "mode"), None)
    if node is None:
        return ""
    return node.value if isinstance(node, ast.Constant) and isinstance(node.value, str) else None


def _scope_nodes(scope):
    """The nodes of one scope, not descending into nested functions (each function is scanned as its own scope)."""
    out, stack = [], [scope]
    while stack:
        n = stack.pop()
        out.append(n)
        for ch in ast.iter_child_nodes(n):
            if not isinstance(ch, (ast.FunctionDef, ast.AsyncFunctionDef)):
                stack.append(ch)
    return out


def stop_ast_problems(src):
    """G15: an AST scan. A STOP path may be reached only by create_stop, with open(path, "x"). Removal is allowed only in remove_partial, by the exact partial path, in a function
    that lists nothing (a glob, a listing or a walk over the cache root is flagged)."""
    if M("g15_scan_off"):
        return []
    tree = ast.parse(src)
    problems = set()
    scopes = [(n.name, n) for n in ast.walk(tree) if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef))] + [("<module>", tree)]
    for fname, scope in scopes:
        nodes = _scope_nodes(scope)
        tainted = set()
        for _ in range(3):
            for n in nodes:
                tgt, val = None, None
                if isinstance(n, ast.Assign):
                    tgt, val = n.targets, n.value
                elif isinstance(n, (ast.AnnAssign, ast.AugAssign)) and n.value is not None:
                    tgt, val = [n.target], n.value
                elif isinstance(n, ast.For):
                    tgt, val = [n.target], n.iter
                if val is not None and _mentions_stop(val, tainted):
                    for t in tgt:
                        for x in ast.walk(t):
                            if isinstance(x, ast.Name):
                                tainted.add(x.id)
        calls = [n for n in nodes if isinstance(n, ast.Call)]
        listers = [n for n in calls if _fname(n) in LISTERS]
        for call in calls:
            nm = _fname(call)
            where = "%s line %d" % (fname, call.lineno)
            if fname == "create_stop" and nm in ("open", "fdopen") and _open_mode(call) != "x" and not M("g15_ignore_open"):
                problems.add("%s: create_stop opens a file in mode %r (only exclusive create, \"x\", is allowed)" % (where, _open_mode(call)))
            reach = any(_mentions_stop(a, tainted) for a in list(call.args) + [k.value for k in call.keywords]) or (
                isinstance(call.func, ast.Attribute) and _mentions_stop(call.func.value, tainted))
            if nm in REMOVERS:
                if fname != "remove_partial":
                    problems.add("%s: %s outside remove_partial" % (where, nm))
                elif listers and not M("g15_glob_ok"):
                    problems.add("%s: removal in a function that lists or globs a directory" % where)
                if reach:
                    problems.add("%s: %s reaches a STOP path" % (where, nm))
            elif nm in MUTATORS and reach and not (nm == "replace" and M("g15_replace_ok")):
                problems.add("%s: %s reaches a STOP path" % (where, nm))
            elif nm in ("open", "fdopen") and reach and not M("g15_ignore_open"):
                mode = _open_mode(call)
                if not (fname == "create_stop" and mode == "x") and (mode is None or any(ch in mode for ch in "wax+")):
                    problems.add("%s: open in mode %r reaches a STOP path" % (where, mode))
    return sorted(problems)


FORBIDDEN_FRAGMENTS = ("series.tsv", "5500", "data/history", "calendar.tsv", "handover.tsv", "the_ask", "research_state", "dropbox")
PATH_CALLS = {"open", "read_file", "read_csv_rows", "join", "exists", "listdir", "isdir", "isfile", "load_workbook", "write_text", "write_bytes", "makedirs", "glob", "walk"}


def isolation_problems(src):
    """G17: the script opens paths from its cache and its output folder only. No path call names a project data file (the series table, a Form 5500 file, the calendar, the
    hand-over log), an absolute path or a parent directory."""
    if M("g17_scan_off"):
        return []
    tree = ast.parse(src)
    out = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Call) and _fname(node) in PATH_CALLS:
            if isinstance(node.func, ast.Attribute) and isinstance(node.func.value, ast.Constant):
                continue                                  # "sep".join(...)
            for a in list(node.args) + [k.value for k in node.keywords]:
                for s in ast.walk(a):
                    if isinstance(s, ast.Constant) and isinstance(s.value, str):
                        v = s.value.lower()
                        if any(fr in v for fr in FORBIDDEN_FRAGMENTS):
                            out.add("line %d: %s(...) names %r" % (node.lineno, _fname(node), s.value))
                        elif _fname(node) in ("open", "read_file", "read_csv_rows", "load_workbook") and (s.value.startswith("/") or ".." in s.value):
                            out.add("line %d: %s(...) opens the absolute or parent path %r" % (node.lineno, _fname(node), s.value))
    return sorted(out)


def checksh_problems(text):
    """G12: the M25 guard in bin/check.sh holds its own literal copy of the seven canonical lines and never reads them from this script."""
    a = text.find("cpp_leak_check() {")
    if a < 0:
        return ["bin/check.sh has no cpp_leak_check function"]
    region = text[a:text.find("\n}\n", a) + 3]
    out = []
    for l in SPEC_ALLOW_TEXT.splitlines():
        name, hdr = l.split(":", 1)
        if ("%s) want='%s'" % (name, hdr)) not in region:
            out.append("check.sh does not hold the literal header of %s" % name)
    if re.search(r"census_public_pensions\.py", region) and not M("g12_checksh_reads_script"):
        out.append("check.sh names the script inside the guard (it must not read the headers from it)")
    return out


# ------------------------------------------------------------------------------------------------------------------------
# selftest: synthetic data only. Every guard G1-G19 of the spec's section 7 has a test; each mutant (a deliberately broken variant, switched on in this
# process only) must make its guard's test fail. A mutant that is not caught is a failure.
# ------------------------------------------------------------------------------------------------------------------------
class Checks:
    def __init__(self):
        self.items = []

    def ok(self, name, cond, detail=""):
        self.items.append((name, bool(cond), "" if cond else detail))

    def eq(self, name, got, want):
        self.ok(name, got == want, "got %r, want %r" % (got, want))

    def raises(self, name, exc, fn):
        try:
            fn()
        except exc:
            self.ok(name, True)
        except Exception as e:
            self.ok(name, False, "raised %s: %s" % (type(e).__name__, e))
        else:
            self.ok(name, False, "did not raise")

    def quiet(self, name, fn):
        try:
            fn()
        except Exception as e:
            self.ok(name, False, "raised %s: %s" % (type(e).__name__, e))
        else:
            self.ok(name, True)

    def failed(self):
        return [(n, d) for n, c, d in self.items if not c]


class FakeClock:
    def __init__(self, t=1_800_000_000.0):
        self.t = t

    def wall(self):
        return self.t

    def mono(self):
        return self.t

    def sleep(self, s):
        self.t += s

    def now_utc(self):
        return dt.datetime.fromtimestamp(self.t, dt.timezone.utc)


# ---- synthetic raw files -------------------------------------------------------------------------------------------------
UNIT_HDR = ["SAMPLE_YEAR", "UNIT_NAME", "PID6", "STATE", "UNIT_TYPE", "ITEM_CODE", "ITEM_VALUE", "EIC", "ISC", "FINAL_WEIGHT"]


def R10(y, uid, name, code, val, w="1"):
    return (str(y), name, uid, "XX", "0", code, str(val), "R", "V", str(w))


def unit_bytes(rows, delim="|", header=False, enc="utf-8", bom=False):
    out = io.StringIO()
    if delim == ",":
        w = csv.writer(out, lineterminator="\n")
        if header:
            w.writerow(UNIT_HDR)
        for r in rows:
            w.writerow(r)
    else:
        if header:
            out.write(delim.join(UNIT_HDR) + "\n")
        for r in rows:
            out.write(delim.join(r) + "\n")
    b = out.getvalue().encode(enc)
    return (codecs.BOM_UTF8 + b) if bom else b


def mk_summary(year, vals, cvs=None, geo=1, census=False, drop=(), losses=None):
    """A synthetic summary workbook shaped like Census's: a geography row, an Item row, one row per label. vals: key -> number as printed (thousands; a count for systems)."""
    import openpyxl
    labels = LABELS_OLD if year <= 2016 else LABELS_CUR
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([""] * geo + ["United States"] + [None] * 4 + ["Alabama"])
    heads = ["State & Local"] + ([] if census else ["CV for State & Local"]) + ["State"] + ([] if census else ["CV for State"]) + ["Local"] + ([] if census else ["CV Local"])
    pad = ["Item"] + ["Filler"] * (geo - 1)
    ws.append(pad + heads + ["State & Local"])
    noise = ["Interest", "Dividends"] if year > 2016 else ["Total payments", "Corporate stocks"]
    for nz in noise:
        ws.append([nz] + ["."] * (geo - 1) + [12345] + ["."] * 5)
    for key, label in labels.items():
        if label in drop:
            continue
        cv = (cvs or {}).get(key)
        row = [label] + ["."] * (geo - 1) + [vals[key]]
        if not census:
            row.append(cv if cv is not None else ".")
        ws.append(row)
    if losses:
        ws.append(["Investment Gains"] + ["."] * (geo - 1) + [losses[0]])
        ws.append(["Investment Losses"] + ["."] * (geo - 1) + [losses[1]])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def q001(x):
    return Decimal(x).quantize(Decimal("0.001"), rounding=decimal.ROUND_HALF_EVEN)


def world_targets(y):
    p = {k: Decimal(v) for k, v in PINS.get(y, {}).items()}
    tot = p.get("contrib_total", Decimal(100000 + 10000 * (y - 2013)))
    emp = p.get("contrib_employee", q001(tot / 4))
    out = {"contrib_total": tot, "contrib_employee": emp, "contrib_government": p.get("contrib_government", tot - emp)}
    out["benefits"] = p.get("benefits", q001(tot * Decimal("1.3")))
    out["withdrawals"] = p.get("withdrawals", q001(tot * Decimal("0.03")))
    out["earnings_ex_gains"] = p.get("earnings_ex_gains", q001(tot * Decimal("0.4")))
    out["gains_losses"] = p.get("gains_losses", q001(tot * Decimal("0.5")))
    out["expenses"] = p.get("expenses", q001(tot * Decimal("0.1")))
    out["assets_end"] = p.get("assets_end", q001(tot * 20))
    return out


def dollars(mnv):
    return int(Decimal(mnv) * 1000000)


def split_ab(T, census):
    """T dollars as one weight-1 unit plus one unit of weight 2 (a sample year), or two weight-1 units (a census year); the weighted sum is exactly T."""
    if census:
        a = T // 2
        return a, T - a, 1, 1
    b = T // 4
    return T - 2 * b, b, 1, 2


def world_unit(y, tg, cp, pad_units=0, n_neg=0):
    census = y in CENSUS_YEARS
    E, G, EA = dollars(tg["contrib_employee"]), dollars(tg["contrib_government"]), dollars(tg["earnings_ex_gains"])
    X, A = dollars(tg["expenses"]), dollars(tg["assets_end"])
    G60 = G // 2
    G70 = G - G60
    totals = collections.OrderedDict([
        ("RX51", E // 2), ("RX52", E - E // 2), ("RX50", E), ("RX60", G60), ("RX61", G60 // 2), ("RX62", G60 - G60 // 2), ("RX70", G70), ("RX71", G70 // 3), ("RX72", G70 // 3),
        ("RX73", G70 - 2 * (G70 // 3)), ("RX40", dollars(tg["benefits"])), ("RX80", dollars(tg["withdrawals"])), ("RX01", EA), ("RX10", EA // 2), ("RX20", EA - EA // 2),
        ("RX30", dollars(tg["gains_losses"]) + 1000 * n_neg), ("RX90", X), ("RX91", X // 3), ("RX92", X - X // 3), ("RZ01", A), ("RZ10", A // 2), ("RZ90", A - A // 2)])
    nameA = "SYNTHETIC EMPLOYEES’ RETIREMENT FUND" if cp else "SYNTHETIC RETIREMENT FUND, THE"
    nameB = "SYNTHETIC LOCAL FUND"
    rows = []
    for code, T in totals.items():
        a, b, wa, wb = split_ab(T, census)
        rows.append(R10(y, "100001", nameA, code, a, wa))
        rows.append(R10(y, "100002", nameB, code, b, wb))
    rows.append(R10(y, "100001", nameA, "RJ01", 500, 1))
    rows.append(R10(y, "100001", nameA, "RP01", 7, 1))
    for i in range(n_neg):
        rows.append(R10(y, str(200000 + i), "SYNTHETIC DUMMY %d" % i, "RX30", -1000, 1))
    for i in range(pad_units):
        rows.append(R10(y, str(300000 + i), "SYNTHETIC PAD %d" % i, "RJ01", 1, 1))
    return rows


def calibrate_summary(y, tg):
    """Summary-table contributions and benefits for 2022-2024, set so that the ratio revisions are the pinned ones."""
    ru = tg["benefits"] / tg["contrib_total"]
    if y == 2022:
        cs, pct = Decimal("282833.491"), Decimal("-1.898")
        bs = q001(ru / (1 + pct / 100) * cs)
    elif y == 2023:
        cs, pct = q001(tg["contrib_total"] / (1 - Decimal("0.00037"))), Decimal("0.102")
        bs = q001(ru / (1 + pct / 100) * cs)
    else:
        pct, bs = Decimal("0.176"), Decimal("405454.166")
        cs = q001(bs / (ru / (1 + pct / 100)))
    return cs, bs


OLD_LM = "2018-08-30T14:00:00Z"
LM_2025 = "2025-05-23T17:00:31Z"
LM_2026 = "2026-05-14T15:15:32Z"


def world_files_raw():
    """{path: (bytes, last_modified_iso)} for the 20 files of the table."""
    out = {}
    for f in FILES:
        y, kind = f["year"], f["kind"]
        tg = world_targets(y)
        if kind == "unit":
            cp = y in (2017, 2021)
            hdr = y >= 2022
            n_neg = {2025: 8, 2024: 24}.get(y, 0)
            pad = (UNITS_PINS.get(y, 0) - 2 - n_neg) if y in UNITS_PINS else 0
            rows = world_unit(y, tg, cp, pad_units=max(pad, 0), n_neg=n_neg)
            ext = os.path.splitext(f["name"])[1]
            if cp:
                raw = unit_bytes(rows, "|", header=False, enc="cp1252")
            else:
                raw = unit_bytes(rows, "|" if ext == ".dsv" else ",", header=hdr, bom=hdr)
            out[f["path"]] = (raw, LM_2026 if y >= 2022 else LM_2025)
        else:
            vals = {}
            t = dict(tg)
            if y in (2022, 2023, 2024):
                cs, bs = calibrate_summary(y, tg)
                t["contrib_total"], t["benefits"] = cs, bs
                t["contrib_employee"] = q001(cs / 4)
                t["contrib_government"] = cs - t["contrib_employee"]
            for k, col in (("contrib_total", "contrib_total"), ("contrib_employee", "contrib_employee"), ("contrib_government", "contrib_government"), ("benefits", "benefits"),
                           ("withdrawals", "withdrawals"), ("assets", "assets_end"), ("earnings", "earnings_ex_gains"), ("gains", "gains_losses"), ("expenses", "expenses")):
                vals[k] = int(t[col] * 1000)
            vals["systems"] = {**SYSTEMS_PINS, 2017: 5529, 2022: 4394}.get(y, 4000 + y - 2013)
            census = y in CENSUS_YEARS
            cvs = None if census else {"contrib_total": 0.002, "benefits": 0.001, "assets": 0.0015}
            if y == 2013:
                cvs = {"contrib_total": 0.0334, "benefits": 0.02, "assets": 0.01}
            out[f["path"]] = (mk_summary(y, vals, cvs, census=census), OLD_LM if y <= 2016 else LM_2025)
    return out


def world_vintage(raw=None, fetch_stamp="2026-10-26T05:44:00Z"):
    raw = raw or world_files_raw()
    meta = {p: {"sha256": sha256_of(b), "bytes": str(len(b)), "last_modified": lm} for p, (b, lm) in raw.items()}
    rows = [(p, m["last_modified"], int(m["bytes"])) for p, m in meta.items()]
    key = vintage_key(rows)
    fetch_rows = [[fetch_stamp[:11] + "05:%02d:%02dZ" % (44 + i // 30, (i * 2) % 60), HOST, p, "200", m["bytes"], m["sha256"], m["last_modified"]] for i, (p, m) in enumerate(sorted(meta.items()))]
    V = Vintage(key, meta, lambda p: raw[p][0], requests=21, notice_years=(), manifest_sha=sha256_of(b"manifest"), retrieved=fetch_stamp)
    return V, fetch_rows


_WORLD = {}


def get_world():
    if _WORLD:
        return _WORLD
    assert not _ACTIVE
    raw = world_files_raw()
    V, fetch_rows = world_vintage(raw)
    p25 = next(f["path"] for f in FILES if f["year"] == 2025 and f["kind"] == "unit")
    sh = V.meta[p25]["sha256"]
    inp = {"vints": [V], "fetch_rows": fetch_rows, "_P": {}, "sha_pin": (sh[:8], sh[-6:])}
    c0 = compute(inp)
    texts, failed = finish(c0)
    _WORLD.update({"raw": raw, "V": V, "fetch_rows": fetch_rows, "inp": inp, "c0": c0, "texts": texts, "failed": failed})
    return _WORLD


def world_compute(**kw):
    return compute(get_world()["inp"], **kw)


def fake_N(overrides=None):
    """A national dict for every survey year with simple numbers, for the ratio tests."""
    N = {}
    for y in SURVEY_YEARS:
        r = {"survey_year": y, "schema": "old" if y <= 2016 else "current", "year_kind": "sample", "contrib_total": Decimal(100), "benefits": Decimal(130), "withdrawals": Decimal(3),
             "expenses": Decimal(10), "earnings_ex_gains": Decimal(40) if y > 2016 else None, "earnings_old": Decimal(40) if y <= 2016 else None,
             "assets_end": Decimal(1000 + 100 * (y - 2013)), "cv_contrib_total": Decimal("0.002"), "cv_benefits": Decimal("0.001"), "cv_assets": Decimal("0.0015"),
             "contrib_employee": Decimal(25), "contrib_government": Decimal(75), "gains_losses": Decimal(50), "systems": 1, "source_kind": "summary", "source_file_date": "2025-05-23",
             "summary_file_date": "2025-05-23", "vintage": "x"}
        r.update((overrides or {}).get(y, {}))
        r["net_cash"] = net_of(r["contrib_total"], r["benefits"], r["withdrawals"], r["expenses"])
        N[y] = r
    return N


# ---- the guards ----------------------------------------------------------------------------------------------------------
def sums_of(up):
    return {k: v for k, v in up.sums.items()}


def t_g1(ck):
    A = [R10(2025, "100001", 'SYNTHETIC FUND, THE', "RX40", 1000, 1), R10(2025, "100002", "SYNTHETIC B", "RX40", 500, 2), R10(2025, "100001", 'SYNTHETIC FUND, THE', "RX51", 25, 1)]
    up = parse_unit(unit_bytes(A, ",", header=True, bom=True), 2025, ".csv")
    ck.eq("a csv with a header, a BOM and a quoted name with a comma: no problem", up.problems, [])
    ck.eq("...the header is recognised and not summed", (up.header, up.rows, up.enc), (True, 3, "utf-8-sig"))
    ck.eq("...and the weighted sums are right", (up.sums["RX40"], up.sums["RX51"]), (Decimal(2000), Decimal(25)))
    up = parse_unit(unit_bytes(A, "|"), 2025, ".dsv")
    ck.eq("a headerless dsv: the first data row is kept", (up.header, up.rows, up.sums["RX40"], up.problems), (False, 3, Decimal(2000), []))
    C = [R10(2017, "100001", "SYNTHETIC EMPLOYEES’ FUND", "RX40", 1000, 1), R10(2017, "100002", "SYNTHETIC B", "RX40", 500, 1)]
    up = parse_unit(unit_bytes(C, "|", enc="cp1252"), 2017, ".dsv")
    ck.eq("a cp1252 file is decoded as cp1252, with the curly apostrophe intact", (up.enc, any("’" in n for n in up.names), up.rows, up.sums["RX40"], up.problems),
          ("cp1252", True, 2, Decimal(1500), []))
    bad = A + [("2025", "N", "9", "XX", "0", "RX40", "5", "R", "V"), ("2025", "N", "9", "XX", "0", "RX40", "5", "R", "V", "1", "extra")]
    up = parse_unit(unit_bytes(bad, "|"), 2025, ".dsv")
    ck.eq("a 9-field and an 11-field row are refused, counted and not summed", (up.bad_fields, up.sums["RX40"]), (2, Decimal(2000)))
    ck.ok("...and the refusal is a problem for C20", any("field count" in p for p in up.problems), str(up.problems))
    up = parse_unit(unit_bytes([R10(2024, "1", "N", "RX40", 5, 1)], "|"), 2025, ".dsv")
    ck.ok("a row whose first field is not the survey year is counted and not summed", up.bad_years == 1 and up.sums.get("RX40", 0) == 0, str((up.bad_years, up.sums.get("RX40"))))


def t_g2(ck):
    rows = [R10(2025, "100001", "SYNTHETIC ONE", "RX40", 100, 1), R10(2025, "100002", "SYNTHETIC TWO", "RX40", 200, 2), R10(2025, "300003", "SYNTHETIC THREE", "RX40", 10, 3),
            R10(2025, "300003", "SYNTHETIC THREE", "RX51", 4, 3), R10(2025, "300003", "SYNTHETIC THREE", "RX70", 5, 4)]
    up = parse_unit(unit_bytes(rows), 2025, ".dsv")
    ck.eq("the weighted sums apply each row's own weight", (up.sums["RX40"], up.sums["RX51"], up.sums["RX70"]), (Decimal(530), Decimal(12), Decimal(20)))
    ck.eq("...the second weight is counted (units, rows, non-blank rows)", (up.units_second, up.rows_second, up.rows_second_nonblank), (1, 1, 1))
    ck.eq("...and the weight facts", (up.all_one, up.max_weight), (False, Decimal(4)))
    P = {y: {} for y in SURVEY_YEARS}
    P[2025]["unit"] = up
    row = check_c12(P)
    ck.ok("C12 prints counts of units and rows", "1 unit(s) with a second weight, 1 non-blank row(s)" in row[3], row[3])
    ck.ok("...and no unit identifier or value", "300003" not in row[3] and "RX70 20" not in row[3], row[3])
    cen = [R10(2017, "100001", "SYNTHETIC ONE", "RX40", 100, 1), R10(2017, "100002", "SYNTHETIC TWO", "RX40", 200, "1.5")]
    P = {y: {} for y in SURVEY_YEARS}
    P[2017]["unit"] = parse_unit(unit_bytes(cen), 2017, ".dsv")
    ck.eq("a census-year fixture with one weight of 1.5 is a FAIL", check_c10(P)[4], "FAIL")
    P[2017]["unit"] = parse_unit(unit_bytes(cen[:1]), 2017, ".dsv")
    ck.eq("...and with every weight 1 it is not", check_c10(P)[4], "PASS")


ALLCODES = {"RX01": 5, "RX10": 2, "RX20": 3, "RX30": -7, "RX40": 11, "RX50": 110, "RX51": 10, "RX52": 100, "RX60": 1000, "RX61": 600, "RX62": 400, "RX70": 20, "RX71": 8, "RX72": 7,
            "RX73": 5, "RX80": 13, "RX90": 17, "RX91": 9, "RX92": 8, "RZ01": 99, "RZ10": 90, "RZ90": 9}


def t_g3(ck):
    rows = [R10(2025, "100001", "SYNTHETIC ONE", c, v * 1000000, 1) for c, v in ALLCODES.items()]
    up = parse_unit(unit_bytes(rows), 2025, ".dsv")
    ck.eq("a fixture with every code parses clean", up.problems, [])
    uc = unit_cols(up)
    ck.eq("employee contributions are RX51 + RX52 (not RX50)", uc["contrib_employee"], Decimal(110))
    ck.eq("government contributions are RX60 + RX70 (not RX61, RX62, RX71-RX73)", uc["contrib_government"], Decimal(1020))
    ck.eq("total contributions are the four codes", uc["contrib_total"], Decimal(1130))
    ck.eq("benefits, withdrawals", (uc["benefits"], uc["withdrawals"]), (Decimal(11), Decimal(13)))
    ck.eq("earnings excluding gains are RX01 alone (not RX30)", uc["earnings_ex_gains"], Decimal(5))
    ck.eq("gains and losses are RX30, signed", uc["gains_losses"], Decimal(-7))
    ck.eq("expenses are RX90 (not RX91, RX92)", uc["expenses"], Decimal(17))
    ck.eq("assets are RZ01", uc["assets_end"], Decimal(99))
    ck.eq("RX90 is not a payout in net cash", net_of(uc["contrib_total"], uc["benefits"], uc["withdrawals"], uc["expenses"]), Decimal(1106))
    ck.eq("a code is matched whole: RX5 would select RX50-RX52, and does not", pick(up.sums, ["RX5"]), Decimal(0))


def t_g4(ck):
    rows = [R10(2025, "100001", "SYNTHETIC ONE", "RX40", 1, 1), R10(2025, "100001", "SYNTHETIC ONE", "RX99", 1, 1), R10(2025, "100001", "SYNTHETIC ONE", "RX98", 1, 1)]
    up = parse_unit(unit_bytes(rows), 2025, ".dsv")
    ck.eq("a 54th code is refused, and every unknown is listed together", sorted(up.unknown), ["RX98", "RX99"])
    ck.ok("...in one problem", len([p for p in up.problems if "outside the 53" in p]) == 1 and "RX98, RX99" in " ".join(up.problems), str(up.problems))
    vals = {"contrib_total": 300, "contrib_employee": 100, "contrib_government": 200, "benefits": 400, "withdrawals": 5, "assets": 9000, "earnings": 50, "gains": 60, "expenses": 7,
            "systems": 12}
    sp = parse_summary(mk_summary(2025, vals, {"contrib_total": 0.002}, drop=("Withdrawals", "Administrative Costs")), 2025)
    ck.ok("missing labels are reported, both together", len(sp.problems) == 1 and "Administrative Costs" in sp.problems[0] and "Withdrawals" in sp.problems[0], str(sp.problems))
    ck.ok("...and are not filled with zero", "withdrawals" not in sp.values and "expenses" not in sp.values, str(sp.values))
    sp = parse_summary(mk_summary(2025, vals, {"contrib_total": 0.002}, geo=3), 2025)
    ck.eq("columns are found by their headings, not by position (State & Local under United States at column 3)", (sp.problems, sp.values.get("contrib_total"), sp.values.get("benefits")),
          ([], Decimal(300), Decimal(400)))
    sp = parse_summary(mk_summary(2025, vals, {"contrib_total": 0.002}), 2025)
    ck.eq("a normal workbook parses clean, with its CV as a fraction", (sp.problems, sp.cv.get("contrib_total")), ([], Decimal("0.002")))


def t_g5(ck):
    rows = [R10(2025, "100001", "SYNTHETIC ONE", "RX30", -5, 1), R10(2025, "100002", "SYNTHETIC TWO", "RX30", 3, 1)]
    up = parse_unit(unit_bytes(rows), 2025, ".dsv")
    ck.eq("a gain row of -5 stays negative", (up.sums["RX30"], up.neg_rx30), (Decimal(-2), 1))
    ck.eq("net cash is contributions minus benefits minus withdrawals", net_of(Decimal(1000), Decimal(300), Decimal(50), Decimal(9)), Decimal(650))
    vals = {"contrib_total": 300, "contrib_employee": 100, "contrib_government": 200, "benefits": 400, "withdrawals": 5, "assets": 9000, "earnings": 50, "gains": 83792371, "expenses": 7,
            "systems": 12}
    sp = parse_summary(mk_summary(2023, vals, None, losses=(219219955, 135427585)), 2023)
    ck.eq("the summary's signed net line is read; Investment Losses is never read as a gain", sp.values["gains"], Decimal(83792371))
    W = get_world()
    N = {y: dict(r) for y, r in W["c0"]["N"].items()}
    ok_row = check_c09(N, W["c0"]["P"])
    ck.eq("C09 passes on the synthetic build", ok_row[4], "PASS")
    N[2025]["net_cash"] = -N[2025]["net_cash"]
    row = check_c09(N, W["c0"]["P"])
    ck.ok("the pin '2025 net cash is negative' fails on a sign-flipped fixture", row[4] == "FAIL" and "net_cash_usd_mn is negative expected True" in row[3], str(row))


def t_g6(ck):
    ck.eq("a summary figure is in thousands of dollars: 1,000 is one million", summary_to_mn(Decimal(1000)), Decimal(1))
    ck.eq("a unit-file figure is in dollars: 1,000,000 is one million", unit_to_mn(Decimal(1000000)), Decimal(1))
    ck.eq("a _usd_mn cell holds millions, not billions", money(Decimal(5)), "5.000")
    W = get_world()
    n = W["c0"]["N"][2025]
    ck.eq("the 2025 total contributions in the synthetic build are 315,020.924 million", money(n["contrib_total"]), "315020.924")
    ck.eq("...and the same figure from the summary route (2020 pinned, from thousands) is 241,972.089 million", money(W["c0"]["N"][2020]["contrib_total"]), "241972.089")


def t_g7(ck):
    c = world_compute()
    N, R, S = c["N"], c["R"], c["summary"]
    rows = summary_csv_rows(S)
    ck.eq("no change row inside the old years or across 2016/2017", [p for p in summary_problems(rows) if "change row" in p], [])
    ck.ok("the changes that exist run between current-schema years", all(r["from_year"] >= 2017 for r in S if r["statistic"].startswith("change_") and r["from_year"]),
          str([(r["from_year"], r["to_year"]) for r in S if r["statistic"].startswith("change_")]))
    ck.ok("a synthetic 2014-to-2015 change row is reported", any("inside the old schema" in p for p in summary_problems(
        rows + [["Q2", "old", "2014", "2015", "change_benefits_per_contrib_usd", "0.1000", "usd_per_usd", "n_a", "no", "x"]])), "not reported")
    ck.ok("a synthetic 2016-to-2018 change row is reported", any("spanning 2016/2017" in p for p in summary_problems(
        rows + [["Q2", "current", "2016", "2018", "change_benefits_per_contrib_usd", "0.1000", "usd_per_usd", "n_a", "no", "x"]])), "not reported")
    ck.ok("old earnings are in the old-schema column, never in earnings_ex_gains", all(N[y]["earnings_ex_gains"] is None and N[y]["earnings_old"] is not None for y in OLD_YEARS),
          str([(y, N[y]["earnings_ex_gains"]) for y in OLD_YEARS]))
    ck.ok("...and gains and expenses are blank in the old years", all(N[y]["gains_losses"] is None and N[y]["expenses"] is None for y in OLD_YEARS), "")
    ck.eq("net cash as a share of opening assets is blank for 2013, 2015 and 2017 (not computed from 2016)", [R[y]["start"] for y in (2013, 2015, 2017)], [None, None, None])
    ck.ok("...and present for 2014, 2016 and 2018", all(R[y]["start"] is not None for y in (2014, 2016, 2018)), "")
    ck.ok("the income ratios are blank for every old year", all(R[y]["income_gap"] is None and R[y]["income_gap_exp"] is None for y in OLD_YEARS), "")
    ck.ok("...and no Q4 row is written for an old year", not [r for r in S if r["question"] == "Q4" and r["segment"] == "old"], "")


def t_g8(ck):
    n = 10000
    rows = [R10(2025, str(100000 + i), "SYNTHETIC %d" % i, "RX40", "98765432101234.56", "1.01") for i in range(n)]
    up = parse_unit(unit_bytes(rows), 2025, ".dsv")
    exact = Decimal("98765432101234.56") * Decimal("1.01") * n
    ck.eq("10,000 cent-weighted rows sum exactly in Decimal", up.sums["RX40"], exact)
    fl = 0.0
    for _ in range(n):
        fl += float("98765432101234.56") * float("1.01")
    ck.ok("(the fixture is sensitive: a float sum is off by more than 0.0005 $mn)", abs(Decimal(fl) - exact) / 1000000 > Decimal("0.0005"), str(abs(Decimal(fl) - exact)))
    ck.ok("...and the Decimal sum is within 0.0005 $mn of the exact figure", abs(unit_to_mn(up.sums["RX40"]) - unit_to_mn(exact)) <= Decimal("0.0005"), "")
    N = fake_N({2023: {"contrib_total": Decimal("0.0015"), "benefits": Decimal("0.0023"), "withdrawals": Decimal("0.0001")}})
    R = ratio_rows(N)
    ck.eq("a ratio is divided before it is rounded", fdec(R[2023]["bpc"], 4), "1.5333")
    W = get_world()
    t = render_data_texts(W["c0"])
    h, rows = csv_body(t["national.csv"])
    money_cols = [i for i, k in enumerate(h) if k.endswith("_usd_mn") and k not in ("earnings_old_schema_usd_mn",)]
    bad = [r[i] for r in rows for i in money_cols if r[i] and not re.fullmatch(r"-?\d+\.\d{3}", r[i])]
    ck.eq("every money cell of national.csv has three decimals", bad, [])
    h, rows = csv_body(t["ratios.csv"])
    bad = [r[2] for r in rows if not re.fullmatch(r"-?\d+\.\d{4}", r[2])]
    ck.eq("every per-dollar ratio has four decimals", bad, [])
    h, rows = csv_body(t["summary.csv"])
    bad = [r for r in rows if r[6] == "usd_per_usd" and not re.fullmatch(r"-?\d+\.\d{4}", r[5])]
    ck.eq("...and so does every usd_per_usd value of summary.csv", bad, [])


def t_g9(ck):
    raw = get_world()["raw"]
    rows = [(p, lm, len(b)) for p, (b, lm) in raw.items()]
    k0 = vintage_key(rows)
    time.sleep(0.002)
    ck.eq("the key is stable for the same headers (not the wall clock)", vintage_key(rows), k0)
    ck.eq("it is 16 hex characters", bool(KEY_RX.match(k0)), True)
    older = [(p, ("2025-05-24T17:00:31Z" if p.endswith("ASPP_Unit_File_2019.dsv") else lm), b) for p, lm, b in rows]
    ck.ok("a change in only an older year's unit file changes the key", vintage_key(older) != k0, "unchanged")
    sized = [(p, lm, (b + 1 if p.endswith("ASPP_2018_Summary.xlsx") else b)) for p, lm, b in rows]
    ck.ok("a change in only a size changes the key", vintage_key(sized) != k0, "unchanged")
    old_tz = os.environ.get("TZ")
    try:
        os.environ["TZ"] = "Pacific/Auckland"
        time.tzset()
        ck.eq("Last-Modified is read in UTC: 23:30 GMT stays on its day", http_date_iso("Fri, 23 May 2025 23:30:00 GMT"), "2025-05-23T23:30:00Z")
    finally:
        if old_tz is None:
            os.environ.pop("TZ", None)
        else:
            os.environ["TZ"] = old_tz
        time.tzset()
    ck.eq("the newest unit file is read from the headers", newest_unit(rows), ("ASPP_Unit_File_2025.csv", "2026-05-14"))


def t_g10(ck):
    W = get_world()
    c = world_compute()
    years = sorted({r["survey_year"] for r in c["revs"]})
    ck.eq("revisions are written for the years whose summary and unit file differ in date (2022-2024), 8 items each", (years, len(c["revs"])), ([2022, 2023, 2024], 24))
    c2 = world_compute(existing_revs=c["revs"])
    ck.eq("revisions.csv is regenerated, not appended (a second build gives the same rows)", len(c2["revs"]), 24)
    ys, yd, gaps, bad = c04_compare(c["P"])
    ck.eq("same-date years are compared and a different-date gap is not a failure", (ys, yd, bad), ([2017, 2018, 2019, 2021], [2022, 2023, 2024], []))
    P = c["P"]
    saved = P[2018]["sum"].values["contrib_total"]
    try:
        P[2018]["sum"].values["contrib_total"] = saved + 5
        ys, yd, gaps, bad = c04_compare(P)
        ck.ok("a same-date summary that differs by 5 thousand dollars is refused", len(bad) == 1 and "2018" in bad[0], str(bad))
    finally:
        P[2018]["sum"].values["contrib_total"] = saved
    with tempfile.TemporaryDirectory(prefix="cpp_g10_", ignore_cleanup_errors=True) as tmp:
        cache = os.path.join(tmp, "cache")
        write_vintage_folder(cache, W["V"], W["raw"])
        part = os.path.join(cache, "partial")
        os.makedirs(part)
        write_manifest(part, {FILES[0]["path"]: {"url_path": FILES[0]["path"], "sha256": "ab", "bytes": "1", "last_modified": LM_2025}})
        names = list_vintage_names(cache)
        ck.eq("the partial folder is not a vintage", names, [W["V"].key])
        ck.quiet("...and the complete one loads", lambda: load_vintage(os.path.join(cache, names[0]), W["fetch_rows"]))


def write_vintage_folder(cache, V, raw, name=None):
    folder = os.path.join(cache, name or V.key)
    os.makedirs(folder, exist_ok=True)
    man = {}
    for f in FILES:
        b, lm = raw[f["path"]]
        write_bytes(os.path.join(folder, f["name"]), b)
        man[f["path"]] = {"url_path": f["path"], "sha256": sha256_of(b), "bytes": str(len(b)), "last_modified": lm}
    write_manifest(folder, man)
    write_text(os.path.join(folder, "requests.txt"), "21\n")
    write_text(os.path.join(folder, "notice.txt"), "listing years above 2025: none\n")
    return folder


def t_g11(ck):
    fb, fs = Decimal("1.898"), Decimal("0.176")
    c = world_compute(frozen=fb, frozen_sample=fs)
    N, R, S = c["N"], c["R"], c["summary"]
    y = 2024
    want_end = Decimal(100) * N[y]["net_cash"] / N[y]["assets_end"]
    want_start = Decimal(100) * N[y]["net_cash"] / N[y - 1]["assets_end"]
    ck.eq("net cash as a percent of end-of-year assets is 100 x net / the same year's assets", fdec(R[y]["end"], 6), fdec(want_end, 6))
    ck.eq("net cash as a percent of opening assets is 100 x net / the previous year's assets", fdec(R[y]["start"], 6), fdec(want_start, 6))
    prim = [r for r in S if r["primary"] == "yes"]
    b = Decimal("1.898")
    ck.eq("a change beyond the band is above_band, below_band, and a smaller one within_band (synthetic changes +5, -5, +1, -1, 0)",
          [band_pattern(Decimal(x), b) for x in ("5", "-5", "1", "-1", "0")], ["above_band", "below_band", "within_band", "within_band", "within_band"])
    ck.eq("...and the edge of the band is within_band on both sides", [band_pattern(b, b), band_pattern(-b, b), band_pattern(None, b)], ["within_band", "within_band", "n_a"])
    ck.eq("exactly one row is primary: the percent change of the ratio, current, 2023 to 2025", [(r["statistic"], r["segment"], r["from_year"], r["to_year"]) for r in prim],
          [("change_benefits_per_contrib_pct", "current", 2023, 2025)])
    want_pct = Decimal(100) * (R[2025]["bpc"] / R[2023]["bpc"] - 1)
    ck.eq("the percent change is 100 x (ratio_b / ratio_a - 1)", fdec(prim[0]["value"], 6), fdec(want_pct, 6))
    ck.eq("...and its label is set against the percent band", prim[0]["pattern"], "within_band" if abs(want_pct) <= fb else ("above_band" if want_pct > 0 else "below_band"))
    ck.ok("the absolute change in dollars carries no pattern", all(r["pattern"] == "n_a" for r in S if r["statistic"] == "change_benefits_per_contrib_usd"), "")
    ck.ok("a change in percentage points carries no pattern", all(r["pattern"] == "n_a" for r in S if r["unit"] == "pp"), "")
    ck.ok("every pattern is one of the four labels (no rose or fell)", all(r["pattern"] in ALLOWED_VALUES["pattern"] for r in S), str({r["pattern"] for r in S}))
    band_rows = {r["statistic"]: r["value"] for r in S if r["statistic"].startswith("revision_band")}
    ck.eq("the sample-years band is the largest of the sample years only (2023, 2024), not 2022", fdec(band_rows["revision_band_sample_years_pct"], 3), "0.176")
    ck.eq("...and the full band is the frozen constant", fdec(band_rows["revision_band_pct"], 3), "1.898")
    ck.eq("the measured bands are the largest of all three revisions and of the sample years only", tuple(fdec(x, 3) for x in measured_bands(
        {2022: Decimal("-1.898"), 2023: Decimal("0.102"), 2024: Decimal("0.176")})), ("1.898", "0.176"))
    c3 = world_compute(frozen=Decimal("0.5"), frozen_sample=fs)
    p3 = [r for r in c3["summary"] if r["primary"] == "yes"][0]
    ck.eq("the band is the frozen constant: with a frozen 0.5 the same change is above_band", p3["pattern"], "above_band")
    exp = N[2025]["expenses"]
    want = N[2025]["earnings_ex_gains"] / (N[2025]["benefits"] + N[2025]["withdrawals"] + exp - N[2025]["contrib_total"])
    ck.eq("income per gap after expenses divides by payouts plus RX90 minus contributions", fdec(R[2025]["income_gap_exp"], 6), fdec(want, 6))
    ck.ok("...and differs from income per gap without expenses", R[2025]["income_gap_exp"] != R[2025]["income_gap"], "")


def t_g12(ck):
    ck.eq("SPEC_ALLOW_SHA256 is the sha256 of the canonical text", hashlib.sha256(allow_text().encode("utf-8")).hexdigest(), SPEC_ALLOW_SHA256)
    ck.eq("the canonical text is 1,072 bytes", len(allow_text().encode("utf-8")), 1072)
    for n in ALLOW_ORDER:
        ck.quiet("the exact header of %s passes" % n, lambda n=n: check_header(n, ALLOW[n]))
    ck.raises("an extra column is refused", DataError, lambda: check_header("national.csv", ALLOW["national.csv"] + ["extra"]))
    ren = list(ALLOW["summary.csv"])
    ren[3] = "year"
    ck.raises("a renamed column is refused", DataError, lambda: check_header("summary.csv", ren))
    h, rows = csv_body("# source: x\n" + ",".join(ALLOW["checks.csv"]) + "\nC01,d,e,o,PASS\n")
    ck.eq("the # source: line is skipped, not compared as a header", h, ALLOW["checks.csv"])
    files = dict(get_world()["texts"])
    files["extra.csv"] = "a,b\n1,2\n"
    ck.ok("an eighth csv is reported", any("not one of the nine" in p for p in leak_problems(files)), str(leak_problems(files)))
    dup = dict(get_world()["texts"])
    dup["checks.csv"] = dup["checks.csv"].rstrip("\n") + "\n" + "C01,d,e,o,PASS\n"
    ck.ok("a duplicate key is reported", any("duplicate key" in p for p in dup_key_problems(dup)), str(dup_key_problems(dup)))
    h, rows = csv_body(get_world()["texts"]["summary.csv"])
    bad_stat = [list(r) for r in rows[:1]]
    bad_stat[0][4] = "my_new_statistic"
    ck.ok("a statistic outside its list is reported", any("statistic" in p for p in summary_problems(bad_stat)), str(summary_problems(bad_stat)))
    bad_q = [list(r) for r in rows[:1]]
    bad_q[0][0] = "Q9"
    ck.ok("a question outside its list is reported", any("question" in p for p in summary_problems(bad_q)), str(summary_problems(bad_q)))
    dupe = [list(rows[0]), list(rows[0])]
    ck.ok("a duplicate summary key is reported", any("duplicate key" in p for p in summary_problems(dupe)), str(summary_problems(dupe)))
    ck.eq("the clean synthetic check.sh holds its own copy of the seven lines", checksh_problems(synthetic_checksh()), [])
    ck.ok("a check.sh that reads the headers from the script is reported", bool(checksh_problems(synthetic_checksh(clean=False))))
    ck.ok("a check.sh missing a header is reported", bool(checksh_problems(synthetic_checksh().replace("checks.csv) want=", "checks.csv) wnt="))))
    p = os.path.join(ROOT, "bin", "check.sh")
    if os.path.exists(p) and "cpp_leak_check() {" in read_file(p, "r"):
        ck.eq("bin/check.sh (guard M25) holds the seven canonical lines and does not read the script", checksh_problems(read_file(p, "r")), [])


def synthetic_checksh(clean=True):
    arms = "\n".join("      %s) want='%s' ;;" % tuple(l.split(":", 1)) for l in SPEC_ALLOW_TEXT.splitlines())
    body = "cpp_leak_check() {\n  case $name in\n%s\n  esac\n%s}\n" % (arms, "" if clean else "  want=$(python3 bin/census_public_pensions.py print-allowlist)\n")
    return "#!/bin/bash\n" + body + "\nnext_function() {\n  :\n}\n"


RAW_ROW_PIPE = '2017|SYNTHETIC, CITY FUND|123456|AK|0|RX40|1000|R|V|1'
RAW_ROW_CSV = '2024,"SYNTHETIC, QUOTED FUND",123456,AK,0,RX40,1000,R,V,1'


def t_g13(ck):
    W = get_world()
    clean = dict(W["texts"])
    names = ("SYNTHETIC PLANTED RETIREMENT SYSTEM",)
    ck.eq("on the clean folder the leak guard passes", leak_problems(clean, names), [])
    bad = dict(clean)
    bad["ASPP_Unit_File_2025.dsv"] = "2025|X|1|AK|0|RX40|1|R|V|1\n"
    bad["pasted.csv"] = RAW_ROW_PIPE + "\n"
    bad["README.md"] = clean["README.md"] + "\n| " + RAW_ROW_CSV + " |\n"
    bad["SOURCES.md"] = clean["SOURCES.md"] + "\nThe planted name SYNTHETIC PLANTED RETIREMENT SYSTEM appears here.\nSee /Users/someone/file.\nSAMPLE_YEAR|UNIT_NAME\ncontact: someone@example.org\n"
    probs = leak_problems(bad, names)
    ck.ok("a raw .dsv is reported", any("raw file" in p and ".dsv" in p for p in probs), str(probs))
    ck.ok("a unit-file row pasted into a .csv is reported (a name with a comma included)", any("raw unit-file row in pasted.csv" in p for p in probs), str(probs))
    ck.ok("a quoted-name csv row inside a Markdown table is reported", any("raw unit-file row in README.md" in p for p in probs), str(probs))
    ck.ok("a planted system name is reported", any("system name in SOURCES.md" in p for p in probs), str(probs))
    ck.ok("a /Users/ path in SOURCES.md is reported", any("local path in SOURCES.md" in p for p in probs), str(probs))
    ck.ok("a unit-file header line is reported", any("header line in SOURCES.md" in p for p in probs), str(probs))
    ck.ok("an at sign (a contact line) in SOURCES.md is reported", any("at sign in SOURCES.md" in p for p in probs), str(probs))
    ck.ok("none of the planned output rows matches the raw-row pattern", not any(RAW_ROW_RX.search(t) for t in clean.values()), "")
    ck.ok("the pattern matches both raw forms of the fixtures", bool(RAW_ROW_RX.search(RAW_ROW_PIPE)) and bool(RAW_ROW_RX.search("| " + RAW_ROW_CSV + " |")), "")


def t_g14(ck):
    opening = ("This table shows how much US state and local government pension systems took in and paid out in each survey year from 2013 to 2025. "
               "The project's figures for private defined-contribution plans stop at 2023 and cover different plans on different dates; nothing in this table is added to, "
               "subtracted from or compared in one number with them. Net cash flow is contributions minus benefits minus withdrawals.")
    ck.eq("the opening and the private-plan sentence pass", wording_problems("README.md", opening), [])
    ck.eq("the denial sentence is exempt (it matches 'sale of equities')", wording_problems("README.md", opening + " " + DENIAL_SENTENCE), [])
    ck.ok("the denial sentence is exempt once only: a second copy fails", bool(wording_problems("README.md", opening + " " + DENIAL_SENTENCE + " " + DENIAL_SENTENCE)), "")
    for bad in ("Public pensions were net sellers last year.", "Benefit payments drove equity sales.", "Money flows out of stocks.", "A forced sale of assets followed.",
                "Retirement flows after 2023 are unclear.", "Net cash flow from members is negative.", "The ratio rose in 2025."):
        ck.ok("'%s' fails" % bad, bool(wording_problems("README.md", opening + " " + bad)), "passed")
    ck.eq("the generated README passes the wording guard", wording_problems("README.md", get_world()["texts"]["README.md"]), [])
    ck.eq("...and the generated CSV files do", [p for n, t in get_world()["texts"].items() if n.endswith(".csv") for p in wording_problems(n, t)], [])


FETCH_GOOD_LISTING = (b'<!DOCTYPE HTML PUBLIC "-//W3C//DTD HTML 3.2 Final//EN"><html><head><title>Index of /programs-surveys/aspp/datasets</title></head><body>'
                      b'<a href="2024/">2024/</a><a href="2025/">2025/</a></body></html>')


def make_handler(calls, clock, override=None):
    """A fake server. override: {path: (status, body, headers) or an exception to raise}."""
    def handler(host, target, headers, timeout, cap):
        calls.append((clock.wall(), host, target, dict(headers)))
        if override and target in override:
            o = override[target]
            if isinstance(o, Exception):
                raise o
            return o
        if target == LISTING_PATH:
            return 200, FETCH_GOOD_LISTING, {}
        return 200, b"PK\x03\x04" + target.encode(), {"last-modified": "Fri, 23 May 2025 17:00:31 GMT"}
    return handler


def rig(tmp, override=None, files_max=None, clock=None, **kw):
    clock = clock or FakeClock()
    calls = []
    cache, out = os.path.join(tmp, "cache"), os.path.join(tmp, "out")
    fe = Fetcher(cache, out, stop_dir=cache, max_files=files_max, opener=make_handler(calls, clock, override), wall=clock.wall, mono=clock.mono, sleeper=clock.sleep,
                 now_utc=clock.now_utc, **kw)
    return fe, calls, clock


def stop_file(fe):
    return fe.stop_path()


def t_g15(ck):
    src = read_file(os.path.abspath(__file__), "r")
    ck.eq("the script's own source passes the STOP scan", stop_ast_problems(src), [])
    bad_sources = {
        "os.remove": "import os\ndef f(cache):\n    os.remove(os.path.join(cache, 'STOP_census'))\n",
        "os.unlink via a name": "import os\ndef f(cache):\n    sp = os.path.join(cache, 'STOP_census')\n    os.unlink(sp)\n",
        "os.replace": "import os\ndef f(tmp):\n    os.replace(tmp, STOP_PATH)\n",
        "os.rename": "import os\ndef f(a):\n    os.rename(a, stop_path)\n",
        "Path.unlink": "from pathlib import Path\ndef f(stop):\n    Path(stop).unlink()\n",
        "Path.write_text": "from pathlib import Path\ndef f(stop):\n    Path(stop).write_text('x')\n",
        "Path.write_bytes": "from pathlib import Path\ndef f(stop):\n    Path(stop).write_bytes(b'x')\n",
        "open w": "def f(stop):\n    open(stop, 'w').write('x')\n",
        "open a": "def f(stop):\n    open(stop, 'a').write('x')\n",
        "shutil.move": "import shutil\ndef f(stop):\n    shutil.move(stop, 'x')\n",
        "rmtree over the cache": "import shutil\ndef f(cache):\n    shutil.rmtree(cache)\n",
        "a glob delete": "import glob, os\ndef remove_partial(cache):\n    for f in glob.glob(cache + '/*'):\n        os.remove(f)\n",
        "create_stop with a": "def create_stop(path, text):\n    open(path, 'a').write(text)\n",
    }
    for name, s in bad_sources.items():
        ck.ok("the scan reports: %s" % name, bool(stop_ast_problems(s)), "not reported")
    good = ("import os, shutil\ndef create_stop(path, text):\n    with open(path, 'x') as fh:\n        fh.write(text)\n"
            "def remove_partial(partial):\n    shutil.rmtree(partial)\n"
            "def check(stop):\n    return os.path.exists(stop) and open(stop).read()\n")
    ck.eq("exclusive create, the exact partial removal and a read of the latch are allowed", stop_ast_problems(good), [])
    with tempfile.TemporaryDirectory(prefix="cpp_g15_", ignore_cleanup_errors=True) as tmp:
        fe, calls, clock = rig(os.path.join(tmp, "a"))
        create_stop(stop_file(fe), "status: 429\nutc_time: x\nurl_path: y\nretry_after: 5\nbody_head: z\n")
        before = sha256_of(read_file(stop_file(fe)))
        say = []
        code = fe.run(say.append)
        ck.eq("with a planted STOP_census, fetch makes no request and exits 2", (code, len(calls)), (2, 0))
        ck.eq("...the STOP file is unchanged", sha256_of(read_file(stop_file(fe))), before)
        ck.ok("...and prints STOP present", any(s.startswith("STOP present") for s in say), str(say))
        create_stop(stop_file(fe), "something else\n")
        ck.eq("create_stop never overwrites an existing latch", sha256_of(read_file(stop_file(fe))), before)
        f0 = FILES[0]["path"]
        cases = [("301", (301, b"", {"location": LISTING_PATH})), ("403", (403, b"Forbidden", {})), ("404", (404, b"Not found", {})), ("429", (429, b"slow down", {"retry-after": "60"})),
                 ("an empty body", (200, b"", {"last-modified": "Fri, 23 May 2025 17:00:31 GMT"})), ("an HTML body", (200, b"<html><body>blocked</body></html>", {"last-modified": "Fri, 23 May 2025 17:00:31 GMT"})),
                 ("an HTML body after a space", (200, b"  \n<!DOCTYPE html><html>", {"last-modified": "Fri, 23 May 2025 17:00:31 GMT"}))]
        for i, (label, resp) in enumerate(cases):
            fe, calls, clock = rig(os.path.join(tmp, "b%d" % i), {f0: resp})
            code = fe.run(lambda s: None)
            ok = os.path.exists(stop_file(fe))
            ck.ok("a file reply of %s writes STOP_census, ends the fetch (exit 2) and requests nothing more" % label, ok and code == 2 and len(calls) == 2, "exit %s, %d calls, stop %s" % (code, len(calls), ok))
            if ok:
                ck.ok("...with the five fields in order", stop_text_ok(read_file(stop_file(fe), "r")), read_file(stop_file(fe), "r")[:200])
        fe, calls, clock = rig(os.path.join(tmp, "l1"), {LISTING_PATH: (403, b"no", {})})
        code = fe.run(lambda s: None)
        ck.ok("a refused listing latches STOP before any file is requested", code == 2 and len(calls) == 1 and os.path.exists(stop_file(fe)), "exit %s, %d calls" % (code, len(calls)))
        fe, calls, clock = rig(os.path.join(tmp, "l2"), {LISTING_PATH: (200, b"<html><title>Please verify you are human</title></html>", {})})
        code = fe.run(lambda s: None)
        ck.ok("a listing that is a web page but not the directory index latches STOP", code == 2 and os.path.exists(stop_file(fe)), "exit %s" % code)
        fe, calls, clock = rig(os.path.join(tmp, "l3"))
        ck.ok("the real directory index (a DOCTYPE page with its title) does not latch", fe.run(lambda s: None) == 0 and not os.path.exists(stop_file(fe)), "")
        fe, calls, clock = rig(os.path.join(tmp, "t1"), {f0: SocketFault("timeout")})
        code = fe.run(lambda s: None)
        ck.ok("a timeout exits 3, writes no STOP and retries nothing", code == 3 and not os.path.exists(stop_file(fe)) and len(calls) == 2, "exit %s, %d calls" % (code, len(calls)))
        _, rows = read_csv_rows(os.path.join(tmp, "t1", "out", "fetch_log.csv"))
        ck.ok("...and the timeout is in the request log", any(r[3] == "socket_fault" for r in rows), str(rows[-1:]))
        # a content failure (R3b) is a FAIL row, never a STOP
        W = get_world()
        raw = dict(W["raw"])
        p24 = FILES[[f["year"] for f in FILES if f["kind"] == "unit"].index(2024) + 4]["path"] if False else UNIT_2024_PATH
        b24, lm24 = raw[p24]
        raw[p24] = (b24 + b"\n2024,N,1,XX,0,RX40,5,R,V\n2023,N,1,XX,0,RX40,5,R,V,1\n", lm24)
        V2, fr = world_vintage(raw)
        sd = os.path.join(tmp, "stopdir")
        inp = {"vints": [V2], "fetch_rows": fr, "stop_dir": sd}
        c = compute(inp)
        st = {r[0]: r[4] for r in c["checks"]}
        ck.eq("a field-count and a year failure are FAIL rows in C20", st["C20"], "FAIL")
        ck.ok("...and write no STOP", not os.path.exists(os.path.join(sd, STOP_NAME)), "a STOP file exists")


def t_g16(ck):
    for host in ("api.census.gov", "data.census.gov", "www.nasra.org", "www.federalreserve.gov", "www.sec.gov", "fred.stlouisfed.org", "www.census.gov"):
        ck.raises("%s is refused before any request" % host, Refused, lambda host=host: check_url("https://%s%s" % (host, LISTING_PATH)))
    ck.raises("http (not https) is refused", Refused, lambda: check_url("http://%s%s" % (HOST, LISTING_PATH)))
    ck.raises("a path that is not one of the 21 is refused", Refused, lambda: check_url("https://%s%sdatasets/2019/aspp-historical-datasets/other.zip" % (HOST, BASE)))
    ck.raises("the 2020 unit file is refused", Refused, lambda: check_url("https://%s%s" % (HOST, UNIT_2020_PATH)))
    ck.quiet("a listed file is allowed", lambda: check_url("https://%s%s" % (HOST, FILES[0]["path"])))
    ck.raises("an address in the identification string is refused", Refused, lambda: assert_ua(UA[:-1] + "; a" + chr(64) + "b.c)"))
    ck.quiet("the neutral identification string is accepted", lambda: assert_ua(UA))
    ck.ok("...and holds no at sign", "@" not in UA, UA)
    with tempfile.TemporaryDirectory(prefix="cpp_g16_", ignore_cleanup_errors=True) as tmp:
        fe, calls, clock = rig(os.path.join(tmp, "full"))
        say = []
        code = fe.run(say.append)
        targets = [c[2] for c in calls]
        ck.eq("a full run exits 0 with 21 requests: the listing, then the 20 files, in table order", (code, targets), (0, [LISTING_PATH] + [f["path"] for f in FILES]))
        gaps = [b[0] - a[0] for a, b in zip(calls, calls[1:])]
        ck.ok("requests are at least 3 seconds apart", all(g >= 3.0 for g in gaps), "smallest gap %s" % (min(gaps) if gaps else None))
        ck.ok("every request carries the neutral User-Agent and identity encoding", all(c[3].get("User-Agent") == UA and c[3].get("Accept-Encoding") == "identity" for c in calls), "")
        ck.ok("only www2.census.gov is requested", {c[1] for c in calls} == {HOST}, str({c[1] for c in calls}))
        cache = os.path.join(tmp, "full", "cache")
        subs = [n for n in os.listdir(cache) if KEY_RX.match(n)]
        ck.ok("the vintage folder exists and partial/ is gone", len(subs) == 1 and not os.path.exists(os.path.join(cache, "partial")), str(os.listdir(cache)))
        if subs:
            man = read_manifest(os.path.join(cache, subs[0]))
            key = vintage_key([(p, r["last_modified"], int(r["bytes"])) for p, r in man.items()])
            ck.eq("...named by the key of its manifest, with the request count kept", (subs[0], len(man), Fetcher.run_requests_in(os.path.join(cache, subs[0]))), (key, 20, 21))
        # slices: resume without a repeated request
        fe, calls, clock = rig(os.path.join(tmp, "sl"), files_max=7)
        c1 = fe.run(lambda s: None)
        fe2, calls2, _ = rig(os.path.join(tmp, "sl"), files_max=7, clock=clock)
        c2 = fe2.run(lambda s: None)
        fe3, calls3, _ = rig(os.path.join(tmp, "sl"), files_max=7, clock=clock)
        c3 = fe3.run(lambda s: None)
        allc = [c[2] for c in calls + calls2 + calls3]
        ck.eq("a sliced run resumes and never requests a saved file again (21 requests, none twice)", (c1, c2, c3, len(allc), len(set(allc))), (0, 0, 0, 21, 21))
        # redirects, retries
        f0 = FILES[0]["path"]
        fe, calls, clock = rig(os.path.join(tmp, "rd"), {f0: (301, b"", {"location": LISTING_PATH})})
        fe.run(lambda s: None)
        ck.eq("a redirect is not followed: one request for the file, then STOP", len(calls), 2)
        fe, calls, clock = rig(os.path.join(tmp, "rt"), {f0: (500, b"oops", {})})
        fe.run(lambda s: None)
        ck.eq("a 500 is not retried", len(calls), 2)
        # the request cap
        fe, calls, clock = rig(os.path.join(tmp, "cap"))
        os.makedirs(os.path.join(tmp, "cap", "cache", "partial"))
        write_text(os.path.join(tmp, "cap", "cache", "partial", "requests.txt"), "30\n")
        write_text(os.path.join(tmp, "cap", "cache", "partial", "notice.txt"), "listing years above 2025: none\n")
        code = fe.run(lambda s: None)
        ck.eq("a 31st request is not made", (code, len(calls)), (3, 0))
        ck.raises("a cap above 30 is refused", Refused, lambda: Fetcher(os.path.join(tmp, "c2"), os.path.join(tmp, "o2"), max_requests=31))
        ck.raises("a pause under 3 seconds is refused", Refused, lambda: Fetcher(os.path.join(tmp, "c3"), os.path.join(tmp, "o3"), sleep=1.0))
        # at most one run a month
        fe, calls, clock = rig(os.path.join(tmp, "mo"))
        os.makedirs(os.path.join(tmp, "mo", "out"))
        stamp = clock.now_utc().strftime("%Y-%m-%dT%H:%M:%SZ")
        write_text(os.path.join(tmp, "mo", "out", "fetch_log.csv"), render_csv("fetch_log.csv", [[stamp, HOST, f0, "200", "5", "ab", "2025-05-23T17:00:31Z"]], "# source: x"))
        ck.raises("a second run in the same calendar month is refused", Refused, lambda: fe.run(lambda s: None))
        # a vintage that already exists: no change, partial removed by its exact path
        fe, calls, clock = rig(os.path.join(tmp, "nc"))
        fe.run(lambda s: None)
        key = [n for n in os.listdir(os.path.join(tmp, "nc", "cache")) if KEY_RX.match(n)][0]
        os.makedirs(os.path.join(tmp, "nc", "out"), exist_ok=True)
        write_text(os.path.join(tmp, "nc", "out", "fetch_log.csv"), render_csv("fetch_log.csv", [], "# source: x"))
        fe, calls, _ = rig(os.path.join(tmp, "nc"), clock=clock)
        say = []
        code = fe.run(say.append)
        ck.ok("a run that lands on a vintage already complete prints no change and removes partial/", code == 0 and any("no change" in s for s in say) and not os.path.exists(os.path.join(tmp, "nc", "cache", "partial")), str(say[-2:]))


def t_g17(ck):
    src = read_file(os.path.abspath(__file__), "r")
    ck.eq("the script's own source opens nothing outside its cache and output folder", isolation_problems(src), [])
    ck.ok("a read of data/series.tsv is reported", bool(isolation_problems("import os\ndef f():\n    return open(os.path.join(ROOT, 'data', 'series.tsv')).read()\n")), "not reported")
    ck.ok("a read of a Form 5500 file is reported", bool(isolation_problems("def f():\n    return read_file('data/form5500/x.csv')\n")), "not reported")
    ck.ok("a read of an absolute path is reported", bool(isolation_problems("def f():\n    return open('/etc/hosts').read()\n")), "not reported")
    texts = get_world()["texts"]
    ck.eq("no committed file but checks.csv holds a figure from another source (R10)", pairing_problems(texts), [])
    bad = dict(texts)
    bad["README.md"] = texts["README.md"] + "\nSurvey year 2024 assets are 85.7% of the $5.13 trillion in the NASRA survey.\n"
    ck.ok("a README that sets ASPP assets against the NASRA figure is reported", bool(pairing_problems(bad)), "not reported")
    ck.ok("checks.csv does hold the reference constants", "5.13 trillion" in texts["checks.csv"] and "6,541.4" in texts["checks.csv"], "")


def t_g18(ck):
    W = get_world()
    c = W["c0"]
    N = {y: dict(r) for y, r in c["N"].items()}
    ck.eq("C18 passes on fractions", check_c18(N)[4], "PASS")
    N[2023]["cv_benefits"] = Decimal("0.20546")
    ck.eq("a CV column in percent (0.20546) is a FAIL", check_c18(N)[4], "FAIL")
    R = ratio_rows(c["N"])
    want = Decimal(100) * (c["N"][2024]["cv_benefits"] + c["N"][2024]["cv_contrib_total"])
    ck.eq("the bound is 100 x (CV of benefits + CV of contributions)", fdec(R[2024]["cv_bound"], 6), fdec(want, 6))
    ck.eq("...and 0 in a census year", R[2022]["cv_bound"], Decimal(0))
    P = parse_vintage(W["V"])
    P[2022]["sum"].cv = {"benefits": Decimal("0.01"), "contrib_total": Decimal("0.01"), "assets": Decimal("0.01")}
    Nn = national_rows(P, "x")
    ck.eq("a census-year CV is blank in the national rows even if a workbook prints one", (Nn[2022]["cv_benefits"], Nn[2022]["cv_assets"]), (None, None))
    ck.eq("the 2025 CV is blank (no 2025 workbook)", (c["N"][2025]["cv_benefits"], R[2025]["cv_bound"]), (None, None))


def t_g19(ck):
    paths = [f["path"] for f in build_files()]
    ck.ok("the 2024 unit file is taken from datasets/2024/aspp-historical-datasets/", UNIT_2024_PATH in paths and not any(p in paths for p in STALE_2024_ROOT_PATHS), str([p for p in paths if "2024" in p]))
    for p in STALE_2024_ROOT_PATHS:
        ck.raises("the stale root file %s is refused" % p.rsplit("/", 1)[1], Refused, lambda p=p: check_url("https://%s%s" % (HOST, p)))
    W = get_world()
    ck.eq("C19 passes on the synthetic build", check_c19(W["V"])[4], "PASS")
    raw = dict(W["raw"])
    b, lm = raw[UNIT_2024_PATH]
    raw[UNIT_2024_PATH] = (b, LM_2025)
    V2, _ = world_vintage(raw)
    ck.eq("a 2024 unit file dated before 2026-05-14, on a vintage whose newest unit file is 2026-05-14, is a FAIL", check_c19(V2)[4], "FAIL")
    meta = dict(W["V"].meta)
    meta[STALE_2024_ROOT_PATHS[0]] = meta.pop(UNIT_2024_PATH)
    V3 = Vintage("x" * 16, meta, lambda p: b"")
    ck.eq("a vintage that holds the root-folder 2024 file instead is a FAIL", check_c19(V3)[4], "FAIL")


def t_x_pins(ck):
    W = get_world()
    c = W["c0"]
    st = {r[0]: r[4] for r in c["checks"]}
    ck.eq("every pin check passes on the synthetic build (the pins are the world's numbers)", [st[k] for k in ("C01", "C02", "C03", "C04", "C04p", "C05", "C06", "C07", "C08", "C09", "C10", "C13", "C14", "C15", "C16", "C17")], ["PASS"] * 16)
    N = {y: dict(r) for y, r in c["N"].items()}
    N[2023]["contrib_total"] += Decimal("0.001")
    c2 = dict(c)
    c2["N"] = N
    rows = {r[0]: r for r in build_checks(c2)}
    ck.eq("a pin off by one thousand dollars on a 2026-05-14 file is a FAIL on the vintage whose newest unit file is 2026-05-14", rows["C03"][4], "FAIL")
    V = c["V"]
    p25 = next(f["path"] for f in FILES if f["year"] == 2025 and f["kind"] == "unit")
    later = Vintage("y" * 16, {p: dict(m, last_modified=("2026-06-01T00:00:00Z" if p == p25 else m["last_modified"])) for p, m in V.meta.items()}, V.read)
    ck.eq("a later vintage has a later newest unit file date", later.newest_unit_file_date, "2026-06-01")
    c3 = dict(c2)
    c3["V"] = later
    rows3 = {r[0]: r for r in build_checks(c3)}
    ck.eq("on a later vintage the same mismatch is a WARN where revisions.csv has a row for the cell (2023 contributions)", rows3["C03"][4], "WARN")
    ck.eq("...and an untouched check (C01) still passes there", rows3["C01"][4], "PASS")


def t_x_cache(ck):
    ck.quiet("a cache outside the repository is accepted", lambda: check_cache(os.path.join(tempfile.gettempdir(), "cpp_cache_x")))
    ck.raises("the committed output folder is refused as a cache", Refused, lambda: check_cache(os.path.join(ROOT, "data", "census_public_pensions", "cachex")))
    ck.raises("a path inside the repository that git does not ignore is refused", Refused, lambda: check_cache(os.path.join(ROOT, "bin", "cache_x")))
    ck.quiet("the default cache under data/vintages/ is accepted", lambda: check_cache(CACHE_DEFAULT))


def t_x_build(ck):
    W = get_world()
    if _ACTIVE:
        texts, failed = finish(world_compute())
    else:
        texts, failed = W["texts"], W["failed"]
    ck.eq("the synthetic build fails no check", failed, [])
    ck.eq("every committed file is present, nothing else", sorted(texts), sorted(COMMITTED))
    ck.eq("the folder guards pass on the rendering", all_folder_problems(texts, ()), [])
    rows = parse_checks(texts["checks.csv"])
    ck.eq("C11 and C21 are PASS in the rendering", sorted((r[0], r[4]) for r in rows if r[0] in ("C11", "C21")), [("C11", "PASS"), ("C21", "PASS")])
    rd = texts["README.md"]
    pos = [rd.find(DENIAL_SENTENCE), rd.find("Survey year 2025 is a first release"), rd.find("## Timing"), rd.find("## What is counted"), rd.find("## Full-count years"), rd.find("## Beside the project"),
           rd.find("## Levels"), rd.find("## Ratios and the primary row"), rd.find("## Revisions and vintages"), rd.find("## Column dictionary"), rd.find("## Citation")]
    ck.ok("the README has its sections in the order of the spec", all(p >= 0 for p in pos) and pos == sorted(pos), str(pos))
    prim = [r for r in W["c0"]["summary"] if r["primary"] == "yes"][0]
    ck.eq("(the synthetic primary row is within_band, so both glosses are due)", prim["pattern"], "within_band")
    ck.ok("the README carries the denial sentence once", rd.count(DENIAL_SENTENCE) == 1, str(rd.count(DENIAL_SENTENCE)))
    ck.ok("the opening glosses within_band as not the same as no change", "; that is not the same as no change" in rd, "")
    ck.ok("the primary-row paragraph says within_band does not mean no change", "`within_band` means smaller than one past revision. It does not mean no change." in rd, "")
    low = rd
    for term, gloss in (("unit file", "a unit file is"), ("summary table", "a summary table is"), ("release", "a release is"), ("vintage", "A vintage is")):
        i, j = low.find(term), low.find(gloss)
        ck.ok("'%s' is glossed at its first use" % term, i >= 0 and j >= 0 and abs(j - i) <= 80, "first use %d, gloss %d" % (i, j))
    ck.ok("the README does not use the header name Last-Modified, and says 'Beside the table'", "Last-Modified" not in rd and "Beside the table are" in rd and "Beside it" not in rd, "")
    ck.ok("the README opens with two sentences that name the primary row", rd.split("\n\n")[1].count(". ") >= 1 and "primary row" in rd.split("\n\n")[1], rd.split("\n\n")[1][:200])
    ck.ok("SOURCES.md has a Builder decisions section", "## Builder decisions" in texts["SOURCES.md"], "")
    t2 = render_data_texts(world_compute())
    ck.eq("a second render over the same inputs is byte-identical", {n: sha256_of(t.encode()) for n, t in t2.items() if n in ("national.csv", "ratios.csv", "summary.csv", "revisions.csv")},
          {n: sha256_of(texts[n].encode()) for n in ("national.csv", "ratios.csv", "summary.csv", "revisions.csv")})
    ck.eq("national.csv has exactly 13 rows, survey years 2013 to 2025", [r[0] for r in csv_body(texts["national.csv"])[1]], [str(y) for y in range(2013, 2026)])


def parse_checks(text):
    return csv_body(text)[1]


GUARDS = [
    ("G1", "parsing (csv module, header, encodings, field counts)", ["g1_header_as_data", "g1_skip_first_row", "g1_decode_ignore", "g1_nine_eleven_ok", "g1_dsv_comma", "g1_split_comma"], t_g1),
    ("G2", "weights (each row's own weight, census-year rule, C12 counts only)", ["g2_unweighted", "g2_modal_weight", "g2_weight_twice", "g2_census_weight_ok", "g2_c12_prints_ids"], t_g2),
    ("G3", "line mapping (exact codes, no sub-lines, RX90 apart)", ["g3_rx50_added", "g3_sublines_added", "g3_rx30_in_earnings", "g3_rx91_added", "g3_rx90_in_net", "g3_substring_codes"], t_g3),
    ("G4", "labels and codes (53 codes, labels together, columns by heading)", ["g4_54th_ok", "g4_missing_zero", "g4_unknown_one_at_a_time", "g4_col_by_position"], t_g4),
    ("G5", "signs (RX30 signed, net cash identity, losses, the C09 pin)", ["g5_abs_rx30", "g5_net_benefits_minus_contrib", "g5_withdrawals_added", "g5_losses_gain", "g5_c09_blind"], t_g5),
    ("G6", "units (summary thousands, unit dollars, millions in _usd_mn)", ["g6_summary_as_dollars", "g6_unit_as_thousands", "g6_bn_in_mn"], t_g6),
    ("G7", "breaks (no change in the old years or across 2016/2017, old earnings apart)", ["g7_change_within_old", "g7_change_span_break", "g7_old_earnings_in_exgains", "g7_start_across_break", "g7_income_for_old"], t_g7),
    ("G8", "numbers (Decimal sums, rounding at write time, decimals)", ["g8_float_sum", "g8_round_before_divide", "g8_wrong_decimals"], t_g8),
    ("G9", "vintage key (headers, all files, sizes, UTC)", ["g9_key_wallclock", "g9_key_newest_only", "g9_key_no_size", "g9_lm_local"], t_g9),
    ("G10", "revisions (regenerated, tolerance, different-date gaps, partial)", ["g10_append", "g10_same_date_5k_ok", "g10_diffdate_fail", "g10_partial_as_vintage", "g10_all_years_revised"], t_g10),
    ("G11", "ratios and band (percent, denominators, frozen band, one primary)", ["g11_pct_no_100", "g11_start_same_year", "g11_end_prev_year", "g11_abs_vs_band", "g11_pattern_on_pp", "g11_rose_fell",
                                                                              "g11_band_recomputed", "g11_sample_band_incl_census", "g11_gap_no_rx90", "g11_two_primary"], t_g11),
    ("G12", "allowlists (headers, hash, source line, nine names, check.sh copy, keys, lists)", ["g12_extra_col_ok", "g12_renamed_ok", "g12_allow_drift", "g12_source_as_header", "g12_eighth_csv_ok", "g12_checksh_reads_script",
                                                                                              "g12_dup_key_ok", "g12_stat_ok"], t_g12),
    ("G13", "leak guard (raw files, raw-row pattern, names, paths)", ["g13_raw_ok", "g13_rowpat_off", "g13_rowpat_anchored", "g13_samplehdr_off", "g13_names_off", "g13_path_ok"], t_g13),
    ("G14", "wording (causal, flow, sale, direction verbs; the one denial sentence)", ["g14_causal_off", "g14_extra_off", "g14_denial_loose", "g14_denial_all", "g11_rose_fell_ok"], t_g14),
    ("G15", "STOP latch (AST scan, statuses, bodies, timeout, content failures)", ["g15_scan_off", "g15_ignore_open", "g15_glob_ok", "g15_replace_ok", "g15_stop_ignored", "g15_overwrite", "g15_stop_short", "g15_html_ok",
                                                                               "g15_empty_ok", "g15_3xx_ignored", "g15_stop_on_timeout", "g15_stop_on_content"], t_g15),
    ("G16", "fetch safety (hosts, user-agent, pacing, redirect, retry, cap, table, month)", ["g16_allow_other_host", "g16_at_in_ua_ok", "g16_pacing_off", "g16_redirect_follow", "g16_retry", "g16_cap_off", "g16_url_any_ok",
                                                                                           "g16_2020_ok", "g16_month_gate_off"], t_g16),
    ("G17", "isolation (AST scan, no pairing with another source)", ["g17_scan_off", "g17_pairing_ok"], t_g17),
    ("G18", "CV scale (fractions, bound, census years)", ["g18_cv_percent_ok", "g18_bound_no_100", "g18_census_cv_nonblank"], t_g18),
    ("G19", "2024 path and date", ["g19_root_2024", "g19_stale_ok"], t_g19),
    ("X1", "pins, statuses on a later vintage", [], t_x_pins),
    ("X2", "cache location", [], t_x_cache),
    ("X3", "end-to-end build on the synthetic world (13 rows, sections, glosses, fixed output)", ["x_render_unstable", "x_drop_gloss_open", "x_drop_gloss_dnm"], t_x_build),
]


def run_test(fn):
    ck = Checks()
    try:
        fn(ck)
    except Exception as e:                      # a test that crashes under a mutant has noticed it
        ck.ok("test ran to completion", False, "%s: %s" % (type(e).__name__, e))
    return ck


def selftest(mutants=True, only=None):
    t0 = time.time()
    get_world()                                  # built first, with no mutant active
    bad_total, n_checks, n_mut = 0, 0, 0
    if only:
        names = {m for _, _, ms, _ in GUARDS for m in ms}
        if only not in names:
            print("unknown mutant %r" % only)
            return 2
        with mutant(only):
            for gid, desc, muts, fn in GUARDS:
                ck = run_test(fn)
                bad = ck.failed()
                print("%-4s %s: %d checks, %s" % (gid, desc, len(ck.items), "FAILED (%d)" % len(bad) if bad else "ok"))
                for name, detail in bad[:3]:
                    print("       FAIL  %s   <- %s" % (name, detail[:140]))
                bad_total += len(bad)
        print("selftest with mutant %s: %d failing check(s) -> %s" % (only, bad_total, "the selftest FAILS (mutant caught)" if bad_total else "the selftest PASSES (mutant NOT caught)"))
        return 1 if bad_total else 0
    for gid, desc, muts, fn in GUARDS:
        ck = run_test(fn)
        n_checks += len(ck.items)
        bad = ck.failed()
        bad_total += len(bad)
        print("%-4s %s: %d checks, %s" % (gid, desc, len(ck.items), "ok" if not bad else "FAILED"))
        for name, detail in bad:
            print("       FAIL  %s   <- %s" % (name, detail[:200]))
        if mutants:
            for m in muts:
                with mutant(m):
                    mk = run_test(fn)
                caught = bool(mk.failed())
                n_mut += 1
                bad_total += 0 if caught else 1
                print("       mutant %-28s %s%s" % (m, "caught" if caught else "NOT CAUGHT", "  (first failing check: %s)" % mk.failed()[0][0][:70] if caught else ""))
    print("selftest: %d guards, %d checks, %d mutants applied one at a time, %d failure(s), %.1f s" % (len(GUARDS), n_checks, n_mut, bad_total, time.time() - t0))
    return 0 if bad_total == 0 else 1


# ==== COMMAND LINE ====
def run_fetch(a):
    os.makedirs(a.out, exist_ok=True)
    fe = Fetcher(a.cache, a.out, sleep=a.sleep, max_files=a.max_files)
    code = fe.run()
    print("requests sent in this invocation: %d" % fe.requests)
    return code


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    pf = sub.add_parser("fetch", help="download the inputs into the cache (the only network command; at most 30 requests a run)")
    pf.add_argument("--cache", default=CACHE_DEFAULT)
    pf.add_argument("--out", default=OUT_DEFAULT, help="where fetch_log.csv is kept")
    pf.add_argument("--sleep", type=float, default=DEFAULT_SLEEP)
    pf.add_argument("--max-files", type=int, default=None, help="download at most this many new files in this invocation (a slice); run again to resume")
    pb = sub.add_parser("build", help="offline: validate the cache and write every output")
    pb.add_argument("--cache", default=CACHE_DEFAULT)
    pb.add_argument("--out", default=OUT_DEFAULT)
    ps = sub.add_parser("selftest", help="synthetic data, mutants applied one at a time; no network")
    ps.add_argument("--no-mutants", action="store_true")
    ps.add_argument("--mutant", help="run every guard with this one mutant switched on; exit 1 (the selftest fails) when it is caught")
    a = ap.parse_args(argv)
    try:
        if a.cmd == "selftest":
            return selftest(mutants=not a.no_mutants, only=a.mutant)
        if a.cmd == "build":
            return run_build(a)
        return run_fetch(a)
    except Refused as e:
        print("refused: %s" % e, file=sys.stderr)
        return 5
    except DataError as e:
        print("STOP (data): %s" % e, file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
