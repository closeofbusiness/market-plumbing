#!/usr/bin/env python3
"""R3b: a duration-weighted ("10-year-equivalent") measure of Japanese government bonds, split between the Bank of Japan
and everyone else ("public"), the same method as bin/r2_duration_supply.py (R2) applies to US Treasuries. R3 had only
face-value flows for Japan (Flow of Funds); duration, not face value, is the interest-rate risk supplied. Whether that
supply moved yields is not tested here (charter; C-122).

Method, issue by issue, at each date (copied from R2):
  - Remaining maturity T = (maturity date - date) / 365.25.
  - A yield on the Ministry of Finance's par curve (1-40 years, semiannual, bin/r3_jgb_egb.py loads the same file) on the
    last trading day on or before the date, linearly interpolated at T; flat extrapolation below 1 year and above 40.
  - Clean price and modified duration from the coupon and that yield: R2's closed form, semiannual coupons, n = 2T
    periods, no coupon calendar, no accrued interest (JGBs pay semiannually, so the convention matches the data).
  - Market value = face x price / 100. 10-year-equivalent = market value x modified duration / modified duration of a
    10-year PAR bond at that date's 10-year yield (R2's par_mod_duration). DV01 = market value x modified duration x 0.0001.
  - Two curves, as in R2: each date's own curve, and a FIXED end-2023 curve (29 Dec 2023) with every date's remaining
    maturities, which separates quantity from yield effects.
  - Public = total - Bank of Japan, computed issue by issue (public_i = total_i - BoJ_i, each priced) and asserted to add
    up. If the BoJ holds more of an issue than is outstanding, that is reported, not clipped.

INPUTS, AND THE ONE THING THAT DIFFERS FROM R2
  - BoJ holdings by issue: the Bank of Japan's "Japanese Government Bonds Held by the Bank of Japan", one xlsx per date
    (face value, 100 million yen, interest-bearing JGBs, delivery basis; no T-bills). Observed at every date used.
  - Market stock by issue: the BoJ's JGB Handbook "outstanding of book-entry JGBs by issue" (binzan.xlsx), a month-end
    snapshot in thousand yen. It is a CURRENT-ONLY file: the live copy is 31 Aug 2026; the Internet Archive holds copies
    of the same BoJ file for 31 May 2022, 31 Oct 2024 and 30 Nov 2024. No free official by-issue stock exists for
    29 Dec 2023 (the MoF's by-issue workbook, maturity.xlsx, is GENERAL BONDS ONLY, i.e. it excludes the FILP bonds that
    sit inside the same issues, and exists only for 31 Mar 2026). So the 29 Dec 2023 stock is ESTIMATED, issue by issue
    (estimate_stock): where observed snapshots lie on both sides (31 May 2022 and 31 Oct 2024), the 29 Dec 2023 stock is the
    earlier one plus the issue's own auction flows (MoF auction-results workbook jgb_historical_data.xls: competitive +
    non-competitive amounts) plus the untabulated remainder (liquidity-enhancement auctions, buy-backs, BoJ rollovers) spread
    evenly over time. Issues seen on one side only are moved by auction flows alone; five 2-year issues (about 15 trillion
    yen) are in no snapshot and are priced from auction totals. A one-sided "nearest snapshot" variant is reported as a
    sensitivity. The identical procedure is back-tested against snapshots that ARE observed (jgb_duration_backtest.csv):
    a bracketed hold-out of 31 Oct 2024 errs by +0.2% in face and +0.4% in 10-year-equivalents; one-sided extrapolations
    over 22-29 months err by 2-3%. An independent control, the Flow of Funds market-value stock for 2023Q4 in R3's
    jp_boj_stock.csv, is within 0.2% (BoJ) and 0.15% (total, approximate).
  - The observed-only comparison (31 Oct 2024 to 31 Aug 2026) needs no estimate and is reported alongside.
  - Curves: MoF jgbcme_all.csv / jgbcme.csv (cached by bin/r3_jgb_egb.py in data/vintages/r3/).

UNIVERSE AND EXCLUSIONS (stated on every output)
  Included: fixed-coupon JGBs of 2, 5, 10, 20, 30 and 40 years and the Japan Climate Transition (GX) 5- and 10-year bonds,
  at face value, FILP bonds included (they sit inside the same issues). Excluded, and reported separately as face where
  an observed source exists: inflation-indexed JGBi (the MoF nominal curve is the wrong curve and indexation is ignored),
  the 15-year floating-rate bond, retail JGBs (fixed 3/5-year, floating 10-year; the BoJ holds none), Treasury discount
  bills, and STRIPS (separate lines in the handbook, 0.26 trillion yen in Oct 2024 and Aug 2026, 0.02% of the stock; whether
  the parent issue's line is net of them was not verified; they are ignored). "Public" means everyone other than the Bank of
  Japan within this universe: banks, insurers, pension funds, overseas investors, households and government accounts.

Fetched into data/vintages/r3/dur/ if missing (gitignored through data/vintages/r3/). User-Agent declares the project and
no person. On HTTP 429 the script waits and retries, never caches a rate-limit page, and never routes around a block.
Outputs: data/r3_rates/jgb_duration_supply.csv, jgb_duration_changes.csv, jgb_duration_issues.csv,
         jgb_duration_boj_only.csv, jgb_duration_backtest.csv, jgb_duration_checks.csv
Requires: Python 3, xlrd and openpyxl. Run from anywhere:
  python3 bin/r3b_jgb_duration.py
"""
import csv, datetime as dt, io, os, re, sys, time, unicodedata, urllib.error, urllib.request
from collections import defaultdict

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
R3RAW = os.path.join(ROOT, "data", "vintages", "r3")
RAW = os.path.join(R3RAW, "dur")
OUT = os.path.join(ROOT, "data", "r3_rates")
UA = {"User-Agent": "ThirdDerivativeResearch/1.0"}

MOF_ALL = "https://www.mof.go.jp/english/policy/jgbs/reference/interest_rate/historical/jgbcme_all.csv"
MOF_CUR = "https://www.mof.go.jp/english/policy/jgbs/reference/interest_rate/jgbcme.csv"
MOF_AUCTIONS = "https://www.mof.go.jp/jgbs/reference/appendix/jgb_historical_data.xls"
BOJ_HOLD = "https://www.boj.or.jp/en/statistics/boj/other/mei/release/20{yy}/mei{ymd}.xlsx"
BOJI_LIVE = "https://www.boj.or.jp/note_tfjgs/jgs/data/binzan.xlsx"
BOJI_WB = "https://web.archive.org/web/{ts}id_/https://www.boj.or.jp/note_tfjgs/jgs/data/binzan.xlsx"

D = dt.date
REPORT = [(D(2023, 12, 29), "estimated stock"), (D(2024, 10, 31), "observed"), (D(2026, 8, 31), "observed")]
FIXED_CURVE_DATE = D(2023, 12, 29)          # R2's rule: every date also priced on the end-2023 curve (quantity only)
# BoJ holdings files fetched (BoJ-only table uses all of them; the three REPORT dates and the back-test dates need theirs)
BOJ_DATES = [D(2022, 5, 31), D(2023, 12, 29), D(2024, 10, 31), D(2024, 11, 29), D(2024, 12, 30), D(2025, 12, 30),
             D(2026, 3, 31), D(2026, 6, 30), D(2026, 8, 31), D(2026, 9, 30)]
# Market stock snapshots: (as-of month-end, file, archive timestamp or None for the live file). The as-of date is
# asserted against the date printed in the file.
SNAPS = [(D(2022, 5, 31), "binzan_2022-05.xlsx", "20220713010840"),
         (D(2024, 10, 31), "binzan_2024-10.xlsx", "20241227160511"),
         (D(2024, 11, 30), "binzan_2024-11.xlsx", "20250130090632"),
         (D(2026, 8, 31), "binzan_2026-08.xlsx", None)]

UNIVERSE = ("2", "5", "10", "20", "30", "40", "GX5", "GX10")
# BoJ-handbook labels -> kind
HB_KIND = {"利付国(2年)": "2", "利付国(5年)": "5", "利付国(10年)": "10", "利付国(20年)": "20", "利付国(30年)": "30",
           "利付国(40年)": "40", "GX国債(5年)": "GX5", "GX国債(10年)": "GX10", "利付国(物価10年)": "JGBi",
           "利付国(変動15年)": "FLT", "個人利国(固3年)": "RET", "個人利国(固5年)": "RET", "個人利国(変10年)": "RET"}
HB_IGNORE = ("分離国",)                       # STRIPS, see docstring
# BoJ holdings-file labels (Japanese line) -> kind
BOJ_KIND = {"2年債": "2", "5年債": "5", "10年債": "10", "20年債": "20", "30年債": "30", "40年債": "40",
            "5年クライメート・トランジション国債": "GX5", "10年クライメート・トランジション国債": "GX10",
            "物価連動債": "JGBi", "変動利付債": "FLT"}
# MoF auction-results workbook sheet -> kind
AUCT_SHEET = {"2": "2年債", "5": "5年債", "10": "10年債", "20": "20年債", "30": "30年債", "40": "40年債",
              "GX5": "GX5年債", "GX10": "GX10年債"}
TENORS = ["1Y", "2Y", "3Y", "4Y", "5Y", "6Y", "7Y", "8Y", "9Y", "10Y", "15Y", "20Y", "25Y", "30Y", "40Y"]


# ---------------------------------------------------------------- fetching
def fetch(url, path, wait=3.0, tries=4):
    """Download url to path once and cache it. On HTTP 429 back off (60 s, 120 s, 240 s) and retry, then give up. A body
    that says 'Too Many Requests' is never cached."""
    if os.path.exists(path):
        return True
    os.makedirs(os.path.dirname(path), exist_ok=True)
    for k in range(tries):
        try:
            time.sleep(wait)
            data = urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=120).read()
            if b"Too Many Requests" in data[:2000]:
                raise urllib.error.HTTPError(url, 429, "rate limited", None, None)
            open(path, "wb").write(data)
            return True
        except urllib.error.HTTPError as e:
            if e.code == 429 and k < tries - 1:
                print(f"  [429] {url} -- waiting {60 * 2 ** k}s")
                time.sleep(60 * 2 ** k)
                continue
            print(f"  [skip] {url}: HTTP {e.code}")
            return False
    return False


# ---------------------------------------------------------------- R2's pricing (copied)
def par_mod_duration(y_pct, n=10):
    y = y_pct / 100.0
    return (1.0 / y) * (1.0 - 1.0 / (1.0 + y / 2.0) ** (2 * n))


def coupon_price_duration(coupon_pct, yield_pct, T):
    """R2's closed form: semiannual coupons, continuous n = 2T, clean price per 100 and modified duration (years)."""
    c, y = coupon_pct / 100.0, yield_pct / 100.0
    n = 2.0 * T
    i = y / 2.0
    if abs(i) < 1e-9:
        i = 1e-9 if i >= 0 else -1e-9
    g = c / 2.0
    df_n = (1.0 + i) ** (-n)
    price = 100.0 * g * (1.0 - df_n) / i + 100.0 * df_n
    denom = g * ((1.0 + i) ** n - 1.0) + i
    if abs(denom) < 1e-12:
        denom = 1e-12
    mac_periods = (1.0 + i) / i - ((1.0 + i) + n * (g - i)) / denom
    return price, (mac_periods / 2.0) / (1.0 + i)


def _selftest():
    for c, y, T in ((3.0, 3.0, 10.0), (0.5, 0.5, 7.0), (2.0, 2.0, 0.25), (1.2, 1.2, 30.0)):
        price, dur = coupon_price_duration(c, y, T)
        assert abs(price - 100.0) < 1e-6 and abs(dur - par_mod_duration(y, n=T)) < 1e-9, (c, y, T, price, dur)


_selftest()


def interp_curve(curve, T):
    pts = sorted(curve.items())
    if T <= pts[0][0]:
        return pts[0][1]
    if T >= pts[-1][0]:
        return pts[-1][1]
    for (x0, y0), (x1, y1) in zip(pts, pts[1:]):
        if x0 <= T <= x1:
            return y0 + (T - x0) / (x1 - x0) * (y1 - y0)
    return pts[-1][1]


def on_or_before(seq, d):
    ks = [k for k in seq if k <= d]
    return max(ks) if ks else None


def load_mof_curve():
    fetch(MOF_ALL, os.path.join(R3RAW, "jgbcme_all.csv"))
    fetch(MOF_CUR, os.path.join(R3RAW, "jgbcme.csv"))
    out = {}
    for f in ("jgbcme_all.csv", "jgbcme.csv"):
        for r in csv.reader(io.StringIO(open(os.path.join(R3RAW, f), encoding="utf-8-sig", errors="replace").read())):
            if r and r[0][:4].isdigit() and "/" in r[0]:
                y, m, d = map(int, r[0].split("/"))
                out[D(y, m, d)] = {int(t[:-1]): float(v) for t, v in zip(TENORS, r[1:16]) if v not in ("-", "")}
    return out


# ---------------------------------------------------------------- parsers
def parse_snapshot(path, expect_asof):
    """BoJ JGB Handbook book-entry balance by issue -> ({(kind, issue_no): face_trn_yen}, {kind: face_trn_yen} for the
    excluded kinds with an issue number, strips_trn). Asserts the as-of month printed in the file."""
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True).active
    head = " ".join(str(c) for row in ws.iter_rows(min_row=1, max_row=3, values_only=True) for c in row if c is not None)
    m = re.search(r"(\d{4})年(\d{2})月末", unicodedata.normalize("NFKC", head))
    assert m and (int(m.group(1)), int(m.group(2))) == (expect_asof.year, expect_asof.month), (path, head)
    out, other, strips = defaultdict(float), defaultdict(float), 0.0
    for r in ws.iter_rows(min_row=5, values_only=True):
        if not r[1] or r[3] is None:
            continue
        n = unicodedata.normalize("NFKC", r[2])
        trn = r[3] / 1e9                                   # thousand yen -> trillion yen
        if n.startswith(HB_IGNORE):
            strips += trn
            continue
        mm = re.match(r"(.+?)第(\d+)回", n)
        assert mm and mm.group(1) in HB_KIND, ("unrecognised handbook line", n)
        kind, no = HB_KIND[mm.group(1)], int(mm.group(2))
        if kind in UNIVERSE or kind == "JGBi":
            out[(kind, no)] += trn
        else:
            other[kind] += trn
    return dict(out), dict(other), strips


def parse_boj(path, asof):
    """BoJ holdings file -> {(kind, issue_no): face_trn_yen} (the file is in 100 million yen)."""
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True).active
    assert ws.cell(3, 1).value.date() == asof, (path, ws.cell(3, 1).value)
    cur, out = None, defaultdict(float)
    for r in ws.iter_rows(min_row=14, values_only=True):
        lab = r[2]
        if isinstance(lab, str) and re.search(r"[぀-ヿ一-鿿]", lab):
            cur = lab
        if isinstance(r[3], (int, float)) and isinstance(r[4], (int, float)):
            assert cur in BOJ_KIND, ("unrecognised BoJ issue type", cur)
            out[(BOJ_KIND[cur], int(r[3]))] += r[4] / 1e4
    return dict(out)


def load_auctions():
    """MoF auction-results workbook -> meta {(kind, no): {maturity, coupon, first_issue}} and
    flows {(kind, no): [(issue_date, face_trn_yen)]}: competitive + non-competitive + non-price I + II."""
    import xlrd
    path = os.path.join(RAW, "mof", "jgb_historical_data.xls")   # fresh copy: R3's cached one stops in June 2026
    fetch(MOF_AUCTIONS, path)
    wb = xlrd.open_workbook(path)
    meta, flows = {}, defaultdict(list)
    for kind, name in AUCT_SHEET.items():
        sh = wb.sheet_by_name(name)
        hdr = [str(sh.cell_value(2, c)) for c in range(sh.ncols)]
        amt_cols = [i for i, h in enumerate(hdr) if any(s in h for s in ("落札・割当額", "定率", "第Ⅰ非価格", "第Ⅱ非価格"))]
        assert "落札・割当額" in "".join(hdr[i] for i in amt_cols), name
        for r in range(5, sh.nrows):
            try:
                no = int(float(sh.cell_value(r, 0)))
            except (ValueError, TypeError):
                continue
            idate = xlrd.xldate_as_datetime(sh.cell_value(r, 2), wb.datemode).date()
            mat = xlrd.xldate_as_datetime(sh.cell_value(r, 3), wb.datemode).date()
            a = sum(sh.cell_value(r, c) for c in amt_cols if isinstance(sh.cell_value(r, c), float)) / 1e4
            k = (kind, no)
            flows[k].append((idate, a))
            m = meta.setdefault(k, {"maturity": mat, "coupon": float(sh.cell_value(r, 4)), "first_issue": idate})
            assert m["maturity"] == mat, ("maturity changes within an issue", k)
            m["first_issue"] = min(m["first_issue"], idate)
    return meta, dict(flows)


def load_inputs():
    os.makedirs(RAW, exist_ok=True)
    snaps = {}
    for asof, fn, ts in SNAPS:
        url = BOJI_LIVE if ts is None else BOJI_WB.format(ts=ts)
        p = os.path.join(RAW, "boji", fn)
        if not fetch(url, p, wait=5):
            sys.exit(f"cannot get {fn}")
        snaps[asof] = parse_snapshot(p, asof)
    boj = {}
    for d in BOJ_DATES:
        p = os.path.join(RAW, "boj", f"mei{d:%y%m%d}.xlsx")
        if not fetch(BOJ_HOLD.format(yy=f"{d:%y}", ymd=f"{d:%y%m%d}"), p, wait=2):
            sys.exit(f"cannot get BoJ holdings {d}")
        boj[d] = parse_boj(p, d)
    meta, flows = load_auctions()
    return snaps, boj, meta, flows


# ---------------------------------------------------------------- estimating a stock on an unobserved date
def flow_sum(flows, key, lo, hi):
    """Auction-flow face (trillion yen) with lo < issue date <= hi."""
    return sum(a for d, a in flows.get(key, ()) if lo < d <= hi)


def estimate_stock(T, snaps, meta, flows, exclude=(), mode="interp"):
    """Market stock by issue on an unobserved date T from the observed snapshots (minus `exclude`), issue by issue.
    Every issue alive on T (first issued on or before T, maturing after it) gets one of:
      interp       an observed snapshot on each side of T: a0 + auction flows (t0, T] + r x (T - t0)/(t1 - t0), where
                   r = a1 - a0 - auction flows (t0, t1] is what the workbooks do not tabulate (liquidity-enhancement
                   auctions, buy-backs, BoJ rollovers), spread evenly over time;
      snap-back /  only a later (earlier) snapshot contains it: that stock moved to T by the issue's own auction flows;
      snap-fwd     mode="nearest" uses this for every issue, with the nearest snapshot (the one-sided sensitivity);
      auction-only no snapshot contains it: auction totals up to T (omits the same untabulated flows).
    Returns ({key: face_trn}, {key: source}). Nothing is clipped; a negative result is returned as is."""
    use = {d: s[0] for d, s in snaps.items() if d not in exclude}
    est, src = {}, {}
    for key, m in meta.items():
        if not (m["first_issue"] <= T < m["maturity"]):
            continue
        before = [d for d in use if d <= T and key in use[d]]
        after = [d for d in use if d > T and key in use[d]]
        if mode == "interp" and before and after:
            t0, t1 = max(before), min(after)
            a0, a1 = use[t0][key], use[t1][key]
            r = a1 - a0 - flow_sum(flows, key, t0, t1)
            est[key] = a0 + flow_sum(flows, key, t0, T) + r * (T - t0).days / (t1 - t0).days
            src[key] = f"interp:{t0.isoformat()}..{t1.isoformat()}"
        elif before or after:
            cands = [(abs((d - T).days), -d.toordinal(), d) for d in before + after]
            d = min(cands)[2]
            a = use[d][key] + (-flow_sum(flows, key, T, d) if d > T else flow_sum(flows, key, d, T))
            est[key], src[key] = a, ("snap-back:" if d > T else "snap-fwd:") + d.isoformat()
        else:
            est[key], src[key] = flow_sum(flows, key, D(1900, 1, 1), T), "auction-only"
    return est, src


# ---------------------------------------------------------------- aggregation
def curves_for(curves, d):
    cd = on_or_before(curves, d)
    return cd, curves[cd]


def price_issue(meta_k, d, curve):
    T = max((meta_k["maturity"] - d).days / 365.25, 0.0)
    return T, coupon_price_duration(meta_k["coupon"], interp_curve(curve, T), T)


def aggregate(stock, meta, d, own_curve, fixed_curve):
    """stock: {key: face_trn}. Returns dict of sums: face, mv, eq10 (own curve), eq10_fx (fixed curve), dv01 (yen bn/bp, own),
    plus the denominators used."""
    d10 = par_mod_duration(own_curve[10], 10)
    d10_fx = par_mod_duration(fixed_curve[10], 10)
    s = defaultdict(float)
    for key, face in stock.items():
        if key[0] not in UNIVERSE:
            continue
        T, (price, dur) = price_issue(meta[key], d, own_curve)
        mv = face * price / 100.0
        _, (price_fx, dur_fx) = price_issue(meta[key], d, fixed_curve)
        mv_fx = face * price_fx / 100.0
        s["face"] += face
        s["mv"] += mv
        s["eq10"] += mv * dur / d10
        s["eq10_fx"] += mv_fx * dur_fx / d10_fx
        s["dv01"] += mv * dur * 0.1                     # trillion yen x years x 1e-4 = trillion; x 1000 = billion yen per bp
        s["pardur"] += face * dur
    s["d10"], s["d10_fx"] = d10, d10_fx
    return dict(s)


def split(total, boj):
    """Issue-by-issue public = total - BoJ. Returns (public {key: face}, violations [(key, boj, total)]). No clipping."""
    pub, viol = {}, []
    for key, t in total.items():
        b = boj.get(key, 0.0)
        pub[key] = t - b
        if b > t + 1e-12:
            viol.append((key, b, t))
    for key, b in boj.items():
        if key[0] in UNIVERSE and key not in total:
            viol.append((key, b, 0.0))
            pub[key] = -b
    return pub, viol


def fmt_key(k):
    return f"{k[0]}#{k[1]}"




def read_fof_control(label_quarter):
    """FoF stocks (market value, trillion yen) from R3's committed table data/r3_rates/jp_boj_stock.csv for one quarter."""
    p = os.path.join(OUT, "jp_boj_stock.csv")
    if not os.path.exists(p):
        return None
    for r in csv.DictReader(open(p)):
        if r["quarter"] == label_quarter:
            return float(r["boj_holdings_mv_tril_yen"]), float(r["outstanding_mv_tril_yen"])
    return None


def write_csv(name, rows):
    with open(os.path.join(OUT, name), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)


def main():
    os.makedirs(OUT, exist_ok=True)
    snaps, boj, meta, flows = load_inputs()
    curves = load_mof_curve()
    _, fixed_curve = curves_for(curves, FIXED_CURVE_DATE)
    checks = []

    def chk(cid, date, label, ours, ref, tol, note=""):
        diff = ours - ref
        checks.append({"check_id": cid, "date": date, "label": label, "ours": round(ours, 6), "reference": round(ref, 6),
                       "diff": round(diff, 6), "tolerance": tol, "result": "PASS" if abs(diff) <= tol else "FAIL", "note": note})

    for d, b in boj.items():
        missing = [k for k in b if k[0] in UNIVERSE and k not in meta]
        chk("0_boj_issue_has_terms", d.isoformat(), "BoJ-held fixed-coupon issues missing coupon/maturity in the MoF auction workbook (count)",
            float(len(missing)), 0.0, 0.0, ", ".join(map(fmt_key, missing)))

    # ---- the scenarios: (label, date, how the market stock is obtained)
    SCEN = [("2023-12-29 estimated", D(2023, 12, 29), "interp"),
            ("2023-12-29 estimated, nearest-snapshot sensitivity", D(2023, 12, 29), "nearest"),
            ("2024-10-31 observed", D(2024, 10, 31), "obs"),
            ("2026-08-31 observed", D(2026, 8, 31), "obs")]
    per, rows, issue_rows = {}, [], []
    for label, d, mode in SCEN:
        if mode == "obs":
            total = {k: v for k, v in snaps[d][0].items() if k[0] in UNIVERSE}
            src = {k: f"snapshot:{d.isoformat()}" for k in total}
            memo = {"jgbi": sum(v for k, v in snaps[d][0].items() if k[0] == "JGBi"), "retail": snaps[d][1].get("RET", 0.0)}
        else:
            est, src = estimate_stock(d, snaps, meta, flows, mode=mode)
            total = {k: v for k, v in est.items() if k[0] in UNIVERSE}
            memo = {"jgbi": None, "retail": None}
        cd, own = curves_for(curves, d)
        b = {k: v for k, v in boj[d].items() if k[0] in UNIVERSE}
        pub, viol = split(total, b)
        A = {"total": aggregate(total, meta, d, own, fixed_curve), "boj": aggregate(b, meta, d, own, fixed_curve),
             "public": aggregate(pub, meta, d, own, fixed_curve)}
        for col in ("face", "mv", "eq10", "eq10_fx", "dv01", "pardur"):
            chk("1_identity_public_plus_boj_eq_total", label, f"public + BoJ = total, {col}", A["public"][col] + A["boj"][col],
                A["total"][col], 1e-9)
        chk("1_total_tie_out", label, "sum of issue-level total face = aggregate total face (trn yen)", sum(total.values()),
            A["total"]["face"], 1e-9)
        if mode == "obs":
            chk("1_total_tie_out_to_file", label, "universe face in the handbook file, summed independently of the pricing loop (trn yen)",
                sum(v for k, v in snaps[d][0].items() if k[0] in UNIVERSE), A["total"]["face"], 1e-9)
        chk("1_boj_tie_out_to_file", label, "BoJ universe face in the BoJ file, summed independently of the pricing loop (trn yen)",
            sum(v for k, v in boj[d].items() if k[0] in UNIVERSE), A["boj"]["face"], 1e-9)
        chk("2_boj_exceeds_issue_outstanding", label, "issues where the BoJ holding exceeds the issue's stock (count)", float(len(viol)), 0.0, 0.0,
            "; ".join(f"{fmt_key(k)} boj={bb:.4f} total={tt:.4f}" for k, bb, tt in viol))
        negs = [k for k, v in total.items() if v < 0]
        chk("2_negative_stock", label, "issues with a negative stock (count)", float(len(negs)), 0.0, 0.0, ", ".join(map(fmt_key, negs)))
        if mode != "obs":
            kinds = {}
            for k, s in src.items():
                if k in total:
                    kinds.setdefault(s.split(":")[0], [0, 0.0])
                    kinds[s.split(":")[0]][0] += 1
                    kinds[s.split(":")[0]][1] += total[k]
            chk("3_estimate_source_mix", label, "face of issues with no snapshot at all (auction totals only), trn yen",
                kinds.get("auction-only", [0, 0.0])[1], 0.0, float("inf"),
                "; ".join(f"{s}: {n} issues, {f:.1f}trn ({f / A['total']['face'] * 100:.1f}%)" for s, (n, f) in sorted(kinds.items())))
        per[label] = {"A": A, "date": d, "mode": mode, "viol": viol}
        for key in sorted(set(total) | set(b), key=lambda k: (UNIVERSE.index(k[0]), k[1])):
            T, (price, dur) = price_issue(meta[key], d, own)
            issue_rows.append({"scenario": label, "date": d.isoformat(), "kind": key[0], "issue": key[1],
                               "coupon_pct": meta[key]["coupon"], "maturity": meta[key]["maturity"].isoformat(),
                               "years_to_maturity": round(T, 3), "stock_source": src.get(key, "not in stock"),
                               "total_face_trn": round(total.get(key, 0.0), 6), "boj_face_trn": round(b.get(key, 0.0), 6),
                               "public_face_trn": round(pub[key], 6), "price": round(price, 3), "mod_duration": round(dur, 4)})
        r = {"scenario": label, "date": d.isoformat(), "market_stock": "observed" if mode == "obs" else "ESTIMATED (" + mode + ")",
             "mof_curve_date": cd.isoformat(), "y10_pct": own[10], "d10_own_curve": round(A["total"]["d10"], 4),
             "d10_fixed_curve": round(A["total"]["d10_fx"], 4)}
        for who in ("total", "boj", "public"):
            a = A[who]
            r[f"{who}_face_trn"] = round(a["face"], 3)
            r[f"{who}_mv_trn"] = round(a["mv"], 3)
            r[f"{who}_10yeq_own_trn"] = round(a["eq10"], 3)
            r[f"{who}_10yeq_fixed_trn"] = round(a["eq10_fx"], 3)
            r[f"{who}_dv01_bn_per_bp"] = round(a["dv01"], 2)
            r[f"{who}_wtd_mod_duration_yrs"] = round(a["pardur"] / a["face"], 3)
        r["boj_share_face_pct"] = round(A["boj"]["face"] / A["total"]["face"] * 100, 2)
        r["boj_share_10yeq_own_pct"] = round(A["boj"]["eq10"] / A["total"]["eq10"] * 100, 2)
        r["boj_share_10yeq_fixed_pct"] = round(A["boj"]["eq10_fx"] / A["total"]["eq10_fx"] * 100, 2)
        r["n_issues_boj_gt_stock"] = len(viol)
        r["memo_excluded_JGBi_market_face_trn"] = "" if memo["jgbi"] is None else round(memo["jgbi"], 3)
        r["memo_excluded_JGBi_boj_face_trn"] = round(sum(v for k, v in boj[d].items() if k[0] == "JGBi"), 3)
        r["memo_excluded_retail_market_face_trn"] = "" if memo["retail"] is None else round(memo["retail"], 3)
        rows.append(r)
    write_csv("jgb_duration_supply.csv", rows)
    write_csv("jgb_duration_issues.csv", issue_rows)

    # ---- independent control for 2023Q4: Flow of Funds stocks (market value, central government securities and FILP bonds, no T-bills)
    fof = read_fof_control("202304")
    if fof:
        lab = "2023-12-29 estimated"
        A = per[lab]["A"]
        jgbi_boj_mv = sum(v for k, v in boj[D(2023, 12, 29)].items() if k[0] == "JGBi")      # face; indexation ignored, price ~ par
        chk("4_boj_mv_vs_flow_of_funds_2023Q4", lab, "BoJ universe market value + JGBi at face vs FoF BoJ holdings 202304 (trn yen)",
            A["boj"]["mv"] + jgbi_boj_mv, fof[0], 0.01 * fof[0], "FoF stock is market value; JGBi added at face (indexation and price ignored)")
        s24 = snaps[D(2024, 10, 31)]
        memo_other = s24[0] and sum(v for k, v in s24[0].items() if k[0] == "JGBi") + s24[1].get("RET", 0.0)
        chk("4_total_mv_vs_flow_of_funds_2023Q4", lab, "estimated universe market value + JGBi and retail at their 31 Oct 2024 face vs FoF outstanding 202304 (trn yen)",
            A["total"]["mv"] + memo_other, fof[1], 0.02 * fof[1],
            "approximate: JGBi and retail are not estimated at end-2023, their Oct 2024 face (%.1ftrn) is used; the FoF total is market value" % memo_other)

    # ---- changes
    PER = [("headline: estimated end-2023 stock to 2026-08-31", "2023-12-29 estimated", "2026-08-31 observed"),
           ("observed only: 2024-10-31 to 2026-08-31", "2024-10-31 observed", "2026-08-31 observed"),
           ("sensitivity: end-2023 stock by nearest snapshot, to 2026-08-31", "2023-12-29 estimated, nearest-snapshot sensitivity", "2026-08-31 observed"),
           ("estimated end-2023 stock to 2024-10-31", "2023-12-29 estimated", "2024-10-31 observed")]
    chg = []
    for name, la, lb in PER:
        A, B = per[la]["A"], per[lb]["A"]
        c = {"period": name, "from": per[la]["date"].isoformat(), "to": per[lb]["date"].isoformat(),
             "start_stock": per[la]["mode"], "end_stock": per[lb]["mode"]}
        for basis, col in (("face_trn", "face"), ("10yeq_own_trn", "eq10"), ("10yeq_fixed_trn", "eq10_fx"), ("dv01_bn_per_bp", "dv01")):
            dT, dB, dP = (B[w][col] - A[w][col] for w in ("total", "boj", "public"))
            c[f"d_total_{basis}"], c[f"d_boj_{basis}"], c[f"d_public_{basis}"] = round(dT, 3), round(dB, 3), round(dP, 3)
            c[f"public_pct_change_{basis}"] = round((B["public"][col] / A["public"][col] - 1) * 100, 2)
            c[f"total_pct_change_{basis}"] = round((B["total"][col] / A["total"][col] - 1) * 100, 2)
            c[f"boj_d_over_total_d_pct_{basis}"] = round(dB / dT * 100, 1) if dT else ""
            chk("1_identity_changes", name, f"d_public + d_boj = d_total, {basis}", dP + dB, dT, 1e-9)
        chg.append(c)
    write_csv("jgb_duration_changes.csv", chg)

    # ---- BoJ-only series: observed at every BoJ date, no market stock needed
    bo = []
    for d in BOJ_DATES:
        cd, own = curves_for(curves, d)
        A = aggregate({k: v for k, v in boj[d].items() if k[0] in UNIVERSE}, meta, d, own, fixed_curve)
        bo.append({"date": d.isoformat(), "mof_curve_date": cd.isoformat(), "boj_face_trn": round(A["face"], 3),
                   "boj_mv_trn": round(A["mv"], 3), "boj_10yeq_own_trn": round(A["eq10"], 3),
                   "boj_10yeq_fixed_trn": round(A["eq10_fx"], 3), "boj_dv01_bn_per_bp": round(A["dv01"], 2),
                   "boj_wtd_mod_duration_yrs": round(A["pardur"] / A["face"], 3),
                   "memo_excluded_JGBi_face_trn": round(sum(v for k, v in boj[d].items() if k[0] == "JGBi"), 3),
                   "memo_excluded_floating15y_face_trn": round(sum(v for k, v in boj[d].items() if k[0] == "FLT"), 3)})
    write_csv("jgb_duration_boj_only.csv", bo)

    # ---- back-test of the estimator on snapshots that ARE observed
    bt = []
    TESTS = [("interior (snapshots on both sides)", D(2024, 10, 31), (D(2024, 10, 31), D(2024, 11, 30)), "interp"),
             ("interior (snapshots on both sides)", D(2024, 11, 30), (D(2024, 11, 30), D(2024, 10, 31)), "interp"),
             ("interior, nearest-snapshot variant", D(2024, 10, 31), (D(2024, 10, 31), D(2024, 11, 30)), "nearest"),
             ("edge (later snapshots only)", D(2022, 5, 31), (D(2022, 5, 31),), "interp"),
             ("edge (earlier snapshots only)", D(2026, 8, 31), (D(2026, 8, 31), D(2024, 11, 30)), "interp")]
    for kind, Tb, ex, mode in TESTS:
        est, src = estimate_stock(Tb, snaps, meta, flows, exclude=ex, mode=mode)
        est = {k: v for k, v in est.items() if k[0] in UNIVERSE}
        act = {k: v for k, v in snaps[Tb][0].items() if k[0] in UNIVERSE}
        _, own = curves_for(curves, Tb)
        Ae, Aa = aggregate(est, meta, Tb, own, fixed_curve), aggregate(act, meta, Tb, own, fixed_curve)
        auc = sum(est[k] for k, s in src.items() if s == "auction-only" and k in est)
        bt.append({"test": kind, "target_date": Tb.isoformat(), "snapshots_available": " ".join(sorted(d.isoformat() for d in snaps if d not in ex)),
                   "mode": mode, "observed_face_trn": round(Aa["face"], 2), "estimated_face_trn": round(Ae["face"], 2),
                   "face_error_pct": round((Ae["face"] / Aa["face"] - 1) * 100, 2),
                   "observed_10yeq_own_trn": round(Aa["eq10"], 2), "estimated_10yeq_own_trn": round(Ae["eq10"], 2),
                   "eq10_own_error_pct": round((Ae["eq10"] / Aa["eq10"] - 1) * 100, 2),
                   "eq10_fixed_error_pct": round((Ae["eq10_fx"] / Aa["eq10_fx"] - 1) * 100, 2),
                   "dv01_error_pct": round((Ae["dv01"] / Aa["dv01"] - 1) * 100, 2),
                   "issue_level_abs_error_trn": round(sum(abs(est.get(k, 0.0) - act.get(k, 0.0)) for k in set(est) | set(act)), 2),
                   "auction_only_face_trn": round(auc, 2),
                   "observed_issues_missing_from_estimate": len([k for k in act if k not in est])})
    write_csv("jgb_duration_backtest.csv", bt)

    write_csv("jgb_duration_checks.csv", checks)
    fails = [c for c in checks if c["result"] == "FAIL"]
    print(f"[done] {len(checks)} checks, {len(fails)} FAIL")
    for c in fails:
        print("  FAIL", c["check_id"], c["date"], c["label"], "|", c["note"][:200])
    assert not [c for c in checks if c["check_id"].startswith("1_") and c["result"] == "FAIL"], "identity check failed"


if __name__ == "__main__":
    main()
